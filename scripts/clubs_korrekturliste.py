#!/usr/bin/env python3
"""Korrekturliste der Clubs als Excel – für die Prüfung durch den Nutzer und zurück.

Vom Nutzer am 19.10.2026 gewünscht: „Sobald du dir bei einem Club unsicher
bist, dann bitte in eine Korrekturliste eintragen und dann schau ich mir das
an." Unsichere Clubs tragen in `clubs.json` `pruefung_offen: true` (samt
`pruefung_grund`) und sind auf der Seite AUSGEBLENDET (`EC.loadClubs()`),
bis sie hier geprüft sind – dieselbe Linie wie bei den Events.

    python3 scripts/clubs_korrekturliste.py export              # -> clubs_korrekturliste.xlsx
    python3 scripts/clubs_korrekturliste.py import datei.xlsx --dry-run
    python3 scripts/clubs_korrekturliste.py import datei.xlsx

Spalte „Status": leer = noch offen, „Stimmt so" / „Korrigiert" = freigeben
(geänderte Zellen werden übernommen), „Löschen" = Club entfernen. Die
Termine stehen als Text, eine Einheit je Zeile:

    wöchentlich Di 19:00
    alle 2 Wochen Sa 09:00 ab 2026-10-03
    monatlich 1. So 10:00          (auch: monatlich letzter Fr 18:30)
    einmalig 2026-10-28 17:30 (Community Run)
    wöchentlich Sa 09:30 (Long Run)

Verglichen wird gegen das ausgeblendete Blatt `_original` (Stand beim
Export) – nur was der Nutzer geändert hat, wird übernommen.
"""
from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("openpyxl fehlt: pip install openpyxl")

WURZEL = Path(__file__).resolve().parent.parent
CLUBS_JSON = WURZEL / "clubs.json"
STANDARD_DATEI = WURZEL / "clubs_korrekturliste.xlsx"

TAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
STATUS = ["Stimmt so", "Korrigiert", "Löschen"]

# (Spaltenkopf, Feld, editierbar)
SPALTEN = [
    ("Status", None, True),
    ("Grund der Prüfung", "pruefung_grund", False),
    ("Name", "name", True),
    ("Sportart", "art", True),
    ("Stadt", "stadt", True),
    ("Land", "land", True),
    ("Treffpunkt", "treffpunkt", True),
    ("Termine", "termine", True),
    ("Distanz km (von–bis)", "distanz_km", True),
    ("Tempo (von–bis)", "pace", True),
    ("Kosten", "kosten", True),
    ("Anmeldung (nein/empfohlen/pflicht)", "anmeldung", True),
    ("Webseite", "webseite", True),
    ("Instagram", "instagram", True),
    ("Strava", "strava", True),
    ("Quelle", "quelle", False),
    ("Anmerkung", None, True),
    ("id", "id", False),
]


# ---------- Text <-> Daten ----------

def termin_text(t: dict) -> str:
    if t["rhythmus"] == "einmalig":
        s = f"einmalig {t['datum']} {t.get('uhrzeit', '')}".strip()
        return s + (f" ({t['bezeichnung']})" if t.get("bezeichnung") else "")
    tag = TAGE[t["wochentag"] - 1]
    if t["rhythmus"] == "zweiwoechentlich":
        s = f"alle 2 Wochen {tag} {t.get('uhrzeit', '')}".strip() + (f" ab {t['ab']}" if t.get("ab") else "")
    elif t["rhythmus"] == "monatlich":
        n = "letzter" if t.get("nter") == -1 else f"{t.get('nter', 1)}."
        s = f"monatlich {n} {tag} {t.get('uhrzeit', '')}".strip()
    else:
        s = f"wöchentlich {tag} {t.get('uhrzeit', '')}".strip()
    if t.get("bezeichnung"):
        s += f" ({t['bezeichnung']})"
    return s


_TERMIN_RE = re.compile(
    r"^(?P<rh>wöchentlich|woechentlich|jeden|alle 2 wochen|zweiwöchentlich|zweiwoechentlich|monatlich)\s+"
    r"(?:(?P<n>letzter|letzte|[1-4])\.?\s+)?(?P<tag>mo|di|mi|do|fr|sa|so)\w*\.?\s+(?P<zeit>\d{1,2}[:.]\d{2})"
    r"(?:\s+ab\s+(?P<ab>\d{4}-\d{2}-\d{2}))?\s*(?:\((?P<bez>.*)\))?$", re.I)


_EINMALIG_RE = re.compile(r"^einmalig\s+(?P<datum>\d{4}-\d{2}-\d{2})\s+(?P<zeit>\d{1,2}[:.]\d{2})\s*(?:\((?P<bez>.*)\))?$", re.I)


def termine_aus_text(text: str, alte: list[dict]) -> list[dict]:
    """Liest die Termin-Zellen zurück; Felder der alten Einheit gleicher
    Stelle (distanz_km, pace) bleiben erhalten. Wirft ValueError."""
    neue = []
    for i, zeile in enumerate(z.strip() for z in re.split(r"\n+", text or "") if z.strip()):
        e = _EINMALIG_RE.match(zeile)
        if e:
            datetime.date.fromisoformat(e["datum"])
            h, mi = re.split(r"[:.]", e["zeit"])
            t = {"rhythmus": "einmalig", "datum": e["datum"], "uhrzeit": f"{int(h):02d}:{mi}"}
            if e["bez"]:
                t["bezeichnung"] = e["bez"].strip()
            neue.append(t)
            continue
        m = _TERMIN_RE.match(zeile)
        if not m:
            raise ValueError(f"Termin nicht lesbar: {zeile!r}")
        rh = m["rh"].lower()
        rhythmus = ("monatlich" if rh.startswith("monat") else
                    "zweiwoechentlich" if ("2" in rh or "zwei" in rh) else "woechentlich")
        h, mi = re.split(r"[:.]", m["zeit"])
        t = dict(alte[i]) if i < len(alte) else {}
        for k in ("ab", "nter", "bezeichnung"):
            t.pop(k, None)
        t.update({"rhythmus": rhythmus, "wochentag": TAGE.index(m["tag"][:2].capitalize()) + 1,
                  "uhrzeit": f"{int(h):02d}:{mi}"})
        if rhythmus == "zweiwoechentlich":
            if not m["ab"]:
                raise ValueError(f"„alle 2 Wochen“ braucht „ab JJJJ-MM-TT“: {zeile!r}")
            if datetime.date.fromisoformat(m["ab"]).isoweekday() != t["wochentag"]:
                raise ValueError(f"Das ab-Datum fällt nicht auf {m['tag']}: {zeile!r}")
            t["ab"] = m["ab"]
        if rhythmus == "monatlich":
            t["nter"] = -1 if (m["n"] or "").lower().startswith("letzt") else int(m["n"] or 1)
        if m["bez"]:
            t["bezeichnung"] = m["bez"].strip()
        neue.append(t)
    if not neue:
        raise ValueError("Keine Termine angegeben")
    return neue


def spanne_text(werte) -> str:
    if not werte:
        return ""
    if isinstance(werte, dict):
        werte = [werte.get("von"), werte.get("bis")]
    a, b = werte[0], werte[1] if len(werte) > 1 else werte[0]
    fmt = lambda v: str(v).replace(".", ",") if isinstance(v, float) else str(v)  # noqa: E731
    return fmt(a) if a == b or b is None else f"{fmt(a)}–{fmt(b)}"


def spanne_aus_text(text: str, art: str | None, als_pace: bool):
    text = (text or "").strip()
    if not text:
        return None
    teile = [t.strip() for t in re.split(r"\s*[–-]\s*|\s+bis\s+", text) if t.strip()]
    if len(teile) == 1:
        teile *= 2

    def wert(t):
        t = re.sub(r"\s*(km/h|min/km|min/100\s*m|km)$", "", t)
        if als_pace and art != "Fahrrad":
            if not re.fullmatch(r"\d{1,2}:[0-5]\d", t):
                raise ValueError(f"Tempo als m:ss: {text!r}")
            return t
        z = float(t.replace(",", "."))
        return int(z) if z.is_integer() else z

    von, bis = wert(teile[0]), wert(teile[1])
    return {"von": von, "bis": bis} if als_pace else [von, bis]


def kosten_text(k) -> str:
    if not k:
        return ""
    if k.get("gratis"):
        return "kostenlos" + (f" ({k['text']})" if k.get("text") else "")
    return k.get("text") or (f"{k['betrag_eur']} €" if k.get("betrag_eur") is not None else "")


def kosten_aus_text(text: str):
    text = (text or "").strip()
    if not text:
        return None
    if re.match(r"^(kostenlos|gratis|frei|free)\b", text, re.I):
        rest = re.sub(r"^(kostenlos|gratis|frei|free)\s*", "", text, flags=re.I).strip(" ()")
        return {"gratis": True, "betrag_eur": None, "text": rest or None}
    m = re.fullmatch(r"(\d+(?:[.,]\d+)?)\s*(€|eur|euro)", text, re.I)
    if m:
        return {"gratis": False, "betrag_eur": float(m[1].replace(",", ".")), "text": None}
    return {"gratis": False, "betrag_eur": None, "text": text}


def zelle(club: dict, feld: str | None) -> str:
    if feld is None:
        return ""
    v = club.get(feld)
    if feld == "termine":
        return "\n".join(termin_text(t) for t in v or [])
    if feld in ("distanz_km", "pace"):
        return spanne_text(v)
    if feld == "kosten":
        return kosten_text(v)
    if feld == "quelle":
        return "\n".join(v or [])
    return "" if v is None else str(v)


# ---------- Export / Import ----------

def lade() -> list[dict]:
    return json.loads(CLUBS_JSON.read_text(encoding="utf-8"))


def speichere(clubs: list[dict]) -> None:
    CLUBS_JSON.write_text(json.dumps(clubs, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def export(ziel: Path) -> int:
    offen = [c for c in lade() if c.get("pruefung_offen")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Clubs prüfen"
    ws.append([s[0] for s in SPALTEN])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B1B33")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    orig = wb.create_sheet("_original")
    orig.append([s[0] for s in SPALTEN])
    grau = PatternFill("solid", fgColor="EEEEEE")
    for club in sorted(offen, key=lambda c: (c.get("land", ""), c.get("stadt", ""), c.get("name", ""))):
        werte = [zelle(club, f) for _, f, _ in SPALTEN]
        ws.append(werte)
        orig.append(werte)
        for i, (_, _, edit) in enumerate(SPALTEN):
            z = ws.cell(row=ws.max_row, column=i + 1)
            z.alignment = Alignment(wrap_text=True, vertical="top")
            if not edit:
                z.fill = grau
    breiten = [12, 40, 30, 11, 14, 14, 34, 30, 12, 14, 20, 14, 34, 34, 34, 40, 30, 28]
    for i, b in enumerate(breiten):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = b
    ws.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{max(ws.max_row, 2)}")
    orig.sheet_state = "hidden"
    hilfe = wb.create_sheet("Anleitung")
    for zeile in [
        "Jede Zeile ist ein Club, bei dem etwas unsicher ist (Spalte „Grund der Prüfung“).",
        "Diese Clubs sind auf der Seite AUSGEBLENDET, bis du sie prüfst.",
        "Status: „Stimmt so“ oder „Korrigiert“ = freigeben (geänderte Zellen werden übernommen), „Löschen“ = Club entfernen, leer = bleibt offen.",
        "Termine: eine Einheit je Zeile, z. B. „wöchentlich Di 19:00“, „alle 2 Wochen Sa 09:00 ab 2026-10-03“, „monatlich 1. So 10:00“, „monatlich letzter Fr 18:30“, „einmalig 2026-10-28 17:30“ (Einzeltermin), Zusatz in Klammern: „wöchentlich Sa 09:30 (Long Run)“.",
        "Tempo: Laufen „5:30–6:30“ (min/km), Rad „25–30“ (km/h), Schwimmen „1:50–2:15“ (min/100 m).",
        "Kosten: „kostenlos“, „5 €“ oder freier Text („CHF 10, Gäste frei“).",
        "Graue Spalten nicht ändern. Was nicht in eine Zelle passt, in „Anmerkung“ schreiben.",
    ]:
        hilfe.append([zeile])
    hilfe.column_dimensions["A"].width = 140
    wb.move_sheet("Anleitung", offset=-2)
    wb.active = 1
    wb.save(ziel)
    print(f"{len(offen)} Clubs zur Prüfung → {ziel}")
    return len(offen)


def importiere(datei: Path, dry_run: bool) -> int:
    wb = load_workbook(datei, data_only=True)
    ws = wb["Clubs prüfen"]
    orig = {r[-1]: r for r in wb["_original"].iter_rows(min_row=2, values_only=True) if r and r[-1]}
    clubs = lade()
    nach_id = {c["id"]: c for c in clubs}
    heute = datetime.date.today().isoformat()
    loeschen, frei, fehler, anmerkungen = [], [], [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[-1]:
            continue
        cid = r[-1]
        status = (r[0] or "").strip()
        club = nach_id.get(cid)
        if club is None:
            fehler.append(f"{cid}: nicht (mehr) in clubs.json")
            continue
        if r[SPALTEN.index(("Anmerkung", None, True))]:
            anmerkungen.append(f"{cid}: {r[SPALTEN.index(('Anmerkung', None, True))]}")
        if not status:
            continue
        if status == "Löschen":
            loeschen.append(cid)
            continue
        if status not in ("Stimmt so", "Korrigiert"):
            fehler.append(f"{cid}: Status {status!r} unbekannt")
            continue
        alt = orig.get(cid)
        neu = dict(club)
        try:
            for i, (kopf, feld, edit) in enumerate(SPALTEN):
                if not edit or feld is None:
                    continue
                wert = "" if r[i] is None else str(r[i]).strip()
                vorher = "" if alt is None or alt[i] is None else str(alt[i]).strip()
                if wert == vorher:
                    continue
                if feld == "termine":
                    neu["termine"] = termine_aus_text(wert, club.get("termine") or [])
                elif feld == "distanz_km":
                    neu[feld] = spanne_aus_text(wert, neu.get("art"), False)
                elif feld == "pace":
                    neu[feld] = spanne_aus_text(wert, neu.get("art"), True)
                elif feld == "kosten":
                    neu[feld] = kosten_aus_text(wert)
                elif feld == "anmeldung" and wert and wert not in ("nein", "empfohlen", "pflicht"):
                    raise ValueError(f"Anmeldung {wert!r}")
                elif feld == "art" and wert not in ("Laufen", "Fahrrad", "Schwimmen"):
                    raise ValueError(f"Sportart {wert!r}")
                elif feld in ("webseite", "instagram", "strava") and wert and not wert.startswith("https://"):
                    raise ValueError(f"{kopf} ohne https://")
                else:
                    neu[feld] = wert or None
                if neu.get(feld) is None:
                    neu.pop(feld, None)
        except ValueError as ex:
            fehler.append(f"{cid}: {ex} – bleibt offen")
            continue
        neu.pop("pruefung_offen", None)
        neu.pop("pruefung_grund", None)
        neu["geprueft"] = heute
        neu.setdefault("quelle", [])
        if "Prüfung durch den Betreiber" not in " ".join(neu["quelle"]):
            neu["quelle"].append(f"Prüfung durch den Betreiber am {heute} (Korrekturliste)")
        nach_id[cid].clear()
        nach_id[cid].update(neu)
        frei.append(cid)
    clubs = [c for c in clubs if c["id"] not in set(loeschen)]
    print(f"freigegeben {len(frei)}, gelöscht {len(loeschen)}, Fehler {len(fehler)}")
    for f in fehler:
        print("  ⚠", f)
    for a in anmerkungen:
        print("  ✎", a)
    if not dry_run:
        speichere(clubs)
    return 1 if fehler else 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--ziel", type=Path, default=STANDARD_DATEI)
    i = sub.add_parser("import")
    i.add_argument("datei", type=Path)
    i.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if a.cmd == "export":
        export(a.ziel)
        return 0
    return importiere(a.datei, a.dry_run)


if __name__ == "__main__":
    sys.exit(main())
