#!/usr/bin/env python3
"""Die Korrekturlisten als Webseite statt als Excel – die Brücke in beide Richtungen.

Vom Nutzer am 19.10.2026 gewünscht: „Wir haben jetzt 2 Korrekturlisten …
Beide sind in Excel, die wir uns immer hin und herschicken müssen. Können
wir … eine Seite bauen, die den gleichen Aufbau hat wie die Excel … Sobald
ein Event dann von mir erledigt ist, dann siehst du sofort meinen Output."

Die Seite ist ein privates Artifact auf claude.ai (`korrektur/korrekturliste.html`,
Adresse in CLAUDE.md, Abschnitt „Korrekturseite“): Die Zeilen kommen als
Datendateien mit der Seite, die Antworten des Nutzers landen in ihrer
Datenbank (`antworten_events/<id>`, `antworten_clubs/<id>`), die Claude über
das Werkzeug `ArtifactData` liest. Eine Seite auf endurance-events.de ginge
nicht: Dort speicherte sie in Firestore, und das kann Claude nicht lesen
(die Regeln sperren es, einen Dienstkonto-Schlüssel gibt es bewusst nicht).

Dieses Skript erfindet keine eigene Regel – es übersetzt nur:

    daten events|clubs         Excel-Export (korrekturliste.py bzw.
                               clubs_korrekturliste.py) -> JSON für die Seite:
                               dieselben Spalten, dieselben Farben (rot = fehlt,
                               orange = fraglich, grau = nicht änderbar),
                               dieselben Auswahllisten.
    einlesen events|clubs DIR  Antworten (ArtifactData-Dateien in DIR) -> eine
                               Excel-Datei im Format des Exports -> der
                               VORHANDENE Import (mit --dry-run nur anzeigen).
                               Schreibt danach `DIR/quittung.json`: EIN neues
                               Dokument für `quittung_<liste>`, das die
                               übernommenen Antworten nennt.

Übernommen wird nur eine Antwort mit Status (Stimmt so / Korrigiert /
Löschen / Fällt aus / Unklar), zu der es noch keine Quittung gibt – eine Zeile, in der
der Nutzer noch tippt, bleibt liegen.

**Warum Quittungen statt eines Felds an der Antwort**: Ein Dokument, das es
schon gibt, schreibt `ArtifactData` nur mit seiner Versionsnummer, und die
steht nicht in den Dateien von `list … out_dir`. Eine Quittung ist dagegen
jedes Mal ein NEUES Dokument (`quittung_<liste>/<Zeitstempel>`, Feld `ids`:
Zeilen-Id -> `geaendert` der Antwort) – ein einziger `set` mit `file_path`.
Die Seite zeigt eine Antwort als „übernommen“, wenn eine Quittung genau
ihren Stand nennt; ändert der Nutzer die Zeile danach noch einmal, ist sie
wieder „bereit“, und der nächste Lauf übernimmt sie erneut. Die Antworten
selbst fasst Claude nie an. Der Ablauf steht in CLAUDE.md („Korrekturseite“).

Clubs: `clubs_korrekturliste.py` liegt auf dem Clubs-Branch
(`claude/clubs-feature-map-kycnmi`). Ohne die Datei hier im Repo braucht
`clubs` den Schalter `--clubs-repo <Checkout des Clubs-Branches>`.
"""
from __future__ import annotations

import argparse
import datetime
import hashlib
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

REPO = Path(__file__).resolve().parent.parent
ZIEL_DATEN = REPO / "korrektur" / "daten"

FARBE_ART = {"F8C9C4": "f", "FCE2B0": "q"}   # rot = fehlt, orange = fraglich
STATUS_WERTE = {"Stimmt so", "Korrigiert", "Löschen", "Fällt aus", "Unklar"}
# Spalten, die über mehrere Zeilen gehen (Textfeld statt Eingabezeile).
MEHRZEILIG = {"Was fehlt / was ist fraglich", "Deine Anmerkung", "Name (nicht änderbar)",
              "Grund der Prüfung", "Termine", "Quelle", "Anmerkung", "Treffpunkt"}

# Auswahllisten, die der Clubs-Export (noch) nicht als Excel-Auswahl
# trägt, die sein Import aber genau so prüft (clubs_korrekturliste.importiere).
AUSWAHL_ZUSATZ = {
    "clubs": {"Sportart": ["Laufen", "Fahrrad", "Schwimmen"],
              "Land": ["Deutschland", "Österreich", "Schweiz", "Italien (Südtirol)"],
              "Anmeldung (nein/empfohlen/pflicht)": ["nein", "empfohlen", "pflicht"]},
}

LISTEN = {
    "events": {"titel": "Events", "blatt": "Events prüfen", "status": "Status",
               "anmerkung": "Deine Anmerkung", "name": "Name (nicht änderbar)",
               "datum": "Datum Start", "ort": "Ort", "hinweis": "Was fehlt / was ist fraglich"},
    "clubs": {"titel": "Clubs", "blatt": "Clubs prüfen", "status": "Status",
              "anmerkung": "Anmerkung", "name": "Name", "datum": None, "ort": "Stadt",
              "hinweis": "Grund der Prüfung"},
}


# ---------------------------------------------------------------------------
# Die beiden Excel-Skripte
# ---------------------------------------------------------------------------

def modul(liste: str, clubs_repo: Path | None):
    if liste == "events":
        sys.path.insert(0, str(REPO / "scripts"))
        import korrekturliste  # noqa: PLC0415
        return korrekturliste
    wurzel = clubs_repo or REPO
    pfad = wurzel / "scripts" / "clubs_korrekturliste.py"
    if not pfad.exists():
        sys.exit(f"{pfad} fehlt – die Clubs liegen auf dem Clubs-Branch; "
                 "--clubs-repo <Checkout von claude/clubs-feature-map-kycnmi> angeben")
    spec = importlib.util.spec_from_file_location("clubs_korrekturliste", pfad)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def editierbar(liste: str, m) -> dict[str, bool]:
    """Spaltenkopf -> änderbar, aus der Spaltenliste des Excel-Skripts."""
    if liste == "events":
        return {kopf: (feld in m.AENDERBAR or feld in ("_status", "_anmerkung"))
                for feld, kopf, _ in m.SPALTEN}
    return {kopf: bool(edit) for kopf, _, edit in m.SPALTEN}


# ---------------------------------------------------------------------------
# daten: Excel -> JSON für die Seite
# ---------------------------------------------------------------------------

def als_text(wert) -> str:
    """Zellwert so, wie die Excel ihn zeigt – und wie der Import ihn wieder
    liest (Datum TT.MM.JJJJ, Dezimalkomma)."""
    if wert is None:
        return ""
    if isinstance(wert, datetime.datetime):
        wert = wert.date()
    if isinstance(wert, datetime.date):
        return wert.strftime("%d.%m.%Y")
    if isinstance(wert, bool):
        return "ja" if wert else ""
    if isinstance(wert, float):
        return (f"{wert:.7f}".rstrip("0").rstrip(".") if not wert.is_integer() else str(int(wert))).replace(".", ",")
    return str(wert)


def zeilen_id(liste: str, werte: list[str], kopf: list[str], orig: str | None) -> str:
    if liste == "clubs":
        return werte[kopf.index("id")]
    # Events: die Werte beim Export (Blatt _original) bestimmen die Zeile –
    # gleicher Stand, gleiche Kennung; die Nr der Excel wechselt mit jedem Export.
    return "e" + hashlib.sha1((orig or "|".join(werte)).encode()).hexdigest()[:14]


def lies_export(liste: str, xlsx: Path, m) -> dict:
    cfg = LISTEN[liste]
    wb = load_workbook(xlsx)
    ws = wb[cfg["blatt"]]
    kopf = [c.value for c in ws[1]]
    edit = editierbar(liste, m)
    auswahl: dict[int, list[str]] = {}
    for dv in ws.data_validations.dataValidation:
        werte = (dv.formula1 or "").strip('"').split(",")
        for bereich in str(dv.sqref).split():
            spalte = column_index_from_string("".join(ch for ch in bereich.split(":")[0] if ch.isalpha()))
            auswahl[spalte - 1] = werte
    orig_json: dict[int, str] = {}
    if liste == "events":
        for r in wb["_original"].iter_rows(min_row=2, values_only=True):
            if r[0] is not None:
                orig_json[int(r[0])] = r[1]

    for i, k in enumerate(kopf):
        if i not in auswahl and k in AUSWAHL_ZUSATZ.get(liste, {}):
            auswahl[i] = AUSWAHL_ZUSATZ[liste][k]
    # Fixierte Spalten wie in der Excel (ws.freeze_panes "F2" = fünf Spalten).
    fixiert = 0
    if ws.freeze_panes:
        fixiert = column_index_from_string("".join(ch for ch in ws.freeze_panes if ch.isalpha())) - 1

    spalten = []
    for i, k in enumerate(kopf):
        breite = ws.column_dimensions[get_column_letter(i + 1)].width or 14
        spalten.append({"kopf": k, "edit": bool(edit.get(k)), "auswahl": auswahl.get(i),
                        "mehrzeilig": k in MEHRZEILIG, "breite": round(breite)})

    zeilen = []
    for r in range(2, ws.max_row + 1):
        zellen = [ws.cell(row=r, column=i + 1) for i in range(len(kopf))]
        werte = [als_text(c.value) for c in zellen]
        if not any(werte):
            continue
        markiert = {}
        for i, c in enumerate(zellen):
            if c.fill is not None and c.fill.fill_type == "solid":
                art = FARBE_ART.get(str(c.fill.fgColor.rgb or "")[-6:].upper())
                if art:
                    markiert[str(i)] = art
        orig = orig_json.get(int(zellen[0].value)) if liste == "events" and zellen[0].value else None
        z = {"id": zeilen_id(liste, werte, kopf, orig), "z": werte}
        if markiert:
            z["m"] = markiert
        if orig:
            z["o"] = orig
        zeilen.append(z)
    return {"liste": liste, "titel": cfg["titel"], "erzeugt": datetime.datetime.now().isoformat(timespec="seconds"),
            "rollen": {k: v for k, v in cfg.items() if k not in ("titel", "blatt")}, "fixiert": fixiert,
            "spalten": spalten, "zeilen": zeilen}


def cmd_daten(args) -> int:
    m = modul(args.liste, args.clubs_repo)
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "liste.xlsx"
        m.export(xlsx)
        daten = lies_export(args.liste, xlsx, m)
    ziel = args.ziel or (ZIEL_DATEN / f"{args.liste}.json")
    ziel.parent.mkdir(parents=True, exist_ok=True)
    ziel.write_text(json.dumps(daten, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    ids = [z["id"] for z in daten["zeilen"]]
    if len(set(ids)) != len(ids):
        print(f"⚠ {len(ids) - len(set(ids))} doppelte Kennungen – Zeilen mit identischen Werten")
    print(f"{len(daten['zeilen'])} Zeilen -> {ziel} ({ziel.stat().st_size // 1024} KB)")
    return 0


# ---------------------------------------------------------------------------
# einlesen: Antworten -> Excel im Exportformat -> vorhandener Import
# ---------------------------------------------------------------------------

def _dokumente(ordner: Path, sammlung: str):
    """Die Dateien, die `ArtifactData list … out_dir` schreibt
    (`<out_dir>/<Sammlung>/<doc_id>.json`, eine je Dokument). Liegen die
    Dateien direkt in `ordner`, gilt der Ordner selbst als Sammlung."""
    unter = ordner / sammlung
    for datei in sorted((unter if unter.is_dir() else ordner).glob("*.json")):
        roh = json.loads(datei.read_text(encoding="utf-8"))
        doc = roh.get("data") if isinstance(roh.get("data"), dict) else roh
        if isinstance(doc, dict):
            yield datei.stem, doc


def lade_antworten(ordner: Path, liste: str) -> list[dict]:
    antworten = []
    for stem, doc in _dokumente(ordner, f"antworten_{liste}"):
        if "zeile" not in doc or doc.get("liste", liste) != liste:
            continue
        doc = dict(doc)
        doc.setdefault("id", stem)
        antworten.append(doc)
    return antworten


def lade_quittungen(ordner: Path, liste: str) -> dict[str, set]:
    """Zeilen-Id -> alle `geaendert`-Stände, die schon übernommen sind."""
    quittiert: dict[str, set] = {}
    for _, doc in _dokumente(ordner, f"quittung_{liste}"):
        for zid, stand in (doc.get("ids") or {}).items():
            quittiert.setdefault(zid, set()).add(stand)
    return quittiert


def baue_excel(liste: str, antworten: list[dict], ziel: Path) -> None:
    """Eine Excel-Datei, wie sie der Export geschrieben hätte – nur mit den
    Zeilen der Antworten, deren Zellen der Nutzer auf der Seite geändert hat."""
    cfg = LISTEN[liste]
    wb = Workbook()
    ws = wb.active
    ws.title = cfg["blatt"]
    orig = wb.create_sheet("_original")
    kopf = antworten[0]["kopf"]
    ws.append(kopf)
    if liste == "events":
        orig.append(["Nr", "Werte beim Export (JSON) - nicht ändern"])
    else:
        orig.append(kopf)
    for nr, a in enumerate(antworten, start=1):
        if a["kopf"] != kopf:
            raise SystemExit(f"{a['id']}: andere Spalten als die übrigen Antworten – Liste neu laden")
        werte = list(a["zeile"])
        if liste == "clubs":
            orig.append(list(werte))
        for k, v in (a.get("werte") or {}).items():
            werte[kopf.index(k)] = v
        werte[kopf.index(cfg["status"])] = a.get("status") or ""
        werte[kopf.index(cfg["anmerkung"])] = a.get("anmerkung") or ""
        if liste == "events":
            werte[0] = nr
            orig.append([nr, a["orig"]])
        ws.append([None if v == "" else v for v in werte])
    wb.save(ziel)


def cmd_einlesen(args) -> int:
    alle = lade_antworten(args.ordner, args.liste)
    quittiert = lade_quittungen(args.ordner, args.liste)
    offen = [a for a in alle if a.get("geaendert") not in quittiert.get(a["id"], set())]
    neu = [a for a in offen if a.get("status") in STATUS_WERTE]
    print(f"{len(alle)} Antworten, {len(alle) - len(offen)} schon übernommen, "
          f"{len(neu)} mit Status zu übernehmen, {len(offen) - len(neu)} noch ohne Status (bleiben liegen)")
    if not neu:
        return 0
    m = modul(args.liste, args.clubs_repo)
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "antworten.xlsx"
        baue_excel(args.liste, neu, xlsx)
        if args.liste == "events":
            rc = m.importiere(xlsx, args.dry_run, not args.ohne_aufraeumen)
        else:
            rc = m.importiere(xlsx, args.dry_run)
    if args.dry_run:
        return rc
    if rc:
        print("\nDer Import meldet Fehler – keine Quittung geschrieben. Erst die Meldungen klären.")
        return rc
    jetzt = datetime.datetime.now(datetime.timezone.utc)
    doc_id = jetzt.strftime("%Y%m%d-%H%M%S")
    quittung = {"liste": args.liste, "am": jetzt.isoformat(timespec="seconds"),
                "ids": {a["id"]: a.get("geaendert") for a in neu}}
    ziel = args.ordner / "quittung.json"
    ziel.write_text(json.dumps(quittung, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"\n{len(neu)} Antworten übernommen – Quittung: ArtifactData set "
          f"collection=quittung_{args.liste} doc_id={doc_id} file_path={ziel}")
    return rc


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("daten", help="Excel-Export -> JSON für die Seite")
    d.add_argument("liste", choices=sorted(LISTEN))
    d.add_argument("--ziel", type=Path)
    d.add_argument("--clubs-repo", type=Path)
    e = sub.add_parser("einlesen", help="Antworten der Seite -> vorhandener Import")
    e.add_argument("liste", choices=sorted(LISTEN))
    e.add_argument("ordner", type=Path, help="Ordner mit den ArtifactData-Dateien")
    e.add_argument("--dry-run", action="store_true")
    e.add_argument("--ohne-aufraeumen", action="store_true")
    e.add_argument("--clubs-repo", type=Path)
    a = ap.parse_args()
    return cmd_daten(a) if a.cmd == "daten" else cmd_einlesen(a)


if __name__ == "__main__":
    sys.exit(main())
