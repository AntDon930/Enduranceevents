#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_scraper_lib.py
======================

Regressionstests für die Textauswertung der Scraper - ohne Netzwerk, ohne
zusätzliche Test-Bibliothek. Aufruf:

    python3 scripts/test_scraper_lib.py

Warum es diese Datei gibt: Die Fehler in diesem Projekt saßen bisher fast
alle in derselben Ecke - im Erkennen von Distanz, Kategorie und Land aus
Freitext. Mehrere davon sind erst aufgefallen, als jemand einzelne Events
in der fertigen Liste stichprobenartig nachgeschlagen hat. Jeder dieser
echten Fehler steht unten als Testfall mit Kommentar, damit er nicht ein
zweites Mal hereinkommt.

Der Exit-Code ist 0 bei Erfolg und 1 bei mindestens einem Fehlschlag,
lässt sich also direkt in CI einhängen.
"""

from __future__ import annotations

import sys

import requests
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from scraper_lib import (  # noqa: E402
    Event,
    SiteConfig,
    art2_from_elevation,
    clean_competition_label,
    expand_competitions,
    guess_art2,
    guess_distance_km,
    guess_land,
    is_portal_link,
    is_same_event,
    parse_competitions,
    parse_duration_h,
    parse_elevation_m,
    round_km,
    update_existing_event,
)

from build_places import (  # noqa: E402
    _hat_ortsbezug,
    _ist_firma,
    _verschmelze_nachbarn,
    normalisiere,
)

CONFIG = SiteConfig(base_url="", calendar_url="")

_failures: list[str] = []


def check(label: str, got, want) -> None:
    if got != want:
        _failures.append(f"{label}: erwartet {want!r}, bekommen {got!r}")
        print(f"  ✗ {label}: erwartet {want!r}, bekommen {got!r}")
    else:
        print(f"  ✓ {label}")


def test_distanz() -> None:
    print("\nDistanz-Erkennung (guess_distance_km):")
    # Echter Bug: Der Nachkommateil war auf zwei Stellen begrenzt, dadurch
    # matchte der Regex bei "42,195" nur "195" als eigenständige km-Angabe.
    check("42,195 km -> 42.2", guess_distance_km("42,195 km", CONFIG), 42.2)
    check("21,0975 km -> 21.1", guess_distance_km("über 21,0975 km", CONFIG), 21.1)
    # Echter Bug: "Strecken: 1 bis 42,195 Kilometer" ergab 195 km.
    check("Spanne mit Kilometer",
          guess_distance_km("Strecken: 0,4 bis 21,1 Kilometer", CONFIG), 21.1)
    # Echter Bug: Das Stichwort "marathon" wurde vor "halbmarathon" geprüft,
    # ein "35. Halbmarathon Altötting" bekam dadurch 42,2 statt 21,1 km.
    check("Halbmarathon-Stichwort",
          guess_distance_km("35. Halbmarathon Altötting", CONFIG), 21.1)
    check("Marathon-Stichwort",
          guess_distance_km("20. Kassel Marathon", CONFIG), 42.2)
    check("ohne Angabe", guess_distance_km("Volkslauf", CONFIG), None)
    # Andere Sprachen (22.09.2026): "Half marathon" traf nur das Teilwort
    # "marathon", der Grand Prix Winterthur stand mit 42,2 km in der Liste.
    check("Half marathon (EN)",
          guess_distance_km("Half marathon (anspruchsvolle Strecke)", CONFIG), 21.1)
    check("Semi-marathon (FR)", guess_distance_km("Semi-marathon de Lausanne", CONFIG), 21.1)
    check("mezza maratona (IT)", guess_distance_km("Mezza Maratona di Merano", CONFIG), 21.1)
    check("½ Marathon", guess_distance_km("½ Marathon", CONFIG), 21.1)
    check("maratona (IT)", guess_distance_km("Maratona dles Dolomites", CONFIG), 42.2)
    # Echter Bug, derselbe Fehlertyp wie oben, nur vor dem Komma: Der
    # Vorkommateil war auf drei Stellen begrenzt, dadurch matchte aus
    # "2067 km" nur "067" - der Transeuropalauf (2067 km) stand mit 67 km
    # in events.json. "1000 km" ergab sogar 0.
    check("2067 km", guess_distance_km("2067 km", CONFIG), 2067.0)
    check("1000 km", guess_distance_km("1000 km", CONFIG), 1000.0)
    check("kein Start mitten in der Zahl",
          guess_distance_km("Etappe 12345 km", CONFIG), 12345.0)


def test_rundung() -> None:
    print("\nRundung (round_km) - immer eine Dezimalstelle:")
    for value, want in [(42.195, 42.2), (21.0975, 21.1), (6.66, 6.7),
                        (9.999, 10.0), (10, 10.0), (None, None), (True, None)]:
        check(f"{value!r}", round_km(value), want)


def test_kategorie() -> None:
    print("\nKategorie-Erkennung (guess_art2) - spezifisch vor generisch:")
    # Echter, vom Nutzer gemeldeter Bug: Das Wort "Marathon" im Namen führte
    # zu "Straße", obwohl "Bergtrail"/"Trail-Marathon" einen Trail beschreibt.
    check("Bergtrail Trail-Marathon",
          guess_art2("5. Beck HochRhön Bergtrail 42k Trail-Marathon", CONFIG), "Trail")
    # Trail, Cross und Berglauf sind EINE Kategorie ("Trail"): Geländelauf,
    # von den Quellen mal so, mal so benannt (vom Nutzer am 19.09.2026 so
    # entschieden - vorher "Trail/Cross" und daneben "Berg").
    check("Höhenmeter -> Trail",
          guess_art2("Panoramalauf (989 Höhenmeter)", CONFIG), "Trail")
    check("Crosslauf", guess_art2("29. Mausebergecrosslauf", CONFIG), "Trail")
    check("Berglauf", guess_art2("47. Gaißacher Berglauf", CONFIG), "Trail")
    check("Alpiner Crosslauf", guess_art2("Alpiner Crosslauf Oberstdorf", CONFIG), "Trail")
    # Die alten Werte dürfen nirgends mehr entstehen.
    from scraper_lib import ART2_KEYWORDS_LAUFEN
    check("kein 'Berg' mehr in der Stichwortliste",
          sorted({k for _, k in ART2_KEYWORDS_LAUFEN}),
          ["Backyard Ultra", "Bahn", "Hindernis", "Marsch", "Straße", "Trail"])
    # Märsche (17.10.2026): eigene Kategorie, nur ohne Laufwort im Namen.
    from scraper_lib import ist_marsch_name, ist_walking
    check("Marsch: Mammutmarsch, Megamarsch, Ultra Walk, Charity Walk, Wandertag, Hollenmarsch",
          [guess_art2(n, CONFIG) for n in ("Mammutmarsch München 42 km", "Megamarsch Hannover",
                                            "Ultra Walk 100 km", "Ahmadiyya Charity Walk Berlin",
                                            "40. Schaumburger Wandertag", "Bödefelder Hollenmarsch",
                                            "Karwendelmarsch Gebirge")],
          ["Marsch"] * 7)
    check("kein Marsch: Lauf mit Walking, Marschland, skywalk, Trailruns",
          [guess_art2(n, CONFIG) for n in ("56. Lorsbacher Winterlauf- und Walkingserie",
                                            "Run & Walk Bern", "Crosslauf Wesermarsch",
                                            "Elbmarsch-Halbmarathon", "skywalk-Lauf",
                                            "Ebbser Koasamarsch Trailruns")],
          ["Straße", "Straße", "Trail", "Straße", "Straße", "Trail"])
    check("Marsch: alle Strecken bleiben, Walking-Beilage eines Laufs fliegt",
          [ist_walking("Laufen", "50 km Wanderung", "Megamarsch Hannover"),
           bool(ist_walking("Laufen", "5 km Walking", "35. Langenberger Volkslauf mit Walking")),
           ist_marsch_name("Winterlauf- und Walkingserie")],
          [None, True, False])
    check("Stadtlauf -> Straße", guess_art2("40. Wolfenbütteler Stadtlauf", CONFIG), "Straße")
    check("Hindernislauf", guess_art2("Spartan Race Hindernislauf", CONFIG), "Hindernis")

    # ---- Hindernisläufe: Marken und der PLURAL ----
    #
    # Vom Nutzer am 21.09.2026 an der "Xletix Challenge - Berlin"
    # gemeldet ("ist ein Hindernis lauf und kein Lauf auf der Straße").
    # Die Nachzählung machte daraus 58 Zeilen. Erkannt wird die Marke
    # oder die PLURALFORM "Hindernisse(n)" - ein Hindernislauf wirbt mit
    # ihrer Zahl.
    for name in ("Xletix Challenge - Berlin", "XLETIX Challenge NRW",
                 "Muddy Angel Run - Dresden", "Mud Masters - Airport Weeze",
                 "CrossDeLuxe Erzgebirge", "Rats-Run - Kupferzell",
                 "Hotfoot Run Warstein Hot-20 (40 Hindernisse)",
                 "Puls 300 Cross- und Hindernis-Lauf",
                 "Bären Run 6 km Strongman Bären, etwa 20 Hindernissen"):
        check(f"{name!r} -> Hindernis", guess_art2(name, CONFIG), "Hindernis")
    # Die Gegenprobe, an der die naheliegende Fassung gescheitert wäre:
    # Ein gewöhnlicher Silvesterlauf nennt im SINGULAR ein Hindernis -
    # und zwar einmal, um es ausdrücklich zu VERNEINEN. Ein Muster auf
    # "hindernis" ohne Plural hätte beide Zeilen zu Hindernisläufen
    # gemacht.
    check("'kein Wasserhindernis' ist kein Hindernislauf",
          guess_art2("TIME2RUN Silvesterlauf in Schwabmünchen "
                     "6,7 km Strecke, kein Wasserhindernis", CONFIG), "Straße")
    check("'mögliches Wasserhindernis' ebenso",
          guess_art2("TIME2RUN Silvesterlauf in Schwabmünchen "
                     "12 km Strecke, mögliches Wasserhindernis", CONFIG), "Straße")

    # ---- Charity ist eine MARKIERUNG, keine Kategorie ----
    #
    # Vom 21.09.2026 an war "Charity" eine Kategorie (art2) und stand in
    # der Stichwortliste ganz vorn - ein Benefiz-Crosslauf war damit
    # Charity STATT Trail. Am selben Tag hat der Nutzer das
    # zurückgenommen ("Aber da bitte wieder die Kategorie einfügen"):
    # Die Kategorie war gerade die Auskunft, die dabei verloren ging.
    # Seitdem trägt das Event ein eigenes Merkmal (ist_charity), und die
    # Kategorie bleibt die Kategorie.
    from scraper_lib import ist_charity
    for name in ("19. Benin Benefiz-Lauf", "Sterntaler Spendenlauf",
                 "Bietlauf für einen Wohltätigen Zweck 9,2 km Crosslauf",
                 "ADAC Charity Treppenlauf", "Sponsorenlauf Brustkrebshilfe Dorsten",
                 "Borne to Run 48-Stunden-Spenden-Lauf",
                 # Vom Nutzer am 21.09.2026 gemeldet - trug keines der
                 # alten Stichwörter und stand deshalb unmarkiert da.
                 "Lauf für einen guten Zweck - Rastenberg"):
        check(f"{name!r} -> charity", ist_charity(name), True)
    check("Lauf ohne Charity-Wort ist nicht markiert",
          ist_charity("Lauf gegen Krebs"), False)
    # Die Kategorie bleibt jetzt erhalten - genau das war der Wunsch.
    check("Benefiz-Crosslauf behält seine Kategorie",
          guess_art2("Bietlauf für einen Wohltätigen Zweck 9,2 km Crosslauf", CONFIG), "Trail")
    check("Charity-Treppenlauf behält seine Kategorie",
          guess_art2("ADAC Charity Treppenlauf", CONFIG), "Trail")
    check("Spendenlauf auf der Straße bleibt Straße",
          guess_art2("Sterntaler Spendenlauf", CONFIG), "Straße")
    check("Benefiz-Seeschwimmen behält Freiwasser",
          guess_art2("Benefiz-Seeschwimmen im Freiwasser", CONFIG, "Schwimmen"), "Freiwasser")
    # Treppenläufe sind Trail (Nutzer, 21.09.2026: "als Trail aufnehmen").
    for name in ("Lotto Thüringen Treppenlauf", "TK Elevator Towerrun",
                 "Bad Wildbader Stäffeleslauf", "Neuwoges-Treppenhauslauf",
                 "Mt. Everest Treppenmarathon"):
        check(f"{name!r} -> Trail", guess_art2(name, CONFIG), "Trail")
    # Die Schwimm-Liste kennt Freiwasser/Becken, ohne Voreinstellung.
    check("Freiwasser", guess_art2("Bodensee Openwater Konstanz", CONFIG, "Schwimmen"), "Freiwasser")
    check("Becken", guess_art2("Hallenbad-Cup 1500 m", CONFIG, "Schwimmen"), "Becken")
    check("Schwimmen ohne Hinweis: keine Kategorie",
          guess_art2("Sommer-Cup", CONFIG, "Schwimmen"), None)
    check("Charity beim Radrennen ist eine Markierung, keine Kategorie",
          (ist_charity("Benefiz-Radrennen"), guess_art2("Benefiz-Radrennen", CONFIG, "Fahrrad")),
          (True, None))

    # Bestehende Daten werden auf die zusammengefasste Kategorie
    # nachgezogen - je Sportart: Beim Laufen wird "Cross" zu "Trail",
    # beim Triathlon bleibt "Cross" eine eigene Kategorie, und beim
    # Fahrrad ist "Cyclecross" eine dritte Sache.
    from clean_events import merge_art2, ART2_MERGED  # lokaler Import, nur hier
    rows = [
        {"name": "Waldtrail", "art1": "Laufen", "art2": "Trail/Cross"},
        {"name": "Mausebergecrosslauf", "art1": "Laufen", "art2": "Cross"},
        {"name": "Gaißacher Berglauf", "art1": "Laufen", "art2": "Berg"},
        {"name": "Tippfehler-Override", "art1": "Laufen", "art2": "Berglauf"},
        {"name": "Hofer Backyard Ultra", "art1": "Laufen", "art2": "Backcountry Ultra"},
        {"name": "Backyardman Würzburg", "art1": "Triathlon", "art2": "Backyard"},
        {"name": "Crosstriathlon", "art1": "Triathlon", "art2": "Cross"},
        {"name": "Stadtlauf", "art1": "Laufen", "art2": "Straße"},
        {"name": "Querfeldein-Rennen", "art1": "Fahrrad", "art2": "Cyclecross"},
    ]
    check("sechs Einträge zusammengefasst", len(merge_art2(rows)), 6)
    check("neue Kategorie steht", [r["art2"] for r in rows],
          ["Trail", "Trail", "Trail", "Trail", "Backyard Ultra", "Backyard Ultra",
           "Cross", "Straße", "Cyclecross"])
    check("idempotent", len(merge_art2(rows)), 0)
    check("kein Zielwert ist zugleich ein Schlüssel",
          [z for m in ART2_MERGED.values() for z in m.values() if z in m], [])

    # Die Kategorie-Liste der Filter (filter-ui.js) muss dieselben Werte
    # anbieten, die die Stichwortliste erzeugt - sonst filtert die Seite
    # einen Wert heraus, den es gibt (genau das wäre beim Umbenennen
    # passiert, wenn man nur eine der beiden Dateien anfasst).
    ui_js = (Path(__file__).resolve().parent.parent / "filter-ui.js").read_text(encoding="utf-8")
    zeile = next(l for l in ui_js.splitlines() if "'Laufen': [" in l)
    from scraper_lib import ART2_KEYWORDS_LAUFEN, DEFAULT_ART2_LAUFEN
    erzeugbar = {kat for _, kat in ART2_KEYWORDS_LAUFEN} | {DEFAULT_ART2_LAUFEN}
    fehlend = sorted(k for k in erzeugbar if f"'{k}'" not in zeile)
    check("Filter kennt alle Laufen-Kategorien", fehlend, [])

    # ---- Mehrsport: Ein Ironman ist kein Lauf ----
    #
    # Alle Quellen sind Laufkalender (default_art1 = "Laufen"), also stand
    # jeder Triathlon als Laufveranstaltung in der Liste - vom Nutzer an
    # "Ironman 70.3 Kraichgau · Laufen · Straße" gemeldet. Bei einem
    # Ironman kann man sich nicht für den Lauf allein anmelden.
    from scraper_lib import (ART2_KEYWORDS_TRIATHLON, DEFAULT_ART2_TRIATHLON,
                             guess_art1)
    print("\nSportart-Erkennung (guess_art1):")
    for name, erwartet in (
            ("Ironman 70.3 Kraichgau", "Triathlon"),
            ("Ironman 5150 Erkner Berlin-Brandenburg", "Triathlon"),
            ("Challenge Roth", "Triathlon"),
            ("Kornwestheimer Triathlon", "Triathlon"),
            ("Baltic X Cross Duathlon", "Triathlon"),
            ("SwimRun Rheinsberg", "Triathlon"),
            # Gegenprobe: Läufe bleiben Läufe. "Marathon" und "Ultra"
            # dürfen nichts umstellen.
            ("Königsforst-Marathon", "Laufen"),
            ("Hofer Backyard Ultra", "Laufen"),
            ("Spartan Race Hindernislauf", "Laufen")):
        check(name, guess_art1(name, CONFIG), erwartet)

    print("\nKategorie beim Mehrsport (guess_art2 mit art1):")
    # Das Format steht VOR dem Gelände: Ein "Cross Duathlon" ist ein
    # Duathlon, der im Gelände stattfindet - "Duathlon" ist die Auskunft,
    # nach der jemand filtert.
    check("Cross Duathlon -> Duathlon",
          guess_art2("Baltic X Cross Duathlon", CONFIG, "Triathlon"), "Duathlon")
    check("Crosstriathlon -> Cross",
          guess_art2("Wuppertaler Sparkassen Crosstriathlon", CONFIG, "Triathlon"), "Cross")
    check("SwimRun", guess_art2("SwimRun Rheinsberg", CONFIG, "Triathlon"), "Swimrun")
    check("Indoor-Triathlon",
          guess_art2("Indoor-Triathlon Aschersleben", CONFIG, "Triathlon"), "Indoor")
    check("Ironman -> Straße",
          guess_art2("Ironman 70.3 Kraichgau", CONFIG, "Triathlon"), "Straße")
    # Ohne art1 bleibt es bei der Lauf-Liste (ältere Aufrufe).
    check("ohne art1 unverändert", guess_art2("Mausebergecrosslauf", CONFIG), "Trail")
    # Eine Kategorie für beide Sportarten: Das Last-Man-Standing-Format
    # heißt beim Laufen und beim Triathlon gleich (Nutzer, 19.09.2026).
    check("Backyard-Triathlon", guess_art2("Backyardman Würzburg", CONFIG, "Triathlon"),
          "Backyard Ultra")

    # Über den ganzen Block statt über eine Zeile: Die Triathlon-Liste
    # ist länger als 80 Zeichen und daher umbrochen.
    start = ui_js.index("'Triathlon': [")
    zeile_tri = ui_js[start:ui_js.index("]", start)]
    erzeugbar_tri = {kat for _, kat in ART2_KEYWORDS_TRIATHLON} | {DEFAULT_ART2_TRIATHLON}
    fehlend_tri = sorted(k for k in erzeugbar_tri if f"'{k}'" not in zeile_tri)
    check("Filter kennt alle Triathlon-Kategorien", fehlend_tri, [])

    # Und die Übersetzung muss jeden Wert kennen - sonst steht im Filter
    # der rohe Schlüssel.
    filters_js = (Path(__file__).resolve().parent.parent / "filters.js").read_text(encoding="utf-8")
    ohne_text = sorted(k for k in erzeugbar_tri if f"'{k}': {{" not in filters_js)
    check("jede Triathlon-Kategorie ist übersetzt", ohne_text, [])

    # Der Bestand wird nachgezogen (clean_events.fix_multisport_art1).
    from clean_events import fix_multisport_art1
    bestand = [
        {"name": "Ironman 70.3 Kraichgau", "art1": "Laufen", "art2": "Straße"},
        {"name": "Cross-Duathlon in Hünsborn", "art1": "Laufen", "art2": "Trail"},
        {"name": "Königsforst-Marathon", "art1": "Laufen", "art2": "Straße"},
    ]
    check("zwei Einträge umgestellt", len(fix_multisport_art1(bestand)), 2)
    check("Sportart und Kategorie stimmen",
          [(r["art1"], r["art2"]) for r in bestand],
          [("Triathlon", "Straße"), ("Triathlon", "Duathlon"), ("Laufen", "Straße")])
    check("idempotent", len(fix_multisport_art1(bestand)), 0)


def test_land() -> None:
    print("\nLand-Erkennung (guess_land):")
    # Echter, vom Nutzer gemeldeter Fall: Die Detailseite nennt
    # "9607 Mosnang (Schweiz)", in events.json fehlte das Land trotzdem.
    check("(Schweiz)", guess_land("9607 Mosnang (Schweiz)"), "Schweiz")
    # Dieselbe Seite nutzt teils dreibuchstabige Kürzel.
    check("(AUT)", guess_land("6020 Innsbruck (AUT)"), "Österreich")
    check("(GER)", guess_land("10115 Berlin (GER)"), "Deutschland")
    check("fünfstellige PLZ", guess_land("83242 Reit im Winkl"), "Deutschland")
    # Vierstellige PLZ trennt AT und CH nicht -> bewusst kein Rateversuch.
    check("vierstellige PLZ bleibt offen", guess_land("6130 Schwaz"), None)
    # Echter Bug: land="Schweiz" für Ebermannstadt (Bayern), weil das Land
    # aus dem EVENT-NAMEN geraten wurde.
    check("Fränkische Schweiz ist Deutschland",
          guess_land("25. Fränkische-Schweiz-Marathon"), None)
    check("Sächsische Schweiz", guess_land("Sächsische Schweiz Trail"), None)
    # Südtirol (21.09.2026): "Italien" gilt nur mit Südtiroler PLZ oder
    # Koordinaten in der Provinz Bozen - sonst bleibt der Zwischenstand
    # "Italien" stehen und fällt später heraus.
    from scraper_lib import in_suedtirol, LAENDER as _regionen
    check("(Italien) mit Südtiroler PLZ", guess_land("39012 Meran (Italien)"), "Italien (Südtirol)")
    check("(ITA) ohne Südtiroler PLZ bleibt Italien", guess_land("20121 Milano (ITA)"), "Italien")
    check("Südtirol im Text", guess_land("Bozen, Südtirol"), "Italien (Südtirol)")
    check("Bozen liegt in Südtirol", in_suedtirol(46.4983, 11.3548), True)
    check("Sterzing liegt in Südtirol", in_suedtirol(46.8967, 11.4300), True)
    check("Innsbruck nicht", in_suedtirol(47.2692, 11.4041), False)
    check("Trient nicht (gleiche Region, andere Provinz)", in_suedtirol(46.0667, 11.1167), False)
    check("Südtirol zählt zu den Regionen", "Italien (Südtirol)" in _regionen, True)
    # Zweibuchstabige Codes dürfen in Freitext NICHT anschlagen.
    check("'de' in Freitext", guess_land("Tour de France"), None)


def test_wettbewerbe() -> None:
    print("\nWettbewerbe (parse_competitions / expand_competitions):")
    check("hm-Klammer entfernt",
          clean_competition_label("Moslig 8000 (229 hm) | 8,5 km"), "Moslig 8000")
    check("races-Layout",
          clean_competition_label("TST 86K | 86 km | + 3500 hm"), "TST 86K")
    # running.life schreibt die Strecken als Satz; alles ab dem Doppelpunkt
    # ist Beschreibung und gehört nicht ins Label.
    check("Satz mit Doppelpunkt",
          clean_competition_label(
              "VR Bank – BraunenBerg-Lauf: 14,6 km , ca. 400 Hm, Strecke endet in Oberalfingen."),
          "VR Bank – BraunenBerg-Lauf")
    check("Label ohne Doppelpunkt bleibt", clean_competition_label("Halbmarathon"), "Halbmarathon")

    # Der vom Nutzer genannte Fall: die Seite listet sieben Wettbewerbe,
    # vorher landete nur der längste in events.json.
    mosnang = [
        "Kids Fun Race | 0,4 km",
        "Plausch-Stafette | 2 km",
        "Moslig 8000 (229 hm) | 8,5 km",
        "Halbmarathon (989 hm) | 21,1 km",
    ]
    comps = parse_competitions(mosnang, CONFIG)
    check("alle vier Strecken", [c.laenge_km for c in comps], [0.4, 2.0, 8.5, 21.1])
    check("Labels", [c.label for c in comps],
          ["Kids Fun Race", "Plausch-Stafette", "Moslig 8000", "Halbmarathon"])

    # Gleiche Distanz mehrfach (Einzel/Staffel/Walking) = ein Eintrag.
    check("Duplikat-Distanz zusammengefasst",
          [c.laenge_km for c in parse_competitions(["10 km", "10 km Staffel"], CONFIG)],
          [10.0])

    base = Event(name="Schnebelhorn Panoramatrail", standort="Mosnang",
                 land="Schweiz", datum_start="2026-09-18")
    variants = expand_competitions(base, comps, CONFIG)
    check("ein Eintrag pro Strecke", len(variants), 4)
    check("Name bleibt unverändert",
          {v.name for v in variants}, {"Schnebelhorn Panoramatrail"})
    check("Wettbewerb gesetzt", variants[2].wettbewerb, "Moslig 8000")

    # art2 pro Strecke: der "hint" der Quelle ("Trailrun") schlägt durch.
    hinted = parse_competitions([("TST 42K | 42 km", "Trailrun")], CONFIG)
    check("art2 aus Streckenart",
          [v.art2 for v in expand_competitions(Event(name="Trophy"), hinted, CONFIG)],
          ["Trail"])

    # Ohne erkannte Wettbewerbe darf nichts verloren gehen.
    check("keine Wettbewerbe -> Basis-Event",
          len(expand_competitions(base, [], CONFIG)), 1)


def test_hoehenprofil() -> None:
    print("\nHöhenprofil (art2_from_elevation):")
    check("Hm lesen", parse_elevation_m("ca. 1100 Hm, Start im Bergwerk"), 1100.0)
    check("Hm mit Tausenderpunkt", parse_elevation_m("+ 3.500 hm"), 3500.0)
    check("Höhenmeter ausgeschrieben", parse_elevation_m("229 Höhenmeter"), 229.0)
    check("keine Angabe", parse_elevation_m("8,2 km"), None)

    # Der vom Nutzer beanstandete Fall: 400 Hm auf 14,6 km ist kein
    # Straßenlauf (27 m/km).
    check("400 hm / 14,6 km -> Trail",
          art2_from_elevation("VR Bank – BraunenBerg-Lauf: 14,6 km, ca. 400 Hm", 14.6, "Straße"),
          "Trail")
    check("248 hm / 8,2 km -> Trail",
          art2_from_elevation("BergBau-Lauf: 8,2 km, ca. 248 Hm", 8.2, "Straße"), "Trail")
    # Flacher Stadtmarathon mit ein paar Brücken bleibt Straße (2,4 m/km).
    check("100 hm / 42,2 km bleibt Straße",
          art2_from_elevation("Stadtmarathon, 100 hm", 42.2, "Straße"), "Straße")
    # Eine spezifischere Kategorie aus dem Namen wird nie überschrieben.
    check("Hindernis bleibt Hindernis",
          art2_from_elevation("Spartan Beast: 21 km, ca. 1100 Hm", 21.0, "Hindernis"),
          "Hindernis")
    check("ohne Distanz keine Aussage",
          art2_from_elevation("ca. 400 Hm", None, "Straße"), "Straße")


def test_offizieller_link() -> None:
    print("\nPortal- vs. offizieller Link (update_existing_event):")
    check("running.life ist Portal", is_portal_link("https://running.life/de/termine/x"), True)
    check("laufen.de ist Portal", is_portal_link("https://laufen.de/laufkalender/details/1"), True)
    check("Veranstalter ist kein Portal", is_portal_link("https://www.braunenberg-lauf.de/"), False)
    # Ein Zeitnehmer ist kein Veranstalter (Kallinchen Triathlon, 19.09.2026).
    check("Zeitnehmer ist Portal",
          is_portal_link("https://www.berlin-timing.de/Kallinchen-Triathlon"), True)
    # Die übrigen Zeitnehmer und Anmeldeplattformen ebenso (19.09.2026,
    # "zieh die anderen Zeitnehmer genauso nach"): Ein Link dorthin ist
    # nie die Veranstalterseite - er wird ersetzt, sobald eine bekannt ist.
    for url in ("https://my.raceresult.com/354931/", "https://runtix.com/sts/10021/3221",
                "https://www.davengo.com/event/overview/x", "https://ladv.de/ausschreibung/detail/1/x.htm",
                "https://www.datasport.de/anmeldeservice/x", "https://rennmeldung.de/cgi-bin/bewerb.cgi?bewerb=1",
                "https://laufen-os.de/", "https://www.sas-online.net/eventportal_bs/837/",
                "https://www.sportprogramme.org/cakespg/events/menue/5037",
                "https://baer-service.de/veranstaltung/AAL/"):
        check(f"Zeitnehmer/Anmeldung ist Portal: {url.split('/')[2]}", is_portal_link(url), True)
    check("Veranstalter mit Zeitnehmer im Namen bleibt", is_portal_link("https://timing-team-lauf.de/"), False)
    # Hostname statt Teilzeichenkette: "tsv-weeze-leichtathletik.de" ist
    # kein Portal, obwohl "leichtathletik.de" darin steckt (Linkprüfung
    # 19.09.2026); Subdomains des Portals zählen dagegen mit.
    check("Vereinsdomain mit Portal-Endung ist kein Portal",
          is_portal_link("https://www.tsv-weeze-leichtathletik.de/staffel.html"), False)
    check("Subdomain des Portals ist Portal",
          is_portal_link("https://www.leichtathletik.de/wettkaempfe/x"), True)

    # Kernfall: Ein gespeicherter Portallink wird durch die offizielle Seite
    # ersetzt - ohne das behalten Altbestände den Portallink für immer.
    target = {"name": "BraunenBerg-Lauf",
              "veranstalter_url": "https://running.life/de/termine/braunenberg-lauf"}
    update_existing_event(target, {"veranstalter_url": "https://www.braunenberg-lauf.de/",
                                   "land": "Deutschland"})
    check("Portallink ersetzt", target["veranstalter_url"], "https://www.braunenberg-lauf.de/")
    check("fehlendes Feld ergänzt", target["land"], "Deutschland")

    # Umgekehrt NICHT: ein bereits direkter Link wird nie durch ein Portal
    # überschrieben.
    target2 = {"veranstalter_url": "https://www.braunenberg-lauf.de/"}
    update_existing_event(target2, {"veranstalter_url": "https://running.life/de/termine/x"})
    check("direkter Link bleibt", target2["veranstalter_url"], "https://www.braunenberg-lauf.de/")


def test_duplikate() -> None:
    print("\nDuplikat-Erkennung (is_same_event):")
    a = {"name": "52. Int. Bodensee-Marathon", "datum_start": "2026-09-20",
         "standort": "Konstanz", "laenge_km": 42.195}
    b = {"name": "Bodensee Marathon", "datum_start": "2026-09-20",
         "standort": "Konstanz", "laenge_km": 42.2}
    check("gleicher Lauf, anderer Name", is_same_event(a, b), True)

    # Unterschiedliche Distanzen desselben Events bleiben getrennte Einträge
    # (ausdrücklicher Wunsch aus dem Chat).
    hm = dict(a, laenge_km=21.1)
    check("Marathon vs. Halbmarathon getrennt", is_same_event(a, hm), False)

    # Achter Weg (07.10.2026): dieselbe Veranstalterseite, derselbe Ort,
    # dieselbe Strecke - ohne ein gemeinsames Namenswort. Die Gegenproben
    # sind der wichtigere Teil, alle am Bestand gezählt.
    print("  achter Weg (gleiche Seite, gleicher Ort, gleiche Strecke):")
    seite = "https://marathonmuenchen.org"
    m1 = {"name": "Bayerische Marathon Meisterschaften", "datum_start": "2026-10-11",
          "standort": "Schwabinger Tor", "lat": 48.1714, "lon": 11.5866, "art1": "Laufen",
          "art2": "Straße", "laenge_km": 10.0, "wettbewerb": "10 km", "veranstalter_url": seite}
    m2 = {"name": "Marathon München by Brooks", "datum_start": "2026-10-11",
          "standort": "München", "lat": 48.16, "lon": 11.58, "art1": "Laufen",
          "art2": "Straße", "laenge_km": 10.0, "wettbewerb": "10 km", "veranstalter_url": seite}
    check("gleiche Seite, gleicher Ort (≤ 3 km), gleiche Strecke -> eins", is_same_event(m1, m2), True)
    check("http://scuchte.de gegen https://www.scuchte.de/ -> dieselbe Seite, eins",
          is_same_event(dict(m1, name="Uchter Volks- und Crosslauf", veranstalter_url="http://scuchte.de"),
                        dict(m2, name="Uchter Wald- und Crosslauf", veranstalter_url="https://www.scuchte.de/")), True)
    check("anderer Pfad auf derselben Domain -> zwei",
          is_same_event(dict(m1, name="Uchter Volks- und Crosslauf", veranstalter_url="https://scuchte.de/volkslauf"),
                        dict(m2, name="Uchter Wald- und Crosslauf", veranstalter_url="https://scuchte.de/crosslauf")), False)
    check("Olympische Distanz gegen 51 km auf derselben Seite -> eins",
          is_same_event(dict(m1, name="Münster City Triathlon", art1="Triathlon", art2=None, laenge_km=51.4,
                             wettbewerb="Olympische Distanz 51,5 km"),
                        dict(m2, name="Sparda Münster City Triathlon", art1="Triathlon", art2=None,
                             laenge_km=51.4, wettbewerb="51 km")), True)
    check("Sprintdistanz gegen Volksdistanz auf derselben Seite -> zwei (zwei Formatwörter)",
          is_same_event(dict(m1, name="Silbersee-Triathlon", art1="Triathlon", art2=None, laenge_km=28.8,
                             wettbewerb="Sprintdistanz 28,8 km (750 m Schwimmen / 23 km Rad / 5 km Laufen)"),
                        dict(m2, name="32. swb Silbersee-Triathlon Stuhr", art1="Triathlon", art2=None,
                             laenge_km=28.5, wettbewerb="Volksdistanz 28,5 km (500 m Schwimmen / 23 km Rad / 5 km Laufen)")), False)
    check("22 km gegen 21,1 km (TOP20RUN gegen Halbmarathon) -> zwei (über 0,5 km)",
          is_same_event(dict(m1, name="Gornergrat Zermatt Marathon", laenge_km=22.0, wettbewerb="TOP20RUN"),
                        dict(m2, name="Zermatt Marathon", laenge_km=21.1, wettbewerb="Halbmarathon")), False)
    check("Dauer gegen Dauer (24 h) -> eins",
          is_same_event(dict(m1, name="3. Chiemgauer100 StundenRundenLauf", laenge_km=None, dauer_h=24, wettbewerb=None),
                        dict(m2, name="Chiemgauer100 StundenRunden - 100mi", laenge_km=None, dauer_h=24, wettbewerb="24 h")), True)
    check("13 km auseinander (Spreewald: Radtour ab Lübbenau UND ab Burg) -> zwei",
          is_same_event(dict(m1, name="Spreewaldmarathon Radmarathon", standort="Lübbenau/Spreewald",
                             lat=51.863, lon=13.961, art1="Fahrrad", art2=None, laenge_km=70.0, wettbewerb="70 km Radtour"),
                        dict(m2, name="Spreewaldmarathon Radtouren Burg", standort="Burg (Spreewald)",
                             lat=51.833, lon=14.15, art1="Fahrrad", art2=None, laenge_km=70.0, wettbewerb="70 km Radtour")), False)
    check("eine Seite für alle parkruns, Orte 300 km auseinander -> zwei",
          is_same_event(dict(m1, name="Englischer Garten parkrun München", laenge_km=5.0, wettbewerb="5 km",
                             veranstalter_url="https://www.parkrun.com.de"),
                        dict(m2, name="Niddapark parkrun Frankfurt", standort="Frankfurt am Main", lat=50.11, lon=8.68,
                             laenge_km=5.0, wettbewerb="5 km", veranstalter_url="https://www.parkrun.com.de")), False)
    check("CTF gegen Gravel Tour über 50 km -> zwei (Gattung)",
          is_same_event(dict(m1, name="Linneser Cross-Country Tag CTF", art1="Fahrrad", art2="Mountainbike",
                             laenge_km=50.0, wettbewerb="CTF 50 km", standort="Gießen", lat=50.58, lon=8.67),
                        dict(m2, name="Linneser Cross-Country-Tag (GRAVEL)", art1="Fahrrad", art2="Gravel",
                             laenge_km=50.0, wettbewerb="Gravel Tour 50 km", standort="Gießen", lat=50.58, lon=8.67)), False)
    check("Schülerlauf gegen Hauptlauf über 5 km -> zwei (Gattung)",
          is_same_event(dict(m1, name="Stadtlauf", laenge_km=5.0, wettbewerb="5 km Schülerlauf"),
                        dict(m2, name="City-Run", laenge_km=5.0, wettbewerb="5 km Hauptlauf")), False)
    check("Zeile ohne Maßzahl -> zwei (passt zu allem)",
          is_same_event(dict(m1, name="Stadtlauf", laenge_km=None, wettbewerb=None),
                        dict(m2, name="City-Run", laenge_km=10.0)), False)
    check("Portallink (raceresult) zählt nicht als Seite",
          is_same_event(dict(m1, name="Stadtlauf", veranstalter_url="https://my.raceresult.com/1/"),
                        dict(m2, name="City-Run", veranstalter_url="https://my.raceresult.com/1/")), False)
    check("verschiedene Kategorie (Bike Marathon gegen Gravel) -> zwei",
          is_same_event(dict(m1, name="Black Forest Ultra Bike Marathon", art1="Fahrrad", art2="Mountainbike",
                             laenge_km=118.0, wettbewerb="Radmarathon 118 km"),
                        dict(m2, name="Black Forest Ultra Gravel", art1="Fahrrad", art2="Gravel",
                             laenge_km=118.0, wettbewerb="Gravel 120 km")), False)

    # Zeitrennen: Stundenlauf und Halbstundenlauf derselben Veranstaltung
    # sind zwei Wettbewerbe (Döbelner Fackellauf); dieselbe Dauer bleibt
    # ein Duplikat, und eine fehlende Dauer schließt nichts aus.
    st = {"name": "Döbelner Fackellauf", "datum_start": "2026-10-28",
          "standort": "Döbeln", "laenge_km": None, "dauer_h": 1.0}
    check("Stundenlauf vs. Halbstundenlauf getrennt",
          is_same_event(st, dict(st, dauer_h=0.5)), False)
    check("gleiche Dauer bleibt Duplikat", is_same_event(st, dict(st)), True)
    check("ohne Dauer weiter Duplikat", is_same_event(st, dict(st, dauer_h=None)), True)
    # Distanz gegen Dauer (seit 21.09.2026): zwei Wettbewerbe. Der
    # 24-h-Lauf des Rokathon ist nicht sein Marathon - vorher verschmolz
    # ein nachgetragenes Zeitrennen mit der Distanz-Zeile und ließ sich
    # gar nicht anlegen. Ohne JEDE Angabe bleibt es ein Duplikat.
    m = dict(st, laenge_km=42.2, dauer_h=None)
    check("Marathon vs. 24-h-Lauf getrennt", is_same_event(m, dict(st, dauer_h=24.0)), False)
    check("Marathon vs. Zeile ohne jede Angabe: Duplikat",
          is_same_event(m, dict(st, dauer_h=None)), True)

    # Vom Nutzer gemeldet: Marathon und Halbmarathon München standen je
    # DOPPELT in der Liste. Ursache war die Wortstellung - als Zeichenfolge
    # nur 0,69 Ähnlichkeit, als Wortmenge identisch.
    D, S = "2026-10-11", "München"
    m1 = {"name": "40. München Marathon by Brooks", "datum_start": D, "standort": S,
          "laenge_km": 42.2, "veranstalter_url": "https://marathonmuenchen.org"}
    m2 = {"name": "MARATHON MÜNCHEN", "wettbewerb": "Marathon", "datum_start": D,
          "standort": S, "laenge_km": 42.2, "veranstalter_url": "https://marathonmuenchen.org/"}
    check("Wortstellung vertauscht", is_same_event(m1, m2), True)

    # Sechster Weg (22.09.2026): zwei Quellen, zwei Labels für dieselbe
    # Strecke - "Marathon" (running.life) neben "Splatterthon" bzw.
    # "42.2 km" (lauftermine.ch). Der Halloween Run Bremen stand fünffach
    # in der Liste. Die Gegenproben sind der wichtigere Teil: Gattung
    # (Walking, Sprint/Volks) und eine halbe Kilometer Unterschied halten
    # zwei Wettbewerbe auseinander.
    hb = {"name": "Halloween Run Bremen", "wettbewerb": "Marathon", "datum_start": "2026-10-31",
          "standort": "Bremen", "laenge_km": 42.2, "art1": "Laufen"}
    check("zwei Labels ohne Gattung: Duplikat",
          is_same_event(hb, dict(hb, name="Halloween-Run-Bremen", wettbewerb="Splatterthon")), True)
    check("Label nur Maßzahl gegen Namen: Duplikat",
          is_same_event(dict(hb, wettbewerb="42.2 km"), hb), True)
    check("Runde in zwei Schreibweisen: Duplikat",
          is_same_event(dict(hb, wettbewerb="Kürbislauf 13,333 K", laenge_km=13.3),
                        dict(hb, wettbewerb="13 km", laenge_km=13.0)), True)
    check("Walking gegen Lauf: getrennt",
          is_same_event(dict(hb, wettbewerb="10 km Walking", laenge_km=10.0),
                        dict(hb, wettbewerb="Stadtwerke 10-km-Lauf", laenge_km=10.0)), False)
    check("Halbmarathon gegen Nordic Walking: getrennt",
          is_same_event(dict(hb, wettbewerb="Halbmarathon", laenge_km=21.1),
                        dict(hb, wettbewerb="Peri Power Nordic Walking", laenge_km=21.1)), False)
    check("Sprint gegen Volksdistanz: getrennt",
          is_same_event(dict(hb, wettbewerb="Sprintdistanz 28,8 km", laenge_km=28.8, art1="Triathlon"),
                        dict(hb, wettbewerb="Volksdistanz 28,5 km", laenge_km=28.5, art1="Triathlon")), False)
    check("halber Kilometer Unterschied: getrennt",
          is_same_event(dict(hb, wettbewerb="12 km Seen-Lauf", laenge_km=12.0),
                        dict(hb, wettbewerb="12,5 km Trailrun", laenge_km=12.5)), False)
    # Gravel neben Straße über dieselbe Länge sind zwei Wettbewerbe (Erkelenzer RTF, 24.09.2026).
    check("'RTF 110 km' und 'Gravel Ride 110 km' bleiben zwei Zeilen",
          is_same_event(dict(hb, art1="Fahrrad", wettbewerb="RTF 110 km", laenge_km=110.0),
                        dict(hb, art1="Fahrrad", wettbewerb="Gravel Ride 110 km", laenge_km=110.0)), False)
    check("Kinderlauf gegen Hauptlauf: getrennt",
          is_same_event(dict(hb, wettbewerb="Mini Marathon", laenge_km=21.1),
                        dict(hb, wettbewerb="21,097 km Generali Halbmarathon", laenge_km=21.1)), False)
    # "Jedermannrennen N km" (radsport-events) gegen den Wettbewerbsnamen
    # einer anderen Quelle (datasport, 30.09.2026): dieselbe Strecke.
    check("'Jedermannrennen 106 km' und '106 KM - Corso Medio': Duplikat",
          is_same_event(dict(hb, art1="Fahrrad", wettbewerb="Jedermannrennen 106 km", laenge_km=106.0),
                        dict(hb, art1="Fahrrad", wettbewerb="106 KM - Corso Medio", laenge_km=106.0)), True)
    check("'Jedermanntriathlon' und '11 km': Duplikat",
          is_same_event(dict(hb, art1="Triathlon", wettbewerb="Jedermanntriathlon", laenge_km=11.0),
                        dict(hb, art1="Triathlon", wettbewerb="11 km", laenge_km=11.0)), True)
    # Ein Gattungswort, das nur die gemeinsame Kategorie wiederholt (30.09.2026):
    # KronplatzKing "MTB-Marathon 64 km" (radsport-events) = "64km KINGMarathon" (datasport).
    check("'MTB-Marathon 64 km' und '64km KINGMarathon' bei art2 Mountainbike: Duplikat",
          is_same_event(dict(hb, art1="Fahrrad", art2="Mountainbike", wettbewerb="MTB-Marathon 64 km", laenge_km=64.0),
                        dict(hb, art1="Fahrrad", art2="Mountainbike", wettbewerb="64km KINGMarathon", laenge_km=64.0)), True)
    check("'Trailmarathon' und '43 km' bei art2 Trail: Duplikat",
          is_same_event(dict(hb, art2="Trail", wettbewerb="Trailmarathon", laenge_km=43.0),
                        dict(hb, art2="Trail", wettbewerb="43 km", laenge_km=43.0)), True)
    check("'MTB 40 km' und 'Rennrad 40 km' mit verschiedenen Kategorien: getrennt",
          is_same_event(dict(hb, art1="Fahrrad", art2="Mountainbike", wettbewerb="MTB 40 km", laenge_km=40.0),
                        dict(hb, art1="Fahrrad", art2="Straße", wettbewerb="Rennrad 40 km", laenge_km=40.0)), False)
    check("'Trailrun 10 km' und '10-km-Stadtlauf' mit verschiedenen Kategorien: getrennt",
          is_same_event(dict(hb, art2="Trail", wettbewerb="Trailrun 10 km", laenge_km=10.0),
                        dict(hb, art2="Straße", wettbewerb="Stadtwerke 10-km-Stadtlauf", laenge_km=10.0)), False)

    # Hier trägt der Wettbewerbs-Name das unterscheidende Wort, und die
    # kürzere Wortmenge steckt komplett in der längeren.
    h1 = {"name": "Marathon München by Brooks", "wettbewerb": "Halbmarathon",
          "datum_start": D, "standort": S, "laenge_km": 21.1}
    h2 = {"name": "München Halbmarathon", "datum_start": D, "standort": S, "laenge_km": 21.1}
    check("Teilmenge mit Wettbewerbs-Name", is_same_event(h1, h2), True)
    check("Marathon bleibt vom Halbmarathon getrennt", is_same_event(m1, h1), False)

    # Ein einzelnes gemeinsames Wort darf NICHT reichen (Teilmengen-Regel
    # verlangt mindestens zwei Wörter).
    check("ein Wort ist zu generisch",
          is_same_event({"name": "Stadtlauf", "datum_start": D, "standort": S, "laenge_km": 10.0},
                        {"name": "Marathon", "datum_start": D, "standort": S, "laenge_km": 10.0}),
          False)

    # Ausprobiert und verworfen: gleiche Veranstalter-Domain als Kriterium.
    # Diese beiden sind verschiedene Veranstaltungen (Alfhausen und
    # Ibbenbüren, 20 km auseinander), verbunden nur durch das
    # Regionalportal laufen-os.de.
    a1 = {"name": "Alfhausener Volkslauf", "datum_start": "2026-09-19", "standort": "Alfhausen",
          "lat": 52.51, "lon": 7.95, "laenge_km": 10.0,
          "veranstalter_url": "https://www.laufen-os.de/wettkaempfe/13-alfhausener-volkslauf"}
    a2 = {"name": "MBH Benefizlauf", "datum_start": "2026-09-19", "standort": "Ibbenbüren",
          "lat": 52.28, "lon": 7.72, "laenge_km": 10.0,
          "veranstalter_url": "https://www.laufen-os.de/wettkaempfe/mbh-benefizlauf"}
    check("gleiche Domain, andere Veranstaltung", is_same_event(a1, a2), False)

    # Einwortiger Veranstaltungsname, eine Zeile ohne Wettbewerb: so stand
    # der SAARathon zweimal mit 42,2 km in den Daten (einmal mit
    # Wettbewerbs-Bezeichnung und offizieller Seite, einmal ohne und mit
    # Portallink). "SAARathon" ist ein Wort - zu kurz für die
    # Teilmengen-Regel -, und gegen die lange Bezeichnung reicht die
    # Ähnlichkeit nicht.
    SD, SS = "2026-10-11", "Saarbrücken"
    sa1 = {"name": "SAARathon", "wettbewerb": "42,195 km Weltkulturerbe-Marathon",
           "datum_start": SD, "standort": SS, "laenge_km": 42.2, "art1": "Laufen",
           "veranstalter_url": "https://www.westspangenlauf.de"}
    sa2 = {"name": "SAARathon", "datum_start": SD, "standort": SS, "laenge_km": 42.2,
           "art1": "Laufen",
           "veranstalter_url": "https://laufen.de/laufkalender/details/26V14000009060002"}
    check("gleicher Name, eine Zeile ohne Wettbewerb", is_same_event(sa1, sa2), True)
    # Nennen BEIDE einen Wettbewerb, ist das Label das Unterscheidende -
    # ein 10-km-Lauf und ein 10-km-Walking sind zwei Einträge.
    sa3 = dict(sa2, wettbewerb="42,2 km Nordic Walking")
    check("beide mit Wettbewerb bleiben getrennt", is_same_event(sa1, sa3), False)
    # Dieselbe Strecke, andere Sportart: auch kein Duplikat (Datenregel 1).
    sa4 = dict(sa2, art1="Wandern")
    check("andere Sportart bleibt getrennt", is_same_event(sa1, sa4), False)

    # Generische Namen am selben Tag in verschiedenen Städten.
    s1 = {"name": "Silvesterlauf", "datum_start": "2026-12-31",
          "standort": "Salzburg", "lat": 47.80, "lon": 13.04, "laenge_km": 10.0}
    s2 = {"name": "Silvesterlauf", "datum_start": "2026-12-31",
          "standort": "München", "lat": 48.14, "lon": 11.58, "laenge_km": 10.0}
    check("gleicher Name, andere Stadt", is_same_event(s1, s2), False)


def test_namensvereinheitlichung() -> None:
    """Die Namens-Vereinheitlichung lebt in clean_events.py, die Regeln
    lassen sich aber ohne Netzwerk und ohne events.json prüfen."""
    print("\nNamen vereinheitlichen (clean_events):")
    from clean_events import (  # lokaler Import, nur für diesen Test
        name_quality_hard, name_quality_soft, repair_umlaut_spelling, tidy_name,
    )

    check("doppeltes Leerzeichen", tidy_name("17.  Preungesheimer Dorflauf"),
          "17. Preungesheimer Dorflauf")
    check("fehlendes Leerzeichen nach Nummer", tidy_name("37.Bessunger Merck-Lauf"),
          "37. Bessunger Merck-Lauf")

    # GROSSSCHREIBUNG darf die Mehrheit nicht durchsetzen.
    check("GROSSGESCHRIEBEN verliert",
          name_quality_hard("MARATHON MÜNCHEN") < name_quality_hard("Marathon München by Brooks"),
          True)

    # Umlaute sind eine Reparatur, keine Auswahl: die Auflagen-Nummer der
    # Mehrheit bleibt, die Schreibweise wird korrigiert.
    check("Umschrift wird repariert",
          repair_umlaut_spelling("22. Baedleslauf", ["22. Baedleslauf", "Bädleslauf"]),
          "22. Bädleslauf")
    check("ß wird repariert",
          repair_umlaut_spelling("Int. Pronsfelder Volks- und Strassenlauf",
                                 ["Int. Pronsfelder Volks- und Strassenlauf",
                                  "39. Int. Pronsfelder Volks- und Straßenlauf"]),
          "Int. Pronsfelder Volks- und Straßenlauf")
    # Kein Umlaut-Transfer zwischen verschiedenen Namen - daran ist ein
    # früherer Versuch gescheitert ("Wiesent Challenge" verlor gegen einen
    # 108 Zeichen langen Namen, nur weil darin "Straßenlauf" vorkam).
    check("kein Transfer zwischen verschiedenen Namen",
          repair_umlaut_spelling("Wiesent Challenge",
                                 ["Wiesent Challenge", "5. Wiesent-Challenge 10 km Straßenlauf"]),
          "Wiesent Challenge")

    # Ein 108-Zeichen-"Name" ist eine Beschreibung und verliert.
    lang = "5. Wiesent-Challenge 10 km Straßenlauf und 12 km Panoramatrail mit Oberfränkischen Meisterschaften im Trail-Lauf"
    check("überlanger Name verliert",
          name_quality_soft(lang) < name_quality_soft("Wiesent Challenge"), True)


def test_zeitrennen() -> None:
    """Zeitlich begrenzte Rennen (24-Stunden-Lauf, 6h, 12h) haben keine
    Distanz - die Dauer steht in `dauer_h` und erscheint in derselben
    Spalte wie die Distanz. Die beiden Fehlerquellen der Erkennung stehen
    hier als Testfall, weil beide in echten Daten vorkommen: "229 hm"
    sind Höhenmeter, und ein "Zeitlimit: 6 Stunden" ist eine
    Zielschlusszeit - die macht aus einem Marathon kein 6-Stunden-Rennen."""
    print("\nZeitrennen (parse_duration_h) und Backyard Ultra:")
    check("24-Stunden-Lauf", parse_duration_h("24-Stunden-Lauf Hamburg"), 24.0)
    check("6h mit Anhang", parse_duration_h("6h Lauf im Stadtpark"), 6.0)
    check("ausgeschrieben", parse_duration_h("12 Stunden von Aachen"), 12.0)
    check("englisch", parse_duration_h("24 hours of Zurich"), 24.0)
    check("Bindestrich-Form", parse_duration_h("6-Stunden-Lauf"), 6.0)

    # Höhenmeter sind keine Stunden.
    check("229 hm ist keine Dauer", parse_duration_h("Berglauf 8,5 km | 229 hm"), None)
    # Zielschlusszeiten dürfen nicht zur Renndauer werden.
    check("Zeitlimit zählt nicht",
          parse_duration_h("Marathon (Zeitlimit: 6 Stunden)"), None)
    check("Karenzzeit zählt nicht",
          parse_duration_h("Halbmarathon, Karenzzeit 3 h"), None)
    check("Startzeit zählt nicht", parse_duration_h("Start 10 Uhr, 10 km"), None)
    check("100 Stunden unplausibel", parse_duration_h("100-Stunden-Rennen"), None)

    # Ein Zeitrennen ist ein eigener Eintrag, obwohl es keine Distanz hat -
    # ohne diese Ausnahme fiel es in expand_competitions() heraus. Und zwei
    # Zeitrennen derselben Veranstaltung sind zwei Einträge, nicht eines
    # (der Dedupe-Schlüssel enthält deshalb die Dauer).
    comps = parse_competitions(
        ["24-Stunden-Lauf", "6h Lauf", "Halbmarathon | 21,1 km", "12 Stunden"], CONFIG)
    check("vier Wettbewerbe erkannt", len(comps), 4)
    check("Dauer am Wettbewerb", [c.dauer_h for c in comps], [24.0, 6.0, None, 12.0])
    base = Event(name="Zeitlauf Testheim", datum_start="2027-05-01", standort="Testheim")
    evs = expand_competitions(base, comps, CONFIG)
    check("vier Einträge", len(evs), 4)
    check("Zeitrennen ohne Distanz", [e.laenge_km for e in evs], [None, None, 21.1, None])
    check("Dauer am Eintrag", [e.dauer_h for e in evs], [24.0, 6.0, None, 12.0])

    # Der Mad Chicken Run (vom Nutzer gemeldet): Die Quelle schreibt
    # "24h Solo auf einer 2km MotoCross-Strecke mit je 60HM (2km)". Die
    # 2 km sind eine RUNDE, keine Renndistanz - als Distanz gespeichert
    # hätte die 5-km-Mindestdistanz den Eintrag anschließend verworfen,
    # und genau deshalb fehlten die vier 24h-Wettbewerbe in der Liste.
    chicken = parse_competitions([
        "24h Solo auf einer 2km MotoCross-Strecke mit je 60HM (2km)",
        "24h Solo auf einer entspannten 2km Runde mit je 15HM (2km)",
        "24h im 5er Team auf der 2km Cross Runde - Einzelergebnisse (2km)",
        "Marathon in schönen 21 flachen Runden (42km)",
        "Halbmarathon in 11 flachen Runden (22km)",
        "10km Einstiegsdroge in 5 flachen Runden (10km)",
    ], CONFIG)
    check("Rundenlänge wird nicht zur Distanz",
          [(c.laenge_km, c.dauer_h) for c in chicken],
          [(None, 24.0), (42.0, None), (22.0, None), (10.0, None)])

    # Ein Umrechnungssatz begrenzt nichts: "24 Stunden ergeben 100 Meilen"
    # beschreibt die Rundenlänge eines Backyard, der läuft aber weiter,
    # bis nur noch eine Person übrig ist.
    check("Umrechnungssatz ist keine Zeitvorgabe",
          parse_duration_h("6,708 km pro Runde beim SWUB; 24 Stunden ergeben 100 Meilen."),
          None)

    # Distanzen aus bestehenden Backyard-Einträgen herausnehmen.
    from clean_events import clear_backyard_lap_km  # lokaler Import
    rows = [
        {"name": "Hofer Backyard Ultra", "art2": "Backyard Ultra", "laenge_km": 7.0},
        {"name": "Backyard SWUB", "art2": "Backyard Ultra", "laenge_km": 6.7},
        {"name": "RET-Team Backyard", "art2": "Backyard Ultra", "laenge_km": 80.0},
        {"name": "Stadtlauf", "art2": "Straße", "laenge_km": 7.0},
    ]
    cleared, to_check = clear_backyard_lap_km(rows)
    # Seit dem 13.10.2026 fällt JEDE Distanz an einem Backyard (vom
    # Nutzer: "die Distanz kann unendlich sein") - auch die 80 km.
    check("drei Distanzen entfernt", len(cleared), 3)
    check("Distanz ist weg", [r["laenge_km"] for r in rows], [None, None, None, 7.0])
    check("nichts mehr nur gemeldet", len(to_check), 0)

    # Die Laufen-Kategorie "Backyard Ultra" (bis 19.09.2026 "Backcountry
    # Ultra") ist das Last-Man-Standing-Format.
    check("Backyard Ultra", guess_art2("Backyard Ultra Berlin", CONFIG),
          "Backyard Ultra")
    check("Schreibweise egal", guess_art2("Murr BackYard 12h", CONFIG),
          "Backyard Ultra")
    check("Last Man Standing", guess_art2("Last Man Standing Hamburg", CONFIG),
          "Backyard Ultra")
    # Höhenmeter machen aus einem Backyard keinen Trail: Das
    # backyard-Stichwort steht VOR der Berg-/Höhenmeter-Zeile.
    check("Backyard mit Höhenmetern bleibt Backyard",
          guess_art2("Backyard Ultra Harz (120 Höhenmeter pro Runde)", CONFIG),
          "Backyard Ultra")
    # ABER: Steht "Trail" im Namen, ist es ein Trailrun, der das Wort nur
    # mitträgt. Deshalb steht das backyard-Stichwort NACH der Trail-Regel.
    check("Backyard Ultra Trail ist ein Trail",
          guess_art2("Backyard Ultra Trail Harz", CONFIG), "Trail")
    # Ein "Backcountry Ultra" ist ein langer Trailrun, kein Rundenformat -
    # seit die Kategorie das Format nennt, ist das Stichwort weg.
    check("Backcountry Ultra Trail ist ein Trail",
          guess_art2("Backcountry Ultra Trail 80 km", CONFIG), "Trail")


def test_vergangene_events() -> None:
    """Vergangene Events gehören nicht in die Liste - weder neu
    hereingeholt (scraper_lib.filter_past) noch in der bestehenden Datei
    (clean_events.drop_past_events). Beide Seiten müssen sich gleich
    verhalten, sonst holt der Scraper heraus, was das Aufräumen entfernt
    hat, oder umgekehrt."""
    from clean_events import drop_past_events  # lokaler Import, nur hier
    from scraper_lib import filter_past

    HEUTE = "2026-09-16"
    roh = [
        Event(name="vorbei", datum_start="2026-09-01", standort="X"),
        Event(name="heute", datum_start=HEUTE, standort="X"),
        # Etappenrennen, das gestern begonnen hat: läuft noch.
        Event(name="mehrtägig läuft", datum_start="2026-09-14",
              datum_ende="2026-09-18", standort="X"),
        Event(name="mehrtägig vorbei", datum_start="2026-09-01",
              datum_ende="2026-09-03", standort="X"),
        # Unlesbares Datum: nicht löschen, nur behalten (dieselbe Linie
        # wie bei den Distanzen - zu viel gelöscht ist unsichtbar).
        Event(name="kaputtes Datum", datum_start="irgendwann", standort="X"),
        Event(name="zukunft", datum_start="2027-01-01", standort="X"),
    ]
    kept, skipped = filter_past(roh, HEUTE)
    check("filter_past verwirft nur Vergangenes", skipped, 2)
    check("filter_past behält heute, laufend, kaputt, künftig",
          [e.name for e in kept],
          ["heute", "mehrtägig läuft", "kaputtes Datum", "zukunft"])

    # Dieselben Fälle als Dicts durch das Aufräumskript.
    dicts = [{"name": e.name, "datum_start": e.datum_start,
              "datum_ende": e.datum_ende} for e in roh]
    kept2, dropped2 = drop_past_events(dicts, HEUTE)
    check("drop_past_events entfernt dieselben zwei", len(dropped2), 2)
    check("drop_past_events behält denselben Rest",
          [e["name"] for e in kept2],
          ["heute", "mehrtägig läuft", "kaputtes Datum", "zukunft"])


def test_kalenderdateien() -> None:
    """Die Dateinamen der .ics-Dateien werden ZWEIMAL berechnet: in
    scripts/build_ics.py (erzeugt die Dateien) und in events.html
    (verlinkt sie). Weichen die beiden ab, zeigt der Kalender-Knopf ins
    Leere - deshalb prüft dieser Test die JS-Umsetzung gegen die
    Python-Umsetzung, indem er node mit den gleichen Fällen aufruft."""
    import json as _json
    import re as _re
    import shutil
    import subprocess

    from build_ics import build_ics, ics_dateiname, slugify

    print("\nKalenderdateien (build_ics + events.html):")
    check("Umlaute werden zerlegt", slugify("Königsforst-Marathon"), "konigsforst-marathon")
    check("ß wird ss", slugify("Straßenlauf Groß-Gerau"), "strassenlauf-gross-gerau")
    check("Sonderzeichen werden Bindestrich",
          slugify("B2Run: Berlin (Firmenlauf!)"), "b2run-berlin-firmenlauf")

    faelle = [
        {"name": "Hofer Backyard Ultra", "datum_start": "2027-06-26",
         "datum_ende": "2027-06-28", "standort": "Hof", "land": "Deutschland",
         "art1": "Laufen", "art2": "Backyard Ultra", "dauer_h": None,
         "laenge_km": None, "wettbewerb": None, "veranstalter_url": None},
        {"name": "Königsforst-Marathon", "datum_start": "2027-03-14",
         "standort": "Bergisch Gladbach", "land": "Deutschland", "art1": "Laufen",
         "art2": "Trail", "laenge_km": 42.2, "wettbewerb": "Marathon 42.2 km",
         "veranstalter_url": "https://example.org/lauf"},
        {"name": "24h Mad Chicken Run", "datum_start": "2026-09-19",
         "datum_ende": "2026-09-20", "standort": "Kolkwitz", "land": "Deutschland",
         "art1": "Laufen", "art2": "Straße", "dauer_h": 24.0, "laenge_km": None,
         "wettbewerb": "24h Solo", "veranstalter_url": None},
        {"name": "Bodensee Openwater", "datum_start": "2027-08-28",
         "standort": "Wallhausen", "land": "Deutschland", "art1": "Schwimmen",
         "art2": "Freiwasser", "laenge_km": 2.5, "wettbewerb": "2,5 km",
         "veranstalter_url": None},
    ]

    check("Dateiname mit Dauer statt Distanz",
          ics_dateiname(faelle[2]),
          "2026-09-19-24h-mad-chicken-run-24h-kolkwitz.ics")
    check("Dateiname ohne beides",
          ics_dateiname(faelle[0]),
          "2027-06-26-hofer-backyard-ultra-x-hof.ics")

    # Mehrtägig: DTEND ist exklusiv, also der Tag NACH dem letzten.
    ics = build_ics(faelle[0], "20260101T000000Z")
    check("DTSTART", "DTSTART;VALUE=DATE:20270626" in ics, True)
    check("DTEND ist Enddatum + 1", "DTEND;VALUE=DATE:20270629" in ics, True)
    check("CRLF-Zeilenenden", ics.count("\r\n") >= 15, True)
    # Ein Label, das nur die Distanz wiederholt, gehört nicht in den Titel.
    ics2 = build_ics(faelle[3], "20260101T000000Z")
    check("Masszahl-Label nicht im Titel",
          "SUMMARY:Bodensee Openwater\r\n" in ics2, True)
    check("Komma maskiert (RFC 5545)",
          "LOCATION:Wallhausen\\, Deutschland" in ics2, True)

    # --- Gegenprobe in JavaScript ---
    node = shutil.which("node")
    if not node:
        print("  (node nicht vorhanden - JS-Gegenprobe übersprungen)")
        return
    quelle = (Path(__file__).resolve().parent.parent / "event-detail.js").read_text(encoding="utf-8")
    # Die vier Funktionen aus event-detail.js (die Detail-Box, von Liste
    # UND Karte genutzt) herausschneiden und in node laufen lassen. Kein
    # Nachbau: es läuft genau der Code, den die Seiten nutzen.
    stuecke = []
    for name in ("function icsSlug(", "function icsMasszahl(", "function eventSlug(",
                 "function icsFileName("):
        i = quelle.index(name)
        j = quelle.index("\n  }\n", i) + len("\n  }\n")
        stuecke.append(quelle[i:j])
    skript = "\n".join(stuecke) + (
        "\nconst faelle = " + _json.dumps(faelle, ensure_ascii=False) + ";"
        "\nconsole.log(JSON.stringify(faelle.map(icsFileName)));"
    )
    ergebnis = subprocess.run([node, "-e", skript], capture_output=True, text=True)
    if ergebnis.returncode != 0:
        check("JS-Gegenprobe lief", ergebnis.stderr.strip()[:200], "")
        return
    js_namen = _json.loads(ergebnis.stdout)
    py_namen = [ics_dateiname(f) for f in faelle]
    check("event-detail.js und build_ics.py erzeugen dieselben Dateinamen",
          js_namen, py_namen)

    # Und: passt der Ordner kalender/ überhaupt zu events.json?
    #
    # Das prüft die CI ohnehin (Schritt "kalender/ passt zu events.json",
    # dort inklusive Dateiinhalt) - aber erst NACH dem Push. Genau daran
    # ist die CI viermal hintereinander rot geworden: Der wöchentliche
    # Datenlauf hatte eine neue events.json committet, ohne die .ics-
    # Dateien mitzunehmen. Lokal fiel das nicht auf, weil kein Test
    # danach gesehen hat. Hier ist der Vergleich billig (nur die
    # Dateinamen, keine Inhalte) und schlägt vor dem Commit an.
    wurzel = Path(__file__).resolve().parent.parent
    events_json = wurzel / "events.json"
    kalender = wurzel / "kalender"
    if events_json.exists() and kalender.is_dir():
        events = _json.loads(events_json.read_text(encoding="utf-8"))
        erwartet = set()
        for e in events:
            if e.get("datum_start") and e.get("name"):
                erwartet.add(ics_dateiname(e))
        vorhanden = {pfad.name for pfad in kalender.glob("*.ics")}
        fehlend = sorted(erwartet - vorhanden)
        ueberzaehlig = sorted(vorhanden - erwartet)
        check("kalender/ hat eine Datei je Event (%d fehlen, %d zu viel)"
              % (len(fehlend), len(ueberzaehlig)),
              (fehlend[:3], ueberzaehlig[:3]), ([], []))
        if fehlend or ueberzaehlig:
            print("  -> python3 scripts/build_ics.py ausführen und mit committen")


def test_meldungen() -> None:
    """Nutzer-Fehlermeldungen (scripts/review_reports.py): bündeln pro
    Strecke, nicht pro Veranstaltung. Der Bündel-Schlüssel ist genau der
    distanzgenaue Override-Schlüssel - sonst würde eine Meldung zur
    10-km-Strecke am Ende die Marathonzeile derselben Veranstaltung
    korrigieren."""
    from review_reports import (  # lokaler Import, nur für diesen Test
        bundle_key, bundle_reports, parse_set,
    )

    ev = {"name": "48. Hochgratlauf", "datum_start": "2026-09-06", "laenge_km": 12.8}
    check("Bündel-Schlüssel distanzgenau",
          bundle_key(ev), "48. Hochgratlauf|2026-09-06|12.8")
    # 42,195 km und 42.2 km müssen im selben Bündel landen - die Anzeige
    # rundet auf eine Dezimalstelle, der Schlüssel muss das auch tun.
    check("Schlüssel rundet auf eine Dezimalstelle",
          bundle_key({"name": "M", "datum_start": "2026-10-11", "laenge_km": 42.195}),
          "M|2026-10-11|42.2")
    check("ohne Distanz kein Distanzteil",
          bundle_key({"name": "M", "datum_start": "2026-10-11", "laenge_km": None}),
          "M|2026-10-11")

    reports = [
        {"id": "a", "kategorie": "laenge", "beschreibung": "zu lang", "event": ev},
        {"id": "b", "kategorie": "laenge", "beschreibung": "12,4 km", "event": ev},
        {"id": "c", "kategorie": "url", "beschreibung": "Link tot",
         "event": {"name": "48. Hochgratlauf", "datum_start": "2026-09-06",
                   "laenge_km": 21.1}},
    ]
    bundles = bundle_reports(reports)["buendel"]
    check("zwei Strecken = zwei Bündel", len(bundles), 2)
    # Am häufigsten gemeldetes Bündel zuerst: zwei unabhängige Meldungen
    # sind ein stärkerer Hinweis als eine.
    check("häufigstes Bündel zuerst", bundles[0]["anzahl"], 2)
    check("Kategorien gezählt", bundles[0]["kategorien"], {"laenge": 2})

    # Komma-Distanzen aus dem Melde-Text dürfen nicht als Text landen.
    check("--set laenge_km=12,4 wird zur Zahl",
          parse_set(["laenge_km=12,4"]), {"laenge_km": 12.4})
    check("--set exclude=true wird bool",
          parse_set(["exclude=true"]), {"exclude": True})


def test_ortsverzeichnis() -> None:
    """Aufbereitung von places.json (scripts/build_places.py).

    Das Ortsverzeichnis füttert die Umkreissuche in events.html. Getestet
    wird, was dort schiefgehen kann: Großempfänger-Postleitzahlen, die
    als Ort durchgehen, und Schreibweisen, unter denen ein Ort nicht mehr
    gefunden wird.
    """
    print("\nOrtsverzeichnis (build_places):")

    # Der Suchschlüssel muss zu normalizePlaceText() in events.html
    # passen - beide Seiten müssen dasselbe aus einer Eingabe machen.
    check("Umlaute", normalisiere("München"), "muenchen")
    check("ß", normalisiere("Weißenburg"), "weissenburg")
    check("Akzent", normalisiere("Genève"), "geneve")
    # "Sankt Anton am Arlberg" heißt in den Daten "St. Anton am Arlberg" -
    # ohne Vereinheitlichung findet die Suche den Ort nicht.
    check("Sankt = St.", normalisiere("Sankt Anton"), normalisiere("St. Anton"))
    check("Bindestrich", normalisiere("Baden-Württemberg"), "baden wuerttemberg")

    # Großempfänger: In der deutschen PLZ-Datei steht bei diesen
    # Postleitzahlen ein Firmenname im Ortsfeld.
    check("Firma erkannt", _ist_firma("Mercedes-Benz Vertrieb NFZ GmbH"), True)
    check("Verein erkannt", _ist_firma("ADAC e. V."), True)
    check("echter Ort", _ist_firma("Aachen"), False)
    # Aalen endet auf "ag"+"en"; die Prüfung darf nur auf ganze Wörter
    # anschlagen, sonst verschwindet der Ort aus der Auswahl.
    check("Ort mit AG im Namen", _ist_firma("Aalen"), False)
    # Eine Rechtsform am Ende zählt nur dort, wo Ortsnamen kein
    # Kantonskürzel tragen: "Deutz AG" ist eine Firma, "Wohlen AG" die
    # Aargauer Gemeinde.
    check("Firma mit Rechtsform am Ende", _ist_firma("Deutz AG", streng=True), True)
    check("Gemeinde mit Kantonskürzel", _ist_firma("Wohlen AG"), False)

    # Was der Firmenfilter nicht erwischt ("Finanzamt Fürth"), fängt der
    # Abgleich mit dem Gazetteer: kein Ort dieses Namens in der Nähe.
    index = {
        "fuerth": [(49.48, 10.97, 132036, True)],
        "hamburg": [(53.55, 9.99, 1973896, True)],
    }
    check("Ort mit Entsprechung", _hat_ortsbezug(index, "Fürth", 49.48, 10.97), True)
    check("Finanzamt ohne Ort", _hat_ortsbezug(index, "Finanzamt Fürth", 49.48, 10.97), False)
    # "Gemeinde Ortsteil": der Ortsteil fehlt oft im Gazetteer, die
    # Gemeinde davor nicht - der Eintrag muss bleiben.
    check("Gemeinde + Ortsteil", _hat_ortsbezug(index, "Hamburg Stellingen", 53.57, 10.01), True)
    # Gleicher Name, aber 600 km entfernt: das ist ein anderer Ort.
    check("zu weit weg", _hat_ortsbezug(index, "Fürth", 53.55, 9.99), False)

    # Postleitzahlbezirke halten sich nicht an Landesgrenzen: die PLZ
    # 22113 liegt teils in Hamburg, teils in Schleswig-Holstein. Ohne
    # Zusammenfassung stand "Hamburg" zweimal in der Trefferliste.
    gruppen = {
        ("hamburg", "HH"): {"name": "Hamburg", "land": "DE", "region": "Hamburg",
                            "lats": [53.57], "lons": [10.01], "plz": {"20095"}},
        ("hamburg", "SH"): {"name": "Hamburg", "land": "DE", "region": "Schleswig-Holstein",
                            "lats": [53.53], "lons": [10.12], "plz": {"22113"}},
        ("neustadt", "BY"): {"name": "Neustadt", "land": "DE", "region": "Bayern",
                             "lats": [49.73], "lons": [11.13], "plz": {"91413"}},
        ("neustadt", "SN"): {"name": "Neustadt", "land": "DE", "region": "Sachsen",
                             "lats": [51.02], "lons": [14.21], "plz": {"01844"}},
    }
    _verschmelze_nachbarn(gruppen)
    check("Hamburg zusammengefasst", sorted(gruppen[("hamburg", "HH")]["plz"]), ["20095", "22113"])
    check("Hamburg nur noch einmal", ("hamburg", "SH") in gruppen, False)
    # Zwei echte "Neustadt" 200 km auseinander bleiben getrennt.
    check("Neustadt bleibt doppelt", ("neustadt", "SN") in gruppen, True)


# Alle HTML-Seiten des Projekts. Zwei Tests laufen darüber (JS-Syntax
# und "keine fremden Dateien"); eine neue Seite hier eintragen, nicht an
# zwei Stellen.
SEITEN = ("index.html", "events.html", "karte.html", "impressum.html", "datenschutz.html", "admin.html")


def test_js_syntax() -> None:
    """`node --check` über die Inline-Skripte und die geteilten Dateien.

    Die drei Seiten tragen ihr Skript im HTML - ein Tippfehler darin fällt
    weder Python noch einem Linter auf, die Seite bleibt einfach leer
    (genau das ist beim Umbau mehrfach passiert). Bisher stand
    "node --check nicht vergessen" nur in CLAUDE.md; jetzt prüft es der
    Test - und damit auch die CI.

    Ohne node wird übersprungen (wie die JS-Gegenprobe oben): Wer nur die
    Daten anfasst, soll nicht an einem fehlenden node hängen.
    """
    import shutil
    import subprocess

    def fehlerzeile(erg) -> str:
        """Die eine Zeile, die den Fehler benennt.

        node schreibt Datei, Codezeile, Pfeil, SyntaxError und zuletzt
        seine eigene Version - `splitlines()[-1]` wäre also "Node.js
        v22" und damit nutzlos.
        """
        if erg.returncode == 0:
            return ""
        zeilen = [z.strip() for z in erg.stderr.splitlines() if z.strip()]
        for z in zeilen:
            if "Error" in z:
                return z
        return zeilen[0] if zeilen else "node --check fehlgeschlagen"

    print("\nJavaScript-Syntax (node --check):")
    node = shutil.which("node")
    if not node:
        print("  (node nicht vorhanden - übersprungen)")
        return
    wurzel = Path(__file__).resolve().parent.parent
    import re as _re
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        for name in SEITEN:
            quelle = (wurzel / name).read_text(encoding="utf-8")
            # Nur Skripte OHNE src: die geladenen Dateien werden unten
            # einzeln geprüft.
            bloecke = _re.findall(r"<script(?![^>]*\bsrc=)[^>]*>(.*?)</script>",
                                  quelle, _re.S)
            check(f"{name}: Inline-Skript vorhanden", len(bloecke) >= 1, True)
            for nr, block in enumerate(bloecke):
                datei = Path(tmp) / f"{name}.{nr}.js"
                datei.write_text(block, encoding="utf-8")
                erg = subprocess.run([node, "--check", str(datei)],
                                     capture_output=True, text=True)
                check(f"{name}: Inline-Skript {nr} ist gültiges JS",
                      fehlerzeile(erg), "")
        for name in ("filters.js", "filter-ui.js", "event-detail.js", "auth.js", "stats.js",
                     "firebase-config.js", "vorschau-sperre.js",
                     "functions/index.js", "functions/stats_aggregate.js"):
            pfad = wurzel / name
            if not pfad.exists():
                continue
            erg = subprocess.run([node, "--check", str(pfad)],
                                 capture_output=True, text=True)
            check(f"{name} ist gültiges JS", fehlerzeile(erg), "")


def test_vorschau_sperre() -> None:
    """Anmeldesperre für /vorschau/ (19.10.2026): pages.yml setzt
    vorschau-sperre.js in JEDE Vorschau-Seite, und vor dem Upload fliegt
    eine Vorschau ohne Sperre ganz heraus. Die Sperre fragt admins/<uid>
    wie admin.html und lädt Firebase aus einem Pfad, den es geben muss."""
    print("\nAnmeldesperre der Vorschau:")
    import re as _re
    wurzel = Path(__file__).resolve().parent.parent
    js = (wurzel / "vorschau-sperre.js").read_text(encoding="utf-8")
    pfad = _re.search(r"FIREBASE_PFAD = '([^']+)'", js).group(1)
    for datei in ("firebase-app-compat.js", "firebase-auth-compat.js"):
        check(f"Firebase-Pfad der Sperre gibt es: {pfad}{datei}", (wurzel / pfad / datei).exists(), True)
    check("Sperre fragt admins/<uid> per REST", "/documents/admins/" in js and "getIdToken" in js, True)
    check("Sperre ist sofort zu (vor <body>)", js.index("root.classList.add('vs-zu')") < js.index("function zu("), True)
    check("verborgen über opacity, nicht display (Leaflet misst sonst 0 px)",
          "opacity:0!important" in js and "display:none" not in js, True)
    regeln = (wurzel / "firestore.rules").read_text(encoding="utf-8")
    check("firestore.rules: das eigene admins-Dokument ist abfragbar",
          "allow get: if request.auth != null && request.auth.uid == uid;" in regeln, True)
    # Als Text, nicht mit yaml: PyYAML gibt es in der CI nicht (13.10.2026).
    workflow = (wurzel / ".github/workflows/pages.yml").read_text(encoding="utf-8")
    teile = _re.split(r"\n      - name: ", workflow)
    schritte = {t.split("\n", 1)[0].strip(): {"run": t} for t in teile[1:]}
    namen = list(schritte)
    marke = 'src="/vorschau/vorschau-sperre.js"'
    vorbereiten = schritte.get("Vorschau vorbereiten", {}).get("run", "")
    check("pages.yml kopiert die Sperre in die Vorschau", "cp vorschau-sperre.js vorschau/" in vorbereiten, True)
    check("pages.yml setzt sie in jede Vorschau-Seite", marke in vorbereiten and "find vorschau -name '*.html'" in vorbereiten, True)
    check("pages.yml räumt die Vorschau bei einem Fehler weg", "trap 'rm -rf vorschau' ERR" in vorbereiten, True)
    waechter = schritte.get("Vorschau nur mit Anmeldesperre", {}).get("run", "")
    check("Wächter prüft dieselbe Marke und wirft sonst die Vorschau raus",
          marke in waechter and "rm -rf vorschau" in waechter, True)
    check("Wächter läuft vor dem Upload",
          "Vorschau nur mit Anmeldesperre" in namen and namen.index("Vorschau nur mit Anmeldesperre") < namen.index("Upload site as Pages artifact"), True)


def test_abo_rhythmen() -> None:
    """Die Abo-Rhythmen stehen an vier Stellen - alle müssen gleich sein.

    `auth.js` legt sie an, `functions/index.js` wertet sie aus,
    `firestore.rules` lässt nur sie durch, und `events.html` hat je
    Rhythmus zwei Texte. Läuft eine Stelle weg, entsteht der leise
    Fehler: Ein Abo wird gespeichert, aber die Function schickt nie eine
    E-Mail (oder die Regel weist es ab, und niemand weiß warum).
    """
    import re as _re

    wurzel = Path(__file__).resolve().parent.parent
    print("\nAbo-Rhythmen (auth.js / functions / rules / events.html):")

    def liste(pfad: str, muster: str) -> list[str]:
        text = (wurzel / pfad).read_text(encoding="utf-8")
        treffer = _re.search(muster, text, _re.S)
        if not treffer:
            return []
        return _re.findall(r"['\"]([a-z]+)['\"]", treffer.group(1))

    aus_auth = liste("auth.js", r"const ABO_RHYTHMEN = \[([^\]]*)\]")
    aus_function = liste("functions/index.js", r"const ABO_RHYTHMEN = \[([^\]]*)\]")
    aus_rules = liste("firestore.rules", r"rhythmus in\s*\n?\s*\[([^\]]*)\]")
    check("auth.js kennt drei Rhythmen", aus_auth, ["sofort", "woechentlich", "monatlich"])
    check("functions/index.js hat dieselben", aus_function, aus_auth)
    check("firestore.rules lässt dieselben durch", sorted(aus_rules), sorted(aus_auth))

    # Je Rhythmus ein Titel und eine Erklärung, in DE und EN.
    seite = (wurzel / "events.html").read_text(encoding="utf-8")
    fehlend = []
    for key in aus_auth:
        for schluessel in (f"abo_r_{key}:", f"abo_r_{key}_note:"):
            if seite.count(schluessel) != 2:      # einmal DE, einmal EN
                fehlend.append(f"{schluessel} ({seite.count(schluessel)}x)")
    check("events.html hat je Rhythmus Titel und Erklärung in DE und EN", fehlend, [])

    # Die Wartezeiten der Function müssen zu den Rhythmen passen.
    fn = (wurzel / "functions/index.js").read_text(encoding="utf-8")
    tage = _re.search(r"const RHYTHMUS_TAGE = \{([^}]*)\}", fn, _re.S)
    check("functions/index.js hat je Rhythmus eine Wartezeit",
          sorted(_re.findall(r"(\w+):", tage.group(1))) if tage else [],
          sorted(aus_auth))


def test_override_schluessel() -> None:
    """Die Schlüssel in manual_overrides.json müssen zu override_keys() passen.

    Die Distanz wird dort mit `:g` formatiert - also "21", nicht "21.0".
    Ein Schlüssel in der falschen Schreibweise wird **stillschweigend nie
    gefunden**: Der Override steht in der Datei, sieht richtig aus und
    tut nichts. Genau das ist beim Eintragen des Fichtel-Duplikats
    passiert (Schlüssel "…|21.0", Wirkung null). Dieser Test macht daraus
    einen roten Test statt einer stillen Lücke.
    """
    import json as _json

    print("\nSchlüssel in manual_overrides.json:")
    wurzel = Path(__file__).resolve().parent.parent
    pfad = wurzel / "scripts" / "manual_overrides.json"
    daten = _json.loads(pfad.read_text(encoding="utf-8"))

    from scraper_lib import override_keys

    falsch = []
    for key in daten:
        if key == "_readme":
            continue
        teile = key.split("|")
        if len(teile) not in (2, 3):
            falsch.append(f"{key}: weder '<Name>|<Datum>' noch '<Name>|<Datum>|<km>'"
                          f" noch '<Name>|<Datum>|@<Wettbewerb>'")
            continue
        if len(teile) == 3 and teile[2].startswith("@"):
            if override_keys(teile[0], teile[1], None, teile[2][1:])[0] != key:
                falsch.append(f"{key}: Wettbewerb mit Leerzeichen am Rand")
            continue
        if len(teile) == 3:
            try:
                km = float(teile[2])
            except ValueError:
                falsch.append(f"{key}: dritter Teil ist keine Zahl")
                continue
            # Genau so baut override_keys() den Schlüssel.
            erwartet = override_keys(teile[0], teile[1], km)[0]
            if key != erwartet:
                falsch.append(f"{key}: müsste {erwartet!r} heißen")
    check("alle Schlüssel in der Form, die find_override sucht", falsch, [])

    # Ein allgemeiner Schlüssel mit Länge beschreibt EINE Strecke
    # (05.10.2026): Eine Schwesterzeile mit ANDERER Distanz behält Länge,
    # Label und Dauer, bekommt aber Link und Koordinaten - vorher setzte
    # der Override 77 nachgetragene Strecken je Lauf auf die falsche
    # Länge, und erst add_manual_events() legte sie wieder an. Ein
    # distanzgenauer Schlüssel korrigiert weiterhin jede Distanz.
    from clean_events import apply_overrides as _apply
    import clean_events as _ce
    _alt = _ce.load_manual_overrides
    _ce.load_manual_overrides = lambda: {
        "Testtriathlon|2027-07-25": {"laenge_km": 25.5, "wettbewerb": "Sprint 25,5 km", "veranstalter_url": "https://test-tri.de/"},
        "Testlauf|2027-07-25|11": {"laenge_km": 10.5, "wettbewerb": "10,5 km"},
    }
    try:
        rows = [
            {"name": "Testtriathlon", "datum_start": "2027-07-25", "laenge_km": None, "wettbewerb": "Triathlon"},
            {"name": "Testtriathlon", "datum_start": "2027-07-25", "laenge_km": 51.5, "wettbewerb": "Olympisch 51,5 km"},
            {"name": "Testlauf", "datum_start": "2027-07-25", "laenge_km": 11.0, "wettbewerb": "11 km"},
        ]
        kept, _, _ = _apply(rows)
    finally:
        _ce.load_manual_overrides = _alt
    check("allgemeiner Schlüssel füllt die Zeile ohne Länge", (kept[0]["laenge_km"], kept[0]["wettbewerb"], kept[0].get("veranstalter_url")), (25.5, "Sprint 25,5 km", "https://test-tri.de/"))
    check("… lässt die Schwesterstrecke in Ruhe, Link kommt trotzdem", (kept[1]["laenge_km"], kept[1]["wettbewerb"], kept[1].get("veranstalter_url")), (51.5, "Olympisch 51,5 km", "https://test-tri.de/"))
    check("distanzgenauer Schlüssel korrigiert die Distanz weiterhin", (kept[2]["laenge_km"], kept[2]["wettbewerb"]), (10.5, "10,5 km"))

    # Die Auflage vorn im Namen wird überlesen (06.10.2026): Der Override
    # "35. Ismaninger Winterlaufserie|2026-12-13|21.1" traf die Zeile
    # "Ismaninger Winterlaufserie" einer zweiten Quelle nicht, und der
    # Halbmarathon stand wieder als "13.12.-21.02." da. Gleicher Tag,
    # gleiche Strecke - dieselbe Veranstaltung, mit oder ohne Nummer.
    from scraper_lib import find_override as _find
    _ov = {
        "35. Testserie|2026-12-13|21.1": {"datum_start": "2027-02-21"},
        "Testcross|2026-11-01": {"art2": "Trail"},
        "7. Testcross|2026-11-01": {"art2": "Hindernis"},
    }
    check("Override mit Auflage trifft den Namen ohne Auflage",
          _find(_ov, "Testserie", "2026-12-13", 21.1), {"datum_start": "2027-02-21"})
    check("… und umgekehrt (Override ohne, Zeile mit Auflage)",
          _find(_ov, "12. Testcross", "2026-11-01", 5.0), {"art2": "Trail"})
    check("ein exakter Schlüssel gewinnt vor dem ohne Auflage",
          _find(_ov, "7. Testcross", "2026-11-01", 5.0), {"art2": "Hindernis"})
    check("anderer Tag trifft nicht",
          _find(_ov, "Testserie", "2027-01-17", 21.1), None)
    check("eine Nummer MITTEN im Namen bleibt unterscheidend",
          _find(_ov, "Testserie 35.", "2026-12-13", 21.1), None)
    check("andere Distanz trifft den distanzgenauen Schlüssel nicht",
          _find(_ov, "Testserie", "2026-12-13", 17.0), None)

    # Allgemeiner und distanzgenauer Eintrag liegen ÜBEREINANDER - der
    # allgemeine (z. B. veranstalter_url für alle Strecken) darf nicht
    # unsichtbar werden, nur weil EINE Strecke einen eigenen Eintrag hat.
    # Sechs Veranstalter-Links aus der Linkprüfung griffen genau deshalb
    # nie (Taubertal 100, Steverlauf, ...), siehe find_override().
    from scraper_lib import find_override

    ov = {
        "_readme": "…",
        "Testlauf|2026-10-03": {"veranstalter_url": "https://testlauf.de/", "art2": "Straße"},
        "Testlauf|2026-10-03|21": {"wettbewerb": "Halbmarathon", "art2": "Trail"},
    }
    beides = find_override(ov, "Testlauf", "2026-10-03", 21.0)
    check("allgemeiner Eintrag scheint durch den distanzgenauen durch",
          beides.get("veranstalter_url"), "https://testlauf.de/")
    check("distanzgenauer Eintrag bringt seine Felder mit", beides.get("wettbewerb"), "Halbmarathon")
    check("bei Widerspruch gewinnt der distanzgenaue Eintrag", beides.get("art2"), "Trail")
    nur_allgemein = find_override(ov, "Testlauf", "2026-10-03", 10.0)
    check("andere Strecke bekommt nur den allgemeinen Eintrag",
          nur_allgemein, {"veranstalter_url": "https://testlauf.de/", "art2": "Straße"})
    check("kein Eintrag -> None", find_override(ov, "Anderer Lauf", "2026-10-03", 10.0), None)
    check("die Datei selbst bleibt unverändert", "wettbewerb" in ov["Testlauf|2026-10-03"], False)

    # Eine Strecke OHNE Distanz über ihren Wettbewerb (Korrekturliste,
    # 05.10.2026): Die nachgetragene Distanz darf nur sie treffen.
    ov["Testlauf|2026-10-03|@Hauptlauf"] = {"laenge_km": 15.0}
    check("Strecke ohne Distanz über den Wettbewerb gefunden",
          find_override(ov, "Testlauf", "2026-10-03", None, "Hauptlauf").get("laenge_km"), 15.0)
    check("andere Strecke ohne Distanz bleibt unberührt",
          "laenge_km" in find_override(ov, "Testlauf", "2026-10-03", None, "Bambini"), False)
    check("der distanzgenaue Eintrag schlägt den Wettbewerb",
          find_override(ov, "Testlauf", "2026-10-03", 21.0, "Hauptlauf").get("art2"), "Trail")


def test_suche_uebersetzungen() -> None:
    """Die Mastersuche sucht in BEIDEN Sprachen - und zwar an zwei Stellen.

    „Germany" muss auch auf der deutschen Seite alle Events in
    Deutschland finden (so vom Nutzer gewünscht). Dafür durchsucht
    `sucheHeuhaufen()` in filters.js die Übersetzungen von Land,
    Sportart, Kategorie und Ort mit. Dieselbe Regel steht ein zweites Mal
    in functions/index.js, damit ein ABO genau das trifft, was die Suche
    gezeigt hat - sonst kämen E-Mails über Events, die man nie gesehen
    hat, oder gar keine.

    Dieser Test vergleicht die beiden Listen Wert für Wert. Läuft eine
    weg, fällt es hier auf und nicht erst beim Empfänger.
    """
    import json as _json
    import re as _re
    import shutil
    import subprocess

    print("\nÜbersetzungen für die Mastersuche (filters.js == functions/index.js):")
    node = shutil.which("node")
    if not node:
        print("  – übersprungen: node nicht gefunden")
        return
    wurzel = Path(__file__).resolve().parent.parent

    # filters.js lässt sich in node laden: Die Datei hängt ihr Objekt an
    # `window` ODER `globalThis` (siehe letzte Zeile dort).
    skript = (f"require({_json.dumps(str(wurzel / 'filters.js'))});"
              "const T = globalThis.EnduranceFilters.VALUE_TRANSLATIONS;"
              "const raus = {};"
              "['land','art1','art2','standort'].forEach(f => {"
              "  raus[f] = {};"
              "  Object.keys(T[f] || {}).forEach(k => { raus[f][k] = [T[f][k].de, T[f][k].en]; });"
              "});"
              "console.log(JSON.stringify(raus));")
    ergebnis = subprocess.run([node, "-e", skript], capture_output=True, text=True)
    if ergebnis.returncode != 0:
        check("filters.js ließ sich laden", ergebnis.stderr.strip()[:200], "")
        return
    aus_filters = _json.loads(ergebnis.stdout)

    # Und die Kopie in der Cloud Function.
    fn = (wurzel / "functions" / "index.js").read_text(encoding="utf-8")
    start = fn.index("const SUCH_UEBERSETZUNGEN")
    ende = fn.index("\n};", start) + 3
    skript2 = (fn[start:ende].replace("const SUCH_UEBERSETZUNGEN", "const S")
               + "console.log(JSON.stringify(S));")
    ergebnis2 = subprocess.run([node, "-e", skript2], capture_output=True, text=True)
    if ergebnis2.returncode != 0:
        check("functions/index.js ließ sich lesen", ergebnis2.stderr.strip()[:200], "")
        return
    aus_function = _json.loads(ergebnis2.stdout)

    for feld in ("land", "art1", "art2", "standort"):
        check(f"{feld}: dieselben Schlüssel",
              sorted(aus_filters.get(feld, {})), sorted(aus_function.get(feld, {})))
        abweichend = sorted(k for k in aus_filters.get(feld, {})
                            if aus_filters[feld][k] != aus_function.get(feld, {}).get(k))
        check(f"{feld}: dieselben Übersetzungen", abweichend, [])

    # Und die Suche selbst: ein deutsches Event muss auf englische
    # Eingaben antworten und umgekehrt.
    probe = _json.dumps({"name": "Marathon München by Brooks", "wettbewerb": "42,2 km",
                         "standort": "München", "land": "Deutschland",
                         "art1": "Laufen", "art2": "Straße"}, ensure_ascii=False)
    skript3 = (f"require({_json.dumps(str(wurzel / 'filters.js'))});"
               "const EF = globalThis.EnduranceFilters;"
               f"const e = {probe};"
               "const st = EF.createState(); const raus = {};"
               "['germany','munich','münchen','running','road','zzznix'].forEach(q => {"
               "  st.suche = q; raus[q] = EF.matchEvent(st, e); });"
               "console.log(JSON.stringify(raus));")
    ergebnis3 = subprocess.run([node, "-e", skript3], capture_output=True, text=True)
    treffer = _json.loads(ergebnis3.stdout) if ergebnis3.returncode == 0 else {}
    check("englische Eingaben treffen deutsche Daten",
          [treffer.get(q) for q in ("germany", "munich", "running", "road")],
          [True, True, True, True])
    check("deutsche Eingaben treffen weiterhin", treffer.get("münchen"), True)
    check("ein Unsinnswort trifft nicht", treffer.get("zzznix"), False)


def test_keine_fremden_dateien() -> None:
    """Keine Skripte und Stylesheets von fremden Servern.

    Leaflet, markercluster und die Firebase-SDKs liegen in `vendor/`.
    Ein CDN-Abruf überträgt die IP-Adresse jedes Besuchers an Dritte -
    bei JEDEM Aufruf, auch wenn niemand sich anmeldet. Rutscht wieder
    ein CDN-Verweis in eine Seite, fällt es hier auf, nicht erst in der
    Datenschutzerklärung.

    Die OpenStreetMap-Kacheln sind die eine bewusste Ausnahme: Eine
    Karte ohne Kartenbilder gibt es nicht. Sie stehen deshalb in der
    Datenschutzerklärung (und nur die Kartenseite lädt sie).
    """
    import re as _re

    wurzel = Path(__file__).resolve().parent.parent
    print("\nKeine fremden Dateien (Selbst-Hosten):")
    erlaubt = ("tile.openstreetmap.org", "www.openstreetmap.org")
    for name in SEITEN:
        quelle = (wurzel / name).read_text(encoding="utf-8")
        # Kommentare weg, sonst zählt die Begründung als Treffer.
        ohne = _re.sub(r"<!--.*?-->", "", quelle, flags=_re.S)
        # Nur GELADENE Dateien zählen: src an beliebigen Tags und href
        # an <link>. Ein <a href> auf geonames.org ist ein Link, den
        # jemand anklicken KANN - dabei werden keine Daten übertragen,
        # solange niemand klickt. (Genau daran ist diese Prüfung beim
        # ersten Versuch gescheitert: die GeoNames-Namensnennung in der
        # Fußzeile galt als Treffer.)
        treffer = [u for u in _re.findall(r'\ssrc="(https?://[^"]+)"', ohne)
                   + _re.findall(r'<link\b[^>]*\shref="(https?://[^"]+)"', ohne)
                   if not any(ok in u for ok in erlaubt)]
        check(f"{name} lädt nichts von fremden Servern", treffer, [])
    for name in ("auth.js", "filters.js", "filter-ui.js", "event-detail.js", "stats.js"):
        quelle = (wurzel / name).read_text(encoding="utf-8")
        ohne = _re.sub(r"//[^\n]*", "", quelle)
        treffer = [u for u in _re.findall(r"['\"](https?://[^'\"]+\.js)['\"]", ohne)
                   if not any(ok in u for ok in erlaubt)]
        check(f"{name} lädt kein fremdes Skript nach", treffer, [])
    # Und die Dateien, auf die verwiesen wird, müssen wirklich da sein.
    fehlend = []
    for name in SEITEN:
        quelle = (wurzel / name).read_text(encoding="utf-8")
        ohne = _re.sub(r"<!--.*?-->", "", quelle, flags=_re.S)
        for pfad in _re.findall(r'(?:src|href)="(vendor/[^"?]+)"', ohne):
            if not (wurzel / pfad).exists():
                fehlend.append(f"{name} -> {pfad}")
    check("alle vendor-Dateien liegen im Repo", fehlend, [])


def test_serien() -> None:
    """Laufserien (06.10.2026): gleicher Name ohne Auflage, gleicher Ort,
    Termine höchstens SERIE_MAX_LUECKE_TAGE auseinander = EINE
    Veranstaltung. Die Regel steht zweimal - EF.markSerien/EF.groupKey in
    filters.js (Liste, Karte, Box, Startseite) und serien_cluster in
    scraper_lib (clean_events, Rauchtest). Hier beide gegeneinander, am
    ganzen Bestand, plus die Gegenproben."""
    import json as _json
    import subprocess as _sp

    print("\nLaufserien (eine Veranstaltung, mehrere Termine):")
    from scraper_lib import SERIE_MAX_LUECKE_TAGE, serien_cluster, serien_schluessel
    from clean_events import unify_auflage_in_serien

    def z(name, datum, ort="Ismaning", km=10.0):
        return {"name": name, "datum_start": datum, "datum_ende": datum, "standort": ort, "laenge_km": km}

    rows = [z("35. Ismaninger Winterlaufserie", "2027-01-17", km=17), z("Ismaninger Winterlaufserie", "2026-12-13", km=13),
            z("35. Ismaninger Winterlaufserie", "2027-02-21", km=21.1),
            z("Stadtlauf", "2026-10-10", "Lauf"), z("Stadtlauf", "2027-10-09", "Lauf"),
            z("1. Lauf Winterserie", "2027-01-10", "Hameln"), z("2. Lauf Winterserie", "2027-01-24", "Hameln"),
            z("Nikolauslauf", "2026-12-06", "Ulm"), z("Nikolauslauf", "2026-12-06", "Ulm", 5.0),
            z("Frühjahrslauf", "2027-03-01", "Köln"), z("Frühjahrslauf", "2027-06-10", "Köln")]   # 101 Tage
    cl = serien_cluster(rows)
    groessen = sorted(len(c) for c in cl)
    check("Ismaning: drei Termine, mit und ohne Auflage, eine Serie",
          [sorted(e["datum_start"] for e in c) for c in cl if c[0]["standort"] == "Ismaning"],
          [["2026-12-13", "2027-01-17", "2027-02-21"]])
    check("Ausgabe 2026 und 2027 (365 Tage) sind zwei Veranstaltungen",
          sum(1 for c in cl if c[0]["standort"] == "Lauf"), 2)
    check("'1. Lauf' und '2. Lauf' derselben Serie sind eine", sum(1 for c in cl if c[0]["standort"] == "Hameln"), 1)
    check("zwei Strecken am selben Tag bleiben eine Veranstaltung", sum(1 for c in cl if c[0]["standort"] == "Ulm"), 1)
    check(f"{SERIE_MAX_LUECKE_TAGE + 1} Tage Lücke trennen", sum(1 for c in cl if c[0]["standort"] == "Köln"), 2)
    check("Größen der Cluster", groessen, [1, 1, 1, 1, 2, 2, 3])

    geaendert = unify_auflage_in_serien(rows)
    check("Auflage innerhalb der Serie angeglichen (ohne -> mit Nummer)",
          sorted({e["name"] for e in rows if e["standort"] == "Ismaning"}), ["35. Ismaninger Winterlaufserie"])
    check("… genau eine Zeile umbenannt", len(geaendert), 1)
    check("… idempotent", unify_auflage_in_serien(rows), [])
    zwei = [z("36. Hamelner Winterlaufserie", "2026-11-22", "Hameln"), z("37. Hamelner Winterlaufserie", "2026-12-05", "Hameln")]
    check("zwei verschiedene Nummern bleiben stehen", (unify_auflage_in_serien(zwei), sorted(e["name"] for e in zwei)),
          ([], ["36. Hamelner Winterlaufserie", "37. Hamelner Winterlaufserie"]))

    # Beide Fassungen am ganzen Bestand: dieselben Schlüssel.
    wurzel = Path(__file__).resolve().parent.parent
    events = _json.loads((wurzel / "events.json").read_text(encoding="utf-8"))
    py = serien_schluessel(events)
    py_keys = sorted(py[id(e)] for e in events if id(e) in py)
    skript = (
        "const fs=require('fs');const window={};"
        f"new Function('window', fs.readFileSync({str(wurzel / 'filters.js')!r},'utf8'))(window);"
        "const EF=window.EnduranceFilters;"
        f"const ev=JSON.parse(fs.readFileSync({str(wurzel / 'events.json')!r},'utf8'));"
        "EF.markSerien(ev);"
        "process.stdout.write(JSON.stringify({luecke: EF.SERIE_MAX_LUECKE_TAGE, keys: ev.map(EF.groupKey).sort()}));"
    )
    aus = _sp.run(["node", "-e", skript], capture_output=True, text=True, check=True)
    js = _json.loads(aus.stdout)
    check("SERIE_MAX_LUECKE_TAGE ist in filters.js und scraper_lib gleich", js["luecke"], SERIE_MAX_LUECKE_TAGE)
    check(f"EF.groupKey und serien_schluessel liefern am Bestand dieselben Schlüssel ({len(set(py_keys))} Veranstaltungen)",
          js["keys"] == py_keys, True)


def test_laender_maske() -> None:
    """laender.json - die Umrisse für die graue Maske auf karte.html.

    Der Prüfstein sind die SCHLÜSSEL: Sie müssen genau die Länder heißen,
    die filters.js in LAENDER führt (und events.json im Feld `land`).
    Wird dort eines umbenannt oder ergänzt, ohne die Maske neu zu bauen,
    liegt das Land stillschweigend unter dem grauen Schleier - sichtbar
    nur, wenn man genau hinschaut.
    """
    import json as _json
    import re as _re

    print("\nLändermaske (build_laender):")
    wurzel = Path(__file__).resolve().parent.parent
    pfad = wurzel / "laender.json"
    check("laender.json liegt im Repo", pfad.exists(), True)
    if not pfad.exists():
        print("  -> python3 scripts/build_laender.py ausführen")
        return
    daten = _json.loads(pfad.read_text(encoding="utf-8"))
    laender = daten.get("laender") or {}

    # Die Namen aus filters.js (LAENDER = [...]) - dieselbe Quelle, aus
    # der die Filterliste entsteht.
    quelle = (wurzel / "filters.js").read_text(encoding="utf-8")
    treffer = _re.search(r"const LAENDER = \[([^\]]*)\]", quelle)
    aus_filters = _re.findall(r"'([^']+)'", treffer.group(1)) if treffer else []
    check("dieselben Länder wie in filters.js",
          sorted(laender), sorted(aus_filters))

    # Der Rahmen der Karte (REGION_BOUNDS in karte.html, seit dem
    # 18.10.2026 der kleinste Zoom) muss die Umrisse umfassen – kommt ein
    # Land dazu, ohne dass er mitwächst, steht es ganz herausgezoomt halb
    # außerhalb des Bildes.
    punkte = [p for ringe in laender.values() for ring in ringe for p in ring]
    karte = (wurzel / "karte.html").read_text(encoding="utf-8")
    rahmen = _re.search(r"REGION_BOUNDS = L\.latLngBounds\(\[([\d.]+), ([\d.]+)\], "
                        r"\[([\d.]+), ([\d.]+)\]\)", karte)
    if rahmen and punkte:
        s_lat, w_lon, n_lat, o_lon = (float(x) for x in rahmen.groups())
        lons = [p[0] for p in punkte]
        lats = [p[1] for p in punkte]
        check("REGION_BOUNDS in karte.html umfasst laender.json (±0,05°)",
              s_lat <= min(lats) + 0.05 and w_lon <= min(lons) + 0.05
              and n_lat >= max(lats) - 0.05 and o_lon >= max(lons) - 0.05, True)
    else:
        check("REGION_BOUNDS in karte.html gefunden", bool(rahmen), True)

    # Jeder Ring ist eine Fläche: mindestens vier Punkte und geschlossen.
    offen = [name for name, ringe in laender.items()
             for ring in ringe if len(ring) < 4 or ring[0] != ring[-1]]
    check("alle Ringe sind geschlossene Flächen", offen, [])

    # Die Koordinaten stehen als [lon, lat] (GeoJSON-Reihenfolge) - karte.html
    # dreht sie für Leaflet. Verdrehte Werte fielen sonst erst am Bild auf.
    verdreht = []
    for name, ringe in laender.items():
        for ring in ringe:
            for lon, lat in ring[:50]:
                if not (-30 <= lon <= 60 and 35 <= lat <= 72):
                    verdreht.append((name, lon, lat))
                    break
    check("Koordinaten liegen in Europa und in der Reihenfolge lon, lat",
          verdreht[:3], [])

    # Größe: Die Datei lädt die Karte bei jedem Aufruf nach. 200 KB wären
    # keine Maske mehr, sondern ein zweites events.json.
    kb = pfad.stat().st_size // 1024
    check("laender.json bleibt klein (%d KB)" % kb, kb < 200, True)


def test_web_data() -> None:
    """Die kompakte Datenfassung (scripts/build_web_data.py → events.web.json)
    muss verlustfrei sein - in Python UND mit dem echten Dekodierer aus
    filters.js (node). Ein Format, das beim Rückweg etwas verliert, fiele
    sonst erst im Browser auf, und dort still: eine fehlende Distanz ist
    nur eine leere Zelle."""
    import json as _json
    import shutil
    import subprocess

    from build_web_data import FORMAT, dekodieren, kodieren

    print("\nKompakte Datenfassung (build_web_data + EF.decodeWebData):")
    # Feld fehlt (Index -1) ist etwas anderes als Feld = null.
    faelle = [
        {"name": "A", "laenge_km": 10.0, "dauer_h": None, "charity": True},
        {"name": "B", "laenge_km": 10.0},
        {"name": "A", "wettbewerb": "10 km", "lat": 48.1, "lon": 11.5},
    ]
    kodiert = kodieren(faelle, stand="2026-09-22")
    check("Format", kodiert["format"], FORMAT)
    check("Wörterbuch fasst gleiche Werte zusammen", kodiert["werte"]["name"], ["A", "B"])
    check("fehlendes Feld ist -1", kodiert["zeilen"]["charity"], [0, -1, -1])
    check("null bleibt ein Wert", kodiert["werte"]["dauer_h"], [None])
    check("Rückweg in Python", dekodieren(kodiert), faelle)

    wurzel = Path(__file__).resolve().parent.parent
    events = _json.loads((wurzel / "events.json").read_text(encoding="utf-8"))
    daten = kodieren(events, stand="2026-09-22")
    check("Rückweg über den ganzen Bestand (Python)", dekodieren(daten) == events, True)
    roh = len(_json.dumps(events, ensure_ascii=False, separators=(",", ":")))
    kompakt = len(_json.dumps(daten, ensure_ascii=False, separators=(",", ":")))
    check("kompakte Fassung ist kleiner als events.json ohne Einrückung", kompakt < roh, True)

    node = shutil.which("node")
    if not node:
        print("  (node fehlt - Dekodierer aus filters.js nicht geprüft)")
        return
    tmp = wurzel / "scripts" / ".web_data_test.json"
    tmp.write_text(_json.dumps(daten, ensure_ascii=False), encoding="utf-8")
    try:
        skript = (
            "global.window = global;"
            f"require({_json.dumps(str(wurzel / 'filters.js'))});"
            f"const d = JSON.parse(require('fs').readFileSync({_json.dumps(str(tmp))}, 'utf8'));"
            "process.stdout.write(JSON.stringify(global.EnduranceFilters.decodeWebData(d)));"
        )
        ergebnis = subprocess.run([node, "-e", skript], capture_output=True, text=True)
    finally:
        tmp.unlink(missing_ok=True)
    if ergebnis.returncode != 0:
        check("filters.js dekodiert in node", ergebnis.stderr.strip()[:300], "")
        return
    zurueck = _json.loads(ergebnis.stdout)
    check("Rückweg mit EF.decodeWebData (node) über den ganzen Bestand", zurueck == events, True)
    print("  ✓ %d Events, %d KB → %d KB (ohne gzip)" % (len(events), roh // 1024, kompakt // 1024))


def test_asset_stempel() -> None:
    """Die ?v=-Stempel an den geteilten Skripten (stamp_assets.py).

    GitHub Pages liefert jede Datei mit max-age=600. Ohne Stempel kann ein
    Browser die neue events.html mit einer zehn Minuten alten filters.js
    kombinieren - dann fehlt eine Funktion, das Inline-Skript bricht in
    seiner ersten Zeile ab und die Seite bleibt LEER. Genau das ist
    einmal passiert. Der Stempel folgt dem Inhalt, dieser Test merkt
    also, wenn ein Modul geändert und der Stempel vergessen wurde.
    """
    from stamp_assets import ASSETS, pruefe, stempel_aller_assets

    print("\nStempel der geteilten Dateien (stamp_assets.py):")
    werte = stempel_aller_assets()
    check("alle geteilten Dateien vorhanden", sorted(werte), sorted(ASSETS))
    probleme = pruefe()
    check("Stempel passen zum Inhalt der Dateien", probleme, [])
    if probleme:
        print("  -> python3 scripts/stamp_assets.py ausführen")


def test_wartung() -> None:
    """Wartungsmodus (13./14.10.2026): Der Schalter ist das Firestore-
    Dokument config/wartung (öffentlich lesbar, nur Admins schreiben -
    firestore.rules), umgelegt im Dashboard admin.html. Der Kopf-Schnipsel
    in Startseite, Liste und Karte fragt es über die REST-Schnittstelle ab
    und führt bei aktiv nach wartung.html; Impressum, Datenschutz und
    admin.html bleiben erreichbar. Dazu das Deploy-Gate in pages.yml."""
    import re as _re
    import shutil
    import subprocess
    import tempfile

    print("\nWartungsmodus (Firestore-Schalter, wartung.html, Kopf-Schnipsel) und Deploy-Gate:")
    wurzel = Path(__file__).resolve().parent.parent
    projekt = _re.search(r'projectId:\s*"([^"]+)"', (wurzel / "firebase-config.js").read_text(encoding="utf-8")).group(1)
    url = f"https://firestore.googleapis.com/v1/projects/{projekt}/databases/(default)/documents/config/wartung"
    check("wartung.json gibt es nicht mehr (Schalter liegt in Firestore)", (wurzel / "wartung.json").exists(), False)
    muster = _re.compile(r"<script>\n\(function \(\) \{\n  if \(!window\.fetch\) return;.*?</script>", _re.S)
    fassungen = {}
    for name in ("index.html", "events.html", "karte.html"):
        html = (wurzel / name).read_text(encoding="utf-8")
        m = muster.search(html)
        fassungen[name] = m.group(0) if m else None
        check(f"{name}: Wartungs-Schnipsel im Kopf", m is not None, True)
        check(f"{name}: Schnipsel fragt das Dokument des eigenen Projekts ab", m is not None and url in m.group(0), True)
        check(f"{name}: Schnipsel steht VOR den Stylesheets",
              m is not None and m.start() < html.index('<link rel="stylesheet"'), True)
    check("alle drei Kopien sind wortgleich", len(set(v for v in fassungen.values() if v)), 1)
    for name in ("impressum.html", "datenschutz.html", "admin.html", "wartung.html"):
        html = (wurzel / name).read_text(encoding="utf-8")
        check(f"{name}: KEIN Schnipsel (bleibt in der Pause erreichbar)", muster.search(html) is None, True)
    seite = (wurzel / "wartung.html").read_text(encoding="utf-8")
    check("wartung.html liest dasselbe Dokument", url in seite, True)
    check("wartung.html verlinkt Impressum und Datenschutz",
          'href="impressum.html"' in seite and 'href="datenschutz.html"' in seite, True)
    check("wartung.html fragt den Schalter erneut ab, geht zurück, kennt ?vorschau",
          "setInterval(pruefen" in seite and "location.replace('index.html')" in seite and "vorschau" in seite, True)
    check("wartung.html: noindex, kein stats.js", 'name="robots" content="noindex"' in seite and 'src="stats.js' not in seite, True)
    ohne = _re.sub(r"<!--.*?-->", "", seite, flags=_re.S)
    check("wartung.html lädt nichts von fremden Servern",
          _re.findall(r'\s(?:src|href)="(https?://[^"]+)"', ohne), [])
    node = shutil.which("node")
    if node:
        with tempfile.TemporaryDirectory() as tmp:
            for nr, block in enumerate(_re.findall(r"<script(?![^>]*\bsrc=)[^>]*>(.*?)</script>", seite, _re.S)):
                datei = Path(tmp) / f"wartung_{nr}.js"
                datei.write_text(block, encoding="utf-8")
                r = subprocess.run([node, "--check", str(datei)], capture_output=True, text=True)
                check(f"wartung.html: Inline-Skript {nr} ist gültiges JavaScript", r.returncode, 0)
    # Regeln: genau dieses eine Dokument öffentlich lesbar, schreiben nur Admins.
    regeln = (wurzel / "firestore.rules").read_text(encoding="utf-8")
    block = _re.search(r"match /config/wartung \{(.*?)\n    \}", regeln, _re.S)
    check("Regeln: config/wartung hat einen eigenen Block", block is not None, True)
    check("Regeln: lesen darf jeder, schreiben nur istAdmin()",
          block is not None and "allow read: if true;" in block.group(1) and "allow write: if istAdmin()" in block.group(1), True)
    check("Regeln: geschlossene Feldliste und Textgrenze",
          block is not None and "hasOnly(" in block.group(1) and "text_de.size() <= 300" in block.group(1), True)
    # Zwei offene Lesezugriffe, beide ohne Personenbezug: der
    # Wartungsschalter und die Instagram-Freigabe (test_insta_freigabe).
    check("Regeln: kein offener Lesezugriff außerhalb von config/wartung und insta_freigaben",
          regeln.count("allow read: if true;"), 2)
    # Dashboard: die Karte mit dem Schalter.
    admin = (wurzel / "admin.html").read_text(encoding="utf-8")
    check("admin.html: Wartungs-Karte und Schreibzugriff auf config/wartung",
          'id="wartung-card"' in admin and "collection('config').doc('wartung')" in admin, True)
    check("admin.html: Texte in beiden Sprachen", admin.count("w_bestaetigen:"), 2)
    ds = (wurzel / "datenschutz.html").read_text(encoding="utf-8")
    check("Datenschutzerklärung nennt den Wartungs-Abruf (DE und EN)", ds.count("t_wartung:"), 2)
    # Das Gate: pages.yml ruft ci.yml, der Deploy hängt daran.
    pages = (wurzel / ".github/workflows/pages.yml").read_text(encoding="utf-8")
    ci = (wurzel / ".github/workflows/ci.yml").read_text(encoding="utf-8")
    check("pages.yml: Deploy braucht die Prüfung (Gate)",
          "uses: ./.github/workflows/ci.yml" in pages and "needs: pruefung" in pages, True)
    check("ci.yml ist aufrufbar (workflow_call)", "workflow_call:" in ci, True)
    # Ohne YAML-Bibliothek (PyYAML steht nicht in requirements.txt, und in
    # der CI fehlt sie - genau daran ist der erste Lauf des Gates
    # gescheitert): der `on:`-Block ist der Text von der Zeile `on:` bis
    # zum nächsten Schlüssel ohne Einrückung, Kommentare abgezogen.
    on_block = _re.search(r"^on:\n(.*?)^\S", _re.sub(r"#[^\n]*", "", ci), _re.S | _re.M)
    check("ci.yml hat einen on:-Block", on_block is not None, True)
    check("ci.yml prüft den Deploy-Branch nicht noch einmal selbst (die Tests laufen in pages.yml)",
          on_block is not None and "claude/endurance-events-website-v1wruf" in on_block.group(1), False)
    check("ci.yml prüft den zweiten Branch weiter direkt",
          on_block is not None and "claude/website-access-9d9tg1" in on_block.group(1), True)

def test_farbschema_skript() -> None:
    """Hell und Dunkel: das Skript im <head>, das `data-theme` setzt, steht
    in allen drei Seiten WORTGLEICH.

    Es muss im Kopf jeder Seite stehen (sonst blitzt die Seite im falschen
    Schema auf), und es ist deshalb dreimal kopiert statt als Datei
    eingebunden. Drei Kopien laufen auseinander, wenn niemand hinsieht -
    dieser Test sieht hin. Dazu: der Knopf in jeder Kopfzeile und site.css
    mit beiden Schemata über `data-theme` (nicht über
    prefers-color-scheme - dann hätte der Knopf keine Wirkung).
    """
    import os
    import re

    print("\nFarbschema-Skript (hell/dunkel) in den Seiten:")
    wurzel = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
    fassungen: dict[str, str | None] = {}
    for name in ("index.html", "events.html", "karte.html", "admin.html"):
        with open(os.path.join(wurzel, name), encoding="utf-8") as f:
            html = f.read()
        m = re.search(r"<script>\n\(function \(\) \{\n  var wahl = null;.*?</script>", html, re.S)
        fassungen[name] = m.group(0) if m else None
        check(f"{name}: Skript im Kopf vorhanden", m is not None, True)
        check(f"{name}: Knopf .theme-btn in der Kopfzeile", 'class="icon-btn theme-btn"' in html, True)
        check(f"{name}: Skript steht VOR den Stylesheets",
              m is not None and m.start() < html.index("<link rel=\"stylesheet\""), True)
    check("alle vier Kopien sind wortgleich", len(set(v for v in fassungen.values() if v)), 1)
    with open(os.path.join(wurzel, "site.css"), encoding="utf-8") as f:
        css = f.read()
    check("site.css: helles Schema über :root[data-theme=\"light\"]", ':root[data-theme="light"]' in css, True)
    check("site.css: kein prefers-color-scheme", "prefers-color-scheme" in css, False)


def test_lauf_im_triathlonkalender() -> None:
    """Ein Lauf im Triathlon-Kalender bleibt ein Lauf.

    running.life führt im Triathlon-Kalender auch die Läufe der
    Triathlon-Veranstalter; der "O-SEE Ultra Trail" stand deshalb als
    Triathlon in der Liste - mit Kinderläufen unter 5 km, die der
    5-km-Regel entgingen (Nutzer 21.09.2026: "hat nichts mit Triathlon
    zu tun", und ein Triathlon kann keine Kategorie Trail haben).
    Die Regel greift NUR bei Triathlon als Voreinstellung und nie gegen
    ein Mehrsport-Stichwort.
    """
    from scraper_lib import SiteConfig, guess_art1
    print("\nLauf im Triathlon-Kalender (guess_art1):")
    tri = SiteConfig(base_url="", calendar_url="", default_art1="Triathlon")
    lauf = SiteConfig(base_url="", calendar_url="", default_art1="Laufen")
    rad = SiteConfig(base_url="", calendar_url="", default_art1="Fahrrad")
    check("O-SEE Ultra Trail im Triathlon-Kalender -> Laufen", guess_art1("O-SEE Ultra Trail", tri), "Laufen")
    check("Zittauer Gebirgslauf im Triathlon-Kalender -> Laufen", guess_art1("Zittauer Gebirgslauf", tri), "Laufen")
    check("XTERRA Germany bleibt Triathlon", guess_art1("XTERRA Germany", tri), "Triathlon")
    check("XTERRA Trail ... ist ein Lauf", guess_art1("XTERRA Trail de la Vallée de Joux", tri), "Laufen")
    check("XTERRA Trail im Laufkalender bleibt Laufen", guess_art1("XTERRA Trail de la Vallée de Joux", lauf), "Laufen")
    check("XTERRA Austria bleibt Triathlon", guess_art1("Xterra Austria", lauf), "Triathlon")
    check("Cross Triathlon Trail Edition bleibt Triathlon", guess_art1("Cross Triathlon Trail Edition", tri), "Triathlon")
    check("Swimrun bleibt Triathlon", guess_art1("Rheinsberger SwimRun", tri), "Triathlon")
    check("Ohne Laufwort bleibt die Voreinstellung", guess_art1("O-SEE Challenge", tri), "Triathlon")
    check("Im Laufkalender wie bisher", guess_art1("O-SEE Ultra Trail", lauf), "Laufen")
    check("MTB Trail Marathon im Radkalender bleibt Fahrrad", guess_art1("MTB Trail Marathon", rad), "Fahrrad")


def test_fremde_sportart() -> None:
    """Ein Radrennen bei einer Laufveranstaltung ist ein Radrennen - eine
    Triathlon-Teilstrecke dagegen nicht.

    Die Unterscheidung kostet Mühe und ist genau die, an der die erste
    Fassung dieser Regel gescheitert ist: Sie hätte aus dem "21,5 km
    Radfahren" des Aluman (einem Triathlon) ein eigenständiges
    Radrennen gemacht. Erkennungszeichen ist die VERBFORM im Label -
    "Radfahren"/"Laufen"/"Run 1" beschreibt eine Etappe, "Rad 100 km"
    ein Rennen, das man einzeln bucht.
    """
    print("\nSportart aus dem Wettbewerbs-Label:")
    from clean_events import fix_fremde_sportart_im_wettbewerb

    def lauf(name, wb, km, art1="Laufen"):
        return {"name": name, "datum_start": "2026-09-19", "wettbewerb": wb,
                "laenge_km": km, "art1": art1, "art2": "Straße"}

    # Laufveranstaltung mit eigenem Radrennen (Drei Talsperren Marathon)
    talsperren = [lauf("Drei Talsperren Marathon", "Marathon", 42.2),
                  lauf("Drei Talsperren Marathon", "Halbmarathon", 21.1),
                  lauf("Drei Talsperren Marathon", "Rad 100 km", 100.0)]
    fix_fremde_sportart_im_wettbewerb(talsperren)
    check("'Rad 100 km' bei einem Marathon -> Fahrrad",
          [e["art1"] for e in talsperren], ["Laufen", "Laufen", "Fahrrad"])
    check("Fahrrad bekommt KEINE Laufen-Kategorie",
          talsperren[2]["art2"], None)

    # Triathlon: beide Zeilen sind Teilstrecken
    roemer = [lauf("RömerMan", "42 km Radfahren mit dem Rennrad", 42.0),
              lauf("RömerMan", "10 km Laufen durch das Grünprojekt", 10.0)]
    fix_fremde_sportart_im_wettbewerb(roemer)
    check("Triathlon-Teilstrecke bleibt unangetastet",
          [e["art1"] for e in roemer], ["Laufen", "Laufen"])

    # Eine einzelne Zeile, die schon als Triathlon geführt wird
    tri = [lauf("Lauchringer Triathlon-Nacht", "ca. 18,6 km Radfahren", 18.6,
                art1="Triathlon")]
    fix_fremde_sportart_im_wettbewerb(tri)
    check("art1 'Triathlon' wird nie überschrieben", tri[0]["art1"], "Triathlon")

    # "Mountainbike" darf nicht über \bbike\b als Teilstrecke gelten
    apfel = [lauf("Frickinger Apfellauf", "Lauf", 10.0),
             lauf("Frickinger Apfellauf", "Mountainbike", 20.0)]
    fix_fremde_sportart_im_wettbewerb(apfel)
    check("'Mountainbike' neben 'Lauf' -> Fahrrad",
          [e["art1"] for e in apfel], ["Laufen", "Fahrrad"])

    # Der Name allein reicht nicht - "Rad" steckt in vielen Laufnamen
    radweg = [lauf("Lauf am Radweg", None, 10.0)]
    fix_fremde_sportart_im_wettbewerb(radweg)
    check("'Rad' im NAMEN ändert nichts", radweg[0]["art1"], "Laufen")


def test_zwei_rennen_in_einer_zeile() -> None:
    """"15 km / 21 km Crosslauf" sind zwei Rennen, nicht eines über 21 km.

    Vor diesem Fix nahm guess_distance_km() die größere Zahl - der
    15-km-Lauf des Limberglaufs Ranis fehlte dadurch komplett in der
    Liste. Ein fehlendes Event ist die unangenehmere Sorte Fehler: eine
    falsche Zahl sieht man, eine fehlende Zeile nicht.
    """
    print("\nZwei Rennen in einer Zeile:")
    from scraper_lib import parse_competitions, SiteConfig
    config = SiteConfig(base_url="", calendar_url="")

    geteilt = parse_competitions(["15 km / 21 km Crosslauf (ab Jahrgang 2008)"], config)
    check("wird in zwei Wettbewerbe geteilt",
          sorted(k.laenge_km for k in geteilt), [15.0, 21.0])
    check("der Zusatz gehört beiden",
          all("Crosslauf" in (k.label or "") for k in geteilt), True)

    # Aufteilungen DERSELBEN Strecke bleiben ein Wettbewerb
    for text, erwartet in (("19 km (14 + 5 km)", [19.0]),
                           ("100 km (10 x 10 km)", [100.0]),
                           ("42,195 km Marathon (Rundkurs 8,33 km, fünfmal)", [42.2]),
                           ("46,5 km, 3 Runden je 15,5 km", [46.5])):
        check(f"{text!r} bleibt ein Wettbewerb",
              [k.laenge_km for k in parse_competitions([text], config)], erwartet)


def test_koordinaten_widerspruch() -> None:
    """Dieselbe Veranstaltung darf nicht an zwei Orten liegen.

    Der Bodensee Marathon stand mit seiner Marathon-Strecke auf
    49.07/10.14 - das ist Franken, 168 km vom Bodensee. Ursache war der
    abgeschnittene Ortsname "Kressbronn" statt "Kressbronn am Bodensee".
    Für die Umkreissuche und die Karte ist das kein Schönheitsfehler:
    Wer im Umkreis von Friedrichshafen sucht, sah den Marathon nicht.
    """
    print("\nWidersprüchliche Koordinaten:")
    from clean_events import report_widerspruechliche_koordinaten

    falsch = [
        {"name": "Bodensee Marathon", "datum_start": "2026-09-19",
         "standort": "Kressbronn", "lat": 49.0656, "lon": 10.1379},
        {"name": "Bodensee Marathon", "datum_start": "2026-09-19",
         "standort": "Kressbronn am Bodensee", "lat": 47.5969, "lon": 9.5982},
    ]
    check("168 km auseinander wird gemeldet",
          len(report_widerspruechliche_koordinaten(falsch)), 1)

    # Nachbarorte derselben Veranstaltung sind normal (Start/Ziel getrennt)
    nah = [
        {"name": "Drei Talsperren Marathon", "datum_start": "2026-09-19",
         "standort": "Eibenstock", "lat": 50.4727, "lon": 12.6115},
        {"name": "Drei Talsperren Marathon", "datum_start": "2026-09-19",
         "standort": "Eibenstock", "lat": 50.4950, "lon": 12.5996},
    ]
    check("3 km auseinander ist kein Fall",
          report_widerspruechliche_koordinaten(nah), [])


def test_override_koordinaten() -> None:
    """Ein Override muss auch lat/lon setzen können - und beide Wege
    (Einsammeln und rückwirkendes Aufräumen) müssen dieselben Felder
    kennen. Zwei getrennte Listen wären ein Fehler, der sich erst Wochen
    später zeigt: Das Feld wirkt je nach Weg oder nicht."""
    print("\nOverride-Felder:")
    from scraper_lib import OVERRIDE_FIELDS
    import clean_events
    import inspect

    check("lat und lon sind erlaubte Override-Felder",
          {"lat", "lon"} <= set(OVERRIDE_FIELDS), True)
    quelle = inspect.getsource(clean_events.apply_overrides)
    check("apply_overrides nutzt die gemeinsame Liste",
          "OVERRIDE_FIELDS" in quelle, True)

    events = [{"name": "Bodensee Marathon", "datum_start": "2026-09-19",
               "laenge_km": 42.2, "standort": "Kressbronn",
               "lat": 49.0656, "lon": 10.1379}]
    behalten, _, geaendert = clean_events.apply_overrides(events)
    check("die Koordinaten werden wirklich verschoben",
          round(behalten[0]["lat"], 3), 47.597)
    check("und die Änderung wird berichtet",
          any("lat" in zeile for zeile in geaendert), True)

    # Ein Portallink über einem Veranstalter-Link bleibt gesperrt - außer
    # der Override sagt `portal_erlaubt` (Entscheidung des Nutzers vom
    # 13.10.2026: Vereinsseite nennt den Lauf nicht mehr, "dann ist die
    # Anmeldeseite besser"). Beide Fassungen (dict und Event) gleich.
    import scraper_lib as _sl
    _alt = clean_events.load_manual_overrides
    _alt_sl = _sl.load_manual_overrides
    fake = lambda: {
        "Gesperrt|2026-10-17": {"veranstalter_url": "https://my.raceresult.com/1/info"},
        "Erlaubt|2026-10-17": {"veranstalter_url": "https://my.raceresult.com/2/info",
                               "portal_erlaubt": True},
    }
    clean_events.load_manual_overrides = fake
    _sl.load_manual_overrides = fake
    try:
        rows = [{"name": "Gesperrt", "datum_start": "2026-10-17", "veranstalter_url": "https://verein-a.de/"},
                {"name": "Erlaubt", "datum_start": "2026-10-17", "veranstalter_url": "https://verein-b.de/"}]
        kept, _, _ = clean_events.apply_overrides(rows)
        check("ohne Flag bleibt die Vereinsseite", kept[0]["veranstalter_url"], "https://verein-a.de/")
        check("mit portal_erlaubt gilt die Anmeldeseite", kept[1]["veranstalter_url"],
              "https://my.raceresult.com/2/info")
        evs = [Event(name="Gesperrt", datum_start="2026-10-17", standort="x", veranstalter_url="https://verein-a.de/"),
               Event(name="Erlaubt", datum_start="2026-10-17", standort="x", veranstalter_url="https://verein-b.de/")]
        evs, _ = _sl.apply_manual_overrides(evs)
        check("Einsammeln: ohne Flag gesperrt", evs[0].veranstalter_url, "https://verein-a.de/")
        check("Einsammeln: mit Flag erlaubt", evs[1].veranstalter_url, "https://my.raceresult.com/2/info")
    finally:
        clean_events.load_manual_overrides = _alt
        _sl.load_manual_overrides = _alt_sl


def test_zwei_sportarten_im_namen() -> None:
    """"Swim&Run" und "Bike+Run" sind Mehrsport - "Decathlon" ist ein Laden.

    Das Trennzeichen war die Lücke: `swim ?run` erlaubte nur ein
    optionales LEERZEICHEN, und damit blieben elf Veranstaltungen
    Laufveranstaltungen, die keine sind ("Wunnebad Swim&Run", "DSW
    Swim & Run", "Kronberger Bike+Run", "Run and Bike Berlin").

    Die Gegenprobe ist der wichtigere Teil dieses Tests: "athlon" als
    Stichwort wäre naheliegend gewesen und hätte die "Decathlon Hybrid
    Series" (den Sporthändler), den "Weinathlon" und den "Eschathlon
    Halbmarathon" zu Triathlons gemacht - alles Wortspiele auf einen
    Laufnamen.
    """
    print("\nZwei Sportarten im Namen:")
    from scraper_lib import guess_art1, guess_art2, SiteConfig
    config = SiteConfig(base_url="", calendar_url="")

    for name, art2 in (("Wunnebad Swim&Run", "Swimrun"),
                       ("DSW Swim & Run", "Swimrun"),
                       ("Swim+Run Winnweiler", "Swimrun"),
                       ("Run and Bike Berlin", "Duathlon"),
                       ("Kronberger Bike+Run", "Duathlon")):
        gefunden = guess_art1(name, config)
        check(f"{name!r} ist Mehrsport", gefunden, "Triathlon")
        check(f"{name!r} bekommt die Kategorie {art2}",
              guess_art2(name, config, gefunden), art2)

    # Gegenprobe: Laufveranstaltungen, die nur so KLINGEN
    for name in ("Decathlon Hybrid Series - Plochingen", "Weinathlon",
                 "Eschathlon Halbmarathon", "Rundlauf am Bikepark",
                 "24. Bordesholmer Seelauf SEE +RUN", "Sunrun Berlin"):
        check(f"{name!r} bleibt ein Lauf", guess_art1(name, config), "Laufen")


def test_stundenlauf() -> None:
    """Ein „Stundenlauf" ist eine Stunde - ein „Halbstundenlauf" eine halbe.

    Das Muster für Zeitrennen verlangte eine ZAHL vor dem Wort, und der
    klassische Ein-Stunden-Lauf heißt einfach „Stundenlauf". 13
    Veranstaltungen standen deshalb ganz ohne Maßzahl in der Liste
    („LCM Stundenlauf", „Warburger Stundenlauf", „Stundenlauf mit
    Musik"). Gefunden bei der Einzelprüfung der ersten 400 Events, die
    Definition gegen Wikipedia/Brockhaus geprüft.

    Die Falle steckt im Wort: „Halb·stundenlauf" enthält „stundenlauf"
    als Teilzeichenkette. Ohne die Sonderzeile davor wäre er doppelt so
    lang wie in Wirklichkeit - deshalb steht sie ZUERST, und dieser Test
    hält die Reihenfolge fest.
    """
    print("\nStundenlauf:")
    from scraper_lib import parse_duration_h

    for text, erwartet in (
            ("LCM Stundenlauf", 1.0),
            ("50. LCM Stundenlauf", 1.0),
            ("Stundenlaufserie - Halle (Saale)", 1.0),
            ("Zwickauer Halb- und Stundenlauf", 1.0),
            ("Halbstundenlauf des TSV Zeulenroda", 0.5),
            # Mit Zahl greift weiterhin das allgemeine Muster.
            ("15. Bokeler 6 Stundenlauf", 6.0),
            ("24-Stunden-Lauf", 24.0),
            ("12h Lauf", 12.0),
            # Gegenproben: kein Zeitrennen.
            ("Halbmarathon", None),
            ("Stundenplan der Wettkämpfe", None),
            ("Zeitlimit 6 Stunden", None)):
        check(f"{text!r} -> {erwartet!r}", parse_duration_h(text), erwartet)


def test_audit_pruefungen() -> None:
    """Die Einzelprüfung meldet echte Fehler - und schweigt bei den
    Fällen, die sie früher fälschlich gemeldet hat.

    Der zweite Teil ist der wichtigere. Von den vier größten
    Fundgruppen der ersten Prüfung waren drei Fehler der REGEL und
    nicht der Daten (siehe CLAUDE.md). Jede dieser Korrekturen steht
    hier als Gegenprobe, damit sie nicht beim nächsten Umbau
    zurückrutscht.
    """
    print("\nEinzelprüfung (audit_events):")
    from audit_events import pruefe_event

    def kategorien(**felder):
        basis = {"name": "Testlauf", "datum_start": "2026-09-19",
                 "datum_ende": "2026-09-19", "standort": "Musterstadt",
                 "land": "Deutschland", "lat": 50.0, "lon": 10.0,
                 "art1": "Laufen", "art2": "Straße", "laenge_km": 10.0,
                 "wettbewerb": "10 km",
                 "veranstalter_url": "https://beispiel-lauf.de/"}
        basis.update(felder)
        return {k for k, _ in pruefe_event(basis)}

    # --- Fälle, die gemeldet werden MÜSSEN ---------------------------
    check("Kinderlauf mit 21,1 km wird gemeldet",
          "Kinderlauf-Label mit Erwachsenendistanz"
          in kategorien(wettbewerb="Kinderlauf", laenge_km=21.1), True)
    check("Label '9 km' bei 8,5 km wird gemeldet",
          "Zahl im Label weicht von laenge_km ab"
          in kategorien(wettbewerb="9 km", laenge_km=8.5), True)
    check("Jahreszahl 2025 an einem Termin 2026 wird gemeldet",
          "Jahreszahl im Namen passt nicht zum Datum"
          in kategorien(name="ONW-Lauf Dannenberg 2025"), True)
    check("Name in Großbuchstaben wird gemeldet",
          "Name komplett in Großbuchstaben"
          in kategorien(name="BFUTR EXTREME UNTERWEGS"), True)
    check("Portallink wird gemeldet",
          "Portallink statt offizieller Seite"
          in kategorien(veranstalter_url="https://my.raceresult.com/123/"), True)
    check("Koordinate außerhalb Deutschlands wird gemeldet",
          "Koordinaten passen nicht zum Land"
          in kategorien(lat=41.9, lon=12.5), True)

    check("ein Probe-Eintrag des Kalenders wird gemeldet",
          "Testeintrag der Quelle?"
          in kategorien(name="TESTVERANSTALTUNG Neujahrslauf"), True)
    # Der Basisname dieser Prüfungen ist „Testlauf" - und genau der darf
    # NICHT anschlagen: Ein Lauf kann so heißen. Gemeldet wird nur das
    # eindeutige „TESTVERANSTALTUNG" (gefunden am 02.01.2030, Dolgesheim).
    check("ein Lauf namens „Testlauf“ ist KEIN Fall",
          "Testeintrag der Quelle?" in kategorien(), False)
    check("und „Härtetest-Lauf“ auch nicht",
          "Testeintrag der Quelle?"
          in kategorien(name="12. Härtetest-Lauf Sauerland"), False)

    # --- Gegenproben: die drei Fehlalarme der ersten Fassung ----------
    check("ein Jugendlauf über 5,6 km ist KEIN Fall",
          "Kinderlauf-Label mit Erwachsenendistanz"
          in kategorien(wettbewerb="Jugendlauf, 5,6 km", laenge_km=5.6), False)
    check("'Bernburger Halbmarathon' mit 12-km-Label ist KEIN Fall",
          "„Halbmarathon\u201c, aber Distanz passt nicht"
          in kategorien(name="Bernburger Halbmarathon",
                        wettbewerb="12 km – ab Altersklasse U16",
                        laenge_km=12.0), False)
    check("'Winterlaufserie 2026/2027' ist KEIN Fall",
          "Jahreszahl im Namen passt nicht zum Datum"
          in kategorien(name="51. Winterlaufserie 2026/2027, 1. Wertungslauf"), False)
    check("ein Freitagslauf ist KEIN Fall",
          "Wochentag Mo-Do" in kategorien(datum_start="2026-09-18",
                                          datum_ende="2026-09-18"), False)
    check("'(L)auf zur Venus' ist KEIN unsauberer Name",
          "Name sieht unsauber aus" in kategorien(name="(L)auf zur Venus"), False)
    check("ein Hindernislauf mit 'Cross' im Namen ist KEIN Fall",
          "Trail/Cross im Namen, andere Kategorie"
          in kategorien(name="Family-CrossDeLuxe Leipzig", art2="Hindernis"), False)
    check("'3 Runden je 15,5 km' ist kein Label-Widerspruch",
          "Zahl im Label weicht von laenge_km ab"
          in kategorien(wettbewerb="46,5 km, 3 Runden je 15,5 km",
                        laenge_km=46.5), False)
    # --- Label, das seine Distanz selbst nennt -------------------------
    check("„Halbmarathon 22,8 km\u201c bei 22,8 km ist KEIN Fall (Trail-Rundkurs)",
          "„Halbmarathon\u201c, aber Distanz passt nicht"
          in kategorien(name="Dörenther Klippen UltraTrail",
                        wettbewerb="Halbmarathon 22,8 km und ca. 560 Hm",
                        laenge_km=22.8), False)
    check("„Halbmarathon\u201c ohne km-Angabe bei 23 km wird gemeldet",
          "„Halbmarathon\u201c, aber Distanz passt nicht"
          in kategorien(wettbewerb="Halbmarathon", laenge_km=23.0), True)

    # --- Rundenlänge gegen Staffel -------------------------------------
    # Die vierte Begegnung mit der Rundenlängen-Falle. Der Unterschied
    # zur Staffel ist der ganze Aufwand: Bei „2x5 km Staffel" läuft man
    # wirklich 5 km (eigener Wettbewerb), bei „4 Runden à 5,27 km"
    # läuft dieselbe Person alle vier.
    check("„4 Runden à ca. 5,27 km\u201c bei 5,3 km wird gemeldet",
          "Rundenlänge statt Renndistanz"
          in kategorien(wettbewerb="Halbmarathon (4 Runden à ca. 5,27 km)",
                        laenge_km=5.3), True)
    check("eine STAFFEL wird nicht gemeldet („2x5 km Staffel\u201c bei 5 km)",
          "Rundenlänge statt Renndistanz"
          in kategorien(wettbewerb="2x5 km Staffel", laenge_km=5.0), False)
    check("ein DUO wird nicht gemeldet",
          "Rundenlänge statt Renndistanz"
          in kategorien(wettbewerb="DUO Marathon 2 x 21,1 km", laenge_km=21.1), False)
    check("stimmt die Summe, ist es kein Fall („3 Runden je 15,5 km\u201c bei 46,5)",
          "Rundenlänge statt Renndistanz"
          in kategorien(wettbewerb="46,5 km, 3 Runden je 15,5 km", laenge_km=46.5), False)

    check("eine saubere Zeile meldet gar nichts", kategorien(), set())


def test_such_vorschlaege() -> None:
    """„Iron" muss „Ironman" anbieten - und nichts Sinnloses.

    Die Vorschläge entstehen AUS DEN DATEN (`EF.buildSuggestions`), nicht
    aus einer gepflegten Markenliste: Eine Liste im Code würde veralten,
    sobald eine Serie dazukommt, und beim großen Datenlauf kommen 16.000
    Events dazu.

    Die Gegenproben sind hier das Eigentliche. Beim Bauen dieser Funktion
    standen in der Liste: „Wings for Life -" (ein Bindestrich als Wort),
    „2026" und „Silvesterlauf 2026" (Jahreszahlen), „München" doppelt
    (einmal als Ort, einmal aus den Namen) und drei Schreibweisen von
    „UltraTrail". Jede dieser Schwächen steht unten als Test.
    """
    print("\nVorschläge der Mastersuche:")
    import json, shutil, subprocess
    node = shutil.which("node")
    if not node:
        print("  (node nicht vorhanden - übersprungen)")
        return
    wurzel = Path(__file__).resolve().parent.parent
    skript = f"""
      global.window = global;
      require({str(wurzel / 'filters.js')!r});
      const EF = global.EnduranceFilters;
      const mach = (name, ort, tag) => ({{ name, standort: ort, datum_start: tag }});
      const daten = [
        mach('Ironman Frankfurt', 'Frankfurt', '2027-06-27'),
        mach('Ironman Hamburg', 'Hamburg', '2027-06-06'),
        mach('Ironman 70.3 Erkner', 'Erkner', '2027-09-12'),
        mach('Ironman 70.3 Leipzig', 'Leipzig', '2027-08-08'),
        // Serie, die NICHT am Namensanfang steht
        mach('Kulmbach Spartan Trifecta Weekend', 'Kulmbach', '2027-06-11'),
        mach('Spartan Berlin', 'Berlin', '2027-05-01'),
        mach('Spartan Muenchen', 'Muenchen', '2027-05-02'),
        // Serie mit Bindestrich-Ort dahinter
        mach('Wings for Life - Frankfurt am Main', 'Frankfurt', '2027-05-09'),
        mach('Wings for Life - Gangelt', 'Gangelt', '2027-05-09'),
        mach('Wings for Life - Bamberg', 'Bamberg', '2027-05-09'),
        // Jahreszahlen im Namen
        mach('Silvesterlauf 2026 Nord', 'Kiel', '2026-12-31'),
        mach('Silvesterlauf 2026 Sued', 'Ulm', '2026-12-31'),
        mach('Silvesterlauf 2026 West', 'Aachen', '2026-12-31'),
        // Schreibvarianten
        mach('Grosser UltraTrail', 'Bonn', '2027-07-01'),
        mach('Kleiner Ultratrail', 'Bonn', '2027-07-02'),
        mach('Dritter ULTRATRAIL', 'Bonn', '2027-07-03')
      ];
      const idx = EF.buildSuggestions(daten);
      const texte = idx.map(v => v.text);
      const treffer = (q) => EF.matchSuggestions(idx, q, 8).map(v => v.text + ':' + v.anzahl);
      console.log(JSON.stringify({{
        iron: treffer('iron'),
        sparta: treffer('sparta'),
        wings: treffer('wings'),
        silv: treffer('silvest'),
        ultra: treffer('ultratr'),
        bonn: treffer('bon'),
        mitBindestrich: texte.filter(t => t.trim().endsWith('-')),
        mitJahr: texte.filter(t => /\\b20\\d\\d/.test(t)),
        // Die Vier-Zeichen-Regel gilt nur für SERIEN. Ein Ort darf
        // kurz sein - "Ulm" hat drei Buchstaben und muss vorkommen.
        kurzeSerie: idx.filter(v => v.art === 'serie' && v.text.length < 4).map(v => v.text),
        kurzerOrt: idx.filter(v => v.art === 'ort' && v.text === 'Ulm').length
      }}));
    """
    ergebnis = subprocess.run([node, "-e", skript], capture_output=True, text=True)
    if ergebnis.returncode != 0:
        check("node konnte filters.js laden", ergebnis.stderr.strip()[-200:], "")
        return
    d = json.loads(ergebnis.stdout)

    check("„iron\u201c schlägt „Ironman\u201c vor (4 Veranstaltungen)",
          "Ironman:4" in d["iron"], True)
    # Der Kernfall: eine Serie, die MITTEN im Namen steht.
    check("„sparta\u201c schlägt „Spartan\u201c vor, auch wenn der Name "
          "mit dem Ort beginnt", "Spartan:3" in d["sparta"], True)
    check("„wings\u201c schlägt „Wings for Life\u201c vor (3)",
          "Wings for Life:3" in d["wings"], True)
    # Der längere Begriff verdrängt den kürzeren bei gleicher Trefferzahl.
    check("„Wings\u201c allein steht NICHT daneben",
          any(x.startswith("Wings:") for x in d["wings"]), False)

    # --- Gegenproben ---------------------------------------------------
    check("kein Vorschlag endet auf einem Bindestrich", d["mitBindestrich"], [])
    check("keine Jahreszahl in einem Vorschlag", d["mitJahr"], [])
    check("keine Serie aus einem Wort unter vier Zeichen", d["kurzeSerie"], [])
    check("ein kurzer ORTSNAME bleibt trotzdem („Ulm\u201c)", d["kurzerOrt"], 1)
    check("„silvest\u201c findet die Serie ohne die Jahreszahl",
          "Silvesterlauf:3" in d["silv"], True)
    check("Schreibvarianten sind EIN Vorschlag mit drei Treffern",
          [x for x in d["ultra"] if x.endswith(":3")] != [], True)
    check("ein Ort steht genau EINMAL in der Liste",
          len([x for x in d["bonn"] if x.startswith("Bonn:")]), 1)


def test_laufen_weiterleitung() -> None:
    """Die laufen.de-Detailseite leitet auf den Veranstalter weiter - und
    genau dieses Ziel ist der Veranstalter-Link (Entscheidungspunkt 15,
    vom Nutzer am 19.09.2026 freigegeben).

    Gegenproben: Ziele auf Anmeldeportalen (lanet3, raceresult, datasport,
    racepedia), sozialen Netzen, Kalenderportalen (PORTAL_DOMAINS) und
    innerhalb von laufen.de (www-Variante) sind KEIN Veranstalter-Link.
    """
    print("\nlaufen.de-Weiterleitung (veranstalter_link_aus_weiterleitung):")
    from laufkalender_scraper import veranstalter_link_aus_weiterleitung as vl
    detail = "https://laufen.de/laufkalender/details/26V09000621060001"
    check("Veranstalter-Host", vl(detail, "https://www.sv-molbergen-leichtathletik.de/"),
          "https://www.sv-molbergen-leichtathletik.de/")
    check("http ohne www", vl(detail, "http://www.svg-poenitz.de"), "http://www.svg-poenitz.de")
    check("Leerzeichen am Rand", vl(detail, "  https://www.taubertal100.de/ "), "https://www.taubertal100.de/")
    for ziel in ("https://lanet3.de/external/dlv/register/13425",
                 "https://my.raceresult.com/354931/",
                 "https://www.datasport.de/anmeldeservice/x",
                 "https://schwabacher-citylauf-2026.racepedia.de",
                 "https://www.facebook.com/events/123", "https://fb.me/e/abc",
                 "https://www.laufen.de/laufkalender/details/26V09000621060001",
                 "https://running.life/de/termine/x",
                 "/laufkalender/details/26V09000621060001",
                 "mailto:ch.olbrich@gmx.de", "", None):
        check(f"kein Veranstalter: {ziel!r}", vl(detail, ziel), None)

    # Der Abruf selbst: eine Weiterleitung setzt den Link und parst nichts,
    # eine echte Detailseite (200) läuft wie bisher durch parse_detail_page.
    import laufkalender_scraper as lk

    class _Resp:
        def __init__(self, status, headers=None, text=""):
            self.status_code, self.headers, self.text = status, headers or {}, text
        def raise_for_status(self):
            if self.status_code >= 400:
                raise requests.exceptions.HTTPError(str(self.status_code))

    class _Session:
        def __init__(self, antworten): self.antworten, self.aufrufe = antworten, []
        def get(self, url, timeout=20, allow_redirects=True):
            self.aufrufe.append((url, allow_redirects))
            return self.antworten[url]

    e1 = lk.Event(name="A", detail_url="https://laufen.de/laufkalender/details/1",
                  veranstalter_url="https://laufen.de/laufkalender/details/1")
    e2 = lk.Event(name="B", detail_url="https://laufen.de/laufkalender/details/2",
                  veranstalter_url="https://laufen.de/laufkalender/details/2")
    e3 = lk.Event(name="C", detail_url="https://laufen.de/laufkalender/details/3",
                  veranstalter_url="https://laufen.de/laufkalender/details/3")
    sess = _Session({
        e1.detail_url: _Resp(302, {"Location": "https://www.wellen-marathon.de"}),
        e2.detail_url: _Resp(302, {"Location": "https://lanet3.de/external/dlv/register/1"}),
        e3.detail_url: _Resp(200, text="<html><body></body></html>"),
    })
    out = lk.enrich_from_details(sess, [e1, e2, e3], delay=0, max_details=0)
    check("Weiterleitung setzt den Veranstalter-Link", out[0].veranstalter_url, "https://www.wellen-marathon.de")
    check("Portal-Ziel lässt den Portallink stehen", out[1].veranstalter_url, e2.detail_url)
    check("Detailseite ohne Weiterleitung bleibt unverändert", out[2].veranstalter_url, e3.detail_url)
    check("kein Event geht verloren", len(out), 3)
    check("ohne allow_redirects abgerufen", all(not ar for _, ar in sess.aufrufe), True)


def test_nicht_ausdauer() -> None:
    """HYROX gehört nicht in die Liste - ein Hindernislauf schon.

    HYROX ist achtmal ein Kilometer Laufen im Wechsel mit acht
    Kraftstationen; man kann sich dafür nicht als Läufer anmelden. Vom
    Nutzer am 18.09.2026 entschieden.

    Die Gegenprobe ist der wichtigere Teil: Ein Hindernislauf (Spartan,
    XLETIX, CrossDeLuxe, Muddy Angel) IST ein Laufformat und bleibt.
    Ein Stichwort wie „Fitness" oder „Hindernis" in NICHT_AUSDAUER
    hätte die alle mitgenommen.
    """
    print("\nKein Ausdauer-Format (NICHT_AUSDAUER):")
    from scraper_lib import ist_nicht_ausdauer
    from clean_events import drop_nicht_ausdauer

    for name in ("HYROX Karlsruhe", "Intersport HYROX Hamburg", "hyrox cologne",
                 "Gymrace Airport Weeze", "Decathlon Hybrid Series - Plochingen",
                 # Nachzügler, vom Nutzer am 21.09.2026 bestätigt
                 "Runworx", "Black Forest Team Battle",
                 # weitere Marken, vom Nutzer am 07.10.2026 bestätigt
                 "Fitletix - Graz", "Raedbox Challenge by FITLETIX", "The Rox",
                 "Deadly Dozen Köln", "Deadly Sprint im Glaspalast", "Deadly Gross Fürth",
                 "ATHX St Gallen", "ATHX Vienna", "StrongAthlon", "4. Strong Athlon",
                 # UCI-Teamrennen 1.x/2.x (Nutzer, 07.10.2026)
                 "16. GP Vorarlberg p/b RadHaus Rankweil UCI 1.2",
                 "Oberösterreich Rundfahrt (UCI Kat. 2.2)", "Tour of Austria UCI 2.Pro",
                 # Gehen und Skilanglauf (Nutzer, 21.09.2026)
                 "Lusatian Race Walking", "Internationaler Kammlauf (Skilanglauf)",
                 "König-Ludwig-Langlauf", "Gehermeeting: Gehertag Naumburg",
                 # virtuelle Läufe (Nutzer, 21.09.2026)
                 "Virtueller Silvesterlauf", "Virtual Run Berlin"):
        check(f"{name!r} fliegt heraus", bool(ist_nicht_ausdauer(name)), True)
    for name in ("Kulmbach Spartan Trifecta Weekend", "XLETIX Challenge - Nürburgring",
                 "Family-CrossDeLuxe Leipzig", "Muddy Angel Run - Berlin",
                 "Tough Mudder Hamburg", "Fitnesslauf Bochum", "Hindernislauf Kiel",
                 "Decathlon Stadtlauf Plochingen", "Hybrid Trail Harz",
                 # "Rox", "Dozen" und "Strong" allein sind keine Marke
                 "Roxheimer Volkslauf", "Dozen Hills Trail", "Strong Viking Run",
                 "Athlon Triathlon", "Deadline Lauf",
                 # UCI C1/C2 und Gran Fondo: jeder meldet sich selbst an
                 "5. Querfeldein GP UCI C2 Gunskirchen", "35. KTM Kamptal Trophy - UCI XCO C1",
                 "UCI Gran Fondo World Series Ötztal", "Säntis Classic 2.2 km",
                 # Walking-Strecken bleiben (offene Frage an den Nutzer),
                 # "Langläufer" ist kein Langlauf, ein Bergläufer kein Geher.
                 "21. Trochtelfinger Nordic Walking Stöckles-Cup", "Ahmadiyya Charity Walk",
                 "5 km Walking", "Langläufer-Cup Oberstdorf", "Marathon der Bergläufer"):
        check(f"{name!r} bleibt", ist_nicht_ausdauer(name), None)
    # Der Ort zählt mit: Die XMAS-Challenge trug "virtuell" nur dort.
    from scraper_lib import nicht_ausdauer_text
    check("Ort 'virtuell' reicht",
          bool(ist_nicht_ausdauer(nicht_ausdauer_text("Blaues Land läuft – XMAS-Challenge",
                                                      "10 km", "virtuell"))), True)
    check("normaler Ort ändert nichts",
          ist_nicht_ausdauer(nicht_ausdauer_text("Stadtlauf", "10 km", "Murnau")), None)

    behalten, entfernt = drop_nicht_ausdauer([
        {"name": "HYROX Berlin", "datum_start": "2027-01-01"},
        {"name": "Spartan Berlin", "datum_start": "2027-01-01"},
        {"name": "XMAS-Challenge", "datum_start": "2027-01-01", "standort": "virtuell"},
    ])
    check("drop_nicht_ausdauer behält den Hindernislauf",
          [e["name"] for e in behalten], ["Spartan Berlin"])
    check("und meldet die Ausschlüsse mit Grund", len(entfernt), 2)


def test_datum_vorlaeufig() -> None:
    """Ein vorläufiger Termin (Kalenderprognose) trägt `datum_vorlaeufig`.

    Vom Nutzer am 21.09.2026 entschieden: Wo kein Tag veröffentlicht ist,
    zeigt die Liste nur "Juni 2027*" mit Fußnote. Das Feld kommt allein
    per Override (OVERRIDE_FIELDS), die Kalenderdatei sagt es im Titel,
    und ein Event ohne das Feld trägt es auch nicht als false in
    events.json (to_dict lässt None weg).
    """
    print("\nVorläufige Termine (datum_vorlaeufig):")
    from scraper_lib import OVERRIDE_FIELDS, Event
    import build_ics
    check("Override-Feld bekannt", "datum_vorlaeufig" in OVERRIDE_FIELDS, True)
    e = Event(name="Borkener Citylauf", datum_start="2027-06-07", datum_ende="2027-06-07",
              standort="Borken", laenge_km=10.0)
    check("ohne Markierung fehlt das Feld in events.json",
          "datum_vorlaeufig" in e.to_dict(), False)
    e.datum_vorlaeufig = True
    check("markiert steht es drin", e.to_dict().get("datum_vorlaeufig"), True)
    ics = build_ics.build_ics(e.to_dict(), "20260101T000000Z")
    check("Kalenderdatei nennt den vorläufigen Termin im Titel",
          "SUMMARY:Borkener Citylauf (Termin vorläufig)" in ics, True)
    check("und in der Beschreibung", "Termin noch zu bestätigen (to be confirmed)" in ics, True)
    ics_ohne = build_ics.build_ics(Event(name="Stadtlauf", datum_start="2027-06-07",
                                         datum_ende="2027-06-07", standort="X").to_dict(),
                                   "20260101T000000Z")
    check("ein normaler Termin bleibt unverändert", "vorläufig" in ics_ohne, False)


def test_staffeln() -> None:
    """Erst einmal keine Staffeln (Nutzer, 21.09.2026).

    Zwei Stufen (siehe ist_staffel): Ein Label, das NUR die Staffel
    beschreibt, nimmt diese eine Zeile; ein Name, der eine
    Staffelveranstaltung nennt, nimmt alle Zeilen. Die Gegenproben sind
    der wichtigere Teil - alle am Bestand gezählt: ein Einzelrennen mit
    Staffel-Option bleibt, eine Laufserie namens "Winterstaffel" bleibt,
    ein Ort namens Staffelsee bleibt.
    """
    print("\nStaffeln (ist_staffel):")
    from scraper_lib import ist_staffel, filter_staffeln, Event
    from clean_events import drop_staffeln

    for name, wb in (("22. Einstein-Marathon", "ZEISS Marathon Staffel"),
                     ("Uni-Lauf Bamberg", "2x5 km Staffel, Wechsel an der Buger Spitze"),
                     ("Lübeck Marathon", "DUO Marathon 2 x 21,1 km"),
                     ("H/21 Halbmarathon Hannover", "H/21 for Two (Staffel)"),
                     ("Landkreislauf Schwandorf", "Läufer-Staffel (10 Läufer, Gesamtstrecke ca. 49,5 km)"),
                     ("Firmenstaffel Sachsen-Anhalt", "5er-Staffel (5 × 3 km = 15 km)"),
                     ("Ostsee Staffel Marathon", "5 km"),
                     ("40. Weezer Staffellauf", None),
                     ("Rasteder Ellernteichstaffellauf", None),
                     ("Stadtpark-Staffel-Marathon", "Marathon"),
                     ("Marathonstaffel Mörfelden", "Marathon"),
                     ("Staffel-Mix-Marathon", None),
                     ("Stralsunder Firmenstaffellauf", "12 km"),
                     # erweitert 07.10.2026 (Nutzer: "Nur Einzelläufe")
                     ("BriggYard Staffel-Run", "42 km"),
                     ("SOLA-Stafette", "115 km"), ("Winti-SOLA", "83 km"),
                     ("Grüntenstafette", None), ("Hulftegg Stafette", "25 km"),
                     ("Lichterpaarlauf Brandenburg", "Paarlauf 60 Minuten (400-m-Runde)"),
                     ("Paarlauf mit Musik des SCC", None),
                     ("Ägeriseelauf", "Paarlauf 2"),
                     ("Napf-Marathon", "Teamlauf mit einer Etappe von 20 km und einer von 22 km"),
                     ("WädiLauf", "6.15 km und 4.35 km Du & Ich Teams, zwei Abschnitte")):
        check(f"{name!r} / {wb!r} fällt", bool(ist_staffel(name, wb)), True)
    for name, wb in (("#ZeroHungerRun Bonn", "10 km Lauf und Staffel"),
                     ("Butterkuchenlauf", "12 km (Einzel oder Staffel)"),
                     ("GaPa Everesting-Festival", "Everesting Solo oder Staffel"),
                     ("Grünwalder Burglauf", "5 km Strecke, ebenfalls als Einzel- oder Duo-Staffel möglich"),
                     ("Inzeller Falkensteinlauf", "Halbmarathon 21,095 km, Einzel und Staffel"),
                     ("GVG-Winterstaffel Pulheim", "Halbmarathon"),
                     ("Meckenheimer Apfelstaffel", None),
                     ("4. Zülpicher Seepark Nikolauslauf mit Fun/Firmenstaffel", None),
                     ("Volks- und Staffeltriathlon TuS Wasserstraße", None),
                     ("Staffelsee Panoramalauf", "5 km"),
                     ("21. Obermain-Marathon Bad Staffelstein", "Sparkassen-Marathon"),
                     ("Rund um den Kellerskopf", "21 km Halbmarathon (2 x 10,5 km Runde)"),
                     # Etappe ohne Team, Teamwertung neben der Einzelwertung
                     ("Die Bergischen 5 Etappe 3", "45 km"),
                     ("Giro delle Dolomiti", "Etappenrennen 124 km"),
                     ("Deutschlandlauf Flensburg - Lörrach", "21 Tagesetappen"),
                     ("Werdauer Waldlauf", "Halbmarathon (Einzelwertung und Teamwertung)"),
                     ("Ägeriseelauf", "14 km"), ("Solarlauf Bocholt", "10 km")):
        check(f"{name!r} / {wb!r} bleibt", ist_staffel(name, wb), None)

    zeilen = [Event(name="Lübeck Marathon", wettbewerb="Marathon", laenge_km=42.2),
              Event(name="Lübeck Marathon", wettbewerb="DUO Marathon 2 x 21,1 km", laenge_km=21.1)]
    behalten, n = filter_staffeln(zeilen)
    check("filter_staffeln nimmt nur die Staffel-Zeile", [e.wettbewerb for e in behalten], ["Marathon"])
    check("und zählt sie", n, 1)
    behalten, entfernt = drop_staffeln([
        {"name": "Weezer Staffellauf", "datum_start": "2027-01-01"},
        {"name": "Stadtlauf", "datum_start": "2027-01-01", "wettbewerb": "10 km"},
    ])
    check("drop_staffeln behält den Stadtlauf", [e["name"] for e in behalten], ["Stadtlauf"])
    check("und meldet den Ausschluss", len(entfernt), 1)


def test_walking() -> None:
    """Reine Walking-Wettbewerbe raus (Nutzer, 07.10.2026: "Nur A raus").

    Nur die Zeile fällt, deren LABEL ausschließlich Walking/Wandern/Marsch
    nennt. Die Gegenproben sind der wichtigere Teil, alle am Bestand
    gezählt: "Lauf und Walking" ist der Lauf, zwei Wettbewerbe in einem
    Label halten die Zeile, "auch als Nordic Walking" ist ein Lauf, und
    Walking im NAMEN (Hollenmarsch, Ultra Walk, Charity Walk) bleibt -
    "der wird nicht nur gelaufen, viele gehen auch beim Ultra".
    """
    print("\nWalking (ist_walking):")
    from scraper_lib import ist_walking, filter_walking, Event
    from clean_events import drop_walking

    for name, wb in (("Laubacher Ramsberglauf", "7,5 km Walking"),
                     ("Kyburglauf", "Nordic Walking 5.7 km"),
                     ("58. Int. Schwarzwald-Marathon", "10 km Nordic Walking/Walking"),
                     ("Harz-Gebirgslauf", "25 km Wandern / Nordic Walking"),
                     ("64. Westenergie Marathon", "Walking Große Schleife"),
                     ("Halloween Run - Hückeswagen", "5-km-Walk"),
                     ("54. Winterlaufserie", "16 km (Nordic-)Walking"),
                     ("Zeiler Waldmarathon", "Zaundirekt Walking und Nordic Walking"),
                     ("SüdwestPark FunRun", "6 km Nordic Walker, offen für alle"),
                     ("Lauf in den Frühling - Freital", "5 km Walking / Nordic Walking, Altersklassen M/W"),
                     ("Schneekopflauf", "11 km Schneekopfwanderung für Nordic Walker und Wanderer"),
                     ("Tour de Tirol", "18 km Hohe Salve Gipfelmarsch, 1.494 hm"),
                     ("Schneeglöckchen-Lauf", "15 km Wanderung"),
                     ("Freundschaftslauf Marpingen", "Walking"),
                     ("Aletsch-Halbmarathon", "Aletsch-Halbmarathon Walking 21.1 km")):
        check(f"{name!r} / {wb!r} fällt", bool(ist_walking("Laufen", wb, name)), True)
    for name, wb in (("MedeRun", "10 km Lauf und Nordic Walking"),
                     ("Zachower Crosslauf", "5 km für Läufer und Walker"),
                     ("Schweizer Frauenlauf Bern", "5 km Running oder (Nordic) Walking"),
                     ("Butterkuchenlauf", "6,5 km (Laufen, Walken, Dogging, Staffel möglich)"),
                     ("Halloween Run Duisburg", "5 km Run/Walk"),
                     ("Rhön-Grabfeld-Cup – Wollbach", "5,5 km Hobby, U18, Walking/Wandern"),
                     ("Berglauf am Brigerberg", "36er und 36er Walk"),
                     ("Martinslauf Düsseldorf", "ca. 6 km Seerunde und Nordic Walking"),
                     ("Bieler Lauftage", "10km fun & walking"),
                     ("Külsheimer TrailRun", "6 km auch als Nordic Walking"),
                     ("Schliersee Lauf", "10,0 km Lauf und Walk"),
                     ("Trail Run Evasion Rhône", "9 km Trail / Walking"),
                     ("Meerlauf", "12,1 km Spendenlauf (Nordic-)Walking"),
                     ("Bödefelder Hollenmarsch", "101 km (Paket XL)"),
                     ("Ultra Walk", "100 km"),
                     ("Ahmadiyya Charity Walk Wetzlar", "7 km"),
                     ("Karwendelmarsch", None),
                     ("Stadtlauf", "10 km")):
        check(f"{name!r} / {wb!r} bleibt", ist_walking("Laufen", wb, name), None)
    check("nur für Laufen", ist_walking("Fahrrad", "7,5 km Walking", "Radtour"), None)

    zeilen = [Event(name="Stettener Volkslauf", wettbewerb="10 km", laenge_km=10.0, art1="Laufen"),
              Event(name="Stettener Volkslauf", wettbewerb="7,5 km Nordic Walking", laenge_km=7.5, art1="Laufen")]
    behalten, n = filter_walking(zeilen)
    check("filter_walking nimmt nur die Walking-Zeile", [e.wettbewerb for e in behalten], ["10 km"])
    check("und zählt sie", n, 1)
    behalten, entfernt = drop_walking([
        {"name": "Kerzerslauf", "datum_start": "2027-03-20", "art1": "Laufen", "wettbewerb": "15km Walk"},
        {"name": "Kerzerslauf", "datum_start": "2027-03-20", "art1": "Laufen", "wettbewerb": "15km"},
    ])
    check("drop_walking behält den Lauf", [e["wettbewerb"] for e in behalten], ["15km"])
    check("und meldet den Ausschluss", len(entfernt), 1)


def test_dauer_minuten() -> None:
    """Dauern unter einer Stunde in Minuten (vom Nutzer am 18.10.2026
    gemeldet: „Was sind denn bitte 0,3 und 0,8 Stunden? Das muss man
    natürlich dann in Minuten ausdrücken."). Dazu, dass die Zahl dafür
    genau genug ist: Bis zum 13.10.2026 stand `dauer_h` auf eine Stelle
    gerundet in events.json (40 min = 0,7 h = 42 min)."""
    import json
    print("\nDauer unter einer Stunde in Minuten:")
    from build_ics import laenge_text
    from clean_events import fix_dauer_minuten
    from scraper_lib import dauer_text
    from insta_inhalte import dauer_wort
    check("Kalender: 30 min", laenge_text({"dauer_h": 0.5}), "30 min")
    check("Kalender: 1 h bleibt Stunden", laenge_text({"dauer_h": 1.0}), "1 h")
    check("Label: 40 min", dauer_text(2 / 3), "40 min")
    check("Label: 1,5 h", dauer_text(1.5), "1.5 h")
    check("Instagram: 40 Minuten", dauer_wort(0.67), "40 Minuten")
    check("Instagram: 6 Stunden", dauer_wort(6.0), "6 Stunden")

    zeilen = [
        {"name": "A", "dauer_h": 0.7, "wettbewerb": "Masters 40 min (Cyclocross Cup Baden-Württemberg)"},
        {"name": "B", "dauer_h": 0.8, "wettbewerb": "Elite 45 min"},
        {"name": "C", "dauer_h": 0.67, "wettbewerb": "Hobby, Masters 40-45 min"},
        {"name": "D", "dauer_h": 0.42, "wettbewerb": "Cyclocross 0.4 h"},
        {"name": "E", "dauer_h": 0.7, "wettbewerb": "Cyclocross 0.7 h"},
        {"name": "F", "dauer_h": 1.0, "wettbewerb": "Cyclocross 59 min"},
        {"name": "G", "dauer_h": 0.5, "wettbewerb": "30-Minuten-Lauf"},
        {"name": "H", "dauer_h": 0.7, "laenge_km": 5.0, "wettbewerb": "Lauf 40 min"},
        {"name": "I", "dauer_h": 0.8, "wettbewerb": "Rennen 20 min"},
    ]
    geaendert, hinweise = fix_dauer_minuten(zeilen)
    dauer = {z["name"]: z["dauer_h"] for z in zeilen}
    check("40 min aus dem Label", dauer["A"], 0.67)
    check("45 min aus dem Label (0,8 war die Rundung)", dauer["B"], 0.75)
    check("Spanne im Label ist keine Angabe", dauer["C"], 0.67)
    check("„0.4 h“ im Label wird „25 min“", zeilen[3]["wettbewerb"], "Cyclocross 25 min")
    check("ungenaues 0,7 behält „0.7 h“ (keine erfundenen 42 min)", zeilen[4]["wettbewerb"], "Cyclocross 0.7 h")
    check("glatte Dauer schlägt ein veraltetes Label", dauer["F"], 1.0)
    check("Halbstunde bleibt", dauer["G"], 0.5)
    check("mit Distanz unberührt", dauer["H"], 0.7)
    check("Label weit daneben: nicht übernommen", dauer["I"], 0.8)
    check("drei Änderungen", len(geaendert), 3)
    check("Hinweise: die beiden ungenauen", sorted(h.split(" ")[0] for h in hinweise), ["E", "I"])
    check("idempotent", fix_dauer_minuten(zeilen)[0], [])

    alt = {"dauer_h": 0.7, "wettbewerb": "Cyclocross 0.7 h"}
    update_existing_event(alt, {"dauer_h": 0.67, "wettbewerb": "Cyclocross 40 min"})
    check("Datenlauf macht die alte Rundung genau", (alt["dauer_h"], alt["wettbewerb"]), (0.67, "Cyclocross 40 min"))
    genau = {"dauer_h": 0.67, "wettbewerb": "Masters 40 min"}
    update_existing_event(genau, {"dauer_h": 0.7, "wettbewerb": "Cyclocross 0.7 h"})
    check("genaue Dauer bleibt gegen eine gerundete", (genau["dauer_h"], genau["wettbewerb"]), (0.67, "Masters 40 min"))
    halb = {"dauer_h": 0.5}
    update_existing_event(halb, {"dauer_h": 0.75})
    check("andere Dauer ist ein anderer Wettbewerb", halb["dauer_h"], 0.5)

    got = _node_event_detail(
        "[global.EnduranceFilters.formatHours(0.67, 'de'), global.EnduranceFilters.formatHours(0.7, 'de'), global.EnduranceFilters.formatHours(0.5, 'en'),"
        " global.EnduranceFilters.formatHours(1, 'de'), global.EnduranceFilters.formatHours(1.5, 'de'), global.EnduranceFilters.formatHours(24, 'en'),"
        " EED.formatLengthSpan([{dauer_h: 0.42}, {dauer_h: 0.67}], 'de'),"
        " EED.formatLengthSpan([{dauer_h: 0.5}, {dauer_h: 1}], 'de'),"
        " EED.formatLengthSpan([{dauer_h: 0.67}, {dauer_h: 2}], 'de')]")
    if got is not None:
        check("Seite: Minuten unter einer Stunde, Stunden darüber", got,
              ["40min", "42min", "30min", "1h", "1,5h", "24h", "25–40min", "30–60min", "40min–2h"])

    ev = json.loads((Path(__file__).resolve().parent.parent / "events.json").read_text(encoding="utf-8"))
    ungenau = [e["name"] for e in ev if isinstance(e.get("dauer_h"), (int, float)) and 0 < e["dauer_h"] < 1
               and e.get("laenge_km") is None and round(e["dauer_h"] * 60) % 5]
    check("Bestand: jede Dauer unter einer Stunde in vollen fünf Minuten", ungenau, [])


def test_rundenlaenge_zeitrennen() -> None:
    """Distanz bis 10 km NEBEN einer Dauer ist die Runde (Korrekturliste
    13.10.2026); eine echte Zielvorgabe mit Zeitlimit bleibt. Dazu: das
    Gedächtnis übernimmt keine Walking- und Kinderzeilen mehr."""
    print("\nRundenlänge bei Zeitrennen:")
    from clean_events import clear_rundenlaenge_bei_zeitrennen
    # Anzeige ohne Distanz (Korrekturliste 13.10.2026, Rückgabe des
    # Nutzers): Minuten unter einer Stunde, "Open End" beim Backyard,
    # Höhenmeter - die Kalenderdatei sagt dasselbe wie die Seite.
    from build_ics import laenge_text
    check("40 Minuten", laenge_text({"dauer_h": 0.67}), "40 min")
    check("unter einer Stunde immer Minuten (18.10.2026)", laenge_text({"dauer_h": 0.7}), "42 min")
    check("Backyard: Open End", laenge_text({"art2": "Backyard Ultra"}), "Open End")
    check("Höhenmeter", laenge_text({"hoehenmeter": 1000}), "1000 hm")
    check("Dauer auf zwei Stellen", Event(name="x", datum_start="2027-01-01", standort="y",
                                           dauer_h=2 / 3).to_dict()["dauer_h"], 0.67)
    rows = [{"name": "Wuhlheider 12 Stundenlauf", "laenge_km": 2.0, "dauer_h": 12.0},
            {"name": "2much4you", "laenge_km": 5.0, "dauer_h": 24.0},
            {"name": "Gravallo", "laenge_km": 85.0, "dauer_h": 6.0},
            {"name": "Brevet 200", "laenge_km": 200.0, "dauer_h": 0.0},
            {"name": "Stundenlauf", "laenge_km": None, "dauer_h": 1.0},
            {"name": "10 km", "laenge_km": 10.0, "dauer_h": None}]
    geaendert = clear_rundenlaenge_bei_zeitrennen(rows)
    check("Runde 2 km neben 12 h fällt", rows[0]["laenge_km"], None)
    check("Runde 5 km neben 24 h fällt", rows[1]["laenge_km"], None)
    check("85 km in 6 h bleibt", rows[2]["laenge_km"], 85.0)
    check("Dauer 0 ist keine Angabe", (rows[3]["dauer_h"], rows[3]["laenge_km"]), (None, 200.0))
    check("ohne Dauer unberührt", rows[5]["laenge_km"], 10.0)
    check("drei Änderungen gemeldet", len(geaendert), 3)
    check("idempotent", clear_rundenlaenge_bei_zeitrennen(rows), [])

    import clean_events
    alt = clean_events.load_manual_events
    try:
        clean_events.load_manual_events = lambda: [
            {"name": "Morat-Fribourg", "datum_start": "2030-10-03", "datum_ende": "2030-10-03",
             "art1": "Laufen", "laenge_km": 17.2, "wettbewerb": "Morat-Fribourg Walking/Nordic Walking 17.2 km",
             "standort": "Murten", "land": "Schweiz"},
            {"name": "Höllenberg Trail-Trophy", "datum_start": "2030-06-19", "datum_ende": "2030-06-19",
             "art1": "Fahrrad", "laenge_km": 5.0, "wettbewerb": "Teens MTB Race",
             "standort": "Spirkelbach", "land": "Deutschland"},
            {"name": "Morat-Fribourg", "datum_start": "2030-10-03", "datum_ende": "2030-10-03",
             "art1": "Laufen", "laenge_km": 9.7, "wettbewerb": "Courtepin-Fribourg 9.7 km",
             "standort": "Murten", "land": "Schweiz"}]
        ergebnis, _ = clean_events.add_manual_events([])
    finally:
        clean_events.load_manual_events = alt
    check("manual_events: Walking und Teens fallen, der Lauf bleibt",
          [e["wettbewerb"] for e in ergebnis], ["Courtepin-Fribourg 9.7 km"])

    # merge_duplicates verkettet nicht über eine Zeile ohne Distanz
    # (LCC Leopoldilauf: 14 und 21,1 km verschwanden, 13.10.2026).
    from clean_events import merge_duplicates
    basis = {"name": "LCC Leopoldilauf", "datum_start": "2030-11-08", "datum_ende": "2030-11-08",
             "art1": "Laufen", "standort": "Wien", "land": "Österreich", "lat": 48.2075, "lon": 16.4205}
    zeilen = [dict(basis), dict(basis, laenge_km=7.0, wettbewerb="7 km"),
              dict(basis, laenge_km=14.0, wettbewerb="14 km"),
              dict(basis, laenge_km=21.1, wettbewerb="Halbmarathon")]
    zusammen, _ = merge_duplicates(zeilen)
    check("Zeile ohne Distanz zieht nicht alle Strecken zusammen",
          sorted(e.get("laenge_km") or 0 for e in zusammen), [7.0, 14.0, 21.1])


def test_kinder() -> None:
    """Reine Kinder-/Jugendwettbewerbe raus (Nutzer, 09.10.2026: „Ja
    Kinderrennen raus"). Treffer und Gegenproben aus dem Bestand."""
    print("\nKinderrennen (ist_kinder):")
    from scraper_lib import ist_kinder, filter_kinder, Event
    from clean_events import drop_kinder

    for name, wb in (("Run & Bike Groisbach", "Sparefroh Kids Run I - 300 m (Jg. 2021 und jünger)"),
                     ("Dirty Race", "Dirty Kids Cross"),
                     ("Chemnitzer Indoor Triathlon", "Swim & Run Schüler C 100m / 400m"),
                     ("Radln & Renna", "Kinder U6/U8"),
                     ("Radln & Renna", "Jugendlauf U16-U18"),
                     ("Finteler Triathlon", "Schülertriathlon (Einzel)"),
                     ("TriZell", "Kinderbewerb (Aquathlon)"),
                     ("Südwestzipfel Lauf", "Jugendlauf, 5,6 km , U14 & U16 (Jg. 2011–2014)"),
                     ("Tauchscher Stadtlauf", "5 km für U14 und U16"),
                     ("Wolfgangseelauf", "Junior-Marathon, Strobl"),
                     ("Solitudelauf", "WLV-Kids-Marathon"),
                     ("4. Backnanger Kindercrossduathlon", None),
                     ("13. NÖ Kinder- und Jugendlauf Strasshof", None)):
        check(f"{name!r} / {wb!r} fällt", bool(ist_kinder(name, wb)), True)
    for name, wb in (("Legdener Mittsommernachtslauf", "5 km für Jugendliche und Erwachsene"),
                     ("46. Krummenseelauf", "7,2 km Lauf für Jugend und Altersklassen"),
                     ("Neuenkirchener Abendlauf", "5 km Straße für Jugend, Aktive und Masters"),
                     ("Halloweenlauf Buchen", "5 km Straße für U16, U18, M30, M40, M50, W30, W40, W50"),
                     ("Dortmunder Westfalenparklauf", "5,0 km (2 Runden) ab M/W U14"),
                     ("Liepnitzseelauf", "8,6 km Lauf ab Jugend U14"),
                     ("Kuckucks-Trail Damscheid", "9,5 km Best of Trail, Jugend ab 14 Jahre, Frauen, Männer"),
                     ("Vogelsanger Mailauf", "5 km Freizeit- und Jugendlauf"),
                     ("ISLANDMAN Norderney", "Schnupper- und Jugend-Triathlon"),
                     ("Kraichgauman Crossduathlon", "Sprint und Jugend"),
                     ("Ostufer Fischhallen-Lauf", "5 km Kurzlauf mit Schülerwertung"),
                     ("Crazy Runners Frankenwald Trail", "Mini Trail"),
                     ("Lindwurm", "Lindwurm mini (24,5 km, 550 hm)"),
                     ("Männedörfler Waldlauf", "U20"),
                     ("MyEifelRide Kids", "RTF 209 km"),
                     ("39. Pesenbachtallauf + Wirt in Pesenbach Kinderlauf", None),
                     ("15. Salzburg Trailrunning Festival - Kidstrail, Festungs- und Panoramatrail", None),
                     ("15. pro kids Lauf mit 5km Jedermannlauf", None),
                     ("Sponsorenlauf „Gemeinsam gegen Kinderkrebs“", None)):
        check(f"{name!r} / {wb!r} bleibt", ist_kinder(name, wb), None)

    zeilen = [Event(name="Dirty Race", wettbewerb="Sprint", art1="Triathlon"),
              Event(name="Dirty Race", wettbewerb="Dirty Kids Cross", laenge_km=1.0, art1="Triathlon")]
    behalten, n = filter_kinder(zeilen)
    check("filter_kinder nimmt nur die Kinderzeile", [e.wettbewerb for e in behalten], ["Sprint"])
    check("und zählt sie", n, 1)
    behalten, entfernt = drop_kinder([
        {"name": "Finteler Triathlon", "datum_start": "2027-09-05", "wettbewerb": "Swim & Run 1 (Kids)"},
        {"name": "Finteler Triathlon", "datum_start": "2027-09-05", "wettbewerb": "Sprintdistanz"},
    ])
    check("drop_kinder behält das Erwachsenenrennen", [e["wettbewerb"] for e in behalten], ["Sprintdistanz"])
    check("und meldet den Ausschluss", len(entfernt), 1)


def test_serientermin_im_label() -> None:
    """Ein Serientermin, der MITTEN im Wettbewerbs-Label steht.

    Die Hammer Winterlaufserie hat drei Termine (31.01. = 10 km,
    14.02. = 15 km, 28.02. = Halbmarathon), und `expand_competitions()`
    hatte jede Distanz an jeden Termin gehängt - neun Zeilen statt drei.
    Zwei davon trugen ihren Termin im Label: „15,0 km (am 14.02.2027 für
    M/W ab 18 bis M/W85)". `_DATE_LABEL` sah das nicht, weil es ein Label
    verlangt, das MIT dem Datum beginnt.

    Die Gegenproben sind der wichtigere Teil - gezählt wurde vorher: Im
    ganzen Bestand tragen genau zwei Zeilen ein „am <Datum>" im Label,
    beide echte Serientermine. Damit die Regel nicht mehr fängt als das,
    gelten zwei Bedingungen:

    - Das Wort „am" muss davorstehen. Ein blankes Datum irgendwo im Text
      kann alles sein.
    - Der Termin muss IM Zeitraum der Veranstaltung liegen, und Wörter
      wie „Anmeldeschluss" schließen ihn aus. Sonst schöbe
      „(Anmeldeschluss am 14.02.2027)" den Lauf auf die Meldefrist.
    """
    print("\nSerientermin im Wettbewerbs-Label (_datum_aus_label):")
    from clean_events import _datum_aus_label, fix_series_dates

    def zeile(wb, start="2027-01-31", ende="2027-02-28"):
        return {"name": "Hammer Winterlaufserie", "datum_start": start,
                "datum_ende": ende, "laenge_km": 15.0, "wettbewerb": wb}

    check("Klammerzusatz wird zum Termin",
          _datum_aus_label(zeile("15,0 km (am 14.02.2027 für M/W ab 18 bis M/W85)")),
          ("2027-02-14", "15,0 km"))
    check("und das Label behält seinen Rest",
          _datum_aus_label(zeile("Halbmarathon (am 28.02.2027 für M/W ab 18)")),
          ("2027-02-28", "Halbmarathon"))
    check("auch ohne Klammern", _datum_aus_label(zeile("10 km am 14.02.2027")),
          ("2027-02-14", "10 km"))

    check("Anmeldeschluss ist kein Termin",
          _datum_aus_label(zeile("10 km (Anmeldeschluss am 14.02.2027)")), None)
    check("ein Datum außerhalb des Zeitraums auch nicht",
          _datum_aus_label(zeile("Lauf am 14.03.2027")), None)
    check("ein blankes Datum ohne „am“ ebenfalls nicht",
          _datum_aus_label(zeile("10 km 14.02.2027 Start 11 Uhr")), None)
    check("ein gewöhnliches Label bleibt unangetastet",
          _datum_aus_label(zeile("15,0 km")), None)

    events = [zeile("15,0 km (am 14.02.2027 für M/W ab 18 bis M/W85)"),
              zeile("10,0 km")]
    meldungen = fix_series_dates(events)
    check("fix_series_dates setzt Start UND Ende",
          (events[0]["datum_start"], events[0]["datum_ende"]),
          ("2027-02-14", "2027-02-14"))
    check("und lässt die Zeile ohne Termin in Ruhe",
          (events[1]["datum_start"], events[1]["datum_ende"]),
          ("2027-01-31", "2027-02-28"))
    check("eine Meldung je verschobener Zeile", len(meldungen), 1)


def test_manuelle_events() -> None:
    """Einzeln recherchierte Strecken nachtragen (manual_events.json).

    Gebraucht wird das, weil ein Override eine Zeile ÄNDERN, aber keine
    ANLEGEN kann - und beim Durchgehen der Streckenlisten fehlten immer
    wieder einzelne Wettbewerbe einer vorhandenen Veranstaltung (die
    Bühlauer Winterlaufserie hat fünf Termine, wir hatten einen).

    Geprüft wird vor allem, dass dabei nichts DOPPELT entsteht:
    Übersprungen wird jeder Eintrag, zu dem `is_same_event()` schon eine
    Zeile findet. Damit ist der Schritt idempotent und verträgt sich mit
    einem späteren Scraper-Lauf, der dieselbe Strecke selbst einsammelt.
    """
    print("\nEinzeln recherchierte Events (manual_events.json):")
    from clean_events import add_manual_events, load_manual_events

    eintraege = load_manual_events()
    check("die Datei ist lesbar und nicht leer", len(eintraege) > 0, True)
    check("Dokumentations-Felder landen nicht in events.json",
          [k for e in eintraege for k in e if k.startswith("_")], [])
    for feld in ("name", "datum_start", "standort", "art1", "veranstalter_url"):
        fehlend = [e.get("name") for e in eintraege if not e.get(feld)]
        check(f"jeder Eintrag hat {feld}", fehlend, [])
    ohne_koordinaten = [e.get("name") for e in eintraege
                        if e.get("lat") is None or e.get("lon") is None]
    check("jeder Eintrag hat Koordinaten", ohne_koordinaten, [])

    beispiel = dict(eintraege[0])
    ergaenzt, meldungen = add_manual_events([beispiel])
    schluessel = (beispiel["name"], beispiel["datum_start"], beispiel.get("laenge_km"))
    check("eine schon vorhandene Strecke wird nicht doppelt angelegt",
          sum(1 for e in ergaenzt
              if (e["name"], e["datum_start"], e.get("laenge_km")) == schluessel), 1)
    check("und alle anderen kommen dazu", len(meldungen), len(eintraege) - 1)

    ergaenzt, meldungen = add_manual_events([])
    check("in eine leere Liste wird alles nachgetragen",
          (len(ergaenzt), len(meldungen)), (len(eintraege), len(eintraege)))
    nochmal, meldungen = add_manual_events(ergaenzt)
    check("ein zweiter Durchlauf trägt nichts mehr nach", meldungen, [])
    check("idempotent, also gleich viele Zeilen", len(nochmal), len(eintraege))


def test_mehrsport_teilstrecken() -> None:
    """Die Länge eines Triathlons ist die SUMME seiner Teilstrecken.

    `guess_distance_km()` nimmt bei mehreren Zahlen die größte - bei
    einem Triathlon ist das die Radstrecke. Deshalb stand die
    Kurzdistanz des Triathlon Höchstadt mit „40 km" in der Liste statt
    mit 51,5 km, und dasselbe bei jedem zweiten Triathlon im Bestand:
    Niederrhein N3T 38 statt 49,5; Frankfurt City 80 statt 102;
    Indeland 88 statt 109,9. Eine Radstrecke als Länge des Rennens ist
    doppelt falsch - die Zahl stimmt nicht, und sie sieht aus wie ein
    Radrennen.

    Die Gegenproben sind hier der wichtigere Teil: Die Regel darf NUR
    dort zuschlagen, wo wirklich Teilstrecken aufgezählt sind.
    """
    print("\nMehrsport: Teilstrecken summieren (summiere_teilstrecken):")
    from scraper_lib import summiere_teilstrecken as summe

    check("1,5 km Schwimmen + 40 km Rad + 10 km Laufen",
          summe("1,5 km Schwimmen, 40 km Radfahren, 10 km Laufen"), 51.5)
    check("englisch und ohne Leerzeichen", summe("400m Swim 20km Bike 5km Run"), 25.4)
    check("Meter mit Tausenderpunkt", summe("1.500m Swim 40km Bike 10km Run"), 51.5)
    check("Wort vor der Zahl", summe("Schwimmen 1,5 km / Rad 40 km / Laufen 10 km"), 51.5)
    check("Langdistanz", summe("3,8 km Schwimmen, 180 km Rad, 42,2 km Laufen"), 226.0)
    check("Duathlon ohne Schwimmen", summe("10 km Laufen und 40 km Radfahren"), 50.0)

    # --- Gegenproben ---------------------------------------------------
    check("ein reiner Lauftext bleibt unberührt",
          summe("5 km Lauf, 10 km Lauf, 21,1 km Lauf"), None)
    check("ein Halbmarathon auch", summe("21,1 km Halbmarathon"), None)
    check("ein Radrennen allein reicht nicht", summe("Rad 100 km"), None)
    check("eine einzelne Teilstrecke auch nicht", summe("ca. 18,6 km Radfahren"), None)
    check("'Mountainbike Rennen 42 km' ist kein Mehrsport",
          summe("Mountainbike Rennen 42 km"), None)
    # Der wichtigste Schutz: Zählt ein Text MEHRERE Wettbewerbe auf,
    # lässt sich nicht sagen, welche Zahlen zusammengehören. Dann lieber
    # nichts - der Aufrufer macht weiter wie bisher.
    check("zwei Wettbewerbe in einem Text ergeben nichts",
          summe("Jedermann 400m Swim 20km Bike 5km Run "
                "Kurzdistanz 1.500m Swim 40km Bike 10km Run"), None)
    check("leerer Text", summe(""), None)

    from scraper_lib import guess_distance_km
    check("guess_distance_km nutzt die Summe",
          guess_distance_km("1,5 km Schwimmen, 40 km Radfahren, 10 km Laufen", CONFIG), 51.5)
    check("und bleibt sonst wie bisher",
          guess_distance_km("5 km, 10 km, 42,195 km", CONFIG), 42.2)


def test_kalender_staging() -> None:
    """`stage_kalender()` fasst den Git-Index nur in GitHub Actions an.

    Der Umweg gibt es, weil GitHub den `schedule`-Trigger nur aus dem
    Standard-Branch liest und die `update-events.yml` auf `main` ein
    älterer Stand ist: Sie committet `git add events.json` OHNE
    `kalender`. `git commit` committet aber den INDEX - was hier gestaget
    ist, geht also mit.

    Geprüft wird vor allem die Gegenprobe: **Lokal darf das Skript den
    Index NIE anfassen.** Ein Werkzeug, das ungefragt `git add` ausführt,
    wäre eine böse Überraschung.
    """
    print("\nKalenderdateien für den Commit vormerken (stage_kalender):")
    import os
    import subprocess
    from update_events import stage_kalender

    skripte = Path(__file__).resolve().parent
    wurzel = skripte.parent

    alt = os.environ.pop("GITHUB_ACTIONS", None)
    try:
        vorher = subprocess.run(["git", "diff", "--cached", "--name-only"],
                                cwd=wurzel, capture_output=True, text=True).stdout
        stage_kalender()
        nachher = subprocess.run(["git", "diff", "--cached", "--name-only"],
                                 cwd=wurzel, capture_output=True, text=True).stdout
        check("ohne GITHUB_ACTIONS bleibt der Index unberührt", nachher, vorher)
    finally:
        if alt is not None:
            os.environ["GITHUB_ACTIONS"] = alt

    quelle = (skripte / "update_events.py").read_text(encoding="utf-8")
    check("der Aufruf steht in run_build_ics()",
          "    stage_kalender()\n    return True" in quelle, True)
    check("und ist an GITHUB_ACTIONS gebunden",
          'os.environ.get("GITHUB_ACTIONS")' in quelle, True)


def _node_event_detail(ausdruck: str):
    """Lässt event-detail.js komplett in node laufen (mit einem Stummel
    für EnduranceFilters) und gibt das JSON des Ausdrucks zurück - kein
    Nachbau, es läuft der Code der Seiten (`EED` ist das Modul)."""
    import json as _json
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        return None
    wurzel = Path(__file__).resolve().parent.parent
    # Erst das echte filters.js (EF), dann event-detail.js - kein
    # Stummel: triathlonFormat() wohnt in filters.js.
    skript = (
        "global.window = global;"
        f"require({_json.dumps(str(wurzel / 'filters.js'))});"
        "global.EnduranceFilters.relativeDays = () => '';"
        f"require({_json.dumps(str(wurzel / 'event-detail.js'))});"
        "const EED = global.EnduranceDetail;"
        f"console.log(JSON.stringify({ausdruck}));"
    )
    ergebnis = subprocess.run([node, "-e", skript], capture_output=True, text=True)
    if ergebnis.returncode != 0:
        check("event-detail.js läuft in node", ergebnis.stderr.strip()[:300], "")
        return None
    return _json.loads(ergebnis.stdout)


def test_wettbewerb_zusatz() -> None:
    """Der Zusatz unter dem Namen (displayWettbewerb): keine Maßzahl mehr.

    Vom Nutzer am 21.09.2026 entschieden: "Wir haben die Länge in der
    Liste schon und in der Detailansicht dann auch. Man muss es nicht 3x
    sehen." Hier die Fälle aus seiner Meldung und die Gegenproben, an
    denen die erste Fassung gescheitert war (leere Klammern, "ca. ,",
    ein führendes "mit", "400-m-Runde" mit Bindestrich)."""
    print("\nZusatz unter dem Namen (displayWettbewerb):")
    faelle = [
        ("Brian Trail (15,5 km, 500 hm)", 15.5, "Brian Trail"),
        ("5.555 m (Berglauf auf den Lousberg)", 5.6, "Berglauf auf den Lousberg"),
        ("10 km Fuchsburg Lauf (ab Jahrgang 2015)", 10, "Fuchsburg Lauf (ab Jahrgang 2015)"),
        ("Marathon", 42.2, None),
        ("21,1 km", 21.1, None),
        ("18 km", 18, None),
        ("Marathon (6 Runden à 7,5 km)", 42.2, None),
        ("Kurzdistanz 51,5 km (1,5 km Schwimmen / 40 km Rad / 10 km Laufen)", 51.5, "Kurzdistanz"),
        ("2x5 km Staffel", 5, "Staffel"),
        ("4-Stunden-Lauf (400-m-Runde)", None, None),
        ("6 km mit 15 Hindernissen (Fun Distanz)", 6, "15 Hindernissen (Fun Distanz)"),
        ("ca. 10.700 m, 25+ Hindernisse", 10.7, "25+ Hindernisse"),
        ("Nordic Walking", 8, "Nordic Walking"),
        ("5 km Walking - Startgebühr 7 Euro", 5, "Walking - Startgebühr 7 Euro"),
        ("Moslig 8000", None, "Moslig 8000"),
        ("Cross lang (10.060 m)", 10.1, "Cross lang"),
        # Wiederholt nur Sportart oder Kategorie: kein Zusatz.
        ("Laufen", 10, None),
        ("Trail", 12, None),
    ]
    import json as _json
    events = [{"wettbewerb": w, "laenge_km": km, "art1": "Laufen", "art2": "Trail"} for w, km, _ in faelle]
    got = _node_event_detail(_json.dumps(events, ensure_ascii=False) + ".map(e => EED.displayWettbewerb(e))")
    if got is None:
        print("  (node nicht vorhanden - übersprungen)")
        return
    for (w, km, erwartet), ist in zip(faelle, got):
        check(f"{w!r} -> {erwartet!r}", ist, erwartet)


def test_triathlon_format() -> None:
    """triathlonFormat(): erst das Label, dann das SCHWIMMEN, dann die Summe.

    Datenregel 21 (21.09.2026): Weinstadt hat 23,6 km Gesamtlänge, aber
    nur 300 m Schwimmen - Super-Sprint, nicht Sprint. Die Regel greift
    nur nach unten, nur mit Teilstrecke im Label und nur ohne
    Format-Stichwort."""
    print("\nTriathlon-Format (triathlonFormat):")
    T = "Triathlon"
    faelle = [
        ({"art1": T, "laenge_km": 23.6, "wettbewerb": "Jedermann-Triathlon 23,6 km (0,3 km Schwimmen / 18,7 km Rad / 4,6 km Laufen)"}, "supersprint"),
        ({"art1": T, "laenge_km": 23.7, "wettbewerb": "Volkstriathlon 23,7 km (0,7 km Schwimmen / 18 km Rad / 5 km Laufen)"}, "sprint"),
        ({"art1": T, "laenge_km": 25.4, "wettbewerb": "Volksdistanz 25,4 km (400 m Schwimmen / 20 km Rad / 5 km Laufen)"}, "supersprint"),
        ({"art1": T, "laenge_km": 25.4, "wettbewerb": "Jedermann Sprint 25,4 km (400 m Schwimmen / 20 km Rad / 5 km Laufen)"}, "sprint"),
        ({"art1": T, "laenge_km": 23.6, "wettbewerb": "Jedermann-Triathlon"}, "sprint"),
        ({"art1": T, "laenge_km": 51.5, "wettbewerb": ""}, "olympisch"),
        ({"art1": T, "laenge_km": 113, "wettbewerb": ""}, "mittel"),
        ({"art1": T, "laenge_km": 226, "wettbewerb": ""}, "lang"),
        ({"art1": T, "laenge_km": 452, "wettbewerb": ""}, "ultra"),
        ({"art1": T, "laenge_km": 25.5, "wettbewerb": "Volksdistanz 25,5 km (500 m Schwimmen / 20 km Rad / 5 km Laufen)"}, "sprint"),
        ({"art1": T, "art2": "Swimrun", "laenge_km": 40, "wettbewerb": ""}, None),
        ({"art1": "Laufen", "laenge_km": 42.2, "wettbewerb": ""}, None),
    ]
    import json as _json
    got = _node_event_detail(_json.dumps([f for f, _ in faelle], ensure_ascii=False)
                             + ".map(e => EED.triathlonFormat(e))")
    if got is None:
        print("  (node nicht vorhanden - übersprungen)")
        return
    for (f, erwartet), ist in zip(faelle, got):
        check(f"{(f.get('wettbewerb') or f['laenge_km'])!r} -> {erwartet!r}", ist, erwartet)


def test_triathlon_format_kopie() -> None:
    """triathlonFormat() steht zweimal - filters.js (Liste, Karte) und
    functions/index.js (Abo-Filter). Beide Fassungen müssen wortgleich
    sein (bis auf Einrückung und Anführungszeichen), und der Längen-
    Filter muss bei JEDEM Triathlon im Bestand dieselbe Kategorie
    treffen, die die Spalte anzeigt (Datenregel 21)."""
    print("\nTriathlon-Format: Kopie in functions/index.js und Filter == Anzeige:")
    wurzel = Path(__file__).resolve().parent.parent
    def block(text: str) -> str:
        a = text.index("TRIATHLON_FORMAT_IM_LABEL = [")
        b = text.index("TRIATHLON_FORMAT_KATEGORIE = {")
        zeilen = []
        for z in text[a:b].splitlines():
            z = z.strip().replace('"', "'")
            if z and not z.startswith("//"):
                zeilen.append(z)
        return "\n".join(zeilen)
    web = block((wurzel / "filters.js").read_text(encoding="utf-8"))
    fn = block((wurzel / "functions" / "index.js").read_text(encoding="utf-8"))
    check("filters.js und functions/index.js: dieselbe Regel", fn == web, True)
    # Dritte Kopie (05.10.2026): audit_events.TRIATHLON_FORMAT_IM_LABEL
    # entscheidet, ob eine Zeile ohne Kilometer "Distanz fehlt" meldet.
    import re as _re
    from audit_events import TRIATHLON_FORMAT_IM_LABEL as _py_fmt
    js_fmt = _re.findall(r"\[/(.+?)/i, '(\w+)'\]", web)
    py_fmt = [(m.pattern, k) for m, k in _py_fmt]
    check("audit_events.py: dieselben Format-Muster wie filters.js", py_fmt, js_fmt)
    if fn != web:
        import difflib
        for d in list(difflib.unified_diff(web.splitlines(), fn.splitlines(), lineterm=""))[:12]:
            print("    " + d)

    import json as _json
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        print("  (node nicht vorhanden - übersprungen)")
        return
    skript = (
        "global.window = global;"
        f"require({_json.dumps(str(wurzel / 'filters.js'))});"
        "const EF = global.EnduranceFilters;"
        f"const events = require({_json.dumps(str(wurzel / 'events.json'))}).filter(e => e.art1 === 'Triathlon' && e.laenge_km != null);"
        "const KAT = { supersprint: 'supersprint', sprint: 'sprint', olympisch: 'olympic', mittel: 'middle', lang: 'long', ultra: 'tultra' };"
        "const keys = Object.values(KAT);"
        "let falsch = [];"
        "for (const e of events) {"
        "  const f = EF.triathlonFormat(e); if (!f) continue;"
        "  const st = EF.createState();"
        "  const getroffen = keys.filter(k => { st.distanceCategories = new Set(['Triathlon:' + k]); return EF.matchEvent(st, e); });"
        "  if (getroffen.length !== 1 || getroffen[0] !== KAT[f]) falsch.push([e.name, e.laenge_km, f, getroffen]);"
        "}"
        "console.log(JSON.stringify({ n: events.length, falsch: falsch.slice(0, 5) }));"
    )
    ergebnis = subprocess.run([node, "-e", skript], capture_output=True, text=True)
    if ergebnis.returncode != 0:
        check("filters.js läuft in node", ergebnis.stderr.strip()[:300], "")
        return
    out = _json.loads(ergebnis.stdout)
    check(f"Filterkategorie == angezeigtes Format bei {out['n']} Triathlon-Zeilen", out["falsch"], [])


def test_charity_merkmal() -> None:
    """fix_charity(): das Merkmal aus dem Namen, die alte Kategorie geleert,
    refresh_art2() holt die echte zurück (Datenregel 17)."""
    print("\nCharity-Merkmal (fix_charity):")
    from clean_events import fix_charity, refresh_art2  # lokaler Import
    rows = [
        {"name": "Bietlauf für einen Wohltätigen Zweck", "datum_start": "2026-10-03", "art1": "Laufen", "art2": "Charity", "wettbewerb": "9,2 km Crosslauf"},
        {"name": "Lauf für einen guten Zweck - Rastenberg", "datum_start": "2026-09-22", "art1": "Laufen", "art2": "Straße"},
        {"name": "Stadtlauf Erding", "datum_start": "2026-09-27", "art1": "Laufen", "art2": "Straße"},
        {"name": "Ahmadiyya Charity Walk Neuwied", "datum_start": "2026-10-11", "art1": "Laufen", "art2": "Charity"},
    ]
    fix_charity(rows)
    check("Benefiz-Crosslauf: charity gesetzt", rows[0].get("charity"), True)
    check("... und die alte Kategorie geleert", rows[0].get("art2"), None)
    check("'guten Zweck' -> charity", rows[1].get("charity"), True)
    check("Stadtlauf bleibt ohne Merkmal", rows[2].get("charity"), None)
    refresh_art2(rows)
    check("refresh_art2 holt die Kategorie aus dem Label zurück", rows[0].get("art2"), "Trail")
    check("Charity Walk ist ein Marsch (Kategorie), das Herz bleibt das Merkmal", rows[3].get("art2"), "Marsch")
    vorher = [dict(r) for r in rows]
    fix_charity(rows); refresh_art2(rows)
    check("fix_charity ist idempotent", rows, vorher)


def test_veranstalter_links() -> None:
    """Die Prüfregel von scripts/veranstalter_links.py: Eine Kandidatenseite
    gilt nur, wenn sie den Lauf am Namen nennt (neunter Durchgang,
    19.09.2026). Festgehalten werden die Gegenproben, die beim Bau
    aufgetreten sind."""
    import veranstalter_links as vl
    # Umlaute im Hostnamen: „Dülmen" heißt im Netz tsg-duelmen.de.
    assert "duelmen" in vl.host_woerter("Nikolauslauf der TSG Dülmen")
    assert vl.nennt_den_lauf("<p>Termine</p>", "https://www.tsg-duelmen.de/de/sport/",
                             ["Nikolauslauf der TSG Dülmen"], ["2026-12-05"], ["Dülmen"]) == ["host:duelmen"]
    assert vl.nennt_den_lauf("<p>x</p>", "http://www.kyffhaeuser-berglauf.de/",
                             ["Kyffhäuser Berglauf"], ["2027-04-10"], ["Bad Frankenhausen"]) == ["host:kyffhaeuser"]
    # Der Ortsname zählt im TEXT nicht - sonst wäre jeder Laufshop der Stadt ein Treffer.
    assert vl.nennt_den_lauf("<p>Laufschuhe in Leipzig kaufen</p>", "https://shop.example.de/",
                             ["Leipzig Run"], ["2026-10-11"], ["Leipzig"]) == []
    # Allgemeine Wörter („Herbstlauf", „Sparkasse") belegen nichts, ein unverwechselbares schon.
    assert vl.nennt_den_lauf("<p>Herbstlauf der Sparkasse</p>", "https://example.de/",
                             ["Sparkassen Herbstlauf"], ["2026-10-11"], ["Bayreuth"]) == []
    assert vl.nennt_den_lauf("<p>Der Ingelheimer Polderlauf startet</p>", "https://www.polderlauf.de/",
                             ["Ingelheimer Polderlauf"], ["2026-10-03"], ["Ingelheim am Rhein"])[0] == "host:polderlauf"
    # Das Datum allein reicht als Beleg.
    assert "datum" in vl.nennt_den_lauf("<p>Start am 03.10.2026</p>", "https://example.de/",
                                        ["Lauf"], ["2026-10-03"], ["Ort"])
    # Umlaut-Domains stehen in der Adresse als Punycode - dekodiert steckt der Name darin
    # (Silvesterlauf Mörschied auf tus-mörschied.de, zehnter Durchgang 20.09.2026).
    assert vl.host_von("https://www.xn--tus-mrschied-8ib.de/index.php/laufen") == "tus-mörschied.de"
    assert vl.nennt_den_lauf("<p>x</p>", "https://www.xn--tus-mrschied-8ib.de/silvesterlauf",
                             ["Silvesterlauf Mörschied"], ["2026-12-31"], ["Mörschied"]) == ["host:morschied"]
    # Der ORT im Hostnamen zählt nur für handverlesene Kandidaten (verifizieren):
    # djk-herzogenrath.de belegt den Volkslauf Herzogenrath - herzogenrath.de wäre die Stadt.
    assert vl.nennt_den_lauf("<p>Ausschreibung</p>", "http://www.djk-herzogenrath.de/seite/ausschreibung.html",
                             ["46. Internationaler Halbmarathon, 56. Internationaler Volkslauf"],
                             ["2026-10-24"], ["Herzogenrath"]) == []
    # Seit dem 22.09.2026 heißt der Treffer „ort:", nicht „host:": `verifizieren` nimmt
    # einen reinen Orts-Treffer nur noch als `unklar` (falkensteinlauf.de war ein anderer
    # Falkenstein-Lauf, tsg-leutkirch.de der Juli-Volkslauf statt der Stadtmeisterschaft).
    assert vl.nennt_den_lauf("<p>Ausschreibung</p>", "http://www.djk-herzogenrath.de/seite/ausschreibung.html",
                             ["46. Internationaler Halbmarathon, 56. Internationaler Volkslauf"],
                             ["2026-10-24"], ["Herzogenrath"], ort_im_host=True) == ["ort:herzogenrath"]
    # Ort im Host PLUS Datum auf der Seite bleibt ein voller Beleg.
    assert vl.nennt_den_lauf("<p>Start am 24.10.2026</p>", "http://www.djk-herzogenrath.de/",
                             ["46. Internationaler Volkslauf"], ["2026-10-24"], ["Herzogenrath"],
                             ort_im_host=True) == ["ort:herzogenrath", "datum"]
    # Ergebnisdienste, Karten und Datenschutzseiten sind keine Veranstalter.
    assert vl.kandidaten_url_normalisieren("https://www.sportstiming.dk/event/17605") is None
    assert vl.kandidaten_url_normalisieren("https://example.de/datenschutz") is None
    print("✓ Veranstalterseiten-Prüfung (Umlaut-Hosts, Ortswort, ort:-Treffer, Allgemeinwörter, Datum, Fremd-Hosts)")


def test_schwimmen_regeln() -> None:
    """Reine Schwimm-Events erst ab 500 m, und keine Meisterschaften
    (vom Nutzer am 24.09.2026 entschieden: "Schwimm events sollten erst
    ab 500m aufgenommen werden. Also reine Schwimmevents. Bei einem
    Triathlon kann es auch weniger sein als 500m. Ich möchte keine
    Schwimm Events aufnehmen, die nicht für jeden sind, also 50m
    deutsche Meisterschaft etc.").

    Die Gegenproben sind der wichtigere Teil: Der 300-m-Schwimmteil
    eines Super-Sprint-Triathlons bleibt, ein 3,5-km-Lauf fällt weiter,
    ein Zeitrennen ist nie zu kurz, und eine LAUF-Meisterschaft bleibt
    (Datenregel 18 - dort entscheidet der Einzelfall)."""
    print("\nSchwimmen: 500-m-Grenze und Meisterschaften:")
    from scraper_lib import (Event, MIN_DISTANCE_BY_ART1, filter_min_distance,
                             filter_nicht_offen_schwimmen, ist_nicht_offen_schwimmen,
                             ist_zu_kurz)
    from clean_events import drop_nicht_offen_schwimmen, drop_too_short

    check("Schwimmen hat 0,5 km", MIN_DISTANCE_BY_ART1.get("Schwimmen"), 0.5)
    check("Laufen hat 5 km", MIN_DISTANCE_BY_ART1.get("Laufen"), 5.0)
    check("400 m Schwimmen ist zu kurz", ist_zu_kurz("Schwimmen", 0.4), True)
    check("500 m Schwimmen bleibt", ist_zu_kurz("Schwimmen", 0.5), False)
    check("3,5 km Lauf ist zu kurz", ist_zu_kurz("Laufen", 3.5), True)
    check("Triathlon kennt keine Grenze", ist_zu_kurz("Triathlon", 0.3), False)
    check("Fahrrad kennt keine Grenze", ist_zu_kurz("Fahrrad", 3.0), False)
    check("unbekannte Distanz ist nie zu kurz", ist_zu_kurz("Schwimmen", None), False)
    check("Zeitrennen ist nie zu kurz", ist_zu_kurz("Schwimmen", 0.1, 24.0), False)

    ev = [Event(art1="Schwimmen", name="Seeschwimmen", laenge_km=0.25),
          Event(art1="Schwimmen", name="Seeschwimmen", laenge_km=0.5),
          Event(art1="Triathlon", name="Weinstadt Triathlon", laenge_km=23.6),
          Event(art1="Laufen", name="Kinderlauf", laenge_km=2.0),
          Event(art1="Schwimmen", name="24h Schwimmen", laenge_km=0.05, dauer_h=24.0)]
    kept, skipped = filter_min_distance(ev)
    check("filter_min_distance wirft 250 m Schwimmen und 2-km-Lauf", skipped, 2)
    check("… und behält 500 m, Triathlon und Zeitrennen", len(kept), 3)
    # Ältere Aufrufer mit ausdrücklicher Grenze: nur Laufen, wie vorher.
    kept, skipped = filter_min_distance(ev, 5.0)
    check("filter_min_distance(min_km) bleibt eine Lauf-Regel", skipped, 1)
    kept, dropped = drop_too_short([
        {"name": "Seeschwimmen", "art1": "Schwimmen", "laenge_km": 0.4},
        {"name": "Seeschwimmen", "art1": "Schwimmen", "laenge_km": 1.0},
        {"name": "Triathlon", "art1": "Triathlon", "laenge_km": 0.3},
        {"name": "24h Halle", "art1": "Schwimmen", "laenge_km": 0.05, "dauer_h": 24}])
    check("clean_events.drop_too_short ebenso", (len(kept), len(dropped)), (3, 1))

    for name in ("Deutsche Meisterschaften Freiwasser", "Bayerische Freiwassermeisterschaft",
                 "DM Kurzbahn Wuppertal", "Landesmeisterschaften Schwimmen",
                 "Open Water Championships Zürich", "Offene Sächsische Freiwassermeisterschaften"):
        check(f"{name!r} ist nicht offen", bool(ist_nicht_offen_schwimmen("Schwimmen", name)), True)
    for name in ("Bodensee Openwater Konstanz", "Alpen Open Water Cup Simssee",
                 "Hafenschwimmen Wilhelmshaven", "24h Friedberg", "Wörthersee Swim",
                 "Jedermann-Schwimmen am Pöhl-Cup"):
        check(f"{name!r} bleibt", ist_nicht_offen_schwimmen("Schwimmen", name), None)
    check("Lauf-Meisterschaft ist nicht betroffen",
          ist_nicht_offen_schwimmen("Laufen", "Bayerische Halbmarathon-Meisterschaften"), None)
    check("Label zählt mit",
          bool(ist_nicht_offen_schwimmen("Schwimmen", "Seefest Wörthsee", "Meisterschaft 1500 m")), True)
    kept, skipped = filter_nicht_offen_schwimmen([
        Event(art1="Schwimmen", name="DM Freiwasser"), Event(art1="Schwimmen", name="Seeschwimmen"),
        Event(art1="Laufen", name="Deutsche Meisterschaft 10 km")])
    check("filter_nicht_offen_schwimmen", (len(kept), skipped), (2, 1))
    kept, entfernt = drop_nicht_offen_schwimmen([
        {"name": "DM Freiwasser", "art1": "Schwimmen", "datum_start": "2027-06-01"},
        {"name": "Seeschwimmen", "art1": "Schwimmen"}])
    check("clean_events.drop_nicht_offen_schwimmen", (len(kept), len(entfernt)), (1, 1))
    # Am Bestand: Keine Schwimm-Zeile darf dadurch fallen, ohne dass
    # es gemeldet wird - hier zählen wir nur, was gemeldet würde.
    import json as _json
    from scraper_lib import EVENTS_JSON_PATH
    bestand = _json.loads(EVENTS_JSON_PATH.read_text(encoding="utf-8"))
    schwimmen = [e for e in bestand if e.get("art1") == "Schwimmen"]
    zu_kurz = [e for e in schwimmen if ist_zu_kurz("Schwimmen", e.get("laenge_km"), e.get("dauer_h"))]
    check("Bestand: keine Schwimm-Zeile unter 500 m", len(zu_kurz), 0)


def test_schwimmkalender() -> None:
    """Der Scraper für schwimmkalender.de (24.09.2026): Listenzeile,
    Detailseite, Distanz/Dauer/Ort aus der Bezeichnung - ohne Netz."""
    print("\nschwimmkalender.de:")
    import schwimmkalender_scraper as sk
    liste = ("<table id='example'><tr><td width='15%' Title='25.09.2026 '>25.09.2026</td>"
             "<td width='60%'><a href='../main/main.php?d=0&Action=views/lstuserevents.php"
             "&fOK=update|userevents|7275&id=185021732606602137475479552159275218493' class='kat64'"
             "   onclick=\"DoScroll2QueryString(event)\" target='_blank'>24h Friedberg</a></td>"
             "<td width='20%'>Hessen</td></tr><tr><td width='15%' Title='03.10.2026 '>03.10.2026</td>"
             "<td width='60%'><a href='../main/main.php?d=0&Action=views/lstuserevents.php"
             "&fOK=update|userevents|7055&id=1' class='kat63' target='_blank'>"
             "Hafenschwimmen, Wilhelmshaven (3k)</a></td><td width='20%'>Niedersachsen</td></tr>"
             "</table>Freiwasser: 2 Einträge gefunden<input type='hidden' name='id' "
             "value='12339903260660213953077955215927521849352' id = 'id'>")
    zeilen = sk.parse_liste(liste)
    check("zwei Listenzeilen", len(zeilen), 2)
    check("Zeile: Nummer, Kategorie, Name, Land",
          (zeilen[1]["nr"], zeilen[1]["kategorie"], zeilen[1]["name"], zeilen[1]["land"]),
          (7055, 63, "Hafenschwimmen, Wilhelmshaven (3k)", "Niedersachsen"))
    check("Sitzungsnummer", sk.sitzungs_id(liste), "12339903260660213953077955215927521849352")
    check("Liste erkannt", sk.ist_liste(liste), True)
    check("Startseite ist keine Liste", sk.ist_liste("<p>Liebe Schwimmerinnen</p>"), False)
    detail = ("<div id='iKategorieNr'>Freiwasser</div><div class='showastext' id='iBezeichnung' >"
              "Hafenschwimmen, Wilhelmshaven (3k)</div><div id='iDatum' >03.10.2026</div>"
              "<div id='iEndDatum' >-</div><div id='iLandNr' >Niedersachsen</div>"
              "<div id='textarea' ><p><a href=\"http://www.hafenschwimmen.de/\" target=\"_blank\">"
              "Hafenschwimmen</a></p></div><div id='simpletextarea' ></div>")
    d = sk.parse_detail(detail)
    check("Detail: Felder", (d["name"], d["datum_start"], d["datum_ende"], d["land"], d["links"]),
          ("Hafenschwimmen, Wilhelmshaven (3k)", "2026-10-03", None, "Niedersachsen",
           ["http://www.hafenschwimmen.de/"]))
    ev = sk.build_events(d, 63, sk.CONFIG)
    check("eine Zeile mit 3 km, Schwimmen/Freiwasser, Ort und Link",
          [(e.name, e.standort, e.land, e.art1, e.art2, e.laenge_km, e.veranstalter_url) for e in ev],
          [("Hafenschwimmen Wilhelmshaven", "Wilhelmshaven", "Deutschland", "Schwimmen", "Freiwasser",
            3.0, "http://www.hafenschwimmen.de/")])
    check("Distanzen aus der Klammer", sk.distanzen_aus_name("Seeschwimmen, Waging am See (2k/5k)"), [2.0, 5.0])
    check("Komma-Dezimale", sk.distanzen_aus_name("Hechtsee X-Treme, Kufstein (3,8k)"), [3.8])
    check("Meter in der Klammer", sk.distanzen_aus_name("Spreeschwimmen (750m)"), [0.75])
    check("Zahl im Namen ist keine Distanz", sk.distanzen_aus_name("100 x 100 Hamburg"), [])
    check("Dauer aus dem Namen", (sk.dauer_aus_name("24h Friedberg"), sk.dauer_aus_name("25h Handorf"),
                                  sk.dauer_aus_name("Hafenschwimmen (3k)")), (24.0, 25.0, None))
    for name, land, ort in (("Hafenschwimmen, Wilhelmshaven (3k)", "Niedersachsen", "Wilhelmshaven"),
                            ("24h Spaichingen-Aldingen", "Baden-Württemberg", "Spaichingen-Aldingen"),
                            ("25h Schwarzenbach a. Wald", "Bayern", "Schwarzenbach a. Wald"),
                            ("Seeschwimmen Waging am See (5k)", "Bayern", "Waging am See"),
                            ("Urban Challenge Berlin", "Berlin", "Berlin"),
                            ("100 x 100 Hamburg", "Hamburg", "Hamburg"),
                            ("Rotary Charity-Schwimmen (24h Freising)", "Bayern", "Freising"),
                            ("Berliner Spreeschwimmen (1k)", "Berlin", "Berlin"),
                            ("Seeschwimmen (2k)", "Bayern", None),
                            ("Vollmondschwimmen (2k)", "Bayern", None)):
        check(f"Ort aus {name!r}", sk.ort_aus_name(name, land), ort)
    check("Name ohne Distanz und Komma", sk.name_ohne_zusatz("Hafenschwimmen, Wilhelmshaven (3k)"),
          "Hafenschwimmen Wilhelmshaven")
    check("Klammer ohne Distanz bleibt", sk.name_ohne_zusatz("Rotary Charity-Schwimmen (24h Freising)"),
          "Rotary Charity-Schwimmen (24h Freising)")
    check("Bundesland → Deutschland", sk.land_aus_feld("Hessen"), "Deutschland")
    check("Staat bleibt", (sk.land_aus_feld("Österreich"), sk.land_aus_feld("Italien")), ("Österreich", "Italien"))
    d24 = {"name": "24h Friedberg", "datum_start": "2026-09-25", "datum_ende": "2026-09-26",
           "land": "Hessen", "links": []}
    ev = sk.build_events(d24, 64, sk.CONFIG)
    check("24h Halle ist ein Zeitrennen im Becken",
          [(e.name, e.standort, e.art1, e.art2, e.dauer_h, e.laenge_km, e.datum_ende) for e in ev],
          [("24h Friedberg", "Friedberg", "Schwimmen", "Becken", 24.0, None, "2026-09-26")])
    dsr = {"name": "Urban Challenge Berlin", "datum_start": "2026-09-27", "datum_ende": None,
           "land": "Berlin", "links": ["https://www.urban-challenge.de/"]}
    ev = sk.build_events(dsr, 69, sk.CONFIG)
    check("SwimRun ist Triathlon/Swimrun", [(e.art1, e.art2, e.standort) for e in ev],
          [("Triathlon", "Swimrun", "Berlin")])
    ev = sk.build_events({"name": "Seeschwimmen (2k)", "datum_start": "2027-06-01", "datum_ende": None,
                          "land": "Bayern", "links": []}, 63, sk.CONFIG)
    check("ohne erkennbaren Ort keine Zeile", ev, [])
    check("Kategorie-URL", sk.kategorie_url(63, "1"),
          "https://www.schwimmkalender.de/sk_kalender/main/main.php?d=0&Action=views/start.php&fOK=userevents|63|63|63|63&id=1")


def test_neue_quellen() -> None:
    """Die Scraper vom 22.09.2026 (endure-cycling, ÖLV, Sparkasse Running,
    Laufkalender NWS, lauftermine.ch): die Parser, die am Bestand der
    Live-Seiten kalibriert wurden, ohne Netz - und die Regeln, die beim
    Bau aufgetreten sind."""
    from scraper_lib import ort_aus_veranstaltungsort as ort
    # Veranstaltungsorte -> Ort (ÖLV- und Sparkassen-Kalender).
    assert ort("OÖ - 4020 Linz, PHDL Linz (Pädagogische Hochschule)") == "Linz"
    assert ort("Heideparkplatz am Ende der Berggasse 2380 Perchtoldsdorf") == "Perchtoldsdorf"
    assert ort("ASKÖ-Stadion Graz-Eggenberg, Schloßstraße 20, 8020 Graz") == "Graz"
    assert ort('Hotel "Fischer am See", Heiterwang 6611') == "Heiterwang"
    assert ort("Wien-Donauinsel") == "Wien"
    assert ort("Sportplatz Weißenbach am Lech") == "Weißenbach am Lech"
    assert ort("Bad Ischl") == "Bad Ischl"
    import endure_scraper as en
    assert en.parse_zeitraum("5. April 2026") == ("2026-04-05", "2026-04-05")
    assert en.parse_zeitraum("18.–20. Juni 2026") == ("2026-06-18", "2026-06-20")
    assert en.parse_zeitraum("30. Mai – 1. Juni 2026") == ("2026-05-30", "2026-06-01")
    # Nur km mit Streckenwort davor sind weitere Strecken - "5 km vom Bahnhof" nicht.
    assert en.strecken_aus_text("Langstrecke 297 km / 2.500 Hm; Light-Variante 145 km. Start 5 km vom Bahnhof") == [297.0, 145.0]
    import oelv_scraper as oe
    assert oe.parse_ort("Ames (ESP)") == ("Ames", "ESP")
    assert oe.parse_ort("Pergine Valsugana (ITA)") == ("Pergine Valsugana", "Italien")
    assert oe.parse_ort("Ebreichsdorf") == ("Ebreichsdorf", "Österreich")
    import nws_scraper as nws
    assert nws.parse_ort("D-Schönau") == ("Schönau", "Deutschland")
    assert nws.parse_ort("Liestal BL") == ("Liestal", "Schweiz")
    assert nws.parse_ort("Basel (Staffellauf)") == ("Basel", "Schweiz")
    assert nws.schoener_name("20. BELCHEN-BERGLAUF") == "20. Belchen-Berglauf"
    assert nws.schoener_name("16. Muttenz Marathon") == "16. Muttenz Marathon"
    import lauftermine_scraper as lt
    kal = ('id=2;mo("Januar 2026","janvier 2026","gennaio 2026","january 2026")\n'
           'v(1,"www.neujahrsmarathon.ch","Neujahrsmarathon Schlieren","42.2 / 18 / 12 / 6","","ZH","")\n'
           '//v(7,"www.x.ch","Abgesagt Irgendwo","10","","ZH","")\n'
           'mo("Dezember","d","d","d")\n'
           'v(31,"www.stauseelauf.ch","Gippinger Stauseelauf","7.53","","AG","")\n'
           'mo("Januar","j","j","j")\n'
           'v(1,"www.neujahrsmarathon.ch","Neujahrsmarathon Schlieren","42.2","","ZH","")\n')
    e = lt.parse_kalender(kal)
    assert [x["datum"] for x in e] == ["2026-01-01", "2026-12-31", "2027-01-01"], e
    assert e[0]["distanzen"] == ["42.2", "18", "12", "6"] and e[0]["kanton"] == "ZH"
    # Der Ort steht im Namen - oder gar nicht (dann kein Eintrag, kein Raten).
    assert lt.ort_aus_name("Neujahrsmarathon Schlieren") == "Schlieren"
    assert lt.ort_aus_name("La Trotteuse-Tissot La Chaux-de-Fonds") == "La Chaux-de-Fonds"
    assert lt.ort_aus_name("Coupe du Vignoble Cortaillod 3/4") == "Cortaillod"
    assert lt.ort_aus_name("Zürcher Silvesterlauf") is None
    assert lt.ort_aus_name("Corrida Bulloise") is None
    assert lt.ort_aus_name("Bierathlon") is None
    # Die neuen Kalender gelten als Portal, bis eine Veranstalterseite bekannt ist.
    from scraper_lib import is_portal_link
    assert is_portal_link("https://events.endure-cycling.com/events/x/")
    assert is_portal_link("https://oelv.athmin.at/event-details.aspx?event=1")
    assert is_portal_link("https://radsport-events.de/x")
    assert not is_portal_link("https://rundumdenharz.com")
    print("  ✓ neue Quellen: Orte, Zeiträume, Kalender-Skript, Portallinks")


def test_turbosport() -> None:
    """Der Scraper für turbo-sport.eu / BRV Timing (24.09.2026): Navigation,
    Rennseite mit Ausschreibungstabelle, nur offene Klassen, Ort aus Name,
    Pfad oder Hostname - ohne Netz."""
    print("\nturbo-sport.eu (BRV Timing):")
    import turbosport_scraper as ts
    nav = ('<nav><ul class="subnav-nav"><li><a href="/veranstaltungen/donnerstagsrennen" class="subnav-link" '
           'title="2026 Donnerstagsrennen Serie">x</a></li><li><a href="/events/augsburg" class="subnav-link" '
           'title="26.09.26 Obergünzburger RR">x</a></li><li><a href="/events/augsburg" class="subnav-link" '
           'title="26.09.26 Obergünzburger RR">x</a></li></ul></nav>')
    liste = ts.parse_events_liste(nav)
    assert [e["datum"] for e in liste] == [None, "2026-09-26"], liste
    assert liste[1]["url"] == "https://turbo-sport.eu/events/augsburg"
    seite_html = (
        '<main><h2 class="element-header"><span>Großer Fritz Neuser Preis am Samstag, 3. Oktober 2026</span></h2>'
        '<p><a href="https://www.rad-net.de/x">Link zur Ausschreibung auf rad-net.de</a></p>'
        '<p><a href="https://rcherpersdorf.de/veranstaltungen/stadtparkrennen-schwabach/">Link zur Event Website</a></p>'
        '<p><a href="https://www.datasport.de/anmeldeservice/x">Link zur Anmeldung bei Datasport Germany</a></p>'
        '<table><tr><th>Kategorie:</th><th>JG von:</th><th>JG bis:</th><th>Wettbewerb:</th><th>Runden:</th>'
        '<th>Distanz:</th><th>Nenngeld:</th><th>Startzeit:</th></tr>'
        '<tr><td>R.4.1 | 4.5 Amateure</td><td>1987</td><td>2007</td><td>Lizenzklasse</td><td>40</td><td>60 km</td><td>21,00 EUR</td><td>12:45:00</td></tr>'
        '<tr><td>R.6.1 | Jedermann (männlich)</td><td>vor</td><td>2008</td><td>Jedermann</td><td>20</td><td>30 km</td><td>25,00 EUR</td><td>9:00:00</td></tr>'
        '<tr><td>R.6.2 | Jedermann (weiblich)</td><td>vor</td><td>2008</td><td>Jedermann</td><td>20</td><td>30 km</td><td>25,00 EUR</td><td>9:00:00</td></tr>'
        '<tr><td>R.1.1 | 4.27 Hobby Männer</td><td>vor</td><td>2009</td><td>Hobbyklasse</td><td>15</td><td>20.2 km</td><td>18,00 EUR</td><td>18:02:55</td></tr>'
        '</table></main>')
    seite = ts.parse_rennseite(seite_html)
    assert seite["name"] == "Großer Fritz Neuser Preis" and seite["datum"] == "2026-10-03", seite
    assert seite["veranstalter_url"] == "https://rcherpersdorf.de/veranstaltungen/stadtparkrennen-schwabach/"
    # Seit dem 24.09.2026 kommen die Lizenzklassen mit ("Bitte auch die mit BDR Lizenz aufnehmen").
    assert seite["lizenzklassen"] == 1 and [k["km"] for k in seite["klassen"]] == [60.0, 30.0, 30.0, 20.2]
    orte = {"schwabach": [("Schwabach", 49.331, 11.024)], "prien am chiemsee": [("Prien am Chiemsee", 47.856, 12.346)],
            "obergünzburg": [("Obergünzburg", 47.846, 10.418)], "burggen": [("Burggen", 47.777, 10.817)],
            "schönberg": [("Schönberg", 1.0, 1.0), ("Schönberg", 2.0, 2.0)], "landshut": [("Landshut", 48.538, 12.146)]}
    # Der Name nennt keinen Ort - der Pfad der Rennseite tut es.
    evs, grund = ts.events_aus_rennseite(seite, "https://turbo-sport.eu/events/schwabacher-stadtparkrennen", None, ts.CONFIG, orte)
    assert grund is None and len(evs) == 3, (grund, evs)
    assert {(e.standort, e.laenge_km, e.wettbewerb, e.art2) for e in evs} == {
        ("Schwabach", 60.0, "Lizenzklasse 60 km", "Straße"),
        ("Schwabach", 30.0, "Jedermann 30 km", "Straße"), ("Schwabach", 20.2, "Hobbyklasse 20.2 km", "Straße")}, evs
    assert evs[0].lat == 49.331 and evs[0].datum_start == "2026-10-03"
    # Lizenz und Jedermann über DIESELBE Distanz: eine Zeile, das Label nennt beide
    # (dedupe_key und ICS-Dateiname kennen kein Label).
    seite_gleich = dict(seite, klassen=[{"klasse": "Lizenzklasse", "kategorie": "U17", "km": 30.0},
                                        {"klasse": "Jedermann", "kategorie": "Jedermann", "km": 30.0},
                                        {"klasse": "Hobbyklasse", "kategorie": "Hobby", "km": 30.0}])
    evs_gleich, _ = ts.events_aus_rennseite(seite_gleich, "https://turbo-sport.eu/events/schwabacher-stadtparkrennen", None, ts.CONFIG, orte)
    assert [e.wettbewerb for e in evs_gleich] == ["Lizenzklasse / Jedermann / Hobbyklasse 30 km"], evs_gleich
    # Mit KLASSEN_FILTER = OFFENE_KLASSEN bleibt es beim alten Verhalten (umkehrbar).
    ts.KLASSEN_FILTER = ts.OFFENE_KLASSEN
    try:
        nur_offen = ts.parse_rennseite(seite_html)
        assert [k["km"] for k in nur_offen["klassen"]] == [30.0, 30.0, 20.2] and nur_offen["lizenzklassen"] == 1
        seite2 = dict(seite, klassen=[], lizenzklassen=3)
        assert ts.events_aus_rennseite(seite2, "https://turbo-sport.eu/events/x", None, ts.CONFIG, orte) == ([], "nur Lizenzklassen (ausgefiltert)")
    finally:
        ts.KLASSEN_FILTER = None
    # Ergebnisseite (keine Tabelle) liefert nichts.
    seite3 = dict(seite, klassen=[], lizenzklassen=0)
    assert ts.events_aus_rennseite(seite3, "https://turbo-sport.eu/events/x", None, ts.CONFIG, orte)[1].startswith("keine Ausschreibungstabelle")
    # Ort aus dem Namen: Adjektivendung, laufende Nummer, Mehrdeutigkeit, Hostname.
    assert ts.ort_aus_name("1.Obergünzburger Rundstreckenrennen", orte)[0] == "Obergünzburg"
    assert ts.ort_aus_name("37. Burggener Straßenpreis", orte)[0] == "Burggen"
    assert ts.ort_aus_name("11. Inklusiver Landshuter Straßenpreis", orte)[0] == "Landshut"
    assert ts.ort_aus_name("5. Schönberger Pfingstradrennen", orte) is None   # zwei Schönberg in Bayern
    assert ts.ort_aus_name("Kampenkönig", orte) is None
    assert ts.ort_aus_host("https://rfv-prien.de/aktivitaeten/kampenkoenig/", orte)[0] == "Prien am Chiemsee"
    assert ts.ort_aus_host("https://www.rad-net.de/x", orte) is None
    assert ts.ort_aus_slug("https://turbo-sport.eu/events/schwabacher-stadtparkrennen", orte)[0] == "Schwabach"
    # Ohne Ort kein Eintrag - geraten wird nie.
    seite4 = dict(seite, name="Kampenkönig", veranstalter_url=None)
    assert ts.events_aus_rennseite(seite4, "https://turbo-sport.eu/events/kampenkoenig", None, ts.CONFIG, orte) == ([], "Ort nicht erkennbar")
    from scraper_lib import is_portal_link
    assert is_portal_link("https://turbo-sport.eu/events/augsburg")
    print("  ✓ turbo-sport.eu: Navigation, Ausschreibungstabelle, alle Klassen (eine Zeile je Distanz), Ort aus Name/Pfad/Host")


def test_cyclingaustria() -> None:
    """Der Scraper für den ÖRV-Rennkalender (24.09.2026): Listenkarten,
    Rennseite mit Rennen-Blöcken, Distanz oder Renndauer, ausgeschlossene
    Disziplinen und Nachwuchsklassen, PLZ vor dem Ort - ohne Netz."""
    print("\ncyclingaustria.at (ÖRV-Rennkalender):")
    import cyclingaustria_scraper as ca
    liste = ('<div data-disziplin="Cycling 4 All" data-date="2026-09-26" class="filter"><a class="om_card" '
             'href="/index.php?option=com_events&amp;view=event&amp;id=66A5150B9B5D4BABB3DD1C836E78761B">'
             '<h3 class="el-title">Teufelsfahrt 2026</h3></a></div>'
             '<div data-disziplin="Enduro, Unlizenziert Kinder" data-date="2026-09-26"><a href="/index.php?option=com_events&amp;view=event&amp;id=1DB0BEFF77C644F09483583293A91C56"><h3>Wexl Kids</h3></a></div>')
    karten = ca.parse_liste(liste)
    assert [k["id"] for k in karten] == ["66A5150B9B5D4BABB3DD1C836E78761B", "1DB0BEFF77C644F09483583293A91C56"], karten
    assert karten[0]["url"].endswith("view=event&id=66A5150B9B5D4BABB3DD1C836E78761B") and karten[0]["datum"] == "2026-09-26"
    assert not karten[0]["abgesagt"]
    abgesagt = ca.parse_liste('<div data-disziplin="Cyclocross" data-date="2026-10-24"><a href="/index.php?option=com_events&amp;view=event&amp;'
                              'id=3A142E458E9042568C7C8459D4B77EC1"><h3>Wienenergie Cyclocross - ABGESAGT !!</h3></a>'
                              '<p>Ausschreibung abgesagt</p></div>')
    assert abgesagt[0]["abgesagt"], abgesagt

    def block(nr, klassen, disziplin, distanz=None, dauer=None):
        felder = f'<h3 class="uk-h5">Disziplin</h3><div class="uk-panel">{disziplin}</div>'
        if distanz:
            felder += f'<h3 class="uk-h5">Distanz</h3><div class="uk-panel">{distanz}</div>'
        if dauer:
            felder += f'<h3 class="uk-h5">Renndauer</h3><div class="uk-panel">{dauer}</div>'
        return (f'<div class="uk-card uk-card-default"><div class="uk-heading-small">Rennen {nr}</div>'
                f'<h2 class="uk-h3">{klassen}</h2>{felder}</div>')
    seite_html = (
        '<main><h1 class="uk-h2">Quer durch\'s Stadion</h1>'
        '<h2 class="uk-h5">Start/Ziel</h2><div class="uk-panel">2763 Neusiedl</div>'
        '<h2 class="uk-h5">Datum</h2><div class="uk-panel">Sa, 21. November 2026–So, 22. November 2026</div>'
        '<h2 class="uk-h5">Sparte</h2><div><ul><li><a href="/kalender?sparten=cyclocross">Cyclocross</a></li></ul></div>'
        '<h2 class="uk-h5">Disziplin</h2><div><ul><li>Unlizenziert Kinder, Cycling 4 All, Cyclocross</li></ul></div>'
        + block(1, "U11 männlich, U11 weiblich, Member Card männlich Jugend", "Touristik", dauer="00:10 (H:min)")
        + block(2, "U13 männlich, U17 weiblich", "Cyclocross", dauer="00:20 (H:min)")
        + block(3, "AMATEURE HERREN, MEN ELITE, MEN MASTERS I", "Cyclocross", dauer="00:40 (H:min)")
        + block(4, "WOMEN ELITE, WOMEN MASTERS", "Cyclocross", dauer="00:40 (H:min)")
        + block(5, "MEN ELITE", "Einzelzeitfahren", distanz="47.2 km")
        + block(6, "Sonderform", "Mannschaftszeitfahren", distanz="47.2 km")
        + '<p><strong>RC Pernitz</strong></p><a class="uk-button" target="_blank" href="https://www.radclub-pernitz.at"> zur Website</a></main>')
    seite = ca.parse_rennseite(seite_html)
    assert seite["name"] == "Quer durch's Stadion" and seite["ort"] == "2763 Neusiedl", seite
    assert (seite["datum_start"], seite["datum_ende"]) == ("2026-11-21", "2026-11-22"), seite
    assert seite["sparte"] == "Cyclocross" and seite["veranstalter_url"] == "https://www.radclub-pernitz.at"
    assert len(seite["rennen"]) == 6 and seite["rennen"][2]["Renndauer"] == "00:40 (H:min)"
    evs, grund = ca.events_aus_rennseite(seite, "https://www.cyclingaustria.at/x", None, ca.CONFIG)
    assert grund is None, grund
    # Kinder- und Jugendrennen (U11/U13/U17), das Mannschaftszeitfahren fallen;
    # 40 min zweimal ist EINE Zeile; das Einzelzeitfahren hat Kilometer.
    assert [(e.standort, e.laenge_km, e.dauer_h, e.wettbewerb, e.art2) for e in evs] == [
        ("Neusiedl", None, 0.67, "Cyclocross 40 min", "Cyclecross"),
        ("Neusiedl", 47.2, None, "Einzelzeitfahren 47.2 km", "Zeitfahren")], evs
    assert evs[0].datum_ende == "2026-11-22" and evs[0].land == "Österreich" and evs[0].art1 == "Fahrrad"
    # Ohne Rennen-Blöcke: eine Zeile ohne Maß, Disziplin als Label, PLZ weg auch bei "Sankt".
    seite2 = dict(seite, rennen=[], ort="4171 Sankt Peter am Wimberg", sparte="Straße", disziplin="Cycling 4 All",
                  veranstalter_url=None)
    evs2, _ = ca.events_aus_rennseite(seite2, "https://www.cyclingaustria.at/x", None, ca.CONFIG)
    assert [(e.standort, e.laenge_km, e.wettbewerb, e.art2, e.veranstalter_url) for e in evs2] == [
        ("Sankt Peter am Wimberg", None, "Cycling 4 All", "Straße", "https://www.cyclingaustria.at/x")], evs2
    # Nur ausgeschlossene Disziplinen -> kein Eintrag, mit Grund.
    seite3 = dict(seite, rennen=[], disziplin="Enduro, Unlizenziert Kinder")
    assert ca.events_aus_rennseite(seite3, "u", None, ca.CONFIG)[1].startswith("Disziplin nicht übernommen")
    seite4 = dict(seite, rennen=[], disziplin="Kunstrad 1er, Radball")
    assert ca.events_aus_rennseite(seite4, "u", None, ca.CONFIG)[0] == []
    seite5 = dict(seite, rennen=[], name="Youngster Grand Prix Windhaag", disziplin="Cross-Country")
    assert ca.events_aus_rennseite(seite5, "u", None, ca.CONFIG) == ([], "Nachwuchsrennen (laut Name)")
    # "https://-" ist kein Veranstalterlink; die Verbandsseite auch nicht.
    kein = seite_html.replace("https://www.radclub-pernitz.at", "https://-")
    assert ca.parse_rennseite(kein)["veranstalter_url"] is None
    kein = seite_html.replace("https://www.radclub-pernitz.at", "https://www.cyclingaustria.at")
    assert ca.parse_rennseite(kein)["veranstalter_url"] is None
    from scraper_lib import is_portal_link
    assert is_portal_link("https://www.cyclingaustria.at/index.php?option=com_events&view=event&id=1")
    print("  ✓ cyclingaustria.at: Karten, Rennseite, Rennen-Blöcke, Ausschlüsse, Ort ohne PLZ")


def test_swimsports() -> None:
    """Der Scraper für den Schwimmkalender Schweiz (24.09.2026): Liste,
    Detailfelder, Distanzen ohne SUP, Absagen, Kantonskürzel - ohne Netz."""
    print("\nswimsports.ch (Schwimmkalender Schweiz):")
    import swimsports_scraper as sw
    liste = ('<article class="swimming-calendar"><h2><a href="/node/2300" hreflang="de">56. Hallwilerseeschwimmen</a></h2>'
             '<time datetime="2026-09-05T12:00:00Z">05.09.2026</time>'
             '<div class="field field--name-field-calendar-location field--item">Beinwil am See</div></article>'
             '<article class="swimming-calendar"><h2><a href="/node/2063">Escape Leman</a></h2></article>'
             '<article class="swimming-calendar"><h2><a href="/node/2076" hreflang="de">Rhyschwümme 2026 - Abgesagt - Das nächste Schwimmen findet am 08.08.2027 statt</a></h2>'
             '<time datetime="2027-08-08T12:00:00Z">08.08.2027</time>'
             '<div class="field field--name-field-calendar-location field--item">Stein am Rhein</div></article>')
    eintraege = sw.parse_liste(liste)
    assert [(e["datum"], e["ort"]) for e in eintraege] == [("2026-09-05", "Beinwil am See"), ("2027-08-08", "Stein am Rhein")], eintraege
    detail = ('<article class="swimming-calendar full"><div class="field field--name-field-calendar-location field--item">Beinwil am See</div>'
              '<div class="field field--name-field-distance"><div class="field--label">Distanz</div><div class="field--item">2 km</div></div>'
              '<div class="field field--name-field-water"><div class="field--item">Hallwilersee</div></div>'
              '<div class="field field--name-field-home-page"><div class="field--item"><a href="https://slrghallwilersee.ch/hallwilerseeschwimmen/">Homepage</a></div></div></article>')
    d = sw.parse_detail(detail)
    assert d["distanz"] == "2 km" and d["homepage"] == "https://slrghallwilersee.ch/hallwilerseeschwimmen/" and d["gewaesser"] == "Hallwilersee", d
    evs, grund = sw.events_aus_eintrag(eintraege[0], d, sw.CONFIG)
    assert grund is None and [(e.name, e.standort, e.land, e.art1, e.art2, e.laenge_km, e.wettbewerb, e.veranstalter_url) for e in evs] == [
        ("56. Hallwilerseeschwimmen", "Beinwil am See", "Schweiz", "Schwimmen", "Freiwasser", 2.0, None, "https://slrghallwilersee.ch/hallwilerseeschwimmen/")], evs
    # Distanzen: SUP zählt nicht, mehrere Angaben sind mehrere Zeilen, "Libre" keine.
    assert sw.distanzen_aus_text("Schwimmen 1.1 km/SUP Fun 1.5 km/SUP Longdistance 8.0 km") == [1.1]
    assert sw.distanzen_aus_text("500 m, 1500 m oder 3000 m") == [0.5, 1.5, 3.0]
    assert sw.distanzen_aus_text("Libre") == [] and sw.distanzen_aus_text("Ca. 11 km, Schwimmdauer ca. 1 ½ Stunden.") == [11.0]
    assert sw.distanzen_aus_text("111 Meter oder länger") == [0.11]
    evs3, _ = sw.events_aus_eintrag(eintraege[0], dict(d, distanz="1 km / 2.5 km"), sw.CONFIG)
    assert [(e.laenge_km, e.wettbewerb) for e in evs3] == [(1.0, "1 km"), (2.5, "2.5 km")], evs3
    # Absage: weg - außer der Name nennt den eigenen (nächsten) Termin.
    assert sw.bereinige_name("Zuger-Seeüberquerung - Abgesagt", "2026-08-23") is None
    assert sw.bereinige_name(eintraege[1]["name"], "2027-08-08") == "Rhyschwümme 2026"
    evs2, _ = sw.events_aus_eintrag(eintraege[1], dict(d, ort=None, distanz="Ca. 11 km", homepage=None), sw.CONFIG)
    assert evs2[0].standort == "Stein am Rhein" and evs2[0].laenge_km == 11.0 and evs2[0].veranstalter_url == "https://www.swimsports.ch/node/2076"
    # Kantonskürzel und Zusätze fallen aus dem Ort.
    for text, ort in [("Greppen/LU", "Greppen"), ("Rapperswil SG", "Rapperswil"), ("Zürich, Strandbad Mythenquai", "Zürich"),
                      ("Küsnacht - Kilchberg", "Küsnacht")]:
        e, _ = sw.events_aus_eintrag(dict(eintraege[0], ort=text), dict(d, ort=None), sw.CONFIG)
        assert e[0].standort == ort, (text, e[0].standort)
    from scraper_lib import is_portal_link
    assert is_portal_link("https://www.swimsports.ch/node/2300")
    print("  ✓ swimsports.ch: Liste, Detailfelder, Distanzen ohne SUP, Absage, Kanton")


def test_fsieben() -> None:
    """Der Scraper für den Triathlon-Kalender Österreich von fsieben.at
    (24.09.2026): JSON in der Seite, nur datierte Bewerbe, ein Eintrag je
    Format, Ort über places.json je Bundesland, unbestätigt = vorläufig."""
    print("\nfsieben.at (Triathlon-Kalender Österreich):")
    import fsieben_scraper as fs
    html = ('<pre id="f7-ev-data" style="display:none">[{"n":"18. Mostiman Triathlon","s":"2027","d":"Samstag, 10. Juli 2027",'
            '"o":"Wallsee an der Donau","b":"Niederösterreich","t":["Olympisch","Sprint","Kids","Staffel"],"u":"https://mosti-man.at/",'
            '"st":"bestätigt","so":"2027-07-10"},'
            '{"n":"ALOHA TRI Linz","d":"Sonntag, 4. Juli 2027","o":"Pleschingersee, Linz","b":"Oberösterreich",'
            '"t":["Olympisch","Cross","Duathlon"],"u":"https://alohasport.at/event/aloha-tri-linz","st":"unbestätigt","so":"2027-07-04"},'
            '{"n":"40. Neufeld Triathlon","d":"19.–20. Juni 2027","o":"Neufeld an der Leitha","b":"Burgenland","t":["Sprint"],"u":"","st":"bestätigt","so":"2027-06-19"},'
            '{"n":"11. Vienna Triathlon","d":"Termin folgt","o":"Wien, Donauinsel","b":"Wien","t":["Sprint"],"u":"https://www.vienna-triathlon.at/","st":"erwartet","so":"9999-12-31"},'
            '{"n":"Bundesschulmeisterschaften Aquathlon","d":"23.–25. Juni 2027","o":"Ort noch offen","b":"Oberösterreich","t":["Aquathlon"],"u":"","st":"bestätigt (Ort offen)","so":"2027-06-23"}]</pre>')
    arr = fs.parse_kalender(html)
    assert len(arr) == 5
    evs, grund = fs.events_aus_eintrag(arr[0])
    assert grund is None and [(e.standort, e.lat, e.wettbewerb, e.art2, e.datum_vorlaeufig) for e in evs] == [
        ("Wallsee", 48.167, "Olympisch", "Straße", None), ("Wallsee", 48.167, "Sprint", "Straße", None)], evs
    assert evs[0].art1 == "Triathlon" and evs[0].land == "Österreich" and evs[0].veranstalter_url == "https://mosti-man.at/"
    # "Pleschingersee, Linz": der See ist kein Ort, Linz schon - und mit dem
    # Bundesland eindeutig (es gibt auch ein Linz in Kärnten). Unbestätigt = vorläufig.
    evs, _ = fs.events_aus_eintrag(arr[1])
    assert [(e.standort, e.wettbewerb, e.art2, e.datum_vorlaeufig) for e in evs] == [
        ("Linz", "Olympisch", "Straße", True), ("Linz", "Cross-Triathlon", "Cross", True), ("Linz", "Duathlon", "Duathlon", True)], evs
    assert evs[0].lat == 48.306, evs[0].lat
    # Zwei Tage, ohne Event-Seite -> Portallink der Liste.
    evs, _ = fs.events_aus_eintrag(arr[2])
    assert (evs[0].datum_start, evs[0].datum_ende, evs[0].veranstalter_url) == ("2027-06-19", "2027-06-20", fs.LISTE_URL), evs
    # "Termin folgt" ohne "zuletzt" - kein Eintrag; Schulmeisterschaft nicht offen.
    assert fs.events_aus_eintrag(arr[3]) == ([], "Termin folgt")
    # Mit "zuletzt" (08.10.2026, vom Nutzer gewünscht): derselbe Wochentag ein
    # Jahr später, vorläufig UND gedaechtnis, Auflage weg, Dauer bleibt.
    from datetime import date as _d
    heute = _d(2026, 10, 8)
    check("fsieben: Vorjahrestermin -> gleicher Wochentag (Sa 22.08.2026 -> Sa 21.08.2027)",
          fs.naechster_termin("zuletzt 22.08.2026", heute), ("2027-08-21", "2027-08-21"))
    check("fsieben: mehrtägig behält die Dauer, auch über den Jahreswechsel",
          (fs.naechster_termin("zuletzt 12.–13.09.2026", heute), fs.naechster_termin("zuletzt 03.01.2026", heute)),
          (("2027-09-11", "2027-09-12"), ("2027-01-02", "2027-01-02")))
    check("fsieben: Schätzung in der Vergangenheit oder ohne Datum - nichts",
          (fs.naechster_termin("zuletzt 01.01.2025", heute), fs.naechster_termin("", heute)), (None, None))
    progn = dict(arr[3], hi="zuletzt 12.–13.09.2026")
    evs, grund = fs.events_aus_eintrag(progn, heute=heute)
    check("fsieben: geschätzter Termin, Name ohne Auflage, markiert",
          [(e.name, e.datum_start, e.datum_ende, e.datum_vorlaeufig, e.gedaechtnis, e.wettbewerb) for e in evs],
          [("Vienna Triathlon", "2027-09-11", "2027-09-12", True, True, "Sprint")])
    check("fsieben: nur Kids/Staffel oder „Kids“ im Namen - keine Schätzung",
          (fs.events_aus_eintrag(dict(progn, t=["Kids", "Staffel"]), heute=heute)[1],
           fs.events_aus_eintrag(dict(progn, n="TRI KIDS Aquathlon"), heute=heute)[1]), ("Termin folgt", "Termin folgt"))
    import clean_events as ce
    schaetzung = evs[0].to_dict()
    echt = dict(schaetzung, name="12. Vienna Triathlon", datum_start="2027-09-18", datum_ende="2027-09-19")
    echt.pop("gedaechtnis"); echt.pop("datum_vorlaeufig")
    check("zurueckziehen_prognosen: echte Ausgabe da -> Schätzung weg",
          [e["datum_start"] for e in ce.zurueckziehen_prognosen([schaetzung, echt], [])[0]], ["2027-09-18"])
    check("zurueckziehen_prognosen: ohne echte Ausgabe, oder anderer Ort, bleibt sie",
          (len(ce.zurueckziehen_prognosen([schaetzung], [])[0]),
           len(ce.zurueckziehen_prognosen([schaetzung, dict(echt, standort="Graz", lat=47.07, lon=15.44)], [])[0])), (1, 2))
    gerasdorf = dict(echt, name="Gerasdorf Triathlon", standort="Gerasdorf bei Wien", lat=48.29, lon=16.47, datum_start="2027-08-08")
    check("zurueckziehen_prognosen: „Triathlon“ allein ist kein gemeinsamer Name (Vienna ≠ Gerasdorf bei Wien)",
          len(ce.zurueckziehen_prognosen([schaetzung, gerasdorf], [])[0]), 2)
    check("zurueckziehen_prognosen: „Int.“ = „Internationaler“, Ort im Namen zählt mit",
          ce._gleicher_name_kern({"name": "Int. Thiersee Triathlon", "standort": "Thiersee"},
                                 {"name": "Internationaler Thiersee Triathlon", "standort": "Thiersee"}), True)
    check("zurueckziehen_prognosen: Kopie aus manual_events.json bleibt der anderen Regel",
          len(ce.zurueckziehen_prognosen([schaetzung, echt], [dict(schaetzung)])[0]), 2)
    assert fs.events_aus_eintrag(arr[4]) == ([], "nicht für jeden offen")
    from scraper_lib import is_portal_link
    assert is_portal_link("https://www.fsieben.at/tools/triathlon-kalender-oesterreich/") and is_portal_link("https://www.triathlon-austria.at/de/service-termine")
    print("  ✓ fsieben.at: JSON, datierte Bewerbe, Formate, Ort je Bundesland, vorläufig")


def test_datasport() -> None:
    """Der Scraper für datasport.com (30.09.2026): sitemap als Liste,
    das Ausgabe-Objekt aus den Next.js-Flight-Daten, je Wettbewerb eine
    Zeile, Mehrsport-Summe, Ausschlüsse (Nachwuchs, Team, E-Bike, fremde
    Sportart), Name ohne Jahr, Veranstalterlink mit Schema."""
    print("\ndatasport.com:")
    import json as _json
    import datasport_scraper as dsp
    from scraper_lib import is_portal_link

    sitemap = ("<urlset><url><loc>https://datasport.com/de/events/morat-fribourg-2026</loc></url>"
               "<url><loc>https://datasport.com/en/events/morat-fribourg-2026</loc></url>"
               "<url><loc>https://datasport.com/de/results/morat-fribourg-2026</loc></url>"
               "<url><loc>https://datasport.com/de/events/gurtenclassic-2026</loc></url>"
               "<url><loc>https://datasport.com/de/events/test-flag</loc></url>"
               "<url><loc>https://datasport.com/de/events/gurtenclassic-2025</loc></url></urlset>")
    assert dsp.slugs_aus_sitemap(sitemap) == ["morat-fribourg-2026", "gurtenclassic-2026", "test-flag", "gurtenclassic-2025"]
    assert dsp.slug_jahr("gurtenclassic-2026") == 2026 and dsp.slug_jahr("bike-revolution-gruyere") is None
    assert dsp.slug_jahr("les-5km-de-terre-des-hommes-suisse2025") == 2025

    def contest(name, sport, disziplin, sections, ages=((18, 100),), tag="2026-10-25"):
        return {"id": 1, "name": name, "shortName": name[:6],
                "discipline": {"id": 3, "name": {"de": disziplin}, "sportType": {"id": 1, "name": {"de": sport}}},
                "from": f"{tag}T00:00:00", "to": f"{tag}T23:59:59",
                "categories": [{"id": i, "name": f"K{i}", "ageFrom": a, "ageTo": b} for i, (a, b) in enumerate(ages)],
                "sections": [{"id": i, "name": n, "distance": m} for i, (n, m) in enumerate(sections)]}

    obj = {"id": 3638, "name": "gurtenCLASSIC 2026", "slug": "gurtenclassic-2026",
           "editionFrom": "2026-10-25", "editionTo": "2026-10-25",
           "organizerLink": "www.gurtenclassic.ch/", "cancelled": False, "visible": True,
           "regions": [{"id": 7, "countryCode": "CHE", "name": {"de": "Kanton Bern"}}],
           "town": {"de": "Wabern", "fr": "Wabern"},
           "preciseLocation": {"latitude": 46.919043, "longitude": 7.440898},
           "contests": [
               contest("Rennvelo-Rundfahrt", "Rennrad", "Strassenrennen", [("Rennvelo-Rundfahrt", 40000)]),
               contest("Classic Run", "Lauf", "Volkslauf", [("Classic Run", 15000)]),
               contest("Gurten Kids Run - 2 Runden", "Lauf", "Volkslauf", [("Kids Run", 1800)], ages=((10, 11),)),
               contest("Firmenstaffel", "Lauf", "Volkslauf", [("", 15000)]),
               contest("Classic Walking/Nordic Walking", "Lauf", "Walking / Nordic Walking", [("", 10000)]),
               contest("49km PRINCEClassic E-Bike", "MTB", "Bike Marathon", [("49ebike", 49000)]),
               contest("Short Distance", "Triathlon", "Andere", [("Swim", 500), ("Bike", 17000), ("Run", 5400)]),
               contest("Marathon", "Langlauf", "Freie Technik (Skating)", [("Marathon", 42000)]),
               contest("12 km Nordic Walking und Walking", "Lauf", "Walking / Nordic Walking", [("", None)]),
               contest("6 Stunden Lauf", "Lauf", "Volkslauf", [("Runde", 1000)]),
               contest("Bergtrail", "Lauf", "Trail Run", [("", 12000)], tag="2026-10-24"),
               contest("Jedermann", "MTB", "Cross Country", [("", 30000)], ages=((16, None),)),
               contest("Classic Run direkt", "Lauf", "Volkslauf", [("", 15000)]),
           ]}
    chunk = _json.dumps({"queries": [{"state": {"data": obj}}]})
    html = ('<html><script>self.__next_f.push([1,"a:[1,2]"])</script>'
            '<script>self.__next_f.push([1,' + _json.dumps("2b:" + chunk) + '])</script></html>')
    gelesen = dsp.event_objekt(html, "gurtenclassic-2026")
    assert gelesen and gelesen["slug"] == "gurtenclassic-2026" and len(gelesen["contests"]) == 13
    assert dsp.event_objekt("<html><script>self.__next_f.push([1,\"x\"])</script></html>") is None

    bericht: dict = {}
    evs = dsp.events_aus_objekt(gelesen, dsp.CONFIG, bericht)
    zeilen = [(e.art1, e.art2, e.laenge_km, e.dauer_h, e.wettbewerb, e.datum_start) for e in evs]
    assert zeilen == [
        ("Fahrrad", "Straße", 40.0, None, "Rennvelo-Rundfahrt 40 km", "2026-10-25"),
        ("Laufen", "Straße", 15.0, None, "Classic Run 15 km", "2026-10-25"),
        ("Laufen", "Marsch", 10.0, None, "Classic Walking/Nordic Walking 10 km", "2026-10-25"),
        ("Triathlon", "Straße", 22.9, None,
         "Short Distance 22.9 km (0.5 km Schwimmen / 17 km Rad / 5.4 km Laufen)", "2026-10-25"),
        ("Laufen", "Marsch", 12.0, None, "12 km Nordic Walking und Walking", "2026-10-25"),
        ("Laufen", "Straße", None, 6.0, "6 Stunden Lauf", "2026-10-25"),
        ("Laufen", "Trail", 12.0, None, "Bergtrail 12 km", "2026-10-24"),
        ("Fahrrad", "Mountainbike", 30.0, None, "Jedermann 30 km", "2026-10-25"),
    ], zeilen
    assert bericht == {"Nachwuchs": 1, "Team/Staffel": 1, "E-Bike": 1, "Sportart langlauf": 1,
                       "gleiche Strecke doppelt": 1}, bericht
    e0 = evs[0]
    assert (e0.name, e0.standort, e0.land, e0.lat, e0.lon, e0.veranstalter_url) == (
        "gurtenCLASSIC", "Wabern", "Schweiz", 46.919043, 7.440898, "https://www.gurtenclassic.ch/"), e0
    # Abgesagt, ohne Ort, Italien (Südtirol entscheidet filter_dach über die Koordinaten).
    assert dsp.events_aus_objekt(dict(obj, cancelled=True), dsp.CONFIG) == []
    assert dsp.events_aus_objekt(dict(obj, town=None), dsp.CONFIG) == []
    bozen = {"latitude": 46.498, "longitude": 11.354}
    it = dsp.events_aus_objekt(dict(obj, regions=[{"countryCode": "ITA", "name": {"de": "Trentino-Südtirol"}}],
                                    preciseLocation=bozen), dsp.CONFIG)
    assert it and it[0].land == "Italien"
    # Italien außerhalb Südtirols (Riva del Garda), fremdes Land, Serie an mehreren Orten, Anmeldearten.
    b2: dict = {}
    assert dsp.events_aus_objekt(dict(obj, regions=[{"countryCode": "ITA", "name": {"de": "Trentino"}}],
                                      preciseLocation={"latitude": 45.886, "longitude": 10.841}), dsp.CONFIG, b2) == []
    assert dsp.events_aus_objekt(dict(obj, regions=[{"countryCode": "LIE", "name": {"de": "Liechtenstein"}}]), dsp.CONFIG, b2) == []
    assert dsp.events_aus_objekt(dict(obj, town={"de": "Andorra / Selva Val Gardena / Riva del Garda"}), dsp.CONFIG, b2) == []
    assert b2 == {"Italien außerhalb Südtirols": 1, "Land LIE": 1, "Serie an mehreren Orten": 1}, b2
    warte = dsp.events_aus_objekt(dict(obj, contests=[
        contest("Warte Liste Sprint Einzel", "Triathlon", "Andere", []),
        contest("Pre-iscrizione / pre-registration / Voranmeldung", "Rennrad", "Strassenrennen", []),
        contest("Triathlon Stafette", "Triathlon", "Andere", [("Swim", 1500), ("Bike", 38200), ("Run", 10500)]),
        contest("1 Erw. & 1 Kind U6 weiblich", "Lauf", "Volkslauf", [("", 1000)], ages=((1, 100),)),
    ]), dsp.CONFIG, b2)
    assert warte == [] and b2["ohne Strecke (Anmeldeart oder ohne Namen)"] == 2 and b2["Team/Staffel"] == 1 and b2["Nachwuchs"] == 1, b2
    assert dsp.events_aus_objekt(dict(obj, town={"de": "Badia (BZ)"}), dsp.CONFIG)[0].standort == "Badia"
    # Ohne organizerLink bleibt die datasport-Seite als Portallink stehen.
    ohne = dsp.events_aus_objekt(dict(obj, organizerLink=""), dsp.CONFIG)
    assert ohne[0].veranstalter_url == "https://datasport.com/de/events/gurtenclassic-2026"
    assert is_portal_link(ohne[0].veranstalter_url)
    assert dsp.veranstalter_link("http://www.triathlon-frauenfeld.ch", "x") == "http://www.triathlon-frauenfeld.ch"
    assert dsp.veranstalter_link("-", "x") == "https://datasport.com/de/events/x"
    # Name ohne Jahr, Nachwuchs über Name oder Altersklassen.
    assert dsp.name_ohne_jahr("Claro Pizzo 2026 - 13° Edizione") == "Claro Pizzo - 13° Edizione"
    assert dsp.name_ohne_jahr("2027 HERO UCI Marathon World Cup") == "HERO UCI Marathon World Cup"
    assert dsp.ist_nachwuchs({"name": "U20 Junior Trail Series", "categories": [{"ageFrom": 18, "ageTo": 19}]})
    assert dsp.ist_nachwuchs({"name": "Jugend", "categories": [{"ageFrom": 14, "ageTo": 17}]})
    assert not dsp.ist_nachwuchs({"name": "Classic Run", "categories": [{"ageFrom": 14, "ageTo": 15}, {"ageFrom": 16, "ageTo": None}]})
    # Ein Lauf in einer Veranstaltung mit Mehrsport-Namen bekommt den Namen
    # seines Wettbewerbs - sonst machte fix_multisport_art1() einen Triathlon daraus.
    inferno = dict(obj, name="Inferno Triathlon + Halbmarathon 2026", slug="inferno-2026", contests=[
        contest("Inferno Halbmarathon", "Lauf", "Berglauf", [("", 21097)], tag="2026-08-15"),
        contest("Trail", "Lauf", "Trail Run", [("", 8400)], tag="2026-08-16"),
        contest("Inferno Triathlon", "Triathlon", "Andere",
                [("Swim", 3100), ("Road Bike", 97000), ("Mountain Bike", 30000), ("Run", 25000)], tag="2026-08-15"),
    ])
    inf = dsp.events_aus_objekt(inferno, dsp.CONFIG)
    assert [(e.name, e.art1, e.art2, e.laenge_km) for e in inf] == [
        ("Inferno Halbmarathon", "Laufen", "Trail", 21.1), ("Inferno Trail", "Laufen", "Trail", 8.4),
        ("Inferno Triathlon + Halbmarathon", "Triathlon", "Straße", 155.1)], inf
    assert inf[2].wettbewerb == "Inferno Triathlon 155.1 km (3.1 km Schwimmen / 97 km Rad / 30 km MTB / 25 km Laufen)"
    # Duathlon/Aquathlon aus den Teilstrecken, wenn der Name nichts sagt.
    du = dsp.events_aus_objekt(dict(obj, contests=[
        contest("Kurz", "Triathlon", "Andere", [("Run", 5000), ("Bike", 20000), ("Run", 2500)]),
        contest("Sprint", "Triathlon", "Andere", [("Swim", 750), ("Run", 5000)])]), dsp.CONFIG)
    assert [(e.art2, e.laenge_km) for e in du] == [("Duathlon", 27.5), ("Aquathlon", 5.8)], du
    # Powerman Zofingen führt "Duathlon" als eigene Sportart - Mehrsport-Schublade (Datenregel 11).
    pm = dsp.events_aus_objekt(dict(obj, contests=[
        contest("Long Distance - OPEN", "Duathlon", "ITU Langdistanz", [("Run", 10000), ("Bike", 150000), ("Run", 30000)])]), dsp.CONFIG)
    assert [(e.art1, e.art2, e.laenge_km) for e in pm] == [("Triathlon", "Duathlon", 190.0)], pm
    print("  ✓ datasport.com: sitemap, Flight-Objekt, Wettbewerbe, Mehrsport-Summe, Ausschlüsse, Namen")


def test_radsportevents() -> None:
    """Der Scraper für die JSON-API von radsport-events.de (24.09.2026):
    Strecken, Rundenlängen, Dauern, Mehrsport-Summen, Ausschlüsse."""
    print("\nradsport-events.de:")
    import radsportevents_scraper as rs
    basis = {"id": 48, "title": "Spreewaldmarathon", "eventDate": "2027-04-24", "endDate": None,
             "city": "Lübbenau", "country": "DE", "latitude": 51.868, "longitude": 13.969,
             "category": "ROAD", "eventType": "RADMARATHON", "organizerUrl": "http://www.spreewaldmarathon.de/",
             "tentative": False, "cancelled": False, "status": "PUBLISHED",
             "distances": [{"distanceKm": 70.0, "durationMinutes": None, "label": None},
                           {"distanceKm": 200.0, "durationMinutes": None, "label": None}]}
    evs = rs.parse_item(basis, rs.CONFIG)
    assert [(e.laenge_km, e.wettbewerb, e.art1, e.art2, e.land, e.standort) for e in evs] == [
        (70.0, "Radmarathon 70 km", "Fahrrad", "Straße", "Deutschland", "Lübbenau"),
        (200.0, "Radmarathon 200 km", "Fahrrad", "Straße", "Deutschland", "Lübbenau")], evs
    assert evs[0].lat == 51.868 and evs[0].veranstalter_url == "http://www.spreewaldmarathon.de/"
    # Rundenlänge fällt weg, 24h-Rennen bekommt 24 h; Dauer in Minuten -> Stunden.
    e24 = rs.parse_item(dict(basis, eventType="RENNEN_24H", distances=[{"distanceKm": 17.0, "label": "Runde"}]), rs.CONFIG)
    assert len(e24) == 1 and e24[0].laenge_km is None and e24[0].dauer_h == 24.0 and e24[0].wettbewerb == "24-Stunden-Rennen 24 h", e24
    kurs = rs.parse_item(dict(basis, title="MTB-Fahrtechnikkurs \"Allround\", Nürnberg #5", category="MTB", eventType="MTB"), rs.CONFIG)
    check("radsport-events: ein Fahrtechnikkurs ist kein Rennen", kurs, [])
    ecx = rs.parse_item(dict(basis, category="MTB", eventType="CYCLOCROSS", distances=[{"distanceKm": 0.0, "durationMinutes": 40}]), rs.CONFIG)
    assert ecx[0].art2 == "Cyclecross" and ecx[0].dauer_h == 0.67 and ecx[0].wettbewerb == "Cyclocross 40 min" and ecx[0].laenge_km is None, ecx
    # Zeitfahren, Gravel, vorläufiger Termin, Portallink ohne organizerUrl.
    ezf = rs.parse_item(dict(basis, eventType="BERGZEITFAHREN", tentative=True, organizerUrl=None,
                             distances=[{"distanceKm": 10.0}]), rs.CONFIG)
    assert ezf[0].art2 == "Zeitfahren" and ezf[0].datum_vorlaeufig is True
    assert ezf[0].veranstalter_url == "https://radsport-events.de/events/48"
    assert rs.parse_item(dict(basis, category="GRAVEL", eventType="GRAVEL_RACE"), rs.CONFIG)[0].art2 == "Gravel"
    # Mehrsport: Summe nur bei beschrifteten Teilstrecken, sonst ohne Länge.
    edu = rs.parse_item(dict(basis, category="MTB", eventType="DUATHLON_CROSS", distances=[
        {"distanceKm": 4.6, "label": "Lauf"}, {"distanceKm": 20.0, "label": "Rad"}, {"distanceKm": 5.0, "label": "Lauf"}]), rs.CONFIG)
    assert len(edu) == 1 and edu[0].art1 == "Triathlon" and edu[0].art2 == "Duathlon" and edu[0].laenge_km == 29.6, edu
    assert edu[0].wettbewerb.startswith("Cross-Duathlon 29.6 km (")
    edu2 = rs.parse_item(dict(basis, eventType="DUATHLON", distances=[{"distanceKm": 3.0}, {"distanceKm": 14.0}]), rs.CONFIG)
    assert len(edu2) == 1 and edu2[0].laenge_km is None and edu2[0].wettbewerb == "Duathlon", edu2
    # Ausschlüsse: abgesagt, virtuell, Camp, Mannschaftszeitfahren, Italien außerhalb Südtirols.
    assert rs.parse_item(dict(basis, cancelled=True), rs.CONFIG) == []
    for typ in ("VIRTUELLE_RTF", "GRAVEL_CAMP", "MANNSCHAFTSZEITFAHREN", "ETAPPENTOUR"):
        assert rs.parse_item(dict(basis, eventType=typ), rs.CONFIG) == [], typ
    assert rs.parse_item(dict(basis, country="IT", latitude=45.69, longitude=9.67), rs.CONFIG) == []   # Bergamo
    bozen = rs.parse_item(dict(basis, country="IT", latitude=46.498, longitude=11.354), rs.CONFIG)   # Bozen
    assert bozen and bozen[0].land == "Italien", bozen  # Südtirol setzt filter_dach() über die Koordinaten
    from scraper_lib import is_portal_link
    assert is_portal_link("https://radsport-events.de/events/48")
    print("  ✓ radsport-events.de: Strecken, Runde/Dauer, Zeitfahren/Gravel/Cyclocross, Mehrsport-Summe, Ausschlüsse")


def test_kleine_radquellen() -> None:
    """Die acht Scraper der zweiten Quellensuche (30.09.2026): Stevens Cup,
    Cyclocross Cup BW, GUNSHA Crosscup, DS Ergebnisdienst, Weser-Ems-Cup,
    Swiss Cycling BOE, Fricktaler Cup, Radsportverband SH - samt
    `ort_im_text()`, ohne Netz."""
    print("\nKleine Radquellen (zweite Suche):")

    def ok(label, bedingung, detail=""):
        check(label if bedingung else f"{label} [{detail}]", bool(bedingung), True)

    from scraper_lib import ort_im_text, orte_aus_places
    de = orte_aus_places("Deutschland")
    be = orte_aus_places("Schweiz", region="Canton de Berne")
    # ort_im_text: Zweiwort-Ort vor Einwort, Adjektiv nur vor Rennwort,
    # Artikel sperrt das Hauptwort, Bindestrich nur hinter Rennwort,
    # mehrdeutig nur mit Flagge.
    faelle = [
        ("WEC Cross Bad Essen / Pr. Oldendorf", de, "Bad Essen"),
        ("Herforder Frühjahrspreis", de, "Herford"),
        ("Finale BL Rund um das Vechtaer Reiterwaldstadion", de, "Vechta"),
        ("Kometen Schmitter Nacht", de, None),
        ("NRW Cross Cup Die Mauer von Kendenich", de, None),
        ("38. Main-Spessart Rundfahrt", de, None),
        ("35. GP von Buchholz", de, None),
        ("Bergrennen Stettlen-Bantiger", be, "Stettlen"),
        ("Einzelzeitfahren Langnau-Zäziwil", be, "Zäziwil"),
        ("Strassenrennen in Deisswil b. Münchenbuchsee", be, "Münchenbuchsee"),
        ("King of Elsigen", be, None),
    ]
    for text, orte, erwartet in faelle:
        t = ort_im_text(text, orte)
        ok(f"ort_im_text: {text!r} -> {erwartet}", (t[0] if t else None) == erwartet, f"{t}")
    t = ort_im_text("35. GP von Buchholz", de, mehrdeutig_ok=True)
    ok("ort_im_text: mehrdeutig nur mit Flagge, dann ohne Koordinaten",
          t is not None and t[1] is None, f"{t}")
    t = ort_im_text("WEC Cross Os - Bornheide", de, ausnahmen={"Os": "Osnabrück"})
    ok("ort_im_text: Kürzel über ausnahmen", t and t[0] == "Osnabrück", f"{t}")

    import stevenscup_scraper as st
    html = ('<div class="mod_eventlist"><div class="event upcoming"><h1>&#35;04 - Mölln</h1>'
            '<h3>„King of Pirates“</h3><p class="info"><time datetime="2026-10-17">Samstag</time></p>'
            '<p class="location"><strong>Adresse:</strong> Motocrossstrecke | Vor dem Bockholt, 23883 Grambek</p>'
            '<div class="racedata"><p><strong>Ausrichter:</strong> <a href="https://www.pirate-hamburg.de/">MC Pirate</a></p>'
            '<p><strong>Meldeschluss:</strong> 11.10.2026</p></div>'
            '<p class="apply_link"><a class="btn" href="https://my.raceresult.com/387088/">Anmelden</a></p></div>'
            '<div class="event"><h1>&#35;09 - Bad Oldesloe</h1><p class="info"><time datetime="2026-11-08">So</time></p>'
            '<p class="location"><strong>Adresse:</strong> Sportplatz, Bad Oldesloe</p>'
            '<div class="racedata"><p><strong>Ausrichter:</strong> RV Trave</p></div>'
            '<p class="apply_link"><a class="btn" href="https://my.raceresult.com/1/">Anmelden</a></p></div></div>')
    rennen = st.parse_rennen(html)
    ok("stevenscup: zwei Blöcke, kommend/vergangen", [r["vergangen"] for r in rennen] == [False, True], f"{rennen}")
    ev, grund = st.event_aus_rennen(rennen[0], de)
    ok("stevenscup: Ort aus der Adresse (Grambek, nicht Mölln)", ev and ev.standort == "Grambek" and ev.lat, f"{ev} {grund}")
    ok("stevenscup: Name, Ausrichter-Link, Meldeschluss, Cyclecross",
          ev and ev.name == "King of Pirates" and ev.veranstalter_url == "https://www.pirate-hamburg.de/"
          and ev.anmeldeschluss == "2026-10-11" and ev.art2 == "Cyclecross" and ev.laenge_km is None, f"{ev}")
    ev2, _ = st.event_aus_rennen(rennen[1], de)
    ok("stevenscup: ohne h3 heißt es nach dem Cup, ohne Ausrichter-Link die Anmeldung",
          ev2 and ev2.name == "Stevens Cyclo-Cross Cup Bad Oldesloe" and ev2.standort == "Bad Oldesloe"
          and ev2.veranstalter_url == "https://my.raceresult.com/1/", f"{ev2}")

    import cyclocrosscup_scraper as cc
    html = ('<div class="entry-content"><h2>1. Lauf – Brumath</h2><h3>Brumath Bike Festival am 03.10.2026 in Brumath (Frankreich)</h3>'
            '<h2>2. Lauf – Baiersbronn</h2><h3>12. Cross im Park am 18.10.2026 in Baiersbronn</h3>'
            '<p>Veranstalter: TV Baiersbronn</p><p>Informationen: <a href="http://www.cross-im-park.com/">x</a></p>'
            '<p>Navi &amp; Parken: Schwimmbadweg, Baiersbronn</p>'
            '<table><tr><td>Uhrzeit</td><td>Klasse</td><td>Renndauer</td><td>Lizenz</td></tr>'
            '<tr><td>10:00</td><td>Hobby m U40/Ü40 (JG 87+)</td><td>30 min</td><td>ohne</td></tr>'
            '<tr><td></td><td>Hobby Frauen</td><td>30 min</td><td>ohne</td></tr>'
            '<tr><td>11:00</td><td>Masters 2, Masters 3, Masters 4</td><td>40 min</td><td>Lizenz</td></tr>'
            '<tr><td>14:30</td><td>U13 m/w</td><td>15 min Rad</td><td>mit/ohne</td></tr>'
            '<tr><td>17:15</td><td>Frauen Elite + Juniorinnen</td><td>40 min</td><td>Lizenz</td></tr>'
            '<tr><td>17:50</td><td>Männer Elite</td><td>50 min</td><td>Lizenz</td></tr>'
            '<tr><td></td><td>U19 m</td><td>40 min</td><td>Lizenz</td></tr>'
            '<tr><td></td><td>Siegerehrung</td><td></td><td></td></tr></table>'
            '<h2>4. Lauf – Heidelberg</h2><h2><strong>5. Heidelberger Crossrennen 21.11.2026</strong></h2>'
            '<p>Informationen: <a href="http://www.rsv-heidelberg.de">x</a></p><h2></h2></div>')
    rennen = cc.parse_rennen(html)
    ok("cyclocrosscup: drei Läufe, leere h2 stört nicht", [r["lauf"][:7] for r in rennen] == ["1. Lauf", "2. Lauf", "4. Lauf"], f"{[r['lauf'] for r in rennen]}")
    evs, grund = cc.events_aus_rennen(rennen[0], de)
    ok("cyclocrosscup: Frankreich fällt weg", not evs and "außerhalb" in (grund or ""), f"{grund}")
    evs, grund = cc.events_aus_rennen(rennen[1], de)
    ok("cyclocrosscup: eine Zeile je Renndauer, Nachwuchs raus",
          [(e.dauer_h, e.wettbewerb) for e in evs] == [(0.5, "Hobby 30 min (Cyclocross Cup Baden-Württemberg)"),
                                                       (0.67, "Masters, Elite 40 min (Cyclocross Cup Baden-Württemberg)"),
                                                       (0.83, "Elite 50 min (Cyclocross Cup Baden-Württemberg)")],
          f"{[(e.dauer_h, e.wettbewerb) for e in evs]} {grund}")
    ok("cyclocrosscup: Ort, Link, Cyclecross", evs and evs[0].standort == "Baiersbronn"
          and evs[0].veranstalter_url == "http://www.cross-im-park.com/" and evs[0].art2 == "Cyclecross", f"{evs[0] if evs else grund}")
    evs, grund = cc.events_aus_rennen(rennen[2], de)
    ok("cyclocrosscup: Titel als zweite h2, Ort aus der Lauf-Überschrift",
          evs and evs[0].name == "5. Heidelberger Crossrennen" and evs[0].datum_start == "2026-11-21"
          and evs[0].standort == "Heidelberg" and evs[0].dauer_h is None, f"{evs} {grund}")

    import crosscup_scraper as cr
    ok("crosscup: Rennlinks aus dem Menü",
          cr.parse_rennen_links('<a href="https://crosscup.org/rennen/borna/">11.10. – Borna</a><a href="/rennen/dornburg/#a">08.11. – Jena</a><a href="/rennen/">z</a>')
          == [("https://crosscup.org/rennen/borna/", "Borna"), ("https://crosscup.org/rennen/dornburg/", "Jena")])
    ev, grund = cr.event_aus_rennseite({"datum": "2026-11-29", "ort": "Volkspark Piesteritz"}, "u", de, "Wittenberg")
    ok("crosscup: Ort aus dem Menü, wenn die Rennseite nur den Park nennt", ev and ev.standort == "Lutherstadt Wittenberg", f"{ev} {grund}")
    seite = cr.parse_rennseite('<html><head><meta property="og:description" content="Termin : 25.10.26 / 10 Uhr Ort : Mühlhäuser Radibor '
                               'GPS-Daten : 51.246088, 14.394373 Umkleidelokal : Sportplatz Streckenlänge : 2,2 km '
                               'Internetseite des Veranstalters : www.rsv-bautzen.de Ausschreibung: rad-net"></head><body></body></html>')
    ev, grund = cr.event_aus_rennseite(seite, "https://crosscup.org/rennen/radibor/", de)
    ok("crosscup: Termin mit zweistelligem Jahr, Koordinaten der Strecke, Rundenlänge keine Distanz, Link ohne Schema",
          ev and ev.datum_start == "2026-10-25" and ev.standort == "Radibor" and ev.lat == 51.246088
          and ev.laenge_km is None and ev.veranstalter_url == "https://www.rsv-bautzen.de" and ev.name == "GUNSHA Crosscup Radibor", f"{ev} {grund}")

    import dsergebnis_scraper as ds
    html = ('<h1>Termine 2027</h1><p>14.03.2027 Herforder Frühjahrspreis</p><p>03./04.07.2027 100 Jahre Rund in Serrig</p>'
            '<p>24.07.2026 43. Uni-Radrennen</p><p>31.07.2027 Radrennen in Offenbach</p><p>11.09.2027 Termin reserviert</p>'
            '<p>10.10.2027 BL Cyclo Cross 5. Galoppcross Bremen</p><p>35. GP von Buchholz</p>')
    zeilen = ds.parse_termine(html)
    ok("dsergebnis: sechs Terminzeilen, Doppeltermin mit Enddatum", len(zeilen) == 6 and zeilen[1]["datum_ende"] == "2027-07-04", f"{zeilen}")
    ergebnisse = [ds.event_aus_zeile(z, de) for z in zeilen]
    ok("dsergebnis: Herford aus dem Adjektiv, Straße", ergebnisse[0][0] and ergebnisse[0][0].standort == "Herford" and ergebnisse[0][0].art2 == "Straße", f"{ergebnisse[0]}")
    ok("dsergebnis: Serrig", ergebnisse[1][0] and ergebnisse[1][0].standort == "Serrig", f"{ergebnisse[1]}")
    ok("dsergebnis: falsches Jahr ist ein Tippfehler der Seite, kein Eintrag", ergebnisse[2][0] is None and "Tippfehler" in ergebnisse[2][1], f"{ergebnisse[2]}")
    ok("dsergebnis: Offenbach nicht eindeutig", ergebnisse[3][0] is None and "Offenbach" in ergebnisse[3][1], f"{ergebnisse[3]}")
    ok("dsergebnis: Termin reserviert ist kein Rennen", ergebnisse[4][0] is None, f"{ergebnisse[4]}")
    ok("dsergebnis: Cross ist Cyclecross", ergebnisse[5][0] and ergebnisse[5][0].art2 == "Cyclecross" and ergebnisse[5][0].standort == "Bremen", f"{ergebnisse[5]}")

    import weserems_scraper as we
    orte_we = we.orte_weser_ems()
    zeilen = we.parse_termine('<p><span>18.10.2026  WEC Cross Bad Essen / Pr. Oldendorf</span></p>'
                              '<p>06.12.2026  WEC u. <strong>LVM NS</strong> Os - Schinkel</p><p>13.12.2026  WEC Cross Engter</p><p>Hier die Termine</p>')
    ok("weserems: drei Zeilen", len(zeilen) == 3, f"{zeilen}")
    evs = [we.event_aus_zeile(z, orte_we)[0] for z in zeilen]
    ok("weserems: erster Ort bei zwei Orten, Name wie bei radsport-events.de, Label schlicht",
          [(e.standort, e.name, e.wettbewerb) for e in evs] == [("Bad Essen", "Weser-Ems-Cup Bad Essen", "Cyclocross"),
                                                  ("Osnabrück", "Weser-Ems-Cup Osnabrück", "Cyclocross"),
                                                  ("Bramsche", "Weser-Ems-Cup Bramsche", "Cyclocross")], f"{[(e.standort, e.name, e.wettbewerb) for e in evs]}")

    import boe_scraper as boe
    html = ('<h2>BOE-Strassenmeisterschaft</h2><div class="cc-m-download-title">13.05.  -  Einzelzeitfahren Langnau-Zäziwil</div>'
            '<div class="cc-m-download-description">Organisation: VC Bärau</div>'
            '<div class="cc-m-download-title">20.09.  -  King of Elsigen</div><div class="cc-m-download-description">Organisation: OK</div>'
            '<h2>BOE-Bikemeisterschaft</h2><div class="cc-m-download-title">17.06. - Bike-Rennen Burgdorf (Fänstu)</div>'
            '<div class="cc-m-download-description">Organisation: RV Ersigen</div>'
            '<div class="cc-m-download-title">21.06. - Bärenried Enduro MTB Race (nur Kids)</div>')
    rennen = boe.parse_kalender(html, 2027)
    ok("boe: vier Termine mit Jahr aus der Adresse", [r["datum"] for r in rennen] == ["2027-05-13", "2027-09-20", "2027-06-17", "2027-06-21"], f"{rennen}")
    ergebnisse = [boe.event_aus_rennen(r, "u", be) for r in rennen]
    ok("boe: Zeitfahren Zäziwil", ergebnisse[0][0] and ergebnisse[0][0].standort == "Zäziwil" and ergebnisse[0][0].art2 == "Zeitfahren", f"{ergebnisse[0]}")
    ok("boe: Berg ist kein Ort", ergebnisse[1][0] is None, f"{ergebnisse[1]}")
    ok("boe: Bikemeisterschaft ist Mountainbike, Organisation gelesen", ergebnisse[2][0] and ergebnisse[2][0].art2 == "Mountainbike"
          and rennen[2]["organisation"] == "RV Ersigen", f"{ergebnisse[2]}")
    ok("boe: nur Kids fällt weg", ergebnisse[3][0] is None, f"{ergebnisse[3]}")

    import fricktal_scraper as fr
    html = ('<div class="event-list"><h3>Sonntag 18.10.2026</h3><div class="cd-tile-h-box"><div class="cd-tile-h-main-heading">'
            'Bikerennen Wittnau - 38. Raiffeisen Fricktaler Cup 2026</div><ul><li><div class="cd-tile-h-detail-label">Ort</div>'
            '<div class="cd-tile-h-detail-value">Turnhalle Wittnau</div></li></ul></div>'
            '<h3>Freitag 06.11.2026</h3><div class="cd-tile-h-box"><div class="cd-tile-h-main-heading">Absenden/Preisverteilung - Cup</div></div></div>')
    eintraege = fr.parse_kalender(html)
    ag = orte_aus_places("Schweiz", region="Kanton Aargau")
    ev, grund = fr.event_aus_eintrag(eintraege[0], ag, orte_aus_places("Schweiz"))
    ok("fricktal: Wittnau im Aargau (nicht das badische), Mountainbike", ev and ev.standort == "Wittnau" and 47.4 < ev.lat < 47.6
          and ev.art2 == "Mountainbike" and ev.datum_start == "2026-10-18", f"{ev} {grund}")
    ok("fricktal: Absenden ist kein Rennen", fr.event_aus_eintrag(eintraege[1], ag, ag)[0] is None)

    import radsportsh_scraper as sh
    orte_sh = sh.orte_sh()
    seite = sh.parse_terminseite('<h1>LTV Almabtrieb (RTF)</h1><div class="news-text-wrap"><p><a href="https://www.ltvkiel-ost.de">x</a></p></div>'
                                 '<div class="-datum ms-2">Samstag, 03.10.2026</div>')
    ev, grund = sh.event_aus_terminseite(seite, "u", orte_sh)
    ok("radsport-sh: Ort aus dem Hostnamen, RTF ist Straße, Format im Label",
          ev and ev.standort == "Kiel" and ev.art2 == "Straße" and ev.wettbewerb == "RTF" and ev.name == "LTV Almabtrieb"
          and ev.veranstalter_url == "https://www.ltvkiel-ost.de", f"{ev} {grund}")
    seite = sh.parse_terminseite('<h1>Sparkasse Mittelholstein AG O-CTF in Westensee (CTF)</h1><div class="-datum">Sonntag, 08.11.2026</div>')
    ev, grund = sh.event_aus_terminseite(seite, "https://www.radsport-sh.de/termin/x", orte_sh)
    ok("radsport-sh: Ort aus dem Namen, CTF ist Mountainbike, ohne Link die Terminseite",
          ev and ev.standort == "Westensee" and ev.art2 == "Mountainbike" and ev.veranstalter_url == "https://www.radsport-sh.de/termin/x", f"{ev} {grund}")
    ev, grund = sh.event_aus_terminseite({"name": "Fachwartetag der Radwanderer", "datum": "2026-11-01"}, "u", orte_sh)
    ok("radsport-sh: Verbandstermin ohne Format fällt weg", ev is None, f"{ev}")
    ev, grund = sh.event_aus_terminseite({"name": "Waffel Ride Kiel 70 km (Gravelride)", "datum": "2027-01-10"}, "u", orte_sh)
    ok("radsport-sh: eine Kilometerzahl im Namen wird die Länge", ev and ev.laenge_km == 70 and ev.standort == "Kiel" and ev.art2 == "Gravel", f"{ev} {grund}")
    ev, grund = sh.event_aus_terminseite({"name": "Viking Bike Winterchallenge, O-CTF und Gravelride 70 km", "datum": "2027-01-10",
                                          "veranstalter_url": "https://www.rv-schleswig.de"}, "u", orte_sh)
    ok("radsport-sh: zwei Formate im Namen - die Kilometer gehören nur zu einem, Länge leer; Ort aus dem Host",
       ev and ev.laenge_km is None and ev.standort == "Schleswig", f"{ev} {grund}")


def test_marsch_quellen() -> None:
    """Mammutmarsch und Megamarsch (17.10.2026): Laufen / Marsch, je
    Distanz eine Zeile, Ort über die PLZ, nur DE/AT/CH."""
    print("\nMammutmarsch / Megamarsch:")
    import json as _j
    import mammutmarsch_scraper as mm
    import megamarsch_scraper as mg
    termin = {"name": "Mammutmarsch Ruhrgebiet - 30/42/55 KM", "slug": "ruhrgebiet", "date": "2027-04-24T00:00:00.000Z",
              "address": "Startorte: 30 KM: Innenhafen Duisburg, Johannes-Corputius-Platz 1, 47051 Duisburg",
              "city": "Ruhr", "eventCountry": "DE", "status": "ACTIVE", "effectiveStatus": "ACTIVE",
              "distances": [30, 42, 55]}
    html = ('<script>self.__next_f.push([1,' + _j.dumps('x"events":' + _j.dumps([termin, dict(termin, eventCountry="ES")]))
            + '])</script>')
    termine = mm.termine_aus_html(html)
    check("mammutmarsch: Terminliste aus den Flight-Daten", len(termine), 2)
    evs, _ = mm.events_aus_termin(termine[0])
    check("mammutmarsch: je Distanz eine Zeile, Ort aus der PLZ, Marsch",
          [(e.name, e.standort, e.laenge_km, e.art1, e.art2, e.veranstalter_url) for e in evs][:1],
          [("Mammutmarsch Ruhrgebiet", "Duisburg", 30.0, "Laufen", "Marsch", "https://mammutmarsch.de/events/ruhrgebiet")])
    check("mammutmarsch: Spanien fällt, Startort offen -> Stadt",
          [mm.events_aus_termin(termine[1])[1],
           mm.events_aus_termin(dict(termin, address="Start/Ziel: wird noch bekannt gegeben", city="Wien",
                                     eventCountry="AT"))[0][0].standort],
          ["Land ES", "Wien"])
    check("mammutmarsch: Name ohne Strecken", mm.name_ohne_strecken("Mammutmarsch Berlin - 30/42/55 KM"),
          "Mammutmarsch Berlin")
    uebersicht = ('<div>\nHANNOVER\n50KM // 12H\n08.05.2027\nWESERBERGLAND bei Nacht\n50KM // 12H\n12.06.2027\n'
                  'DRESDEN\n25KM // 6H\n06.03.2027\n</div><a href="https://www.megamarsch.de/hannover">x</a>'
                  '<a href="https://www.megamarsch.de/weserbergland-bei-nacht">x</a>'
                  '<a href="https://www.megamarsch.de/dresden-25-km">x</a><a href="https://www.megamarsch.de/dresden">x</a>')
    t = mg.termine(uebersicht)
    seiten = mg.eventseiten(uebersicht)
    check("megamarsch: Termine und Eventseiten",
          [(x["ort"], x["km"], x["datum"], mg.seite_fuer(x["ort"], x["km"], seiten)) for x in t],
          [("HANNOVER", 50, "2027-05-08", "hannover"),
           ("WESERBERGLAND bei Nacht", 50, "2027-06-12", "weserbergland-bei-nacht"),
           ("DRESDEN", 25, "2027-03-06", "dresden-25-km")])
    check("megamarsch: Startort über PLZ, Gemeinde aus der Anlage, offen -> Stadt",
          [mg.startort("<p>ORT</p><p>STARTORT:</p><p>Messeparkplatz,</p><p>30880 Laatzen</p><p>ZIELORT:</p>")[0],
           mg.startort("<p>STARTORT:</p><p>Campingbeach Annenheim Seeuferstraße 109</p>"
                       "<p>9520 Treffen am Ossiacher See Österreich</p><p>ZIELORT:</p>")[::3],
           mg.startort("<p>STARTORT:</p><p>Wird noch bekannt gegeben</p><p>ZIELORT:</p>", "Luzern")[::3],
           mg.startort("<p>STARTORT:</p><p>Calle 1, 07001 Palma</p><p>ZIELORT:</p>", "Mallorca")],
          ["Laatzen", ("Annenheim", "Österreich"), ("Luzern", "Schweiz"), None])
    check("megamarsch: Anzeigename", mg.anzeigename("WESERBERGLAND bei Nacht"), "Weserbergland bei Nacht")


def test_dealgrid() -> None:
    """Der Scraper für den Gravel-Kalender von dealgrid.de (08.10.2026):
    JSON-LD plus Listenzeile, Strecken aus der Meta-Zeile (Preise,
    Spannen, Runden, Vorjahreswerte), Ort gegen places.json, Camps,
    Serien und Ausland raus - ohne Netz."""
    print("\ndealgrid.de:")
    import dealgrid_scraper as dg
    s = dg.strecken_aus_meta
    check("dealgrid: Zahlen ohne Einheit zählen erst mit km, Preise nicht",
          [k for k, _ in s("Waltrop, NRW · 43 · 67 · 101 km · 8 · 10 €")], [43, 67, 101])
    check("dealgrid: Label-Wörter, CTF-Spanne nicht",
          s("Paderborn, NRW · 48 km Gravel · CTF bis 63 km · 6–10 €"), [(48, "Gravel 48 km")])
    check("dealgrid: Spanne fällt, die Zahl davor bleibt",
          [k for k, _ in s("Kiel, SH · 65 · 100–120 km")], [65])
    check("dealgrid: Runden ergeben die Renndistanz, Kinderstrecke raus",
          [k for k, _ in s("Rheinzabern, RLP · 10 × 5,8 km · 20 €")] + [k for k, _ in s("Sonnewalde, BB · 30 · 55 km · Kids 3,5 km")],
          [58, 30, 55])
    check("dealgrid: Vorjahreswerte und Rundenlänge nicht",
          [s("Gamlitz, Stmk · Strecken 2027 folgen (2026: 40 · 96 · 126 km)"),
           s("Bad Feilnbach, BY · Gravel 56 · 94 km (Stand 2026) · 139 €"),
           s("X, DE · 17,7-km-Runde, jede Stunde neu"), s("X, DE · 9 Routen, 20 bis 100 km")], [[], [], [], []])
    check("dealgrid: Tausenderpunkt, „oder“, Höhenmeter",
          [k for k, _ in s("Pfalz, RLP · 1.120 km · 22.500 hm")] + [k for k, _ in s("Ibbenbüren, NRW · Gravel 46, 80 oder 115 km · 10 €")],
          [1120, 46, 80, 115])
    html = ('<script type="application/ld+json">{"@context":"https://schema.org","@graph":['
            '{"@type":"SportsEvent","name":"Rad am Ring, Gravel-Rennen","startDate":"2027-07-24","endDate":"2027-07-24",'
            '"eventStatus":"https://schema.org/EventScheduled","location":{"address":{"addressLocality":"Nürburgring","addressCountry":"DE"},'
            '"geo":{"latitude":50.335,"longitude":6.947}},"url":"https://radamring.de/"},'
            '{"@type":"SportsEvent","name":"Riedgravel-Camp","startDate":"2027-08-27","location":{"address":{"addressLocality":"Riedstadt","addressCountry":"DE"}},"url":"https://x.de/"},'
            '{"@type":"SportsEvent","name":"Elsass Gravel","startDate":"2027-07-04","location":{"address":{"addressLocality":"Wangenbourg","addressCountry":"FR"}},"url":"https://y.fr/"},'
            '{"@type":"SportsEvent","name":"Rund um den Harz","startDate":"2027-04-03","location":{"address":{"addressLocality":"Harz","addressCountry":"DE"},"geo":{"latitude":51.7,"longitude":10.6}},"url":"https://z.de/"}'
            ']}</script><ul>'
            '<li class="dgkal-row" id="termin-50"><span class="dgkal-kicker-s">Rennen</span><a class="dgkal-name" href="https://radamring.de/">Rad am Ring, Gravel-Rennen</a>'
            '<span class="dgkal-meta">Nürburgring, Rheinland-Pfalz · 50 · 100 km · 65 €</span></li>'
            '<li class="dgkal-row" id="termin-51"><span class="dgkal-kicker-s">Camp</span><a class="dgkal-name" href="https://x.de/">Riedgravel-Camp</a>'
            '<span class="dgkal-meta">Riedstadt, Hessen · 60 km</span></li></ul>')
    termine = dg.parse_kalender(html)
    check("dealgrid: vier Termine aus dem JSON-LD", [t["name"] for t in termine],
          ["Rad am Ring, Gravel-Rennen", "Riedgravel-Camp", "Elsass Gravel", "Rund um den Harz"])
    evs, grund = dg.events_aus_termin(termine[0])
    check("dealgrid: Name ohne „, Gravel-Rennen“, Gattung im Label, Ort mit Koordinaten der Quelle",
          [(e.name, e.standort, e.laenge_km, e.wettbewerb, e.art2, e.veranstalter_url, e.lat) for e in evs],
          [("Rad am Ring", "Nürburgring", 50, "Gravel-Rennen 50 km", "Gravel", "https://radamring.de/", 50.335),
           ("Rad am Ring", "Nürburgring", 100, "Gravel-Rennen 100 km", "Gravel", "https://radamring.de/", 50.335)])
    check("dealgrid: Camp, Ausland und Gegend fallen",
          [dg.events_aus_termin(t)[1][:12] for t in termine[1:]], ["kein Renntermin"[:12], "außerhalb de", "Ort nicht ei"])
    orte = {"freiburg im breisgau": [("Freiburg im Breisgau", 47.999, 7.839)],
            "wald": [("Wald", 47.933, 9.167)], "stuttgart": [("Stuttgart", 48.78, 9.18)]}
    check("dealgrid: Ortsname kürzer als in places.json, nur mit Koordinaten",
          (dg.standort_finden("Freiburg", 47.99, 7.85, orte), dg.standort_finden("Freiburg", None, None, orte)),
          (("Freiburg im Breisgau", 47.99, 7.85), None))
    check("dealgrid: „Bayerischer Wald“ ist nicht der Ort Wald, Ortsteil zählt als Ort",
          (dg.standort_finden("Bayerischer Wald", 49.1, 12.9, orte), dg.standort_finden("Stuttgart-Wangen", 48.778, 9.241, orte)),
          (None, ("Stuttgart", 48.778, 9.241)))


def test_kilometerliebe() -> None:
    """Der Scraper für kilometerliebe.de (30.09.2026): Monatsseite (Datum +
    Link je Karte, vergangene weg), Eventseite (JSON-LD, Faktenliste,
    Event-Website), Distanzen, Zeitrennen, Triathlon-Formate, Wandern raus,
    Latin-1-Mojibake - ohne Netz."""
    print("\nkilometerliebe.de:")
    import kilometerliebe_scraper as kl
    monat = ('<nav><a href="/events/map">Karte</a></nav>'
             '<article><time datetime="2026-10-03T00:00:00">3. Okt</time><a href="/events/kap-arkona-lauf/">Kap-Arkona-Lauf</a></article>'
             '<article><time datetime="2026-10-01">1. Okt</time><a href="/events/alt-2026/">Alt</a></article>'
             '<article><time datetime="2026-10-03">3. Okt</time><a href="https://www.kilometerliebe.de/events/kap-arkona-lauf/">nochmal</a></article>')
    links = kl.event_links(monat, ab="2026-10-02")
    check("kilometerliebe: Karten mit Datum, vergangene und doppelte weg, Karte-Link nicht",
          links, [("2026-10-03", "Kap-Arkona-Lauf", "https://www.kilometerliebe.de/events/kap-arkona-lauf/")])
    seite = ('<html><head><script type="application/ld+json">{"@context":"https://schema.org","@graph":[{"@type":"SportsEvent",'
             '"name":"13. Belgenbach-Trail","description":"13. Belgenbach-Trail am 7. März 2027 in Monschau. Angebotene Distanzen: 5 km, 9,5 km, 18 km.",'
             '"startDate":"2027-03-07","endDate":"2027-03-07","location":{"@type":"Place","name":"Monschau","address":{"addressLocality":"Monschau",'
             '"addressRegion":"Nordrhein-Westfalen","addressCountry":"DE"},"geo":{"latitude":50.55,"longitude":6.24}},"eventStatus":"https://schema.org/EventScheduled"}]}</script></head>'
             '<body><dl><dt>Termin</dt><dd>So, 07.03.2027</dd><dt>Distanzen</dt><dd>5 km, 9,5 km, 18 km</dd><dt>Kategorie</dt><dd>Trail</dd><dt>Status</dt><dd>Geplant</dd></dl>'
             '<a href="data:text/calendar" download="x.ics" class="ev-bib__action">Kalender</a>'
             '<a href="https://belgenbachtrail.de/" class="ev-bib__action">Event-Website</a></body></html>')
    e = kl.parse_eventseite(seite)
    check("kilometerliebe: Eventseite gelesen", (e.get("name"), e.get("datum_start"), e.get("ort"), e.get("kategorie"), e.get("distanzen"), e.get("veranstalter_url")),
          ("13. Belgenbach-Trail", "2027-03-07", "Monschau", "Trail", "5 km, 9,5 km, 18 km", "https://belgenbachtrail.de/"))
    evs, grund = kl.events_aus_seite(e, "https://www.kilometerliebe.de/events/x/", kl.CONFIG)
    check("kilometerliebe: je Distanz eine Zeile, Trail, Koordinaten, Veranstalterseite",
          [(x.laenge_km, x.art2, x.lat, x.veranstalter_url) for x in evs],
          [(5.0, "Trail", 50.55, "https://belgenbachtrail.de/"), (9.5, "Trail", 50.55, "https://belgenbachtrail.de/"), (18.0, "Trail", 50.55, "https://belgenbachtrail.de/")])
    e2 = dict(e, name="Drehwurm", distanzen="6 h", kategorie="Lauf", veranstalter_url=None)
    evs, _ = kl.events_aus_seite(e2, "https://www.kilometerliebe.de/events/y/", kl.CONFIG)
    check("kilometerliebe: 6 h ist ein Zeitrennen, ohne Event-Website die Kalenderseite",
          [(x.laenge_km, x.dauer_h, x.veranstalter_url) for x in evs], [(None, 6.0, "https://www.kilometerliebe.de/events/y/")])
    e3 = dict(e, name="1. Friedberger Triathlon", kategorie="Triathlon", distanzen=None,
              beschreibung="1. Friedberger Triathlon am 19. September 2027 in Friedberg. Angebotene Wertungen: Olympische Distanz (Schwimmen – Radfahren – Laufen), Mitteldistanz (…).")
    evs, _ = kl.events_aus_seite(e3, "u", kl.CONFIG)
    check("kilometerliebe: Triathlon je Format eine Zeile", [(x.art1, x.wettbewerb, x.art2) for x in evs],
          [("Triathlon", "Olympisch", "Straße"), ("Triathlon", "Mitteldistanz", "Straße")])
    e4 = dict(e, name="Mammutmarsch Hamburg", kategorie="Wandern")
    check("kilometerliebe: Mammutmarsch/Megamarsch kommen aus der eigenen Quelle",
          kl.events_aus_seite(e4, "u", kl.CONFIG), ([], "Wandern (eigene Quelle)"))
    evs, _ = kl.events_aus_seite(dict(e, name="Rennsteig Wandermarathon", kategorie="Wandern",
                                      distanzen="42 km"), "u", kl.CONFIG)
    check("kilometerliebe: andere Wander-Events als Laufen/Marsch",
          [(x.art1, x.art2, x.laenge_km) for x in evs], [("Laufen", "Marsch", 42.0)])
    e5 = dict(e, status_text="Abgesagt")
    check("kilometerliebe: abgesagt fällt weg", kl.events_aus_seite(e5, "u", kl.CONFIG)[1], "abgesagt")
    check("kilometerliebe: Latin-1-Mojibake repariert", kl._utf8("KÃ¶ln Marathon"), "Köln Marathon")


def test_siebter_weg() -> None:
    """Siebter Weg von _same_name (30.09.2026): Sponsor, Auflage, Jahr und
    Ortsadjektiv im Namen - samt Gegenproben."""
    print("\nNamensregel, siebter Weg (Sponsor/Ort im Namen):")
    from scraper_lib import is_same_event
    def e(name, standort, km=None, wb=None, art1="Laufen", lat=None, lon=None):
        return dict(name=name, standort=standort, laenge_km=km, wettbewerb=wb, art1=art1, datum_start="2027-01-01", lat=lat, lon=lon)
    faelle = [
        ("SAARathon", "Saarbrücken", 5, "5 km Lauf für alle", "Sparkassen-SAARathon", "Saarbrücken", 5, "5 km", True),
        ("Frauenlauf Berlin", "Berlin", 10, "10 km Lauf", "Vitamin Well Frauenlauf Berlin 2027", "Berlin", 10, "10 km", True),
        ("Marathon Bonn", "Bonn", 21.1, "Halbmarathon", "24. Deutsche Post Marathon Bonn", "Bonn", 21.1, "21,1 km", True),
        ("Trollinger Marathon", "Heilbronn", 21.1, "Halbmarathon", "Heilbronner Trollinger-Marathon 2027", "Heilbronn", 21.1, "21,1 km", True),
        ("Keltenlauf Ditzingen", "Ditzingen", 13, "13 km", "2. Ditzinger Keltenlauf", "Ditzingen", 13, "13 km", True),
        ("Runners Festival Klouster", "Rinchnach", 12, "Hauptlauf", "4. Runner's Festival Klouster", "Rinchnach", 12, "12 km", True),
        ("Alstervorland parkrun Hamburg", "Hamburg", 5, "5 km", "Krupunder See parkrun Hamburg", "Hamburg", 5, "5 km", False),
        ("Hamburg Halbmarathon", "Hamburg", 21.1, "Halbmarathon", "Alstertallauf Hamburg", "Hamburg", 21.1, "21,1 km", False),
        ("Frauenlauf Berlin", "Berlin", 10, "10 km Walking", "Vitamin Well Frauenlauf Berlin 2027", "Berlin", 10, "10 km Lauf", False),
        ("Essener-Silvesterlauf 2026", "Essen", 5, "5 km Jedermannlauf", "Silvesterlauf Gießen", "Gießen", 5, "5 km Lauf", False),
        ("Stadtlauf Erding", "Erding", 10, "10 km", "Stadtlauf Erding", "Erding", 5, "5 km", False),
        # Eine Seite ohne Distanz, beide Labels nennen dasselbe Format
        # (05.10.2026, Munich Triathlon / Kilometerliebe) - samt Gegenproben.
        ("Munich Triathlon", "München", 25.5, "Sprintdistanz 25,5 km (500 m Schwimmen / 20 km Rad / 5 km Laufen)", "Munich Triathlon 2027", "München", None, "Sprint", True),
        ("Munich Triathlon", "München", 51.5, "Kurzdistanz 51,5 km (1,5 km Schwimmen / 40 km Rad / 10 km Laufen)", "Munich Triathlon 2027", "München", None, "Sprint", False),
        ("Allgäu Triathlon", "Immenstadt im Allgäu", 57.5, "Olymp 57,5 km (1,5 km Schwimmen / 46 km Rad / 10 km Laufen)", "Allgäu Triathlon 2027", "Immenstadt im Allgäu", None, "Olympisch", True),
        ("Triathlon Tübingen", "Tübingen", 27.4, "Sprintdistanz 27,4 km (750 m Schwimmen / 22 km Rad / 4,7 km Laufen)", "Mey Generalbau Triathlon Tübingen 2027", "Tübingen", None, "Super-Sprint", False),
        ("Stadtlauf Erding", "Erding", 10, "10 km Hauptlauf", "Stadtlauf Erding 2027", "Erding", None, "Nachtlauf", False),
        # Die Teilstrecken-Klammer zählt nicht als Gattung: "Rad" darin ist
        # kein Radrennen (Siebzehnter Durchgang, bis dahin per Override).
        ("Sprint Triathlon Hof", "Hof", 26, "Sprintdistanz 26 km (750 m Schwimmen / 20 km Rad / 5 km Laufen)", "Sprint Triathlon Hof", "Hof", 26, "Sprintdistanz 26 km", True),
    ]
    for n1, o1, k1, w1, n2, o2, k2, w2, erwartet in faelle:
        art1 = "Triathlon" if "Triathlon" in n1 else "Laufen"
        check(f"{n1!r} <> {n2!r}", is_same_event(e(n1, o1, k1, w1, art1), e(n2, o2, k2, w2, art1)), erwartet)
    from scraper_lib import _label_gattung
    check("Gattung: Teilstrecken-Klammer zählt nicht", _label_gattung("Sprintdistanz 26 km (750 m Schwimmen / 20 km Rad / 5 km Laufen)"), frozenset({"sprint"}))
    check("Gattung: Super-Sprint ist nicht Sprint", _label_gattung("Super-Sprint"), frozenset({"super-sprint"}))
    check("Gattung: Staffel-Klammer bleibt", _label_gattung("Staffel (4 x 10 km)"), frozenset({"staffel"}))


def test_lvpfalz_lck() -> None:
    """LV-Pfalz-Terminliste und Kaltern-Kalender (30.09.2026), ohne Netz."""
    print("\nLV Pfalz und Läuferclub Kaltern:")
    import lvpfalz_scraper as lp
    import lck_scraper as lck
    from scraper_lib import orte_aus_places, SUEDTIROL
    html = ('<div class="event"><h2 itemprop="name">44. Gemüselauf</h2><p class="time"><time datetime="2026-10-03T15:00:00+02:00">03.10.2026</time></p>'
            '<div><p>Veranstalter: TSG Maxdorf</p><p>10 km Gemüselauf; 0,8 km Junges Gemüse</p><p><a href="mailto:x">x</a></p>'
            '<p><a href="http://www.tsg-maxdorf.de">http://www.tsg-maxdorf.de</a></p></div></div>'
            '<div class="event"><h2 itemprop="name">Holzland-Crosslauf</h2><p class="time"><time datetime="2026-10-03">03.10.2026</time></p>'
            '<div><p>Veranstalter: TuS Heltersberg</p><p>Crosslauf; Crosslauf lang</p><p><a href="https://tus06heltersberg.de/">x</a></p></div></div>')
    laeufe = lp.parse_liste(html)
    check("lvpfalz: zwei Läufe mit Veranstalter, Strecken und Link",
          [(l["name"], l["datum"], l["veranstalter"], l["strecken"], l["url"]) for l in laeufe],
          [("44. Gemüselauf", "2026-10-03", "TSG Maxdorf", "10 km Gemüselauf; 0,8 km Junges Gemüse", "http://www.tsg-maxdorf.de"),
           ("Holzland-Crosslauf", "2026-10-03", "TuS Heltersberg", "Crosslauf; Crosslauf lang", "https://tus06heltersberg.de/")])
    rlp = orte_aus_places("Deutschland", region="Rheinland-Pfalz")
    evs, grund = lp.events_aus_lauf(laeufe[0], lp.CONFIG, rlp)
    check("lvpfalz: Ort aus dem Verein, je Strecke eine Zeile", [(e.standort, e.laenge_km) for e in evs], [("Maxdorf", 10.0), ("Maxdorf", 0.8)])
    evs, grund = lp.events_aus_lauf(laeufe[1], lp.CONFIG, rlp)
    check("lvpfalz: Crosslauf ohne Kilometer bleibt eine Zeile, Trail", [(e.standort, e.laenge_km, e.art2) for e in evs][:1], [("Heltersberg", None, "Trail")])
    frag = ('<h5>Berglauf, Trail</h5><div class="sm:flex"><div>ausgetragen</div><p class="my-3">Samstag, 03.10.2026</p><h2>Radlsee Hüttenlauf</h2>'
            '<p><strong>Distanz: 5,1km (874m)</strong></p><p class="mt-2">Berglauf zur Radlseehütte vom Garner Wetterkreuz / Feldthurns.</p></div>'
            '<h5>VSS - Dorflauf, Kinder- und Jugendlauf</h5><div class="sm:flex"><p>Samstag, 03.10.2026</p><h2>Dorflauf Niederdorf</h2><p><strong>Distanz: 5km - 10km</strong></p>'
            '<p>x</p><a href="http://www.vss.bz.it/x">vss</a></div>'
            '<h5>Berglauf, Trail</h5><div class="sm:flex"><p>Samstag, 03.10.2026</p><h2>Fiemme Ultra Sky</h2><p><strong>Distanz: 120km - 80km</strong></p><p>Trentino</p><a href="https://www.fiemmeultrasky.com/">x</a></div>')
    eintraege = lck.parse_fragment(frag)
    check("lck: drei Einträge gelesen", [(e["name"], e["datum"], e.get("distanz"), e["kategorie"]) for e in eintraege],
          [("Radlsee Hüttenlauf", "2026-10-03", "5,1km (874m)", "Berglauf, Trail"), ("Dorflauf Niederdorf", "2026-10-03", "5km - 10km", "VSS - Dorflauf, Kinder- und Jugendlauf"),
           ("Fiemme Ultra Sky", "2026-10-03", "120km - 80km", "Berglauf, Trail")])
    st = orte_aus_places(SUEDTIROL)
    evs, _ = lck.events_aus_eintrag(eintraege[0], st)
    check("lck: Ort aus dem Text (Feldthurns), Höhenmeter keine Distanz, Trail", [(e.standort, e.laenge_km, e.art2, e.land) for e in evs], [("Feldthurns", 5.1, "Trail", SUEDTIROL)])
    evs, _ = lck.events_aus_eintrag(eintraege[1], st)
    check("lck: Dorflauf mit zwei Distanzen, Ort aus dem Namen, VSS-Link", [(e.standort, e.laenge_km, e.veranstalter_url) for e in evs],
          [("Niederdorf", 5.0, "http://www.vss.bz.it/x"), ("Niederdorf", 10.0, "http://www.vss.bz.it/x")])
    check("lck: Trentiner Lauf ohne Südtiroler Ort fällt weg", lck.events_aus_eintrag(eintraege[2], st), ([], "kein Südtiroler Ort"))
    check("lck: Seitenleisten-Ids", lck.sidebar_ids('<a href=?id=15&kategorie=>x</a><a href=?id=9&kategorie=>y</a><a href=?id=15&kategorie=>z</a>'), [9, 15])


def test_gedaechtnis() -> None:
    """Gedächtnis je Veranstaltung ohne Datum (veranstalter_seiten.py,
    05.10.2026): Schlüssel ohne Jahr/Auflage, Seite und Strecken merken,
    Seite/Koordinaten auf die neue Ausgabe anwenden, Abgleich mit dem
    Vorjahr, plausible nächste Termine, Übernahme nach manual_events."""
    print("\nGedächtnis je Veranstaltung (veranstalter_seiten.py):")
    import json as _json
    import tempfile
    import veranstalter_seiten as vs

    def z(name, datum, km, wb="", url="https://www.munich-triathlon.com/", lat=48.14, lon=11.58, standort="München"):
        return dict(name=name, standort=standort, land="Deutschland", lat=lat, lon=lon, art1="Triathlon",
                    art2="Straße", datum_start=datum, datum_ende=datum, laenge_km=km, dauer_h=None,
                    wettbewerb=wb, veranstalter_url=url)
    check("Schlüssel ohne Jahr und Auflage", vs.schluessel_von(z("49. Munich Triathlon 2027", "2027-07-25", 25.5)),
          vs.schluessel_von(z("Munich Triathlon", "2026-07-26", 25.5)))
    check("Name ohne Jahr", (vs.name_ohne_jahr("Munich Triathlon 2027"), vs.name_ohne_jahr("Winterlaufserie 2026/2027"), vs.name_ohne_jahr("Lauf 2027 - Nord")),
          ("Munich Triathlon", "Winterlaufserie", "Lauf - Nord"))
    check("Name ohne Auflage vorn", (vs.name_ohne_jahr("6. Appelhülsener Landlauf"), vs.name_ohne_jahr("Lauf zum 4. Advent")),
          ("Appelhülsener Landlauf", "Lauf zum 4. Advent"))
    ged = vs.Gedaechtnis({})
    alt = [z("Munich Triathlon", "2026-07-26", 25.5, "Sprint 25,5 km"), z("Munich Triathlon", "2026-07-26", 51.5, "Kurzdistanz 51,5 km")]
    vs.lernen(alt, ged, overrides={})
    e = ged.finde(alt[0])
    check("Ausgabe 2026 gelernt: Seite und zwei Strecken", (e["url"], [s["laenge_km"] for s in e["strecken"]], e["datum"]),
          ("https://www.munich-triathlon.com/", [25.5, 51.5], "2026-07-26"))
    zweimal = vs.lernen([z("Munich Triathlon", "2026-07-26", 25.5, "Sprint 25,5 km", url="https://munich-triathlon.com/")], ged, overrides={})
    check("gleiche Seite in zweiter Schreibweise: nichts hin und her", (zweimal, e["url"]), ([], "https://www.munich-triathlon.com/"))
    # handpruefung (13.10.2026): Fälle, deren Veranstaltung schon wieder in
    # events.json steht, fallen aus der Liste - der Bericht vergisst nichts.
    gh = vs.Gedaechtnis({})
    vs.lernen([z("Munich Triathlon", "2026-07-26", 25.5, "Sprint"),
               z("Stadtlauf Erding", "2026-10-03", 10.0, "10 km", url="https://sv-erding.de/", standort="Erding", lat=48.3, lon=11.9)], gh, overrides={})
    k_m = vs.schluessel_von(z("Munich Triathlon", "2026-07-26", 25.5))
    k_e = vs.schluessel_von(z("Stadtlauf Erding", "2026-10-03", 10.0, standort="Erding"))
    bericht = {k_m: {"flag": "NEU?", "kandidaten": ["2027-07-24", "2027-07-25"], "name": "Munich Triathlon"},
               k_e: {"flag": "NEU?", "kandidaten": ["2027-10-02"], "name": "Stadtlauf Erding"},
               "nichts|x": {"flag": "NICHTS"}}
    aktuell = vs.aktuelle_schluessel(gh, [z("Munich Triathlon 2027", "2027-07-25", 25.5, "Sprint")])
    faelle, weg = vs.handpruefung_faelle(bericht, aktuell)
    check("handpruefung: erledigter Fall (Zeile in events.json) fällt weg, offener bleibt",
          ([k for k, _ in faelle], weg), ([k_e], 1))
    # Eine Veranstaltung OHNE eigene Seite steht trotzdem im Gedächtnis.
    ohne = [z("Stadtlauf Erding", "2026-10-03", 10.0, "10 km", url="https://my.raceresult.com/1/", standort="Erding")]
    vs.lernen(ohne, ged, overrides={})
    e2 = ged.finde(ohne[0])
    check("ohne Seite: Eintrag mit Strecken, url leer", (e2 is not None and e2.get("url"), [s["laenge_km"] for s in e2["strecken"]]), (None, [10.0]))
    # Die neue Ausgabe kommt mit Portallink und ohne Koordinaten.
    neu = [z("Munich Triathlon 2027", "2027-07-25", None, "Sprint", url="https://www.kilometerliebe.de/x", lat=None, lon=None)]
    vergleich = vs.abgleich_vorjahr(neu, ged)
    check("Abgleich: Vorjahr hatte Strecken, neue Ausgabe ohne Maßzahl", len(vergleich) == 1 and "25.5 km" in vergleich[0] and "51.5 km" in vergleich[0], True)
    angewendet = vs.anwenden(neu, ged, overrides={})
    check("anwenden: Seite und Koordinaten der vorigen Ausgabe", (neu[0]["veranstalter_url"], neu[0]["lat"], len(angewendet)),
          ("https://www.munich-triathlon.com/", 48.14, 2))
    frueher = [z("Munich Triathlon", "2026-05-01", 25.5, url="https://my.raceresult.com/2/")]
    vs.anwenden(frueher, ged, overrides={})
    check("anwenden: nie von einer SPÄTEREN Ausgabe (Serie)", frueher[0]["veranstalter_url"], "https://my.raceresult.com/2/")
    eigene = [z("Munich Triathlon 2027", "2027-07-25", 25.5, url="https://neue-seite.de/")]
    vs.anwenden(eigene, ged, overrides={})
    check("anwenden: eine eigene Seite wird nie überschrieben", eigene[0]["veranstalter_url"], "https://neue-seite.de/")
    vs.lernen(neu, ged, overrides={})
    check("lernen: die neue Ausgabe gilt", ged.finde(alt[0])["datum"], "2027-07-25")
    vs.lernen(alt, ged, overrides={})
    check("lernen: die ältere Ausgabe überschreibt die jüngere nicht", (ged.finde(alt[0])["datum"], [s["laenge_km"] for s in ged.finde(alt[0])["strecken"]]), ("2027-07-25", [None]))
    # Plausible nächste Termine: 10-14 Monate nach der vorigen Ausgabe.
    daten = {"2026-11-01", "2027-07-24", "2027-07-25", "2028-07-23", "2027-01-15"}
    check("kandidaten_termine: nur das Fenster 10-14 Monate", vs.kandidaten_termine(daten, "2026-07-26", "2026-10-05"), ["2027-07-24", "2027-07-25"])
    check("kandidaten_termine: ohne vorige Ausgabe alles Künftige bis 14 Monate", vs.kandidaten_termine(daten, None, "2026-10-05"),
          ["2026-11-01", "2027-01-15", "2027-07-24", "2027-07-25"])
    check("spanne: zwei Tage sind eine Ausgabe", vs.spanne(["2027-07-25", "2027-07-24"]), ("2027-07-24", "2027-07-25"))
    check("spanne: zwei Termine im Abstand von Wochen sind mehrdeutig", vs.spanne(["2027-06-01", "2027-07-25"]), None)
    check("spanne: ein Termin", vs.spanne(["2027-07-25"]), ("2027-07-25", "2027-07-25"))
    # Übernahme nach manual_events.json: je Strecke eine Zeile, ohne Jahr im Namen.
    ged2 = vs.Gedaechtnis({})
    vs.lernen(alt, ged2, overrides={})
    eintrag = ged2.finde(alt[0])
    zeilen = vs.manual_events_aus(eintrag, "2027-07-25", "2027-07-25", "2026-10-05")
    check("manual_events_aus: zwei Strecken, Datum, Seite, Quelle, Notiz",
          ([(r["name"], r["datum_start"], r["laenge_km"], r["veranstalter_url"]) for r in zeilen], zeilen[0]["_quelle"], "Gedächtnis" in zeilen[0]["_note"]),
          ([("Munich Triathlon", "2027-07-25", 25.5, "https://www.munich-triathlon.com/"), ("Munich Triathlon", "2027-07-25", 51.5, "https://www.munich-triathlon.com/")],
           "https://www.munich-triathlon.com/", True))
    with tempfile.TemporaryDirectory() as tmp:
        pfad = Path(tmp) / "manual_events.json"
        pfad.write_text(_json.dumps({"_readme": "x", "events": []}), encoding="utf-8")
        b1 = vs.uebernehmen([("k", eintrag, "2027-07-25", "2027-07-25")], "2026-10-05", pfad)
        b2 = vs.uebernehmen([("k", eintrag, "2027-07-25", "2027-07-25")], "2026-10-05", pfad)
        roh = _json.loads(pfad.read_text(encoding="utf-8"))
        check("uebernehmen: einmal angelegt, beim zweiten Mal übersprungen", (len(b1), len(b2), len(roh["events"]), roh["_readme"]), (1, 0, 2, "x"))
        # Dieselbe Ausgabe unter einem zweiten Gedächtnis-Eintrag mit Sponsornamen
        # (19.10.2026, RIBE Citylauf Schwabach): keine zweite Zeile für dieselbe
        # Strecke – sonst stünde sie doppelt in manual_events.json.
        sponsor = dict(eintrag, name="Allianz Munich Triathlon")
        b3 = vs.uebernehmen([("k2", sponsor, "2027-07-25", "2027-07-25")], "2026-10-05", pfad)
        roh = _json.loads(pfad.read_text(encoding="utf-8"))
        check("uebernehmen: Sponsorname derselben Ausgabe legt nichts doppelt an",
              (len(roh["events"]), "anderen Namen" in " ".join(b3)), (2, True))
        # Dieselbe Strecke zweimal IM Eintrag (26.10.2026, Naturpark
        # Schönbuchlauf: 25 km mit und ohne Label): nur einmal, die mit Label.
        doppelt = dict(eintrag, name="Schönbuchlauf", strecken=[
            {"laenge_km": 5.0, "dauer_h": None, "wettbewerb": "5 km", "art2": "Straße"},
            {"laenge_km": 25.0, "dauer_h": None, "wettbewerb": None, "art2": "Straße"},
            {"laenge_km": 25.0, "dauer_h": None, "wettbewerb": "25 km Schönbuchlauf", "art2": "Straße"}])
        b4 = vs.uebernehmen([("k3", doppelt, "2027-10-17", "2027-10-17")], "2026-10-05", pfad)
        roh = _json.loads(pfad.read_text(encoding="utf-8"))
        neu4 = [(e["laenge_km"], e.get("wettbewerb")) for e in roh["events"] if e["name"] == "Schönbuchlauf"]
        check("uebernehmen: doppelte Strecke im Gedächtnis nur einmal, mit Label",
              (neu4, "doppelt im Gedächtnis" in " ".join(b4)), ([(5.0, "5 km"), (25.0, "25 km Schönbuchlauf")], True))
        g = Path(tmp) / "ged.json"
        ged2.speichern(g)
        check("speichern/laden: gültiges JSON, Eintrag unverändert", vs.Gedaechtnis.laden(g).finde(alt[0]), eintrag)
    # Zurückgezogen: Dieselbe Ausgabe trägt die gemerkte Seite nicht mehr
    # (jemand hat sie entfernt - tv-albig.de war der Verein, nicht der
    # Lauf). Dann darf sie nicht an die nächste Ausgabe weitergehen.
    g2 = vs.Gedaechtnis({})
    albig = [z("Albiger Doppelzwölfer", "2026-10-24", None, "12 h", url="https://tv-albig.de/", standort="Albig")]
    vs.lernen(albig, g2, overrides={})
    check("zurückgezogen: Seite zunächst gemerkt", g2.finde(albig[0])["url"], "https://tv-albig.de/")
    albig_ohne = [z("Albiger Doppelzwölfer", "2026-10-24", None, "12 h", url="https://my.raceresult.com/421768/", standort="Albig")]
    rueck = vs.lernen(albig_ohne, g2, overrides={})
    e3 = g2.finde(albig_ohne[0])
    check("zurückgezogen: url leer, quelle zurueckgezogen", (e3["url"], e3["quelle"], any("zurückgezogen" in r for r in rueck)), (None, "zurueckgezogen", True))
    albig_2027 = [z("Albiger Doppelzwölfer", "2027-10-23", None, "12 h", url="https://my.raceresult.com/5/", standort="Albig")]
    vs.anwenden(albig_2027, g2, overrides={})
    check("zurückgezogen: die nächste Ausgabe bekommt sie nicht", albig_2027[0]["veranstalter_url"], "https://my.raceresult.com/5/")
    # ... und dieselbe Ausgabe ohne Seite bekommt sie nie zurück (die
    # Seite der Ausgabe 2026 gilt erst ab 2027).
    g3 = vs.Gedaechtnis({})
    vs.lernen(albig, g3, overrides={})
    vs.anwenden(albig_ohne, g3, overrides={})
    check("anwenden: nie auf dieselbe Ausgabe", albig_ohne[0]["veranstalter_url"], "https://my.raceresult.com/421768/")
    # ---- Seit dem 05.10.2026 (Frage des Nutzers: „auch wenn im nächsten
    # Jahr die Webseite leicht geändert oder neu aufgesetzt wurde, oder aus
    # einem 10 km ein Halbmarathon wurde") ------------------------------
    # Sponsorwechsel im Namen: derselbe Eintrag, wenn der Kern im anderen
    # steckt, der Ort passt und ein Jahr dazwischenliegt.
    g4 = vs.Gedaechtnis({})
    saar = [z("Sparkassen-SAARathon", "2026-10-11", 42.2, url="https://saarathon.de/", standort="Saarbrücken", lat=49.23, lon=6.99)]
    g4.lerne(saar[0], "quelle"); g4.lerne_ausgabe(vs.schluessel_von(saar[0]), saar)
    saar27 = [z("SAARathon 2027", "2027-10-10", 42.2, url="https://my.raceresult.com/9/", standort="Saarbrücken", lat=49.23, lon=6.99)]
    vs.anwenden(saar27, g4, overrides={})
    check("Sponsorwechsel: SAARathon bekommt die Seite von Sparkassen-SAARathon", saar27[0]["veranstalter_url"], "https://saarathon.de/")
    vs.lernen(saar27, g4, overrides={})
    check("Sponsorwechsel: EIN Eintrag, Schlüssel auf den kleineren Kern umgezogen", (len(g4), list(g4.daten)), (1, ["saarathon|saarbrucken"]))
    # Zwei Einträge aus der Zeit vor der unscharfen Suche (Sponsor-Variante
    # unter eigenem Schlüssel) - lernen() legte so etwas bis zum 05.10.2026 an.
    g4b = vs.Gedaechtnis({"asvo king lake of the|attersee": {"url": "https://kotl.at/", "name": "Asvö King of the Lake am Attersee", "standort": "Attersee",
                                                              "lat": 47.87, "lon": 13.55, "datum": "2026-09-26", "quelle": "quelle", "art1": "Fahrrad"},
                          "king lake of the|attersee": {"url": "https://kotl.at/", "name": "King of the Lake", "standort": "Attersee",
                                                         "lat": 47.87, "lon": 13.55, "datum": "2027-09-18", "quelle": "quelle", "art1": "Fahrrad"}})
    kotl = z("King of the Lake", "2027-09-18", 47.2, standort="Attersee", lat=47.87, lon=13.55); kotl["art1"] = "Fahrrad"
    check("auch_unscharf: der genaue UND der unscharfe Eintrag (für die Fälligkeit in pruefen)",
          (len(g4b._kandidaten(kotl)), len(g4b._kandidaten(kotl, auch_unscharf=True))), (1, 2))
    geschwister = [z("Ironman 70.3 Erkner", "2027-09-12", 113, url="https://my.raceresult.com/8/", standort="Erkner", lat=52.42, lon=13.75)]
    g5 = vs.Gedaechtnis({})
    g5.lerne(z("Ironman 5150 Erkner Berlin-Brandenburg", "2027-09-11", 51.5, url="https://ironman5150.de/", standort="Erkner", lat=52.42, lon=13.75), "quelle")
    vs.anwenden(geschwister, g5, overrides={})
    check("Gegenprobe: Geschwister einer Serie am Folgetag sind kein Sponsorwechsel", geschwister[0]["veranstalter_url"], "https://my.raceresult.com/8/")
    g6 = vs.Gedaechtnis({})
    g6.lerne(z("Vienna Triathlon", "2026-09-12", 51.5, url="https://vienna-triathlon.at/", standort="Wien", lat=48.2, lon=16.37), "quelle")
    lauf = [z("Run Vienna", "2027-09-11", 10, url="https://my.raceresult.com/7/", standort="Wien", lat=48.2, lon=16.37)]
    lauf[0]["art1"] = "Laufen"
    vs.anwenden(lauf, g6, overrides={})
    check("Gegenprobe: andere Sportart, Sportwort im Unterschied - kein Treffer", lauf[0]["veranstalter_url"], "https://my.raceresult.com/7/")
    # Neu aufgesetzte Seite: die JÜNGERE Ausgabe gewinnt auch gegen einen Override-Link.
    g7 = vs.Gedaechtnis({})
    g7.lerne(z("Ruhr Trail Run", "2026-06-06", 21.1, url="https://brooks-ruhr-trail-run.de/", standort="Essen"), "override")
    g7.lerne(z("Ruhr Trail Run", "2027-06-05", 21.1, url="https://altra-ruhr-trail-run.de/", standort="Essen"), "quelle")
    check("neue Domain für die neue Ausgabe schlägt den Override-Link der alten", g7.finde(z("Ruhr Trail Run", "2028-06-03", 21.1, standort="Essen"))["url"], "https://altra-ruhr-trail-run.de/")
    g7.lerne(z("Ruhr Trail Run", "2027-06-05", 21.1, url="https://verein-essen.de/", standort="Essen"), "quelle")
    check("… dieselbe Ausgabe, gleicher Rang, anderer Host: Konflikt", g7.finde(z("Ruhr Trail Run", "2028-06-03", 21.1, standort="Essen")).get("konflikt") is not None, True)
    # Datumsangaben in vier Sprachen, zweistelliges Jahr, Schrägstrich.
    from veranstalter_links import norm as _norm
    for text, erwartet in [("Sonntag, 12. Oktober 2027", "2027-10-12"), ("am 12.10.27 um 9 uhr", "2027-10-12"),
                           ("12/10/2027", "2027-10-12"), ("dimanche 12 octobre 2027", "2027-10-12"),
                           ("domenica 12 ottobre 2027", "2027-10-12"), ("October 12, 2027", "2027-10-12"),
                           ("12 October 2027", "2027-10-12"), ("24. März 2027", "2027-03-24")]:
        check(f"daten_aus: {text!r}", vs.daten_aus(_norm(text)), {erwartet})
    check("daten_aus: keine Versionsnummer und keine Uhrzeit", vs.daten_aus(_norm("Version 1.2.27. und 12.10.27.30")), set())
    check("daten_aus: Tagesspanne mit Gedankenstrich (fällt in norm weg)", vs.daten_aus(_norm("Der Termin steht fest: 24.–26. September 2027")), {"2027-09-24", "2027-09-26"})
    check("daten_aus: Tagesspanne mit Schrägstrich und numerisch", vs.daten_aus(_norm("08./09. Mai 2027 und 24.-26.09.2027")), {"2027-05-08", "2027-05-09", "2027-09-24", "2027-09-26"})
    check("daten_aus: 12.10.2027 bleibt EIN Tag", vs.daten_aus(_norm("am 12.10.2027")), {"2027-10-12"})
    check("distanzen_aus: km, Meter, Halbmarathon, Marathon, Italienisch",
          vs.distanzen_aus(_norm("Halbmarathon, 10 km und 5.000 m, Mezza Maratona, 1500 m")), {21.1, 10.0, 5.0, 1.5})
    check("distanzen_aus: Mezza Maratona ist kein Marathon", 42.2 in vs.distanzen_aus(_norm("Mezza Maratona")), False)
    # Strecken gegen die Seite: aus 10 km wurde ein Halbmarathon.
    e10 = {"url": "https://lauf.de/", "name": "Seelauf", "standort": "Erding", "lat": 48.3, "lon": 11.9, "datum": "2026-10-03",
           "art1": "Laufen", "land": "Deutschland",
           "strecken": [{"laenge_km": 10.0, "dauer_h": None, "wettbewerb": "10 km", "art2": "Straße"},
                        {"laenge_km": 5.0, "dauer_h": None, "wettbewerb": "5 km", "art2": "Straße"},
                        {"laenge_km": None, "dauer_h": 1.0, "wettbewerb": "Stundenlauf", "art2": "Straße"}]}
    zl = vs.manual_events_aus(e10, "2027-10-02", "2027-10-02", "2026-10-05", seiten_km={21.1, 5.0})
    check("Streckenprüfung: 10 km fällt (Seite nennt 21,1 und 5), 5 km und das Zeitrennen bleiben",
          [(r.get("laenge_km"), r.get("dauer_h")) for r in zl], [(5.0, None), (None, 1.0)])
    check("Streckenprüfung: die Notiz nennt das Weggefallene und das Neue", ("10 km" in zl[0]["_note"], "21.1 km" in zl[0]["_note"]), (True, True))
    check("Streckenprüfung: jede Zeile trägt gedaechtnis", all(r.get("gedaechtnis") is True for r in zl), True)
    zl2 = vs.manual_events_aus(e10, "2027-10-02", "2027-10-02", "2026-10-05", seiten_km={21.1})
    check("Streckenprüfung: nichts passt mehr - eine Zeile ohne Maßzahl statt erfundener Strecken",
          [(r.get("laenge_km"), r.get("dauer_h")) for r in zl2], [(None, 1.0)])
    zl3 = vs.manual_events_aus(e10, "2027-10-02", "2027-10-02", "2026-10-05", seiten_km=set())
    check("Streckenprüfung: Seite ohne Kilometer - Vorjahr unverändert", len(zl3), 3)
    # Die Prüfung einer Seite: Weiterleitung, tote Unterseite, tote Seite.
    class FakeAbrufer:
        def __init__(self, antworten, umleitung=None):
            self.antworten, self.endadresse = antworten, dict(umleitung or {})
        def hole(self, url):
            return self.antworten.get(url, (404, ""))
    seite = "<html><body>Fuchsburglauf Erding - nächster Termin: Sonntag, 3. Oktober 2027. Strecken: 10 km, 5 km</body></html>"
    v = {"url": "https://lauf.de/2026/ausschreibung", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "override", "art1": "Laufen"}
    r, fund = vs.pruefe_eintrag("k", v, FakeAbrufer({"https://lauf.de/2026/ausschreibung": (404, ""), "https://lauf.de/": (200, seite)}), "2026-10-05")
    check("tote Unterseite: Startseite nennt den Lauf -> Adresse umgezogen, Termin gefunden",
          (v["url"], r.get("umgezogen"), r.get("kandidaten"), r.get("seiten_km")), ("https://lauf.de/", "https://lauf.de/", ["2027-10-03"], [5.0, 10.0]))
    vl = {"url": "https://lauf.de/ausschreibung", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "quelle", "art1": "Laufen"}
    rl, fl = vs.pruefe_eintrag("k", vl, FakeAbrufer({"https://lauf.de/ausschreibung": (200, seite)}), "2026-10-05")
    check("lebende Seite des Laufs nennt den Termin -> NEU mit geprüften Strecken", (rl["flag"], fl[2], sorted(fl[4])), ("NEU", "2027-10-03", [5.0, 10.0]))
    seite_mo = seite.replace("Sonntag, 3. Oktober 2027", "Montag, 11. Oktober 2027")
    vm = {"url": "https://lauf.de/ausschreibung", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "quelle", "art1": "Laufen"}
    rm, fm = vs.pruefe_eintrag("k", vm, FakeAbrufer({"https://lauf.de/ausschreibung": (200, seite_mo)}), "2026-10-05")
    check("anderer Wochentag als die vorige Ausgabe -> nur NEU?", (rm["flag"], fm, rm.get("grund")), ("NEU?", None, "anderer Wochentag als die vorige Ausgabe"))
    check("Feiertag: derselbe Kalendertag zählt wie derselbe Wochentag - nur an festen Tagen",
          (vs._gleicher_wochentag("2027-10-03", "2026-10-03"), vs._gleicher_wochentag("2027-10-02", "2026-10-04"), vs._gleicher_wochentag("2027-09-26", "2026-09-26")), (True, False, False))
    check("tote Unterseite: der Termin von der Startseite ist nur ein Hinweis (NEU?)", (r["flag"], fund, "grund" in r, v.get("seite_ist_startseite")), ("NEU?", None, True, True))
    r_s, fund_s = vs.pruefe_eintrag("k", v, FakeAbrufer({"https://lauf.de/": (200, seite)}), "2026-10-26")
    check("… und bleibt es bei der nächsten Prüfung der Startseite", (r_s["flag"], fund_s), ("NEU?", None))
    v2 = {"url": "https://alt.de/lauf", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "quelle", "art1": "Laufen"}
    r, fund = vs.pruefe_eintrag("k", v2, FakeAbrufer({"https://alt.de/lauf": (200, seite)}, {"https://alt.de/lauf": "https://neu.de/seelauf/"}), "2026-10-05")
    check("Weiterleitung auf eine neue Domain: das Gedächtnis merkt sich die neue Adresse", (v2["url"], r.get("umgezogen")), ("https://neu.de/seelauf/", "https://neu.de/seelauf/"))
    v3 = {"url": "https://weg.de/", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "quelle", "art1": "Laufen"}
    tot = FakeAbrufer({"https://weg.de/": (None, "fehler:ConnectionError")})
    r1, _ = vs.pruefe_eintrag("k", v3, tot, "2026-10-05")
    check("tote Seite: beim ersten Mal nur FEHLER und Zähler", (r1["flag"], v3.get("tot_zaehler"), v3["url"]), ("FEHLER", 1, "https://weg.de/"))
    r2, _ = vs.pruefe_eintrag("k", v3, tot, "2026-10-26")
    check("tote Seite: beim zweiten Mal tot - url leer, Adresse in url_tot", (r2["flag"], v3["url"], v3["url_tot"], v3["quelle"]), ("TOT", None, "https://weg.de/", "tot"))
    g8 = vs.Gedaechtnis({"fuchsburglauf|erding": v3})
    neu_seite = [z("Fuchsburglauf Erding", "2027-10-02", 10, url="https://neue-seite.de/", standort="Erding")]
    check("tote Seite: anwenden gibt nichts weiter", vs.anwenden([z("Fuchsburglauf Erding", "2027-10-02", 10, url="https://my.raceresult.com/5/", standort="Erding")], g8, overrides={}), [])
    vs.lernen(neu_seite, g8, overrides={})
    check("tote Seite: die erste eigene Seite einer Quelle ersetzt sie", (g8.daten["fuchsburglauf|erding"]["url"], g8.daten["fuchsburglauf|erding"]["quelle"]), ("https://neue-seite.de/", "quelle"))
    v4 = {"url": "https://bot.de/", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "quelle", "art1": "Laufen"}
    vs.pruefe_eintrag("k", v4, FakeAbrufer({"https://bot.de/": (403, "")}), "2026-10-05")
    check("403 ist eine Bot-Sperre, keine tote Seite", (v4.get("tot_zaehler"), v4["url"]), (None, "https://bot.de/"))
    vo = {"url": "https://lauf.de/ausschreibung", "name": "Fuchsburglauf Erding", "standort": "Erding", "datum": "2026-10-04", "quelle": "quelle"}
    ro, fo = vs.pruefe_eintrag("k", vo, FakeAbrufer({"https://lauf.de/ausschreibung": (200, seite)}), "2026-10-05")
    check("ohne Sportart im Gedächtnis -> nur NEU? (jede Zeile braucht art1)", (ro["flag"], fo), ("NEU?", None))
    # Übernahme: Vorjahreskopien räumt uebernehmen() wieder weg, wenn sie vorbei sind.
    with tempfile.TemporaryDirectory() as tmp:
        pfad = Path(tmp) / "manual_events.json"
        pfad.write_text(_json.dumps({"events": [{"name": "Alt", "datum_start": "2026-01-01", "datum_ende": "2026-01-01", "gedaechtnis": True},
                                                {"name": "Hand", "datum_start": "2026-01-01", "datum_ende": "2026-01-01"}]}), encoding="utf-8")
        b = vs.uebernehmen([("k", e10, "2027-10-02", "2027-10-02", {21.1, 5.0})], "2026-10-05", pfad)
        roh = _json.loads(pfad.read_text(encoding="utf-8"))
        check("uebernehmen: vergangene Kopie weg, Handeintrag bleibt, neue Ausgabe mit geprüften Strecken",
              ([e["name"] for e in roh["events"]], len(b)), (["Hand", "Seelauf", "Seelauf"], 2))
    # Beim Lernen zählt die eigene Vorjahreskopie nur, solange keine Quelle liefert.
    g9 = vs.Gedaechtnis({})
    kopie = z("Seelauf", "2027-10-02", 10, "10 km", url="https://lauf.de/", standort="Erding"); kopie["gedaechtnis"] = True
    g9.lerne_ausgabe("seelauf|erding", [kopie])
    check("lerne_ausgabe: Kopie gelernt, als übernommen markiert", (g9.daten["seelauf|erding"]["strecken_uebernommen"], [s["laenge_km"] for s in g9.daten["seelauf|erding"]["strecken"]]), (True, [10]))
    g9.lerne_ausgabe("seelauf|erding", [z("Seelauf", "2027-10-02", 21.1, "Halbmarathon", url="https://lauf.de/", standort="Erding")])
    check("lerne_ausgabe: die Quelle ersetzt die Kopie (keine Vereinigung mit 10 km)",
          ([s["laenge_km"] for s in g9.daten["seelauf|erding"]["strecken"]], g9.daten["seelauf|erding"].get("strecken_uebernommen")), ([21.1], None))
    # clean_events: Vorjahreskopien verschwinden, sobald eine Quelle die Ausgabe liefert.
    from clean_events import zurueckziehen_vorjahreskopien, merge_duplicates
    from scraper_lib import update_existing_event
    kopie10 = dict(kopie)
    quelle21 = z("Seelauf Erding 2027", "2027-10-03", 21.1, "Halbmarathon", url="https://lauf.de/", standort="Erding")
    ev, man, ber = zurueckziehen_vorjahreskopien([kopie10, quelle21], [dict(kopie10)])
    check("Vorjahreskopie zurückgezogen: Quelle liefert die Ausgabe (einen Tag daneben, anderer Name)",
          ([e["laenge_km"] for e in ev], man, len(ber)), ([21.1], [], 1))
    ev, man, ber = zurueckziehen_vorjahreskopien([kopie10], [dict(kopie10)])
    check("Vorjahreskopie bleibt, solange keine Quelle liefert", ([e["laenge_km"] for e in ev], len(man)), ([10], 1))
    bestaetigt = dict(kopie10); update_existing_event(bestaetigt, z("Seelauf", "2027-10-02", 10, "10 km", url="https://lauf.de/", standort="Erding"))
    check("update_existing_event: eine Quelle bestätigt die Strecke - Markierung weg", "gedaechtnis" in bestaetigt, False)
    prim = dict(kopie10); prim["veranstalter_url"] = "https://lauf.de/"
    res, _ = merge_duplicates([prim, z("Seelauf", "2027-10-02", 10, "10 km", url="https://my.raceresult.com/1/", standort="Erding")])
    check("merge_duplicates: bestätigt durch eine zweite Zeile - Markierung weg, eine Zeile", (len(res), "gedaechtnis" in res[0]), (1, False))
    # Die Datei im Repo: jeder Eintrag trägt Schlüssel in der erwarteten Form.
    if vs.PFAD.exists():
        d = _json.loads(vs.PFAD.read_text(encoding="utf-8"))
        d.pop("_readme", None)
        check("veranstalter_seiten.json: alle Schlüssel '<kern>|<ort>'", [k for k in d if "|" not in k][:3], [])
        check("veranstalter_seiten.json: jede Seite ist eine eigene (kein Portal)",
              [v["url"] for v in d.values() if v.get("url") and not vs.eigene_seite(v["url"])][:3], [])


def test_pruefung_offen() -> None:
    """Livegang (13.10.2026): Was in der Korrekturliste steht, ist auf der
    Seite ausgeblendet - EINE Regel (audit_events.offene_pruefungen) für
    die Excel-Liste, das Feld in events.json und den Loader der Seite."""
    print("\nAusgeblendet bis geprüft (pruefung_offen):")
    from audit_events import offene_pruefungen, pruefe_event
    from clean_events import mark_pruefung_offen
    wurzel = Path(__file__).resolve().parent.parent
    sauber = {"name": "Testlauf Musterstadt", "datum_start": "2027-05-01", "datum_ende": "2027-05-01",
              "laenge_km": 10.0, "wettbewerb": "10 km", "art1": "Laufen", "art2": "Straße",
              "standort": "Musterstadt", "land": "Deutschland", "lat": 50.1, "lon": 8.6,
              "veranstalter_url": "https://testlauf-musterstadt.de/"}
    check("die saubere Zeile hat keine Audit-Meldung", pruefe_event(sauber), [])
    ohne_km = dict(sauber, name="Lauf ohne Distanz", laenge_km=None, wettbewerb=None)
    geprueft_ohne_km = dict(ohne_km, name="Geprüfter Lauf ohne Distanz")
    vorlaeufig = dict(sauber, name="Prognoselauf", datum_vorlaeufig=True)
    unklar = dict(sauber, name="Unklarer Lauf")
    vergangen = dict(ohne_km, name="Vergangener Lauf", datum_start="2026-01-09", datum_ende="2026-01-09")
    doppelt_a = dict(sauber, name="Doppellauf"); doppelt_b = dict(sauber, name="Doppellauf")
    rows = [sauber, ohne_km, geprueft_ohne_km, vorlaeufig, unklar, vergangen, doppelt_a, doppelt_b]
    protokoll = {"geprüfter lauf ohne distanz|2027-05-01": {"ergebnis": "quelle_ok"},
                 "unklarer lauf|2027-05-01": {"ergebnis": "unklar", "notiz": "Seite tot"}}
    offen = {e["name"]: [a for a, _ in g] for e, g in offene_pruefungen(rows, "2026-10-13", protokoll)}
    check("sauber ist nicht offen", "Testlauf Musterstadt" in offen, False)
    check("ohne Distanz ist offen (audit)", offen.get("Lauf ohne Distanz"), ["audit"])
    check("geprüft ohne Distanz ist NICHT offen (geprueft.json)", "Geprüfter Lauf ohne Distanz" in offen, False)
    check("vorläufiger Termin ist offen", offen.get("Prognoselauf"), ["vorlaeufig"])
    check("frühere Prüfung unklar ist offen", offen.get("Unklarer Lauf"), ["unklar"])
    check("vergangen zählt nicht", "Vergangener Lauf" in offen, False)
    check("exakt doppelte Zeilen sind offen", offen.get("Doppellauf"), ["doppelt"])

    import audit_events as _ae
    _alt = _ae.lade_geprueft
    _ae.lade_geprueft = lambda: protokoll
    try:
        markiert, frei = mark_pruefung_offen(rows, "2026-10-13")
        # Ein geschätzter Termin ALLEIN blendet nicht aus (Sternchen reicht,
        # Nutzer 13.10.2026) - die vier anderen schon.
        check("markiert: vier Zeilen (ohne Distanz, unklar, 2x doppelt) - vorläufig nicht", len(markiert), 4)
        check("das Feld steht nur an offenen Zeilen",
              [e.get("pruefung_offen") for e in rows], [None, True, None, None, True, None, True, True])
        vorlaeufig_ohne_km = dict(vorlaeufig, name="Prognoselauf ohne Distanz", laenge_km=None, wettbewerb=None)
        mark_pruefung_offen(rows + [vorlaeufig_ohne_km], "2026-10-13")
        check("vorläufig PLUS Distanz fehlt ist ausgeblendet", vorlaeufig_ohne_km.get("pruefung_offen"), True)
        ohne_km["laenge_km"] = 10.0; ohne_km["wettbewerb"] = "10 km"
        markiert, frei = mark_pruefung_offen(rows, "2026-10-13")
        check("Distanz nachgetragen - wieder sichtbar", (len(markiert), len(frei), "pruefung_offen" in ohne_km),
              (0, 1, False))
    finally:
        _ae.lade_geprueft = _alt

    # Der Loader der Seite filtert die Zeilen - alle drei Seiten laden darüber.
    quelle = (wurzel / "filters.js").read_text(encoding="utf-8")
    check("EF.loadEvents() wirft pruefung_offen weg (beide Wege)",
          quelle.count("dropUnchecked(") >= 3 and "!e.pruefung_offen" in quelle, True)
    import abo_meldung as _abo
    check("die Abo-Meldung zählt ausgeblendete Zeilen nicht als sichtbar",
          _abo.sichtbar({"name": "X", "datum_start": "2099-01-01", "pruefung_offen": True}, "2026-10-20"), False)


def test_kopierschutz() -> None:
    """Schutz der Datenbank (14.10.2026): robots.txt, Vorbehalt für Text
    und Data Mining, Nutzungsbedingungen, LICENSE und die Fangzeile."""
    print("\nKopierschutz (robots.txt, TDM-Vorbehalt, Fangzeile):")
    import datetime as _dt
    import json
    import fangzeile
    from audit_events import offene_pruefungen, pruefe_event
    wurzel = Path(__file__).resolve().parent.parent

    # robots.txt: Gruppen auswerten wie ein Crawler (jede Gruppe = Agenten + Regeln).
    gruppen, agenten, regeln = {}, [], None
    for zeile in (wurzel / "robots.txt").read_text(encoding="utf-8").splitlines():
        zeile = zeile.split("#", 1)[0].strip()
        if not zeile or ":" not in zeile:
            continue
        feld, wert = (t.strip() for t in zeile.split(":", 1))
        if feld.lower() == "user-agent":
            if regeln is not None:
                agenten, regeln = [], None
            agenten.append(wert.lower())
        elif feld.lower() == "disallow":
            if regeln is None:
                regeln = []
                for a in agenten:
                    gruppen[a] = regeln
            regeln.append(wert)
    check("robots.txt: Googlebot darf die Datendatei (Liste rendern)",
          "/events.web.json" in gruppen.get("googlebot", ["?"]), False)
    check("robots.txt: Googlebot darf die Seiten", "/" in gruppen.get("googlebot", ["/"]), False)
    for bot in ("gptbot", "claudebot", "ccbot", "google-extended", "bytespider"):
        check(f"robots.txt: {bot} gesperrt", gruppen.get(bot), ["/"])
    for pfad in ("/events.json", "/events.web.json", "/kalender/", "/places.json"):
        check(f"robots.txt: alle anderen nicht an {pfad}", pfad in gruppen.get("*", []), True)
    check("robots.txt: alle anderen dürfen die Seiten", "/" in gruppen.get("*", ["/"]), False)

    tdm = json.loads((wurzel / ".well-known" / "tdmrep.json").read_text(encoding="utf-8"))
    check("tdmrep.json: Vorbehalt für die ganze Seite", tdm, [{"location": "/", "tdm-reservation": 1}])
    for seite in ("index.html", "events.html", "karte.html", "impressum.html",
                  "datenschutz.html", "wartung.html", "admin.html"):
        check(f"{seite}: Meta-Tag tdm-reservation",
              '<meta name="tdm-reservation" content="1">' in (wurzel / seite).read_text(encoding="utf-8"), True)
    imp = (wurzel / "impressum.html").read_text(encoding="utf-8")
    check("Impressum: Abschnitt #nutzung, DE und EN",
          ('id="nutzung"' in imp, imp.count("t_nutzung3:"), "87a" in imp, "44b" in imp), (True, 2, True, True))
    lizenz = (wurzel / "LICENSE").read_text(encoding="utf-8")
    check("LICENSE: alle Rechte vorbehalten, Ausnahmen genannt",
          all(w in lizenz for w in ("Alle Rechte vorbehalten", "vendor/", "bilder/", "places.json", "GeoNames")), True)

    # Fangzeile: wandernder Termin, unsichtbar, sauber, genau einmal.
    check("Termin: zweiter Samstag acht Monate voraus",
          [fangzeile.termin(_dt.date(*d)).isoformat() for d in ((2026, 10, 14), (2026, 5, 31), (2027, 1, 1))],
          ["2027-06-12", "2027-01-09", "2027-09-11"])
    z = fangzeile.zeile(_dt.date(2026, 10, 14))
    check("Fangzeile ist ausgeblendet (pruefung_offen)", z.get("pruefung_offen"), True)
    check("Fangzeile hat keine Audit-Meldung (kein Eintrag in der Korrekturliste)",
          (pruefe_event(z), offene_pruefungen([z], "2026-10-14")), ([], []))
    rows = [{"name": "A", "standort": "X", "datum_start": "2027-01-01"},
            {"name": "B", "standort": "X", "datum_start": "2027-12-01"}]
    einmal = fangzeile.einsetzen(rows, _dt.date(2026, 10, 14))
    zweimal = fangzeile.einsetzen(einmal, _dt.date(2026, 10, 14))
    check("einsetzen: in der Datumsfolge, genau einmal, idempotent",
          ([e["name"] for e in einmal], einmal == zweimal), (["A", fangzeile.NAME, "B"], True))
    check("entfernen nimmt sie wieder heraus", fangzeile.entfernen(einmal), rows)
    ev = json.loads((wurzel / "events.json").read_text(encoding="utf-8"))
    treffer = [e for e in ev if fangzeile.ist_fangzeile(e)]
    check("events.json trägt die Fangzeile genau einmal, ausgeblendet",
          [e.get("pruefung_offen") for e in treffer], [True])


def test_statistik() -> None:
    """Besucherstatistik (13.10.2026): stats.js, die Cloud Functions
    `collect`/`aggregateStats`, admin.html und die Regeln.

    Was hier festgehalten wird, läuft sonst still auseinander: die Liste
    der Ereignisse steht in stats.js UND in functions/index.js (was dort
    fehlt, verwirft die Function ohne Meldung), jede Seite mit stats.js
    muss der Function als Seite bekannt sein, der Abschalt-Eintrag heißt
    in stats.js und in der Datenschutzerklärung gleich - und die
    Rechenregeln (Besuche, Absprung, Dauer, Besucher je Tag) werden an
    einem festen Beispiel mit node nachgerechnet.
    """
    import json as _json
    import re as _re
    import shutil
    import subprocess

    print("\nBesucherstatistik (stats.js, Functions, admin.html):")
    wurzel = Path(__file__).resolve().parent.parent
    stats_js = (wurzel / "stats.js").read_text(encoding="utf-8")
    index_js = (wurzel / "functions" / "index.js").read_text(encoding="utf-8")

    def liste(quelle, name):
        m = _re.search(name + r"\s*=\s*(?:new Set\()?\[([^\]]*)\]", quelle)
        return _re.findall(r"['\"]([a-z_]+)['\"]", m.group(1)) if m else None

    ereignisse = liste(stats_js, "var EREIGNISSE")
    check("Ereignisliste in stats.js und functions/index.js gleich",
          liste(index_js, "const STATS_EREIGNISSE"), ereignisse)
    check("Ereignisliste beginnt mit dem Seitenaufruf", ereignisse[:1] if ereignisse else None, ["pv"])
    check("Werbeplatz-Aktionen sind Ereignisse (14.10.2026)",
          all(e in (ereignisse or []) for e in ("anzeige_sicht", "anzeige_klick")), True)

    # Zusatzangaben: Was die Function je Aktion annimmt (STATS_PROPS), muss
    # stats.js überhaupt senden können (PROP_LAENGE) - und jede Aktion mit
    # Zusatzangaben muss ein Ereignis sein.
    m = _re.search(r"const STATS_PROPS = \{(.*?)\n\};", index_js, _re.S)
    props_fn = dict(_re.findall(r"\n\s*([a-z_]+): \{([^}]*)\}", m.group(1))) if m else {}
    m = _re.search(r"var PROP_LAENGE = \{([^}]*)\}", stats_js)
    props_js = set(_re.findall(r"([a-z]+):", m.group(1))) if m else set()
    fehlt = sorted({f for felder in props_fn.values() for f in _re.findall(r"([a-z]+):", felder)} - props_js)
    check("jede Zusatzangabe der Function kennt auch stats.js", fehlt, [])
    check("jede Aktion mit Zusatzangaben ist ein Ereignis", sorted(set(props_fn) - set(ereignisse or [])), [])
    check("Kalender und Teilen tragen Event und Sportart (Bericht je Veranstaltung)",
          all("ev" in props_fn.get(a, "") and "sp" in props_fn.get(a, "") for a in ("kalender", "teilen")), True)

    # Die Listen der Tagesdokumente: AKTIONS_DETAIL in stats_aggregate.js
    # gegen LISTEN_FELDER in admin.html - eine Liste, die die Seite nicht
    # kennt, wird still nicht angezeigt.
    agg_js = (wurzel / "functions" / "stats_aggregate.js").read_text(encoding="utf-8")
    admin_html = (wurzel / "admin.html").read_text(encoding="utf-8")
    m = _re.search(r"const AKTIONS_DETAIL = \{(.*?)\n\};", agg_js, _re.S)
    detail = _re.findall(r"\n\s*([a-z_]+): \[(.*)\]", m.group(1)) if m else []
    check("jede Aktion in AKTIONS_DETAIL ist ein Ereignis", sorted({a for a, _ in detail} - set(ereignisse or [])), [])
    listen_agg = sorted({l for _, paare in detail for _, l in _re.findall(r'\["([a-z]+)", "([a-z_]+)"\]', paare)} | {"sport_region"})
    m = _re.search(r"const LISTEN_FELDER = \[([^\]]*)\]", admin_html)
    listen_admin = set(_re.findall(r"'([a-z_]+)'", m.group(1))) if m else set()
    check("admin.html kennt jede Liste der Tagesdokumente", sorted(set(listen_agg) - listen_admin), [])
    # Drei Listen haben eine eigene Karte (Kreuztabelle, Werbeplatz-Tabelle)
    # statt einer Balkenliste - die brauchen kein l_<name>.
    eigene_karte = {"sport_region", "anzeige_sicht", "anzeige_klick"}
    check("admin.html beschriftet jede Balkenliste (l_<name>) in DE und EN",
          sorted(l for l in listen_agg if l not in eigene_karte and admin_html.count(f"l_{l}:") != 2), [])
    seiten_fn = liste(index_js, "const STATS_SEITEN")
    fehlt = []
    for name in SEITEN:
        html = (wurzel / name).read_text(encoding="utf-8")
        if 'src="stats.js?v=' not in html:
            fehlt.append(f"{name}: lädt stats.js nicht")
        if name[:-5] not in (seiten_fn or []):
            fehlt.append(f"{name}: der Function unbekannt")
    check("jede Seite lädt stats.js und ist der Function bekannt", fehlt, [])
    check("die Function kennt keine Seite, die es nicht gibt",
          sorted(set(seiten_fn or []) - {n[:-5] for n in SEITEN}), [])

    projekt = _re.search(r'projectId:\s*"([^"]+)"', (wurzel / "firebase-config.js").read_text(encoding="utf-8")).group(1)
    check("stats.js schickt an die Function des eigenen Projekts",
          f"https://europe-west1-{projekt}.cloudfunctions.net/collect" in stats_js, True)
    for stelle in ("'localhost'", "'127.0.0.1'", "navigator.webdriver", "globalPrivacyControl", "doNotTrack"):
        check(f"stats.js schweigt bei {stelle}", stelle in stats_js, True)
    ds = (wurzel / "datenschutz.html").read_text(encoding="utf-8")
    check("Abschalt-Eintrag heißt in stats.js und Datenschutz gleich",
          "var SPEICHER_AUS = 'endurance-stats'" in stats_js and "const SPEICHER_STATS = 'endurance-stats'" in ds, True)
    check("Datenschutz nennt die Statistik in DE und EN",
          ds.count("h_statistik"), 3)
    check("Datenschutz hat den Schalter", 'id="stats-schalter"' in ds, True)

    regeln = (wurzel / "firestore.rules").read_text(encoding="utf-8")
    check("Regeln: stats_daily nur für Admins lesbar",
          _re.search(r"match /stats_daily/\{tag\} \{\s*allow read: if istAdmin\(\);\s*allow write: if false;", regeln) is not None, True)
    check("Regeln: hits für Clients gesperrt",
          _re.search(r"match /hits/\{id\} \{\s*allow read, write: if false;", regeln) is not None, True)
    check("Regeln: admins nur das eigene Dokument",
          "match /admins/{uid}" in regeln and "request.auth.uid == uid" in regeln, True)
    check("Regeln: keine E-Mail-Adresse (das Repo ist öffentlich)", "@" in regeln.split("service cloud.firestore")[1], False)

    from stamp_assets import ASSETS, HTML_FILES
    check("stamp_assets.py stempelt stats.js", "stats.js" in ASSETS, True)
    check("stamp_assets.py kennt alle Seiten", sorted(HTML_FILES), sorted(SEITEN))

    node = shutil.which("node")
    if not node:
        print("  (node nicht vorhanden - Rechenprobe übersprungen)")
        return
    skript = r"""
      const a = require(process.argv[1]);
      const T = 1760000000000;
      const pv = (t, h, p, v, s, extra) => Object.assign({ t, h, e: 'pv', p, v, s, r: '', ld: 'DE', la: 'de-DE', sl: 'de', th: 'dark', dv: 'handy', br: 'Safari', os: 'iOS' }, extra || {});
      const hits = [
        pv(T, 9, 'index', 'a', 's1', { r: 'google.com' }),
        pv(T + 60000, 9, 'events', 'a', 's1'),
        pv(T + 3 * 3600000, 12, 'karte', 'a', 's1'),         // 3 h Pause = neuer Besuch
        pv(T + 100, 9, 'events', 'b', 's2', { ld: 'AT', dv: 'rechner' }),
        { t: T + 200, e: 'detail', p: 'events', s: 's2', pr: { ev: 'Berlin Marathon', sp: 'Laufen' } },
        { t: T + 300, e: 'filter', p: 'events', s: 's2', pr: { k: 'sportart', w: 'Laufen' } },
        { t: T + 400, e: 'suche', p: 'events', s: 's2', pr: { q: 'marathon' } },
        { t: T + 500, e: 'kalender', p: 'events', s: 's2', pr: { art: 'ics', ev: 'Berlin Marathon', sp: 'Laufen' } },
        // Zielgruppe (14.10.2026): Sportart je Region kommt aus der Sitzung
        // (s1 = DE), ohne Seitenaufruf der Sitzung "??"; Werbeplatz-Zahlen.
        { t: T + 600, e: 'detail', p: 'karte', s: 's1', pr: { ev: 'Ötztaler Radmarathon', sp: 'Fahrrad' } },
        { t: T + 700, e: 'veranstalter', p: 'karte', s: 's1', pr: { ev: 'Ötztaler Radmarathon', sp: 'Fahrrad' } },
        { t: T + 800, e: 'detail', p: 'events', s: 's9', pr: { ev: 'Seeschwimmen', sp: 'Schwimmen' } },
        { t: T + 900, e: 'anzeige_sicht', p: 'events', s: 's1', pr: { platz: 'liste-oben' } },
        { t: T + 901, e: 'anzeige_sicht', p: 'events', s: 's2', pr: { platz: 'liste-oben' } },
        { t: T + 902, e: 'anzeige_klick', p: 'events', s: 's2', pr: { platz: 'liste-oben' } },
      ];
      const r = a.aggregiere(hits, '2026-10-13');
      console.log(JSON.stringify({
        pv: r.pv, besucher: r.besucher, besuche: r.besuche, absprung: r.absprung, dauer_sek: r.dauer_sek, dauer_n: r.dauer_n,
        einstieg: r.einstieg, referrer: r.referrer, laender: r.laender, geraete: r.geraete, stunden: [r.stunden[9], r.stunden[12]],
        aktionen: r.aktionen, events: r.events, filter: r.filter, suche: r.suche, kalender: r.kalender,
        sport_events: r.sport_events, sport_veranstalter: r.sport_veranstalter, sport_region: r.sport_region,
        events_kalender: r.events_kalender, anzeige_sicht: r.anzeige_sicht, anzeige_klick: r.anzeige_klick,
        land: [a.landAus('Europe/Vienna', 'de-DE'), a.landAus('Europe/Berlin', 'de-AT'), a.landAus('', 'de-CH'),
               a.landAus('America/New_York', 'en-US'), a.landAus('Europe/Oslo', 'nb'), a.landAus('Mars/Olympus', 'xx'), a.landAus(null, null)]
      }));
    """
    erg = subprocess.run([node, "-e", skript, str(wurzel / "functions" / "stats_aggregate.js")],
                         capture_output=True, text=True)
    check("Rechenprobe läuft", erg.returncode, 0)
    if erg.returncode != 0:
        print("  " + erg.stderr.strip()[:300])
        return
    r = _json.loads(erg.stdout)
    check("Seitenaufrufe", r["pv"], 4)
    check("Besucher (verschiedene Kennungen)", r["besucher"], 2)
    check("Besuche (Pause über 30 min trennt)", r["besuche"], 3)
    check("Absprünge (Besuche mit einem Aufruf)", r["absprung"], 2)
    check("Besuchsdauer nur aus Besuchen mit zwei Aufrufen", [r["dauer_sek"], r["dauer_n"]], [60, 1])
    check("Einstiegsseiten je Besuch", r["einstieg"], [{"k": "events", "n": 1}, {"k": "index", "n": 1}, {"k": "karte", "n": 1}])
    check("Referrer je Besuch, ohne Referrer = direkt", r["referrer"], [{"k": "(direkt)", "n": 2}, {"k": "google.com", "n": 1}])
    check("Regionen und Geräte je Seitenaufruf", [r["laender"], r["geraete"]],
          [[{"k": "DE", "n": 3}, {"k": "AT", "n": 1}], [{"k": "handy", "n": 3}, {"k": "rechner", "n": 1}]])
    check("Stunden", r["stunden"], [3, 1])
    check("Aktionen gezählt", r["aktionen"],
          [{"k": "detail", "n": 3}, {"k": "anzeige_sicht", "n": 2}, {"k": "anzeige_klick", "n": 1}, {"k": "filter", "n": 1},
           {"k": "kalender", "n": 1}, {"k": "suche", "n": 1}, {"k": "veranstalter", "n": 1}])
    check("Zusatzangaben der Aktionen", [r["events"], r["filter"], r["suche"], r["kalender"]],
          [[{"k": "Berlin Marathon", "n": 1}, {"k": "Seeschwimmen", "n": 1}, {"k": "Ötztaler Radmarathon", "n": 1}],
           [{"k": "sportart=Laufen", "n": 1}], [{"k": "marathon", "n": 1}], [{"k": "ics", "n": 1}]])
    check("Zielgruppe: geöffnete Events und Veranstalter-Klicks je Sportart",
          [r["sport_events"], r["sport_veranstalter"]],
          [[{"k": "Fahrrad", "n": 1}, {"k": "Laufen", "n": 1}, {"k": "Schwimmen", "n": 1}], [{"k": "Fahrrad", "n": 1}]])
    check("Zielgruppe: Sportart je Region der Sitzung (ohne Seitenaufruf: ??)",
          r["sport_region"], [{"k": "Fahrrad|DE", "n": 1}, {"k": "Laufen|AT", "n": 1}, {"k": "Schwimmen|??", "n": 1}])
    check("Kalender je Event (Bericht je Veranstaltung)", r["events_kalender"], [{"k": "Berlin Marathon", "n": 1}])
    check("Werbeplätze: Einblendungen und Klicks je Platz",
          [r["anzeige_sicht"], r["anzeige_klick"]], [[{"k": "liste-oben", "n": 2}], [{"k": "liste-oben", "n": 1}]])
    check("Region aus Zeitzone, sonst Sprache, sonst Erdteil",
          r["land"], ["AT", "DE", "CH", "AM", "NO", "??", "??"])

    # Kanal Instagram (14.10.2026): utm_source, Referrer instagram.com oder
    # der Browser der App - je BESUCH am ersten Aufruf; dazu, was diese
    # Besucher taten. Und admin.html muss die neuen Listen kennen.
    skript = r"""
      const a = require(process.argv[1]);
      const T = 1760000000000;
      const pv = (t, p, v, s, x) => Object.assign({ t, h: 9, e: 'pv', p, v, s, r: '', br: 'Safari' }, x || {});
      const r = a.aggregiere([
        pv(T, 'index', 'a', 's1', { q: 'instagram', m: 'bio' }), pv(T + 60000, 'events', 'a', 's1'),
        { t: T + 70000, e: 'detail', p: 'events', s: 's1', pr: { ev: 'X', sp: 'Laufen' } },
        pv(T, 'events', 'b', 's2', { r: 'l.instagram.com' }),
        pv(T, 'index', 'c', 's3', { br: 'Instagram' }),
        pv(T, 'index', 'd', 's4', { r: 'google.de' }), pv(T, 'index', 'e', 's5', { r: 'facebook.com' }),
        pv(T, 'index', 'f', 's6', { q: 'newsletter' }), pv(T, 'index', 'g', 's7', { r: 'laufen.de' }), pv(T, 'index', 'h', 's8'),
        pv(T, 'index', 'a', 's9', { q: 'IG' }),
      ], '2026-10-14');
      console.log(JSON.stringify({ k: r.kanaele, ig: r.instagram, m: r.instagram_medien, e: r.instagram_einstieg, a: r.instagram_aktionen }));
    """
    erg = subprocess.run([node, "-e", skript, str(wurzel / "functions" / "stats_aggregate.js")],
                         capture_output=True, text=True)
    check("Rechenprobe Instagram läuft", erg.returncode, 0)
    if erg.returncode == 0:
        r = _json.loads(erg.stdout)
        check("Kanäle je Besuch", sorted((x["k"], x["n"]) for x in r["k"]),
              [("direkt", 1), ("instagram", 4), ("kampagne", 1), ("sozial", 1), ("suche", 1), ("verweis", 1)])
        check("Instagram: Besucher, Besuche, Aufrufe, Absprung, Dauer", r["ig"],
              {"besucher": 3, "besuche": 4, "pv": 5, "absprung": 3, "dauer_sek": 60, "dauer_n": 1})
        check("Instagram: Link (utm_medium) und Einstieg", [r["m"], r["e"]],
              [[{"k": "(ohne)", "n": 3}, {"k": "bio", "n": 1}], [{"k": "index", "n": 3}, {"k": "events", "n": 1}]])
        check("Instagram: Aktionen der Instagram-Sitzungen", r["a"], [{"k": "detail", "n": 1}])
    fehlt = [l for l in ("kanaele", "instagram_medien", "instagram_kampagnen", "instagram_einstieg", "instagram_aktionen")
             if f"'{l}'" not in admin_html or admin_html.count(f"l_{l}:") != 2]
    check("admin.html kennt und beschriftet die Instagram-Listen (DE/EN)", fehlt, [])
    check("stats.js erkennt den Browser der Instagram-App", "br = 'Instagram'" in stats_js, True)


def _insta_testevents(heute):
    """Testdaten für die Instagram-Tests: 60 Zehner in München über zwei
    Wochen, dazu was die Formate seit dem 16.10.2026 brauchen - ein
    130-km-Lauf (krass), ein Backyard (wuerdest), fünf Marathons vier
    Monate voraus (saison)."""
    import datetime as _dt
    ev = []
    for i in range(60):
        tag = heute + _dt.timedelta(days=i % 14)
        ev.append({"name": f"Testlauf {i}", "standort": "München", "land": "Deutschland", "art1": "Laufen",
                   "art2": "Straße", "datum_start": tag.isoformat(), "datum_ende": tag.isoformat(),
                   "lat": 48.14, "lon": 11.57, "laenge_km": 10.0, "veranstalter_url": "https://x.de"})
    tag = (heute + _dt.timedelta(days=20)).isoformat()
    for km in (42.2, 88.0, 130.0):
        ev.append({"name": "Isartal Ultra", "standort": "Wolfratshausen", "land": "Deutschland", "art1": "Laufen",
                   "art2": "Trail", "datum_start": tag, "datum_ende": tag, "lat": 47.91, "lon": 11.42,
                   "laenge_km": km, "veranstalter_url": "https://isartal-ultra.de/"})
    tag = (heute + _dt.timedelta(days=30)).isoformat()
    ev.append({"name": "Backyard Testhof", "standort": "Dachau", "land": "Deutschland", "art1": "Laufen",
               "art2": "Backyard Ultra", "datum_start": tag, "datum_ende": tag, "lat": 48.26, "lon": 11.43,
               "veranstalter_url": "https://backyard-testhof.de/"})
    for i, t in enumerate(("2027-02-06", "2027-02-13", "2027-02-14", "2027-02-20", "2027-02-27")):
        ev.append({"name": f"Wintermarathon {i}", "standort": "Augsburg", "land": "Deutschland", "art1": "Laufen",
                   "art2": "Straße", "datum_start": t, "datum_ende": t, "lat": 48.37, "lon": 10.90,
                   "laenge_km": 42.2, "veranstalter_url": f"https://wintermarathon{i}.de/"})
    return ev


def test_instagram_konten() -> None:
    """Das eigene Instagram-Konto großer Veranstaltungen (seit 27.10.2026):
    die Regel (nur Konten, deren Name aus dem Namen der Veranstaltung
    besteht - nie Verein, Agentur, Sponsor, Tourismus, Schwesterrennen),
    das Sammeln (eigene Seite, Startseite als Rückfall, Fehler behält
    das Alte) und die Zuordnung in clean_events (Override setzt und
    schaltet ab, Portallink zählt nicht, idempotent)."""
    print("\ntest_instagram_konten")
    import datetime as _dt
    import json
    SCRIPTS_DIR = Path(__file__).resolve().parent
    import instagram_konten as ik
    from clean_events import fix_instagram

    k = ik.event_konto
    # Fälle aus der Stichprobe vom 27.10.2026 (41 Veranstalterseiten).
    for konten, name, ort, soll in [
        (["leipzigmarathon"], "Leipzig Marathon", "Leipzig", "leipzigmarathon"),
        (["wachaumarathon"], "WACHAUmarathon", "Krems an der Donau", "wachaumarathon"),
        (["oetztalerradmarathon"], "Ötztaler Radmarathon", "Sölden", "oetztalerradmarathon"),
        (["viking.triathlon"], "6. Viking-Triathlon", "Schleswig", "viking.triathlon"),
        (["neujahrsmarathon.ch"], "Neujahrsmarathon Schlieren", "Schlieren", "neujahrsmarathon.ch"),
        (["evl.halbmarathon"], "EVL-HalbMarathon", "Leverkusen", "evl.halbmarathon"),
        (["cityundtrail"], "Alpenstadt City&Trail", "Bad Reichenhall", "cityundtrail"),
        (["3kings3hills"], "3Kings3Hills Trail Festival 2027", "Haidmühle", "3kings3hills"),
        (["rennsteiglauf.official"], "Rennsteiglauf", "Schmiedefeld am Rennsteig", "rennsteiglauf.official"),
        (["berlinmarathon2027"], "BMW Berlin-Marathon", "Berlin", "berlinmarathon2027"),
        (["brooksrunning", "muenchenmarathon"], "Brooks München Marathon", "München", "muenchenmarathon"),
        (["temposport.erlentriathlon", "temposport.einsiedelntriathlon", "erlentriathlon"],
         "Einsiedeln Triathlon", "Höfe", "temposport.einsiedelntriathlon"),
        (["lhivernalle", "lagueuleduloup_"], "L'Hivern'Alle", "Alle", "lhivernalle"),
        (["halloweenrunbremen"], "Halloween Run Bremen", "Bremen", "halloweenrunbremen"),
        (["algemarin_sv3sathlon"], "Algemarin SV3sathlon", "Mühltal-Traisa", "algemarin_sv3sathlon"),
        (["schwebebahnlauf"], "Schwebebahnlauf Wuppertal 2027", "Wuppertal", "schwebebahnlauf"),
        (["der_frauenfelder"], "Frauenfelder (Halb-)Marathon Wil", "Wil", "der_frauenfelder"),
    ]:
        check(f"Event-Konto: {name}", k(konten, name, ort), soll)
    for konten, name, ort, warum in [
        (["lsvham"], "Lufthansa-Technik-Halbmarathon", "Bramfelder See", "Verein"),
        (["tsvpenzberg"], "TSV Penzberg Volkslauf", "Penzberg", "Verein mit Ort"),
        (["hnt_hamburg"], "Süderelbe Halbmarathon durch das Alte Land", "Hamburg Neugrab", "nur der Ort"),
        (["alohasport.at"], "34. Welser Halbmarathon", "Wels", "Agentur"),
        (["pitztal"], "Pitz Alpine Snow Trail", "Jerzens", "Tourismus"),
        (["brooksrunning"], "Brooks München Marathon", "München", "Sponsor"),
        (["erlentriathlon"], "Einsiedeln Triathlon", "Höfe", "Schwesterrennen (nur 'triathlon')"),
        (["stadtlauf"], "Stadtlauf Erding", "Erding", "ein Allerweltswort"),
        (["exil_ev"], "Lauf gegen Rassismus", "Osnabrück", "Verein"),
        # Durchsicht der ersten 128 Treffer (27.10.2026): der Verein im Namen des Laufs ...
        (["tuswasserstrasse"], "Oktoberlauf mit Mühlenkreisserie - TUS Wasserstraße", "Petershagen", "Verein im Namen"),
        (["alvmainz"], "Nikolauslauf des ALV Mainz", "Mainz", "Vereinskürzel im Namen"),
        (["scc.la"], "62. SCC Cross Country", "Berlin", "Vereinskürzel im Namen"),
        (["rc1886villingen"], "Gravelrennen des RC 1886 Villingen", "Mönchweiler", "Vereinskürzel im Namen"),
        (["llg_wustweiler"], "Freundschaftslauf Wustweiler", "Illingen (Saar)", "Kürzel als erstes Glied"),
        (["tsgaltenbuseck"], "Alten-Busecker Winterlaufserie", "Buseck", "TSG vorn"),
        (["athletik_waldniel"], "Athletik Waldniel Jahresendlauf", "Schwalmtal", "Vereinswort"),
        (["jeunesse_de_fey"], "Cross de la Jeunesse de Fey", "Fey", "Vereinswort"),
        # ... und dieselbe Serie in einer anderen Stadt.
        (["halloweenrunbremen"], "1. Halloween-Run-Rotenburg", "Rotenburg (Wümme)", "fremder großer Ort"),
        (["laufen", "muenchen"], "Halbmarathon München", "München", "zu kurz/nur Ort"),
        ([], "Leipzig Marathon", "Leipzig", "keine Konten"),
        (None, "Leipzig Marathon", "Leipzig", "Seite nie gelesen"),
        (["Bad Name!"], "Leipzig Marathon", "Leipzig", "kein gültiger Kontoname"),
    ]:
        check(f"kein Event-Konto ({warum}): {konten}", k(konten, name, ort), None)
    check("Gleichstand zweier Konten heißt keines",
          k(["leipzigmarathon", "marathonleipzig"], "Leipzig Marathon", "Leipzig"), None)

    # Sammeln mit einer Attrappe statt des Netzes.
    seiten = {
        "https://a.de/lauf": '<a href="https://www.instagram.com/alauf/">',
        "https://b.de/sub": "<p>nichts</p>",
        "https://b.de/": '<footer><a href="https://instagram.com/bmarathon">x</a></footer>',
        "https://d.de/": '<a href="https://instagram.com/p/xyz">Beitrag</a>',
    }

    class Attrappe:
        def _darf(self, url):
            return "c.de" not in url

        def holen(self, url):
            return seiten.get(url)

    heute = _dt.date(2026, 10, 27)
    evs = [{"name": "A-Lauf", "datum_start": "2026-11-01", "veranstalter_url": "https://a.de/lauf"},
           {"name": "B Marathon", "datum_start": "2026-12-01", "veranstalter_url": "https://b.de/sub"},
           {"name": "C", "datum_start": "2026-12-01", "veranstalter_url": "https://c.de/"},
           {"name": "D", "datum_start": "2026-12-02", "veranstalter_url": "https://d.de/"},
           {"name": "E", "datum_start": "2026-12-02", "veranstalter_url": "https://e.de/tot"},
           {"name": "alt", "datum_start": "2026-01-01", "veranstalter_url": "https://alt.de/"},
           {"name": "Portal", "datum_start": "2026-12-01", "veranstalter_url": "https://my.raceresult.com/1/"}]
    daten = {"https://e.de/tot": {"geprueft": "2026-01-01", "konten": ["ealt"]}}
    z = ik.pruefen(evs, daten, heute, 100, 60, abrufer_fabrik=Attrappe, parallel=3)
    check("gesammelt: Konto der Seite", daten["https://a.de/lauf"]["konten"], ["alauf"])
    check("gesammelt: Startseite als Rückfall",
          (daten["https://b.de/sub"]["konten"], daten["https://b.de/sub"].get("quelle")), (["bmarathon"], "startseite"))
    check("robots.txt verbietet: leer, kein Abruf", daten["https://c.de/"]["konten"], [])
    check("Beitrag ist kein Konto", daten["https://d.de/"]["konten"], [])
    check("nicht lesbar: altes Ergebnis bleibt, als Fehler markiert",
          (daten["https://e.de/tot"]["konten"], daten["https://e.de/tot"].get("fehler")), (["ealt"], True))
    check("vergangene Events und Portale werden nicht gelesen",
          ("https://alt.de/" in daten, any("raceresult" in u for u in daten)), (False, False))
    check("Zählung", (z["geprueft"], z["mit_konto"], z["fehler"], z["offen"]), (5, 2, 1, 0))
    z2 = ik.pruefen(evs, daten, heute, 100, 60, abrufer_fabrik=Attrappe, parallel=1)
    check("frisch geprüft: nichts fällig, nur der Fehler kommt früher wieder",
          z2["geprueft"], 0)
    check("Fehler nach drei Wochen wieder fällig",
          ik.faellig(daten["https://e.de/tot"], heute + _dt.timedelta(days=21)), True)
    check("gelesene Seite nach 90 Tagen wieder fällig",
          (ik.faellig(daten["https://a.de/lauf"], heute + _dt.timedelta(days=89)),
           ik.faellig(daten["https://a.de/lauf"], heute + _dt.timedelta(days=90))), (False, True))

    # Zuordnung (clean_events.fix_instagram).
    konten = {"https://www.leipzigmarathon.de/": {"konten": ["leipzigmarathon"], "geprueft": "2026-10-27"},
              "https://tsv.de/": {"konten": ["tsvpenzberg"], "geprueft": "2026-10-27"}}
    rows = [{"name": "Leipzig Marathon", "datum_start": "2027-04-18", "laenge_km": 42.2,
             "standort": "Leipzig", "veranstalter_url": "https://www.leipzigmarathon.de/"},
            {"name": "Leipzig Marathon", "datum_start": "2027-04-18", "laenge_km": 10,
             "standort": "Leipzig", "veranstalter_url": "https://www.leipzigmarathon.de/"},
            {"name": "TSV Penzberg Volkslauf", "datum_start": "2027-05-01", "laenge_km": 10,
             "standort": "Penzberg", "veranstalter_url": "https://tsv.de/", "instagram": "veraltet"},
            {"name": "Kleiner Lauf", "datum_start": "2027-05-01", "laenge_km": 10,
             "standort": "Penzberg", "veranstalter_url": "https://tsv.de/"},
            {"name": "Portal-Lauf", "datum_start": "2027-05-01", "laenge_km": 10, "standort": "X",
             "veranstalter_url": "https://my.raceresult.com/1/"}]
    ov = {"Kleiner Lauf|2027-05-01": {"instagram": "@KleinerLauf", "_note": "von Hand"},
          "Leipzig Marathon|2027-04-18|10": {"instagram": None, "_note": "abgeschaltet"}}
    aend = fix_instagram(rows, konten, ov, grosse=[])
    check("je Zeile: Event-Konto, Override aus, Verein nie, von Hand, Portal nie",
          [r.get("instagram") for r in rows], ["leipzigmarathon", None, None, "kleinerlauf", None])
    check("Änderungen werden gemeldet (gesetzt, entfernt, von Hand)", len(aend), 3)
    check("idempotent", fix_instagram(rows, konten, ov, grosse=[]), [])
    # Die recherchierte Liste der großen Events (grosse_events.json) schlägt
    # die Erkennung, der Override schlägt beide.
    gross = [{"name": "Portal-Lauf", "suche": ["portal-lauf"], "land": None, "sport": None,
              "instagram": "portallauf_official"},
             {"name": "Leipzig", "suche": ["leipzig marathon"], "land": None, "sport": None,
              "instagram": "leipzig_recherchiert"}]
    fix_instagram(rows, konten, ov, grosse=gross)
    check("große Events: recherchiertes Konto, auch hinter einem Portallink; Override gewinnt",
          [r.get("instagram") for r in rows],
          ["leipzig_recherchiert", None, None, "kleinerlauf", "portallauf_official"])
    check("ohne Konto kein Feld in events.json", "instagram" in rows[2], False)
    ik_datei = json.loads((SCRIPTS_DIR / "instagram_konten.json").read_text(encoding="utf-8")) \
        if (SCRIPTS_DIR / "instagram_konten.json").exists() else {}
    falsch = [u for u, v in ik_datei.items()
              if not isinstance(v, dict) or "geprueft" not in v
              or not all(ik.NAME_RE.match(n) for n in v.get("konten") or [])]
    check("instagram_konten.json: jeder Eintrag mit Datum und gültigen Namen", falsch[:3], [])
    ev = json.loads((SCRIPTS_DIR.parent / "events.json").read_text(encoding="utf-8"))
    ungueltig = [e["name"] for e in ev if "instagram" in e and not ik.NAME_RE.match(e["instagram"] or "")]
    check("events.json: jedes instagram-Feld ist ein gültiger Kontoname", ungueltig[:3], [])
    js = (SCRIPTS_DIR.parent / "event-detail.js").read_text(encoding="utf-8")
    check("Box: Zeichen nur mit gültigem Namen, Link auf instagram.com, DE und EN",
          ("/^[a-z0-9._]{1,30}$/.test(e.instagram" in js, "https://www.instagram.com/${encodeURIComponent(insta)}/" in js,
           js.count("detail_instagram:")), (True, True, 2))


def test_insta_freigabe() -> None:
    """Instagram-Freigabe mit einem Klick (14.10.2026): die Kennung des
    Inhalts, der Freigabe-Plan, die Entscheidung des Abendlaufs mit einer
    Firestore-Attrappe - und dass Workflow, Pages-Workflow, Regeln,
    gitignore und Dashboard zusammenpassen."""
    import datetime as _dt
    import json as _json
    import re as _re

    print("\nInstagram-Freigabe (insta_freigabe, plan.json, Dashboard):")
    import insta_inhalte as I
    import insta_freigabe as F
    wurzel = Path(__file__).resolve().parent.parent
    heute = _dt.date(2026, 10, 19)
    ev = _insta_testevents(heute)
    momente = I.lade_momente()
    b = I.plane(heute, ev, momente=momente, clips=[])
    h = I.beitrag_hash(b)
    check("Kennung: 16 Hex-Stellen", bool(_re.fullmatch(r"[0-9a-f]{16}", h)), True)
    check("Kennung: derselbe Beitrag, dieselbe Kennung", I.beitrag_hash(I.plane(heute, ev, momente=momente, clips=[])), h)
    check("Kennung ändert sich mit der Caption", I.beitrag_hash(dict(b, caption=b["caption"] + " x")) == h, False)
    check("Kennung ändert sich mit den Folien", I.beitrag_hash(dict(b, folien=b["folien"][1:])) == h, False)
    check("Kennung ignoriert Felder außerhalb von Format/Folien/Caption/Stories",
          I.beitrag_hash(dict(b, video={"x": 1}, veroeffentlicht=None, erwaehnt=["x"])), h)
    check("Kennung ändert sich mit den Stories", I.beitrag_hash(dict(b, stories=b["stories"][1:])) == h, False)

    plan = I.plan_liste(heute, 14, ev, momente=momente, clips=[])
    check("Plan: 14 Tage", len(plan), 14)
    check("Plan: Felder je Tag", sorted(k for k in plan[0] if k not in ("moment_id", "clip_id")),
          ["caption", "datum", "folien", "format", "hash", "stories", "wochentag"])
    check("Plan: erster Tag = Start, Montag, krass", (plan[0]["datum"], plan[0]["wochentag"], plan[0]["format"]),
          (heute.isoformat(), "Mo", "krass"))
    check("Plan: Stories je Tag mit Art und Folie", all({"art", "folie"} <= set(st) for st in plan[0]["stories"]), True)
    check("Dashboard: Stories-Zeile", ("storiesText(b.stories)" in (wurzel / "admin.html").read_text(encoding="utf-8")), True)
    check("Plan: Kennung je Tag = beitrag_hash des geplanten Beitrags", plan[0]["hash"], h)
    check("Plan: ist JSON", bool(_json.dumps(plan, ensure_ascii=False)), True)

    # Entscheidung des Abendlaufs
    doc = lambda status, hash_: (200, _json.dumps({"fields": {"status": {"stringValue": status}, "hash": {"stringValue": hash_}}}))  # noqa: E731
    check("REST: 404 = keine Freigabe", F.lese_freigabe("2026-10-19", lambda u: (404, "{}")), None)
    check("REST: Dokument gelesen", F.lese_freigabe("2026-10-19", lambda u: doc("freigegeben", h)),
          {"status": "freigegeben", "hash": h, "format": None})
    check("REST: Serverfehler wird genannt", F.lese_freigabe("x", lambda u: (500, "")), {"fehler": "Firestore antwortet 500"})
    check("URL zeigt auf insta_freigaben/<Datum> des Projekts",
          F.url_fuer("2026-10-19").endswith("/documents/insta_freigaben/2026-10-19") and "endurance-5177a" in F.url_fuer("x"), True)
    faelle = [
        ("ohne Freigabe, ohne Autopost: nein", None, False, False),
        ("ohne Freigabe, Autopost: ja", None, True, True),
        ("freigegeben, Kennung passt: ja", {"status": "freigegeben", "hash": h}, False, True),
        ("freigegeben, Kennung veraltet: nein", {"status": "freigegeben", "hash": "alt"}, False, False),
        ("freigegeben, Kennung veraltet, Autopost: ja", {"status": "freigegeben", "hash": "alt"}, True, True),
        ("gesperrt schlägt Autopost", {"status": "gesperrt", "hash": h}, True, False),
        ("Fehler ohne Autopost: nein", {"fehler": "x"}, False, False),
        ("Fehler mit Autopost: ja", {"fehler": "x"}, True, True),
    ]
    for label, frei, auto, erwartet in faelle:
        ok, grund = F.entscheide(h, frei, auto)
        check(f"Entscheidung: {label}", (ok, bool(grund)), (erwartet, True))
    check("Autopost-Variable: ja/yes/true/1, Rest nein",
          [F.autopost_aktiv({"INSTAGRAM_AUTOPOST": v}) for v in ("ja", " Ja ", "yes", "true", "1", "nein", "", "0")],
          [True, True, True, True, True, False, False, False])
    check("darf_posten reicht alles durch", F.darf_posten("2026-10-19", h, lambda u: doc("freigegeben", h), {})[0], True)

    # Die Teile passen zusammen
    wf = (wurzel / ".github" / "workflows" / "instagram.yml").read_text(encoding="utf-8")
    check("instagram.yml: Morgen- und Abendlauf", ("cron: '47 6 * * *'" in wf, "cron: '47 15 * * *'" in wf), (True, True))
    check("instagram.yml: Abendlauf postet mit --freigabe, Handlauf ohne",
          ("--posten --freigabe" in wf, "ABENDLAUF" in wf, "HANDPOSTEN" in wf, "INSTAGRAM_AUTOPOST: ${{ vars.INSTAGRAM_AUTOPOST }}" in wf),
          (True, True, True, True))
    pages = (wurzel / ".github" / "workflows" / "pages.yml").read_text(encoding="utf-8")
    check("pages.yml erzeugt den Plan", "insta_inhalte.py --plan 14 --json marketing/instagram/plan.json" in pages, True)
    check(".gitignore kennt plan.json", "marketing/instagram/plan.json" in (wurzel / ".gitignore").read_text(encoding="utf-8"), True)
    regeln = (wurzel / "firestore.rules").read_text(encoding="utf-8")
    # Bis zur delete-Regel schneiden - ein schlichtes "}" träfe die
    # Klammern des Datumsmusters ([0-9]{4}).
    start = regeln.index("match /insta_freigaben/{datum} {")
    ende = regeln.index("allow delete: if istAdmin();", start)
    block = regeln[start:ende + len("allow delete: if istAdmin();")]
    check("Regeln: insta_freigaben lesen alle, schreiben Admins, Status begrenzt",
          ("allow read: if true;" in block, "istAdmin()" in block, "['freigegeben', 'gesperrt']" in block,
           "hasOnly(['status', 'hash', 'format', 'geaendert'])" in block, "allow delete: if istAdmin();" in block),
          (True, True, True, True, True))
    admin = (wurzel / "admin.html").read_text(encoding="utf-8")
    check("Dashboard: Karte, Plan-Adresse, Sammlung", ('id="insta-freigabe-card"' in admin,
          "'marketing/instagram/plan.json'" in admin, "collection('insta_freigaben')" in admin), (True, True, True))
    for key in ("f_titel", "f_frei", "f_sperren", "f_zurueck", "f_s_veraltet", "f_kein_plan", "f_fmt_wochenvorschau"):
        check(f"Dashboard: Text {key} in DE und EN", admin.count(f"      {key}:"), 2)
    for fmt in I.FORMATE:
        check(f"Dashboard: Formatname {fmt}", f"f_fmt_{fmt}:" in admin, True)
    lauf = (wurzel / "scripts" / "insta_lauf.py").read_text(encoding="utf-8")
    check("insta_lauf.py: --freigabe fragt insta_freigabe, Kennung vor dem Rendern",
          ("--freigabe" in lauf, "freigabe.darf_posten(" in lauf, 'beitrag["hash"] = inhalte.beitrag_hash(beitrag)' in lauf), (True, True, True))


def test_instagram() -> None:
    """Instagram-Account „AusdauerMaschine" (14.10.2026): Wochenplan,
    Captions, Momente, Lizenzprüfung der Clips, der API-Ablauf beim Posten
    (mit einer Attrappe, ohne Netz) und dass der Link in der Bio genau das
    trägt, woran die Statistik Instagram erkennt."""
    import datetime as _dt
    import json as _json
    import random as _random

    print("\nInstagram (insta_inhalte, insta_video, insta_posten):")
    import insta_inhalte as I
    wurzel = Path(__file__).resolve().parent.parent
    heute = _dt.date(2026, 10, 19)            # ein Montag
    ev = _insta_testevents(heute)
    ev.append(dict(ev[0], name="Geheimer Lauf", pruefung_offen=True, laenge_km=42.195))
    ev.append(dict(ev[0], name="Wolfgangsee Junior-Marathon", laenge_km=42.195))
    momente = I.lade_momente()
    formate = [I.plane(heute + _dt.timedelta(days=i), ev, momente=momente, clips=[])["format"] for i in range(7)]
    check("Wochenplan Mo-So (seit 16.10.2026; ohne Clips: Do = Humor)", formate,
          ["krass", "naehe", "wuerdest", "humor", "wochenende", "saison", "moment"])
    woche = [I.plane(heute + _dt.timedelta(days=i), ev, momente=momente, clips=[]) for i in range(7)]
    check("jedes Format beginnt mit Zahl oder Satz (sofort sichtbar), endet mit dem Aufruf",
          [(w["folien"][0]["typ"] in ("hook", "satz", "moment"), w["folien"][-1]["typ"]) for w in woche],
          [(True, "cta")] * 7)
    check("krass: das längste Laufrennen, 130 KM", (woche[0]["folien"][0]["zahl"], woche[0]["folien"][1]["name"]),
          ("130 KM", "Isartal Ultra"))
    mi2 = I.plane(heute + _dt.timedelta(days=9), ev, momente=momente, clips=[])
    check("wuerdest: wechselt wöchentlich die Besonderheit (Ultra, Backyard mit seiner Runde)",
          sorted([(woche[2]["folien"][0]["zahl"], woche[2]["folien"][1]["name"]),
                  (mi2["folien"][0]["zahl"], mi2["folien"][1]["name"])]),
          [("130 KM", "Isartal Ultra"), ("6,7 KM", "Backyard Testhof")])
    check("saison: Marathons im Februar, vier Monate voraus", woche[5]["folien"][0]["titel"], "MARATHONS IM FEBRUAR")
    check("saison: Rennen einzeln und schnell", all(f.get("schnell") for f in woche[5]["folien"] if f["typ"] == "event"), True)
    check("naehe: rund um München, Teilen-Aufruf", (woche[1]["stadt"], "SCHICK DAS" in woche[1]["folien"][-2]["titel"]),
          ("München", True))
    check("Rennen tragen die Veranstalterseite (für die @-Erwähnung)", woche[0]["folien"][1].get("url"),
          "https://isartal-ultra.de/")
    check("Stories: höchstens zwei, bekannte Arten",
          all(len(w["stories"]) <= 2 and {st["art"] for st in w["stories"]} <= {"rennen", "countdown", "reel"} for w in woche), True)
    witze = I.lade_humor()
    check("humor.json: Pflichtfelder, eindeutige Ids, kurze Texte",
          [all(w.get(k) for w in witze for k in ("id", "oben", "text_oben", "unten", "text_unten")),
           len({w["id"] for w in witze}) == len(witze), max(len(w["text_oben"] + w["text_unten"]) for w in witze) <= 90,
           len(witze) >= 12],
          [True, True, True, True])
    check("Caption ohne doppelte Satzzeichen (Frage + Doppelpunkt)", all("?:" not in w["caption"] for w in woche), True)
    check("format_am: Wochenplan ohne Events", [I.format_am(heute + _dt.timedelta(days=i), momente) for i in range(7)], formate)
    b = I.plane(heute, ev, momente=momente, clips=[])
    namen = [e.get("name") for f in b["folien"] for e in [f] + f.get("events", [])]
    check("ausgeblendete Zeilen (pruefung_offen) kommen nie ins Reel", "Geheimer Lauf" in namen, False)
    check("Junior-/Kinderläufe sind kein Highlight", "Wolfgangsee Junior-Marathon" in namen, False)
    check("derselbe Tag ergibt denselben Beitrag", I.plane(heute, ev, momente=momente, clips=[]), b)
    caps = [I.plane(heute + _dt.timedelta(days=i), ev, momente=momente, clips=[])["caption"] for i in range(14)]
    check("Captions: höchstens 2.200 Zeichen, 30 Hashtags, nennen die Seite",
          all(len(c) <= 2200 and c.count("#") <= 30 and I.SEITE in c for c in caps), True)
    check("Einzahl: 1 Radrennen / 1 Schwimm-Event / 2 Läufe",
          [I.sport_anzahl(1, "Fahrrad"), I.sport_anzahl(1, "Schwimmen"), I.sport_anzahl(2, "Laufen")],
          ["1 Radrennen", "1 Schwimm-Event", "2 Läufe"])
    # Jahrestag: Kipchoge 12.10.2019 - 2027 ein Dienstag (ersetzt den
    # Plan), 2026 ein Montag (Montag bleibt bei den Terminen).
    b27 = I.plane(_dt.date(2027, 10, 12), ev, momente=momente, clips=[])
    check("Jahrestag ersetzt den Plan (Di) und sagt „Heute vor 8 Jahren“",
          [b27["format"], b27["folien"][0]["ueber"]], ["moment", "Heute vor 8 Jahren"])
    mo_jahrestag = momente + [dict(momente[0], id="test-montag", datum="2000-10-19")]
    check("Montag bleibt beim Plan, auch am Jahrestag",
          I.plane(heute, ev, momente=mo_jahrestag, clips=[])["format"], "krass")

    pflicht = ("id", "datum", "sport", "titel", "text", "quelle")
    fehler = [m.get("id") for m in momente if any(not m.get(k) for k in pflicht)
              or m["sport"] not in I.SPORTARTEN or not m["quelle"].startswith("https://")]
    for m in momente:
        try:
            _dt.date.fromisoformat(m["datum"])
        except ValueError:
            fehler.append(m["id"])
    check("momente.json: Pflichtfelder, Sportart, Quelle, Datum", fehler, [])
    check("momente.json: Ids eindeutig", len({m["id"] for m in momente}), len(momente))
    clips = _json.loads((wurzel / "marketing" / "instagram" / "clips.json").read_text(encoding="utf-8"))
    check("clips.json ist eine Liste", isinstance(clips.get("clips"), list), True)

    import insta_video as V
    erlaubt, verboten = [], []
    for lizenz in ("CC0", "CC BY 4.0", "CC BY-SA 3.0", "Pexels", "eigenes Material", "Genehmigung Veranstalter 12.10.2026"):
        try:
            V.pruefe_lizenz({"id": "t", "lizenz": lizenz, "urheber": "x"})
            erlaubt.append(lizenz)
        except V.LizenzFehler:
            pass
    for lizenz in ("CC BY-NC 4.0", "CC BY-ND 4.0", "ZDF", "YouTube", ""):
        try:
            V.pruefe_lizenz({"id": "t", "lizenz": lizenz, "urheber": "x"})
        except V.LizenzFehler:
            verboten.append(lizenz)
    check("Lizenzprüfung: freie Lizenzen durch", len(erlaubt), 6)
    check("Lizenzprüfung: NC/ND, Sender, ohne Lizenz abgewiesen", len(verboten), 5)
    for c in clips.get("clips", []):
        try:
            V.pruefe_lizenz(c)
        except V.LizenzFehler as e:
            check(f"clips.json: {c.get('id')}", str(e), "freigegeben")
    check("Emoji fliegen aus dem Videotext", V.ohne_emoji("🏃 12 Läufe 🇩🇪 in Deutschland"), "12 Läufe in Deutschland")
    plan = V.zeitplan([{"typ": "titel"}, {"typ": "liste"}, {"typ": "cta"}], 15.0)
    check("Zeitplan gestreckt auf die Clip-Länge", round(plan[-1][1], 6), 15.0)
    check("letzte Folie bleibt stehen, die erste steht ab dem ersten Bild, spätere blenden ein",
          [V.sichtbarkeit(14.99, plan[-1][0], plan[-1][1], True)[0], V.sichtbarkeit(0.0, 0, 3, False)[0],
           V.sichtbarkeit(3.0, 3, 6, False)[0]], [1.0, 1.0, 0.0])
    import insta_musik as M
    schlag = M.schlag_s(124)
    takt = V.zeitplan(woche[1]["folien"], schlag=schlag)
    check("Zeitplan im Takt: jeder Schnitt auf einem Schlag",
          all(abs(s_ / schlag - round(s_ / schlag)) < 1e-9 for s_, _, _ in takt), True)
    proben, bpm = M.erzeuge(1.0, 7)
    check("Beat: Stereo, 44,1 kHz, Tempo aus der Auswahl, derselbe Tag klingt gleich",
          [len(proben), bpm in M.BPM_AUSWAHL, M.erzeuge(1.0, 7)[0] == proben, M.erzeuge(1.0, 8)[0] == proben],
          [2 * M.RATE, True, True, False])
    for f in woche[0]["folien"] + woche[3]["folien"] + [st["folie"] for st in woche[0]["stories"]]:
        bild = V.folie_bild(f, "Laufen")
        check(f"Folie {f['typ']} wird gezeichnet (1080 × 1920, nicht leer)", [bild.size, bild.getbbox() is not None],
              [(1080, 1920), True])
    import tempfile as _tf0
    with _tf0.TemporaryDirectory() as tmp:
        st_pfad = V.rendere_story(woche[0]["stories"][0], Path(tmp) / "s.jpg", 1)
        from PIL import Image as _Img
        check("Story: JPEG 1080 × 1920", _Img.open(st_pfad).size, (1080, 1920))
    # Hintergründe wechseln (vom Nutzer am 17.10.2026: „immer das gleiche
    # Hintergrundbild"): mehrere Fotos je Sportart, nie zweimal hintereinander.
    hg = _json.loads((wurzel / "marketing" / "instagram" / "hintergruende" / "hintergruende.json")
                     .read_text(encoding="utf-8"))["fotos"]
    for f in hg:
        try:
            V.pruefe_lizenz({"id": f["datei"], "lizenz": f["lizenz"], "urheber": f["urheber"]})
            ok = (wurzel / "marketing" / "instagram" / "hintergruende" / f["datei"]).exists() and \
                f["quelle"].startswith("https://commons.wikimedia.org/")
        except V.LizenzFehler:
            ok = False
        check(f"Hintergrund {f['datei']}: Datei, freie Lizenz, Commons-Quelle", ok, True)
    for sport in ("Laufen", "Fahrrad", "Schwimmen", "Triathlon", None):
        tage = [V.foto_wahl(sport, 740000 + i)[0] for i in range(10)]
        check(f"Hintergrund {sport}: mind. 4 Fotos, nie zweimal hintereinander",
              [len(set(tage)) >= 4, all(a_ != b_ for a_, b_ in zip(tage, tage[1:]))], [True, True])

    import insta_veranstalter as IV
    html = ('<a href="https://www.instagram.com/isartal.ultra/">IG</a> <a href="https://instagram.com/p/abc123/">x</a>'
            '<a href="//www.instagram.com/ausdauermaschine">wir</a> <a href="https://www.instagram.com/explore/tags/lauf/">t</a>')
    check("Instagram-Namen im HTML: Konten ja, Beiträge/Explore/wir selbst nein", IV.namen_im_html(html), ["isartal.ultra"])
    check("ein Name gilt, mehrere nur mit Namenswort",
          [IV.waehle(["isartal.ultra"], "Isartal Ultra", "x.de"),
           IV.waehle(["sponsor_bank", "isartal.ultra"], "Isartal Ultra", "isartal-ultra.de"),
           IV.waehle(["sponsor_bank", "radhaus"], "Isartal Ultra", "isartal-ultra.de")],
          ["isartal.ultra", "isartal.ultra", None])
    domains = IV.portal_domains()
    check("Portale aus scraper_lib.py gelesen (ohne bs4)", ["raceresult.com" in domains, "laufen.de" in domains, len(domains) > 30],
          [True, True, True])

    class AbruferAttrappe:
        def __init__(self):
            self.urls = []

        def holen(self, url):
            self.urls.append(url)
            return '<footer><a href="https://www.instagram.com/isartal.ultra">Instagram</a></footer>' if url.endswith("/") else ""
    ab, cache = AbruferAttrappe(), {}
    check("Portallink: kein Abruf, kein Name",
          [IV.name_fuer("https://my.raceresult.com/1/", "X", ab, cache, heute, domains), ab.urls], [None, []])
    check("Unterseite ohne Link: Startseite gelesen",
          [IV.name_fuer("https://isartal-ultra.de/anmeldung.html", "Isartal Ultra", ab, cache, heute, domains), ab.urls],
          ["isartal.ultra", ["https://isartal-ultra.de/anmeldung.html", "https://isartal-ultra.de/"]])
    ab.urls = []
    check("Cache: 30 Tage kein neuer Abruf",
          [IV.name_fuer("https://isartal-ultra.de/anmeldung.html", "Isartal Ultra", ab, cache, heute, domains), ab.urls],
          ["isartal.ultra", []])
    check("Erwähnung vor den Hashtags",
          IV.caption_mit_erwaehnung("Text\n\nLink\n\n#a #b", ["isartal.ultra"]),
          "Text\n\nLink\n\n📣 Veranstalter: @isartal.ultra\n\n#a #b")
    lauf_q = (wurzel / "scripts" / "insta_lauf.py").read_text(encoding="utf-8")
    check("Erwähnung NACH der Kennung (Freigabe bleibt gültig)",
          lauf_q.index('beitrag["hash"] = inhalte.beitrag_hash(beitrag)') < lauf_q.index("veranstalter.erwaehnungen(beitrag)"), True)

    import insta_posten as P
    import base64 as _b64
    import re as _re
    import os as _os
    import requests as _rq

    class Antwort:
        def __init__(self, daten, status=200):
            self.status_code, self._d = status, daten
            self.text = _json.dumps(daten)
            self.content = self.text.encode()

        def json(self):
            return self._d

    SM = "https://secretmanager.googleapis.com/v1/"
    V = "projects/939878897826/secrets/instagram-token/versions/"

    class Sitzung:
        """Attrappe für Google (Tresor, Speicher) und Instagram - ohne Netz."""

        def __init__(self, erstellt="2026-10-01T08:00:00.123456789Z", tresor_leer=False):
            self.aufrufe, self.status = [], ["IN_PROGRESS", "FINISHED"]
            self.erstellt, self.tresor_leer, self.abgelegt, self.vernichtet = erstellt, tresor_leer, None, []

        def _kurz(self, url):
            for vorn in (SM + "projects/endurance-5177a/secrets/instagram-token", SM + V[:-1],
                         "https://storage.googleapis.com", "https://graph.instagram.com/v25.0",
                         "https://graph.instagram.com"):
                if url.startswith(vorn):
                    return url[len(vorn):] or "/"
            return url

        def get(self, url, headers=None, params=None, **kw):
            k = self._kurz(url)
            self.aufrufe.append(("GET", k))
            if k == "/versions/latest":
                return Antwort({"error": {"code": 404}}, 404) if self.tresor_leer else \
                    Antwort({"name": V + "3", "createTime": self.erstellt, "state": "ENABLED"})
            if url == SM + V + "3:access":
                return Antwort({"payload": {"data": _b64.b64encode(b"ALT-SCHLUESSEL").decode()}})
            if k == "/versions":
                check("Tresor: nur aktive Fassungen auflisten", params, {"filter": "state:ENABLED"})
                return Antwort({"versions": [{"name": V + "4"}, {"name": V + "3"}]})
            if k == "/refresh_access_token":
                check("Erneuern: grant_type ig_refresh_token mit dem alten Schlüssel", params,
                      {"grant_type": "ig_refresh_token", "access_token": "ALT-SCHLUESSEL"})
                return Antwort({"access_token": "NEU-SCHLUESSEL", "token_type": "bearer", "expires_in": 5184000})
            if k == "/me":
                check("Konto: Bearer-Schlüssel, Felder user_id,username",
                      [headers["Authorization"], params["fields"]], ["Bearer NEU-SCHLUESSEL", "user_id,username"])
                return Antwort({"id": "app-scoped", "user_id": "17841", "username": "ausdauermaschine"})
            if "content_publishing_limit" in k:
                return Antwort({"data": [{"quota_usage": 0}]})
            if k == "/m1":
                return Antwort({"permalink": "https://www.instagram.com/reel/abc/"})
            if k == "/c1":
                return Antwort({"status_code": self.status.pop(0)})
            return Antwort({"error": {"message": "unerwartet " + k}}, 400)

        def post(self, url, headers=None, params=None, json=None, data=None, **kw):
            k = self._kurz(url)
            self.aufrufe.append(("POST", k if not k.startswith("/upload/") else "/upload"))
            if k == ":addVersion":
                check("Tresor: neuer Schlüssel als neue Fassung",
                      _b64.b64decode(json["payload"]["data"]), b"NEU-SCHLUESSEL")
                return Antwort({"name": V + "4"})
            if url.endswith(":destroy"):
                self.vernichtet.append(url[len(SM):-len(":destroy")])
                return Antwort({"state": "DESTROYED"})
            if k.startswith("/upload/storage/v1/b/endurance-5177a-instagram/o"):
                check("Speicher: einfacher Upload als video/mp4",
                      [params["uploadType"], headers["Content-Type"], data], ["media", "video/mp4", b"1234"])
                self.abgelegt = params["name"]
                return Antwort({"name": params["name"]})
            if k == "/17841/media":
                check("Container: REELS mit öffentlicher video_url, Caption, Vorschaubild",
                      [json["media_type"], json["video_url"], json["caption"], json["thumb_offset"], json["share_to_feed"]],
                      ["REELS", "https://storage.googleapis.com/endurance-5177a-instagram/" + self.abgelegt, "Hallo",
                       2000, True])
                return Antwort({"id": "c1"})
            if k == "/17841/media_publish":
                check("Veröffentlichen mit creation_id", json, {"creation_id": "c1"})
                return Antwort({"id": "m1"})
            return Antwort({"error": {"message": "unerwartet " + k}}, 400)

        def delete(self, url, **kw):
            self.aufrufe.append(("DELETE", "/storage/v1/b/…/o/" + url.rsplit("/", 1)[-1]))
            return Antwort({}, 204)

    z = {"gcp_token": "g", "projekt": "endurance-5177a", "bucket": "endurance-5177a-instagram",
         "version": "v25.0", "uid": None}
    jetzt = _dt.datetime(2026, 10, 16, 9, 0, tzinfo=_dt.timezone.utc)

    s = Sitzung(erstellt="2026-10-12T08:00:00Z")
    pf = P.pflege_schluessel(z, s, jetzt)
    check("Schlüssel 4 Tage alt: nicht erneuern, nur lesen",
          [pf["token"], pf["alter_tage"], pf["erneuert"], [a for a in s.aufrufe]],
          ["ALT-SCHLUESSEL", 4, False, [("GET", "/versions/latest"), ("GET", "/3:access")]])

    s = Sitzung()
    pf = P.pflege_schluessel(z, s, jetzt)
    check("Schlüssel 15 Tage alt: erneuern, neue Fassung, alte vernichten",
          [pf["token"], pf["erneuert"], pf["gueltig_tage"], s.vernichtet, [a[1] for a in s.aufrufe]],
          ["NEU-SCHLUESSEL", True, 60, [V + "3"],
           ["/versions/latest", "/3:access", "/refresh_access_token", ":addVersion", "/versions",
            "/3:destroy"]])

    try:
        P.pflege_schluessel(z, Sitzung(tresor_leer=True), jetzt)
        check("leerer Tresor meldet sich", "kein Fehler", "Fehler")
    except P.PostenFehler as e:
        check("leerer Tresor nennt den Weg (Einrichtungsskript --schluessel)",
              "instagram_google_einrichten.sh --schluessel" in str(e), True)

    class Kaputt:
        def get(self, url, **kw):
            raise _rq.ConnectionError(f"Max retries exceeded with url: {url}?access_token=GEHEIM")
    try:
        P.erneuere("GEHEIM", Kaputt())
        check("Netzfehler wird gemeldet", "kein Fehler", "Fehler")
    except P.PostenFehler as e:
        check("Netzfehler: nur der Typ, nie die Adresse mit dem Schlüssel", ["GEHEIM" in str(e), "ConnectionError" in str(e)],
              [False, True])

    alt_env = _os.environ.get("GITHUB_ACTIONS")
    import io as _io
    import contextlib as _cl
    puffer = _io.StringIO()
    _os.environ["GITHUB_ACTIONS"] = "true"
    with _cl.redirect_stdout(puffer):
        P.maskiere("GEHEIM")
    _os.environ.pop("GITHUB_ACTIONS")
    with _cl.redirect_stdout(puffer):
        P.maskiere("LOKAL")
    if alt_env is not None:
        _os.environ["GITHUB_ACTIONS"] = alt_env
    check("Schlüssel werden in Actions geschwärzt (::add-mask::), lokal nicht",
          puffer.getvalue(), "::add-mask::GEHEIM\n")

    import tempfile as _tf
    with _tf.TemporaryDirectory() as tmp:
        vid = Path(tmp) / "r.mp4"
        vid.write_bytes(b"1234")
        s = Sitzung()
        with _cl.redirect_stdout(_io.StringIO()):
            erg = P.poste_reel(vid, "Hallo", 2000, z, "NEU-SCHLUESSEL", session=s, warten=lambda _s: None,
                               kennung="reel-2026-10-16")
        check("Posten: Konto → Video ablegen → Container → Status bis FINISHED → Veröffentlichen → Video löschen → Permalink",
              [a for a in s.aufrufe if "limit" not in a[1]],
              [("GET", "/me"), ("POST", "/upload"), ("POST", "/17841/media"), ("GET", "/c1"), ("GET", "/c1"),
               ("POST", "/17841/media_publish"), ("DELETE", "/storage/v1/b/…/o/" + s.abgelegt), ("GET", "/m1")])
        check("Posten: Ergebnis mit Permalink", erg["permalink"], "https://www.instagram.com/reel/abc/")
        check("Video-Name: Datum plus Zufall (16 Hex-Stellen)",
              bool(_re.fullmatch(r"reel-2026-10-16-[0-9a-f]{16}\.mp4", s.abgelegt)), True)

        s = Sitzung()
        s.status = ["ERROR"]
        try:
            with _cl.redirect_stdout(_io.StringIO()):
                P.poste_reel(vid, "Hallo", 2000, z, "NEU-SCHLUESSEL", session=s, warten=lambda _s: None)
            check("Verarbeitungsfehler meldet sich", "kein Fehler", "Fehler")
        except P.PostenFehler:
            check("Verarbeitungsfehler: Video wird trotzdem gelöscht", s.aufrufe[-1][0], "DELETE")

    # Vorschau für den „View"-Knopf (16.10.2026): Reel + Titelbild je
    # zweimal (Kennung und „aktuell"), Namen wie in admin.html.
    class VorschauSitzung:
        def __init__(self):
            self.abgelegt = []

        def post(self, url, headers=None, params=None, data=None, **kw):
            self.abgelegt.append((params["name"], headers["Content-Type"]))
            return Antwort({"name": params["name"]})

    with _tf.TemporaryDirectory() as tmp:
        (Path(tmp) / "reel.mp4").write_bytes(b"mp4")
        (Path(tmp) / "cover.jpg").write_bytes(b"jpg")
        vs = VorschauSitzung()
        adressen = P.lade_vorschau(Path(tmp), "2026-10-17", "abc123", z, session=vs)
        check("Vorschau: Reel und Titelbild unter Kennung und 'aktuell'", sorted(vs.abgelegt), sorted([
            ("vorschau/2026-10-17-abc123.mp4", "video/mp4"), ("vorschau/2026-10-17-aktuell.mp4", "video/mp4"),
            ("vorschau/2026-10-17-abc123.jpg", "image/jpeg"), ("vorschau/2026-10-17-aktuell.jpg", "image/jpeg")]))
        check("Vorschau: öffentliche Adresse", adressen[0],
              "https://storage.googleapis.com/endurance-5177a-instagram/vorschau/2026-10-17-abc123.mp4")
    admin = (Path(__file__).resolve().parent.parent / "admin.html").read_text(encoding="utf-8")
    check("admin.html: Vorschau-Adresse passt zum Speicher",
          "'https://storage.googleapis.com/" + P.VORGABE_BUCKET + "/" + P.VORSCHAU_ORDNER + "/'" in admin, True)

    class StorySitzung:
        """Attrappe für eine Story und die Auswertung (ohne Netz)."""

        def __init__(self, mit_insights=True):
            self.aufrufe, self.abgelegt, self.mit_insights = [], None, mit_insights

        def get(self, url, headers=None, params=None, **kw):
            k = url.replace("https://graph.instagram.com/v25.0", "")
            self.aufrufe.append(("GET", k))
            if k == "/me":
                return Antwort({"user_id": "17841", "username": "ausdauermaschine", "followers_count": 12, "media_count": 2})
            if k == "/s1":
                return Antwort({"status_code": "FINISHED"})
            if k == "/17841/media":
                return Antwort({"data": [{"id": "m1", "timestamp": "2026-10-16T15:02:40+0000", "media_product_type": "REELS",
                                          "permalink": "https://www.instagram.com/reel/abc/", "like_count": 5, "comments_count": 1}]})
            if k == "/m1/insights":
                if not self.mit_insights:
                    return Antwort({"error": {"message": "permission", "code": 10}}, 400)
                return Antwort({"data": [{"name": "reach", "values": [{"value": 300}]}, {"name": "saved", "values": [{"value": 4}]}]})
            return Antwort({"error": {"message": "unerwartet " + k}}, 400)

        def post(self, url, headers=None, params=None, json=None, data=None, **kw):
            k = url.replace("https://graph.instagram.com/v25.0", "")
            if "/upload/storage/" in url:
                check("Story: Bild als image/jpeg abgelegt", [headers["Content-Type"], params["uploadType"]], ["image/jpeg", "media"])
                self.abgelegt = params["name"]
                self.aufrufe.append(("POST", "/upload"))
                return Antwort({"name": params["name"]})
            self.aufrufe.append(("POST", k))
            if k == "/17841/media":
                check("Story: Container STORIES mit öffentlicher image_url",
                      [json["media_type"], json["image_url"]],
                      ["STORIES", "https://storage.googleapis.com/endurance-5177a-instagram/" + self.abgelegt])
                return Antwort({"id": "s1"})
            if k == "/17841/media_publish":
                return Antwort({"id": "sm1"})
            return Antwort({"error": {"message": "unerwartet " + k}}, 400)

        def delete(self, url, **kw):
            self.aufrufe.append(("DELETE", "bild"))
            return Antwort({}, 204)

    with _tf.TemporaryDirectory() as tmp:
        bild = Path(tmp) / "s.jpg"
        bild.write_bytes(b"jpg")
        s = StorySitzung()
        with _cl.redirect_stdout(_io.StringIO()):
            erg = P.poste_story(bild, dict(z, uid=None), "NEU-SCHLUESSEL", session=s, warten=lambda _s: None,
                                kennung="story-2026-10-16-1")
        check("Story: Konto → Bild ablegen → Container → Status → Veröffentlichen → Bild löschen",
              [a for a in s.aufrufe], [("GET", "/me"), ("POST", "/upload"), ("POST", "/17841/media"), ("GET", "/s1"),
                                       ("POST", "/17841/media_publish"), ("DELETE", "bild")])
        check("Story: Ergebnis und Name mit Zufall", [erg["media_id"], bool(_re.fullmatch(r"story-2026-10-16-1-[0-9a-f]{16}\.jpg", s.abgelegt))],
              ["sm1", True])
    aw = P.auswertung(z, "NEU-SCHLUESSEL", session=StorySitzung())
    check("Auswertung: Follower, Likes, Reichweite, Gespeichert",
          [aw["follower"], aw["beitraege"][0]["likes"], aw["beitraege"][0]["reach"], aw["beitraege"][0]["saved"], aw["insights"]],
          [12, 5, 300, 4, True])
    aw = P.auswertung(z, "NEU-SCHLUESSEL", session=StorySitzung(mit_insights=False))
    check("Auswertung ohne Insights-Berechtigung: Likes ja, Hinweis auf die Berechtigung",
          [aw["beitraege"][0]["likes"], aw["insights"], "instagram_business_manage_insights" in (aw["hinweis"] or "")],
          [5, False, True])
    import insta_lauf as L
    text = L.auswertung_text(P.auswertung(z, "NEU-SCHLUESSEL", session=StorySitzung()))
    check("Auswertung als Tabelle mit Format aus dem Wochenplan (16.10.2026 = Fr = wochenende)",
          ["12 Follower" in text, "| wochenende |" in text, "Schnitt je Format" in text], [True, True, True])

    alt_gcp = _os.environ.pop("GCP_ACCESS_TOKEN", None)
    check("ohne Google-Zugang: kein Zugang (nur rendern)", P.zugang(), None)
    if alt_gcp is not None:
        _os.environ["GCP_ACCESS_TOKEN"] = alt_gcp

    # Workflow, Einrichtungsskript und Skript nennen dieselben Namen.
    wf = (wurzel / ".github" / "workflows" / "instagram.yml").read_text(encoding="utf-8")
    sh = (wurzel / "scripts" / "instagram_google_einrichten.sh").read_text(encoding="utf-8")
    check("instagram.yml: Google-Zugang per Workload Identity (id-token, auth@v3, ohne Schlüsseldatei, darf vorher fehlen)",
          ["id-token: write" in wf, "google-github-actions/auth@v3" in wf, "create_credentials_file: false" in wf,
           "continue-on-error: true" in wf, "GCP_ACCESS_TOKEN: ${{ steps.google.outputs.access_token }}" in wf],
          [True, True, True, True, True])
    check("instagram.yml: kein Instagram-Schlüssel mehr in GitHub", ["IG_ACCESS_TOKEN" in wf, "IG_USER_ID" in wf],
          [False, False])
    nummer = _re.search(r"projects/(\d+)/locations/global/workloadIdentityPools/github/providers/enduranceevents", wf)
    fc = (wurzel / "firebase-config.js").read_text(encoding="utf-8")
    check("Anbieter im Workflow: Projektnummer = messagingSenderId aus firebase-config.js",
          bool(nummer) and f'messagingSenderId: "{nummer.group(1)}"' in fc, True)
    check("Einrichtungsskript = Workflow = insta_posten (Projekt, Speicher, Tresor, Dienstkonto, Pool, Anbieter, Version)",
          [f'PROJEKT="{P.VORGABE_PROJEKT}"' in sh, f'BUCKET="{P.VORGABE_BUCKET}"' in sh, f'SECRET="{P.SECRET}"' in sh,
           f'IG_VERSION="{P.VORGABE_VERSION}"' in sh, 'POOL="github"' in sh, 'ANBIETER="enduranceevents"' in sh,
           'DIENSTKONTO_NAME="instagram-reel"' in sh, "instagram-reel@endurance-5177a.iam.gserviceaccount.com" in wf],
          [True] * 8)
    check("Einrichtungsskript: Zugang nur für instagram.yml aus diesem Repository",
          "assertion.workflow_ref.startsWith('${REPO}/.github/workflows/instagram.yml@')" in sh
          and 'REPO="AntDon930/Enduranceevents"' in sh, True)
    check("Einrichtungsskript: Speicher nicht auflistbar, Reste nach einem Tag weg, Schlüssel unsichtbar eingeben",
          ["roles/storage.legacyObjectReader" in sh, "roles/storage.objectViewer" in sh, '"age": 1' in sh,
           "read -r -s" in sh], [True, False, True, True])

    admin = (wurzel / "admin.html").read_text(encoding="utf-8")
    agg = (wurzel / "functions" / "stats_aggregate.js").read_text(encoding="utf-8")
    check("Bio-Link trägt utm_source=instagram (Statistik erkennt ihn) und steht so im Dashboard",
          ["utm_source=instagram" in I.BIO_LINK, f"'{I.BIO_LINK}'" in admin, 'q === "instagram"' in agg],
          [True, True, True])


def test_korrektur_seite() -> None:
    """Korrekturseite (19.10.2026): Die Antworten der Seite gehen über eine
    Excel-Datei im Exportformat in den VORHANDENEN Import; übernommen wird
    nur, was einen Status und noch keine Quittung hat; die Quittung ist ein
    neues Dokument, das die Seite auf „übernommen“ schaltet."""
    print("\nKorrekturseite (Antworten -> Import, Quittungen):")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("  (übersprungen: openpyxl fehlt)")
        return
    import argparse
    import json
    import tempfile
    import korrektur_seite as ks
    from openpyxl import load_workbook
    wurzel = Path(__file__).resolve().parent.parent
    kopf = ["Nr", "Name (nicht änderbar)", "Distanz (km)", "Status", "Deine Anmerkung"]
    def antwort(zid, status, werte=None, geaendert="2026-10-19T10:00:00.000Z"):
        return {"liste": "events", "id": zid, "status": status, "werte": werte or {}, "anmerkung": "",
                "kopf": kopf, "zeile": ["7", "Testlauf", "21,1", "", ""], "geaendert": geaendert,
                "orig": '{"name": "Testlauf", "laenge_km": 21.1}'}
    with tempfile.TemporaryDirectory() as tmp:
        ordner = Path(tmp)
        (ordner / "antworten_events").mkdir()
        (ordner / "quittung_events").mkdir()
        docs = {"e1": antwort("e1", "Korrigiert", {"Distanz (km)": "10,5"}),
                "e2": antwort("e2", "Stimmt so"),
                "e3": antwort("e3", "", {"Distanz (km)": "5"}),            # in Arbeit: bleibt liegen
                "e4": antwort("e4", "Stimmt so", geaendert="2026-10-19T09:00:00.000Z"),
                "e5": antwort("e5", "Unklar", geaendert="2026-10-19T11:00:00.000Z")}
        for zid, d in docs.items():
            (ordner / "antworten_events" / f"{zid}.json").write_text(json.dumps(d), encoding="utf-8")
        # e4 ist mit genau diesem Stand übernommen, e5 nur mit einem ÄLTEREN
        (ordner / "quittung_events" / "20261019-090500.json").write_text(json.dumps(
            {"liste": "events", "ids": {"e4": "2026-10-19T09:00:00.000Z", "e5": "2026-10-19T08:00:00.000Z"}}),
            encoding="utf-8")
        check("Antworten gelesen", sorted(a["id"] for a in ks.lade_antworten(ordner, "events")),
              ["e1", "e2", "e3", "e4", "e5"])
        check("Clubs-Ordner leer", ks.lade_antworten(ordner, "clubs"), [])
        aufrufe = []
        class FakeImport:
            @staticmethod
            def importiere(xlsx, dry_run, aufraeumen=True):
                wb = load_workbook(xlsx)
                ws, orig = wb["Events prüfen"], wb["_original"]
                aufrufe.append({"zeilen": [[c for c in r] for r in ws.iter_rows(min_row=2, values_only=True)],
                                "orig": [list(r) for r in orig.iter_rows(min_row=2, values_only=True)],
                                "kopf": [c.value for c in ws[1]], "dry_run": dry_run})
                return 0
        alt_modul = ks.modul
        ks.modul = lambda liste, clubs_repo=None: FakeImport
        try:
            args = argparse.Namespace(liste="events", ordner=ordner, dry_run=True, ohne_aufraeumen=True, clubs_repo=None)
            ks.cmd_einlesen(args)
            check("Probelauf schreibt keine Quittung", (ordner / "quittung.json").exists(), False)
            args.dry_run = False
            ks.cmd_einlesen(args)
        finally:
            ks.modul = alt_modul
        lauf = aufrufe[-1]
        check("Kopfzeile wie im Export", lauf["kopf"], kopf)
        check("übernommen: e1, e2 und e5 (neuer Stand) – nicht e3 (ohne Status), nicht e4 (quittiert)",
              [(z[0], z[2], z[3]) for z in lauf["zeilen"]],
              [(1, "10,5", "Korrigiert"), (2, "21,1", "Stimmt so"), (3, "21,1", "Unklar")])
        check("_original trägt die Werte beim Export", lauf["orig"][0], [1, docs["e1"]["orig"]])
        quittung = json.loads((ordner / "quittung.json").read_text(encoding="utf-8"))
        check("Quittung nennt genau die übernommenen Stände",
              quittung["ids"], {z: docs[z]["geaendert"] for z in ("e1", "e2", "e5")})
        # Ein zweiter Lauf mit der neuen Quittung übernimmt nichts mehr
        (ordner / "quittung_events" / "20261019-120000.json").write_text(json.dumps(quittung), encoding="utf-8")
        aufrufe.clear()
        ks.modul = lambda liste, clubs_repo=None: FakeImport
        try:
            ks.cmd_einlesen(argparse.Namespace(liste="events", ordner=ordner, dry_run=False,
                                               ohne_aufraeumen=True, clubs_repo=None))
        finally:
            ks.modul = alt_modul
        check("zweiter Lauf: nichts mehr zu übernehmen", aufrufe, [])
    seite = (wurzel / "korrektur" / "korrekturliste.html").read_text(encoding="utf-8")
    for teil in ("'antworten_' + l", "'quittung_' + l", "'daten/' + l + '.json'", "q.has(a.geaendert)"):
        check(f"Seite: {teil}", teil in seite, True)
    check("Seite: kein Feld `stand` mehr (Quittungen entscheiden)", "a.stand" in seite, False)
    # Vom Nutzer am 19.10.2026: „nur eine Excel-Liste, wo ich umschalten kann
    # von Events zu Clubs“ – Blattreiter, Tabelle, Statuszeile, sonst nichts.
    check("Seite: Blattreiter für Events und Clubs",
          all(f'data-liste="{l}"' in seite for l in ("events", "clubs")), True)
    check("Seite: kein Kopf, keine Zähler, keine Suche, keine Hilfe",
          [t for t in ('<header', 'id="zaehler"', 'id="suche"', 'id="hilfe"') if t in seite], [])
    # 20.10.2026: Nr und „Termin vorläufig“ ausgeblendet (vorläufig steht im
    # Datum, „13.10.26 (vorläufig)“), keine festen Spalten mehr – nur die
    # Kopfzeile bleibt oben; der Import liest das zweistellige Jahr.
    check("Seite: Nr und Termin vorläufig ausgeblendet",
          "new Set(['Nr', 'Termin vorläufig', 'id'])" in seite, True)
    check("Seite: keine festen Spalten", ("th.fix" in seite, "position: sticky; top: 0" in seite), (False, True))
    import korrekturliste as kl
    check("Import: Datum mit zweistelligem Jahr",
          (kl.lies_wert("datum_start", "13.10.26"), kl.lies_wert("datum_start", "13.10.2026")),
          ("2026-10-13", "2026-10-13"))


def test_namensrand() -> None:
    """Trenn- und Aufzählungszeichen am Rand eines Namens fliegen
    (Fehlerscan 19.10.2026: "-25. Laufwochenende Brugg"), Klammern,
    Anführungszeichen und Bindestriche IM Namen bleiben."""
    print("\nSatzzeichen am Rand des Namens:")
    import scraper_lib
    import clean_events
    b = scraper_lib.bereinige_name
    check("führender Bindestrich", b("-25. Laufwochenende Brugg"), "25. Laufwochenende Brugg")
    check("Gedankenstrich am Ende", b("23. Lauf um den Allendorfer Bahnhof - Volksbank Sauerland Lauf Cup -"),
          "23. Lauf um den Allendorfer Bahnhof - Volksbank Sauerland Lauf Cup")
    check("Bindestrich im Namen bleibt", b("Run-Bike-Run Ostfildern-Nellingen"), "Run-Bike-Run Ostfildern-Nellingen")
    check("Klammer bleibt", b("(L)auf zur Venus"), "(L)auf zur Venus")
    check("Anführungszeichen bleiben", b("„Rund um“"), "„Rund um“")
    check("nur Satzzeichen: unverändert", b("--"), "--")
    check("None", b(None), None)
    check("Event.to_dict", scraper_lib.Event(name=" • Stadtlauf, ", datum_start="2027-05-01",
                                               standort="Ulm").to_dict()["name"], "Stadtlauf")
    ev = [{"name": "-Testlauf"}, {"name": "Testlauf"}]
    check("fix_namensrand meldet einen", len(clean_events.fix_namensrand(ev)), 1)
    check("fix_namensrand idempotent", clean_events.fix_namensrand(ev), [])


def test_faellt_aus() -> None:
    """Status „Fällt aus" (vom Nutzer am 20.10.2026 am 25h Handorf
    gewünscht: „einen Status der das Event rausnimmt aus der Liste aber
    auch für nächstes Jahr merkt, weil löschen will ich es nicht"): Der
    Import schreibt `exclude` + `abgesagt`, das Gedächtnis behält die
    Veranstaltung mit Seite und Strecken und gibt sie an `pruefen`, die
    Markierung fällt mit der nächsten Ausgabe weg; schwimmkalender.de
    liest eine Absage im Namen nicht als Ort."""
    print("\nStatus „Fällt aus“ (Ausgabe abgesagt, Veranstaltung gemerkt):")
    import json as _json
    import tempfile
    import clean_events
    import schwimmkalender_scraper as sk
    import veranstalter_seiten as vs
    wurzel = Path(__file__).resolve().parent.parent

    # Beide Statuslisten kennen den Wert (die Seite übernimmt ihn aus der
    # Datenprüfung der Excel, korrektur_seite.py lässt ihn durch).
    for datei in ("korrekturliste.py", "korrektur_seite.py"):
        text = (wurzel / "scripts" / datei).read_text(encoding="utf-8")
        zeile = next(z for z in text.splitlines() if z.startswith("STATUS_WERTE"))
        check(f"{datei}: „Fällt aus“ in STATUS_WERTE", "\"Fällt aus\"" in zeile, True)

    def z(name, datum, standort="Handorf", lat=51.99, lon=7.70, url="https://bad-handorf.example/"):
        return dict(name=name, standort=standort, land="Deutschland", lat=lat, lon=lon, art1="Schwimmen",
                    art2="Becken", datum_start=datum, datum_ende=datum, laenge_km=None, dauer_h=25.0,
                    wettbewerb=None, veranstalter_url=url)
    ged = vs.Gedaechtnis({})
    vs.lernen([z("25h Handorf", "2026-10-24")], ged, overrides={})
    k = vs.schluessel_von(z("25h Handorf", "2026-10-24"))
    zeile = z("25h Handorf", "2026-10-24")
    bericht = vs.merke_abgesagt([zeile], ged)
    check("abgesagte Ausgabe gemerkt", (ged.daten[k].get("abgesagt"), len(bericht)), ("2026-10-24", 1))
    check("Eintrag behält Seite und Strecken", (ged.daten[k]["url"], [s["dauer_h"] for s in ged.daten[k]["strecken"]]),
          ("https://bad-handorf.example/", [25.0]))
    check("zweimal gemerkt = einmal", vs.merke_abgesagt([zeile], ged), [])
    # Ohne Zeile in events.json gehört der Eintrag zu denen, die `pruefen`
    # nach der nächsten Ausgabe fragt.
    check("pruefen fragt die abgesagte Veranstaltung ab", k in vs.aktuelle_schluessel(ged, []), False)
    check("nächster Termin im Fenster nach der abgesagten Ausgabe",
          vs.kandidaten_termine({"2027-10-23", "2026-10-24"}, ged.daten[k]["datum"], "2026-10-25"), ["2027-10-23"])
    # Die nächste Ausgabe ist da: die Markierung ist Geschichte.
    vs.lernen([z("25h Handorf", "2027-10-23")], ged, overrides={})
    check("nächste Ausgabe nimmt die Markierung weg", "abgesagt" in ged.daten[k], False)
    # Eine verdorbene Zeile (die Absage im Namen, als Ort gelesen) legt
    # keinen neuen Eintrag an - der Bericht meldet sie.
    vorher = len(ged)
    kaputt = z("25h Handorf FÄLLT AUS!!!", "2026-10-24", standort="FÄLLT AUS!!!", lat=None, lon=None)
    bericht = vs.merke_abgesagt([kaputt], ged)
    check("ohne Eintrag: Warnung, nichts angelegt", (len(ged), bericht[0].startswith("⚠")), (vorher, True))

    # clean_events findet die Zeilen über den Override.
    alt_laden = clean_events.load_manual_overrides
    clean_events.load_manual_overrides = lambda: {
        "25h Handorf|2026-10-24": {"exclude": True, "abgesagt": True},
        "Anderer Lauf|2026-10-24": {"exclude": True}}
    try:
        treffer = clean_events.abgesagte_zeilen([zeile, z("Anderer Lauf", "2026-10-24")])
    finally:
        clean_events.load_manual_overrides = alt_laden
    check("abgesagte_zeilen: nur abgesagt, nicht gelöscht", [e["name"] for e in treffer], ["25h Handorf"])

    # schwimmkalender.de schreibt die Absage in den Namen.
    detail = {"name": "25h Handorf FÄLLT AUS!!!", "datum_start": "2026-10-24", "datum_ende": "2026-10-25",
              "land": "Nordrhein-Westfalen", "links": ["https://bad-handorf.example/"]}
    check("schwimmkalender: abgesagt im Namen = keine Zeile", sk.build_events(detail, 64, sk.CONFIG), [])
    check("schwimmkalender: Absage-Wörter", [bool(sk.ABGESAGT.search(n)) for n in (
        "25h Handorf FÄLLT AUS!!!", "Seeschwimmen abgesagt", "Fällt 2026 aus", "Entfällt", "Kaltwasser-Ausfallschwimmen")],
          [True, True, True, True, False])

    # Der Import schreibt den Override und das Protokoll (braucht openpyxl).
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("  (Import übersprungen: openpyxl fehlt)")
        return
    import korrektur_seite as ks
    import korrekturliste as kl
    kopf = [k for _, k, _ in kl.SPALTEN]
    werte = {k: "" for k in kopf}
    werte.update({"Nr": 1, "Datum Start": "24.10.2026", "Datum Ende": "25.10.2026", "Name (nicht änderbar)": "25h Handorf",
                  "Sportart": "Schwimmen", "Kategorie": "Becken", "Dauer (h)": "25", "Ort": "Handorf",
                  "Land": "Deutschland", "Veranstalter-Link": "https://bad-handorf.example/"})
    orig = {"datum_start": "2026-10-24", "datum_ende": "2026-10-25", "datum_vorlaeufig": None, "name": "25h Handorf",
            "wettbewerb": None, "art1": "Schwimmen", "art2": "Becken", "laenge_km": None, "dauer_h": 25.0,
            "standort": "Handorf", "land": "Deutschland", "veranstalter_url": "https://bad-handorf.example/",
            "lat": 51.99, "lon": 7.70}
    antwort = {"liste": "events", "id": "e1", "status": "Fällt aus", "werte": {}, "anmerkung": "",
               "kopf": kopf, "zeile": [werte[k] for k in kopf], "orig": _json.dumps(orig, ensure_ascii=False)}
    pfade = ("MANUAL_OVERRIDES_PATH", "GEPRUEFT_PATH", "EVENTS_JSON_PATH")
    alt = {n: getattr(kl, n) for n in pfade}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for n, inhalt in zip(pfade, ({}, {}, [dict(orig)])):
            (tmp / f"{n}.json").write_text(_json.dumps(inhalt), encoding="utf-8")
            setattr(kl, n, tmp / f"{n}.json")
        try:
            ks.baue_excel("events", [antwort], tmp / "a.xlsx")
            rc = kl.importiere(tmp / "a.xlsx", False, False)
            ov = _json.loads((tmp / "MANUAL_OVERRIDES_PATH.json").read_text(encoding="utf-8"))
            prot = _json.loads((tmp / "GEPRUEFT_PATH.json").read_text(encoding="utf-8"))
        finally:
            for n, v in alt.items():
                setattr(kl, n, v)
    eintrag = ov.get("25h Handorf|2026-10-24", {})
    check("Import: exclude + abgesagt", (rc, eintrag.get("exclude"), eintrag.get("abgesagt")), (0, True, True))
    check("Import: Protokoll „abgesagt“", prot.get("25h Handorf|2026-10-24", {}).get("ergebnis"), "abgesagt")


def test_check_again() -> None:
    """Status „Check again" (vom Nutzer am 20.10.2026: „Ich schreibe meine
    Kommentare dann hinein, und du müsstest das Event noch einmal genau
    anschauen"): Die Zeile geht NICHT in den Import - weder auf der Seite
    noch in der Excel -, sie steht als Auftrag in `nochmal.json`; nach der
    Prüfung quittiert `erledigt` sie mit einer Notiz, die die Seite zeigt."""
    print("\nStatus „Check again“ (Auftrag an Claude statt Import):")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("  (übersprungen: openpyxl fehlt)")
        return
    import argparse
    import json as _json
    import tempfile
    import korrekturliste as kl
    import korrektur_seite as ks
    from openpyxl import load_workbook
    wurzel = Path(__file__).resolve().parent.parent
    check("Excel-Auswahl kennt „Check again“", kl.NOCHMAL in kl.STATUS_WERTE, True)
    check("Seite: „Check again“ geht nie in den Import", ks.NOCHMAL in ks.STATUS_WERTE, False)

    kopf = [k for _, k, _ in kl.SPALTEN]
    orig = {"datum_start": "2027-05-01", "datum_ende": "2027-05-01", "datum_vorlaeufig": None, "name": "Testlauf",
            "wettbewerb": None, "art1": "Laufen", "art2": "Straße", "laenge_km": 10.0, "dauer_h": None,
            "standort": "Ulm", "land": "Deutschland", "veranstalter_url": "https://testlauf.example/",
            "lat": 48.4, "lon": 10.0}
    werte = {k: "" for k in kopf}
    werte.update({"Nr": 1, "Datum Start": "01.05.2027", "Datum Ende": "01.05.2027", "Name (nicht änderbar)": "Testlauf",
                  "Sportart": "Laufen", "Kategorie": "Straße", "Distanz (km)": "10", "Ort": "Ulm", "Land": "Deutschland",
                  "Veranstalter-Link": "https://testlauf.example/"})
    def antwort(zid, status, anm="", w=None, geaendert="2026-10-20T10:00:00.000Z"):
        return {"liste": "events", "id": zid, "status": status, "werte": w or {}, "anmerkung": anm,
                "kopf": kopf, "zeile": [werte[k] for k in kopf], "geaendert": geaendert,
                "orig": _json.dumps(orig, ensure_ascii=False)}
    with tempfile.TemporaryDirectory() as tmp:
        ordner = Path(tmp)
        (ordner / "antworten_events").mkdir()
        docs = {"e1": antwort("e1", "Stimmt so"),
                "e2": antwort("e2", "Check again", "Ich glaube, es gibt auch 21 km", {"Distanz (km)": "21,1"})}
        for zid, d in docs.items():
            (ordner / "antworten_events" / f"{zid}.json").write_text(_json.dumps(d), encoding="utf-8")
        aufrufe = []
        class FakeImport:
            GEPRUEFT_PATH = ordner / "geprueft.json"
            @staticmethod
            def importiere(xlsx, dry_run, aufraeumen=True):
                ws = load_workbook(xlsx)["Events prüfen"]
                aufrufe.append([r[kopf.index("Status")] for r in ws.iter_rows(min_row=2, values_only=True)])
                return 0
        FakeImport.GEPRUEFT_PATH.write_text("{}", encoding="utf-8")
        alt_modul = ks.modul
        ks.modul = lambda liste, clubs_repo=None: FakeImport
        try:
            args = argparse.Namespace(liste="events", ordner=ordner, dry_run=False, ohne_aufraeumen=True, clubs_repo=None)
            ks.cmd_einlesen(args)
            auftrag = _json.loads((ordner / "nochmal.json").read_text(encoding="utf-8"))
            quittung = _json.loads((ordner / "quittung.json").read_text(encoding="utf-8"))
            check("Import sieht nur „Stimmt so“", aufrufe, [["Stimmt so"]])
            check("Quittung nur für die übernommene Zeile", sorted(quittung["ids"]), ["e1"])
            e = auftrag["eintraege"][0]
            check("Auftrag: Name, Link, Kommentar, Zellen, Schlüssel",
                  (len(auftrag["eintraege"]), e["id"], e["name"], e["link"], e["anmerkung"], e["werte"], e["schluessel"]),
                  (1, "e2", "Testlauf", "https://testlauf.example/", "Ich glaube, es gibt auch 21 km",
                   {"Distanz (km)": "21,1"}, "Testlauf|2027-05-01"))
            # Claude beantwortet den Auftrag; ein zweites `einlesen` behält die Antwort.
            e.update(antwort="21,1 km gibt es – nachgetragen.", ergebnis="korrigiert")
            (ordner / "nochmal.json").write_text(_json.dumps(auftrag), encoding="utf-8")
            ks.cmd_einlesen(args)
            check("zweites einlesen behält die Antwort",
                  _json.loads((ordner / "nochmal.json").read_text(encoding="utf-8"))["eintraege"][0]["antwort"],
                  "21,1 km gibt es – nachgetragen.")
            args.dry_run = True
            ks.cmd_erledigt(args)
            check("erledigt --dry-run schreibt nichts",
                  ((ordner / "quittung-nochmal.json").exists(), _json.loads(FakeImport.GEPRUEFT_PATH.read_text())), (False, {}))
            args.dry_run = False
            rc = ks.cmd_erledigt(args)
            q = _json.loads((ordner / "quittung-nochmal.json").read_text(encoding="utf-8"))
            prot = _json.loads(FakeImport.GEPRUEFT_PATH.read_text(encoding="utf-8"))
            check("Quittung mit Stand und Notiz",
                  (rc, q["ids"], q["notizen"]), (0, {"e2": docs["e2"]["geaendert"]}, {"e2": "21,1 km gibt es – nachgetragen."}))
            check("Protokoll: Ergebnis und Kommentar",
                  (prot["Testlauf|2027-05-01"]["ergebnis"], "Check again" in prot["Testlauf|2027-05-01"]["notiz"]),
                  ("korrigiert", True))
            e.update(antwort="", ergebnis="korrigiert")
            (ordner / "nochmal.json").write_text(_json.dumps(auftrag), encoding="utf-8")
            check("Ergebnis ohne Antwort wird abgewiesen", ks.cmd_erledigt(args), 1)
            e.update(antwort="x", ergebnis="passt")
            (ordner / "nochmal.json").write_text(_json.dumps(auftrag), encoding="utf-8")
            check("unbekanntes Ergebnis wird abgewiesen", ks.cmd_erledigt(args), 1)
        finally:
            ks.modul = alt_modul

    # Der Excel-Weg: Die Zeile wird übersprungen, auch mit geänderter Zelle.
    pfade = ("MANUAL_OVERRIDES_PATH", "GEPRUEFT_PATH", "EVENTS_JSON_PATH")
    alt = {n: getattr(kl, n) for n in pfade}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for n, inhalt in zip(pfade, ({}, {}, [dict(orig)])):
            (tmp / f"{n}.json").write_text(_json.dumps(inhalt), encoding="utf-8")
            setattr(kl, n, tmp / f"{n}.json")
        try:
            ks.baue_excel("events", [antwort("e2", "Check again", "bitte prüfen", {"Distanz (km)": "21,1"})], tmp / "a.xlsx")
            rc = kl.importiere(tmp / "a.xlsx", False, False)
            ov = _json.loads((tmp / "MANUAL_OVERRIDES_PATH.json").read_text(encoding="utf-8"))
            prot = _json.loads((tmp / "GEPRUEFT_PATH.json").read_text(encoding="utf-8"))
        finally:
            for n, v in alt.items():
                setattr(kl, n, v)
    check("Excel-Import: „Check again“ schreibt weder Override noch Protokoll", (rc, ov, prot), (0, {}, {}))

    seite = (wurzel / "korrektur" / "korrekturliste.html").read_text(encoding="utf-8")
    for teil in ("const NOCHMAL = 'Check again'", "tr.s-nochmal", "v.notizen", "claude-notiz"):
        check(f"Seite: {teil}", teil in seite, True)


def test_abo_meldung() -> None:
    """E-Mail-Abos erst nach dem Deploy und mit freigegebenen Zeilen (vom
    Nutzer am 20.10.2026: „Ja, bitte beide Punkte ändern"): Neu ist, was
    jetzt SICHTBAR ist und beim letzten Melden nicht. Gemeldet wird einmal
    je Datenlauf (sonst wanderte die Wochen-Mail mit jedem Korrektur-
    Deploy), erst im Job nach dem Deploy; ein Rückfall meldet nichts."""
    print("\nE-Mail-Abos (abo_meldung.py):")
    import argparse
    import datetime as _dt
    import json as _json
    import os
    import socket as _socket
    import tempfile
    import urllib.error
    import abo_meldung as a
    import fangzeile
    heute = "2026-10-20"

    def z(name, tag="2026-11-07", km=10.0, **mehr):
        e = {"name": name, "datum_start": tag, "datum_ende": tag, "laenge_km": km, "standort": "Ulm"}
        e.update(mehr)
        return e

    vorher = [z("Bleibt"), z("Ausgeblendet", pruefung_offen=True), z("Bleibt versteckt", pruefung_offen=True),
              z("Gestrichen")]
    nachher = [z("Bleibt"), z("Ausgeblendet"), z("Bleibt versteckt", pruefung_offen=True), z("Ganz neu"),
               z("Ganz neu"), z("Neu, aber vorbei", tag="2026-10-01"), z("Neu und versteckt", pruefung_offen=True),
               z("Bleibt", km=21.1), fangzeile.zeile(_dt.date(2026, 10, 20)) | {"pruefung_offen": False}]
    namen = [e["name"] for e in a.neue_sichtbare(vorher, nachher, heute)]
    check("neue Zeile aus dem Datenlauf wird gemeldet", "Ganz neu" in namen, True)
    check("je Kennung nur einmal", namen.count("Ganz neu"), 1)
    check("freigegebene Zeile (vorher pruefung_offen) wird gemeldet", "Ausgeblendet" in namen, True)
    check("weiter ausgeblendete Zeile nicht", "Bleibt versteckt" in namen or "Neu und versteckt" in namen, False)
    check("schon sichtbare Zeile nicht", namen.count("Bleibt"), 1)  # nur die neue 21,1-km-Strecke
    check("vergangene Zeile nicht", "Neu, aber vorbei" in namen, False)
    check("Fangzeile nie", fangzeile.NAME in namen, False)
    check("Kennung wie dedupe_key (Distanz gerundet)", a.kennung(z("X", km=42.195)), ("x", "2026-11-07", 42.2))

    gross = [z(f"Lauf {i}") for i in range(a.MIN_VORHER)]
    A, B, C = "a" * 40, "b" * 40, "c" * 40   # gemeldet, live, dieser Deploy
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "events.json").write_text(_json.dumps(gross + [z("Ganz neu")]), encoding="utf-8")
        aus = tmp / "abo" / "abo-neu.json"
        dj = tmp / "deploy.json"
        alt_stand, alt_git = a.live_stand, a.git

        def lauf(stand, vorfahr=True, datenlauf=True, inhalt=gross, bot=A):
            a.live_stand = lambda seite, cache: stand
            def git(*g):
                if g[0] == "cat-file":
                    return ""
                if g[0] == "merge-base":
                    return "" if vorfahr else None
                if g[0] == "show":
                    return _json.dumps(inhalt)
                if g[0] == "log":
                    if g[1] == "-1":
                        return bot + "\n" if bot else ""
                    return "d" * 40 + "\n" if datenlauf else ""
                return None
            a.git = git
            if dj.exists():
                dj.unlink()
            a.cmd_vergleichen(argparse.Namespace(seite="https://x.example", cache="1", commit=C, deploy_json=str(dj),
                                                 nachher=str(tmp / "events.json"), aus=str(aus), heute=heute))
            d = _json.loads(aus.read_text(encoding="utf-8"))
            meldestand = _json.loads(dj.read_text(encoding="utf-8")) if dj.exists() else None
            return [e["name"] for e in d["events"]], d.get("grund"), meldestand

        try:
            namen, grund, ds = lauf(("ok", {"commit": B, "gemeldet": A}))
            check("neuer Datenlauf: alles seit der letzten Meldung", (namen, grund), (["Ganz neu"], None))
            check("… und dieser Stand gilt danach als gemeldet", ds, {"commit": C, "gemeldet": C})
            namen, grund, ds = lauf(("ok", {"commit": B, "gemeldet": A}), datenlauf=False)
            check("Korrektur-Deploy ohne Datenlauf: nichts melden, Meldestand bleibt",
                  (namen, grund, ds), ([], None, {"commit": C, "gemeldet": A}))
            namen, grund, ds = lauf(("ok", {"commit": B, "gemeldet": A}), vorfahr=False)
            check("Rückfall (Live-Commit kein Vorfahr): nichts melden, kein Fehler",
                  (namen, grund, ds), ([], None, {"commit": C, "gemeldet": A}))
            namen, grund, ds = lauf(("ok", {"commit": B, "gemeldet": A}), inhalt=gross[:10])
            check("kaputter gemeldeter Stand: nichts melden, mit Grund", (namen, grund is not None), ([], True))
            namen, grund, ds = lauf(("fehlt", {}), datenlauf=False)
            check("ohne deploy.json: der letzte Datenlauf gilt als gemeldet",
                  (namen, grund, ds), ([], None, {"commit": C, "gemeldet": A}))
            namen, grund, ds = lauf(("fehler", {}))
            check("deploy.json nicht lesbar: nichts melden, mit Grund, ohne Meldestand",
                  (namen, grund is not None, ds), ([], True, {"commit": C}))
        finally:
            a.live_stand, a.git = alt_stand, alt_git

        # senden: ohne Datei / mit Grund rot, leer grün, ohne Adresse grün.
        alt_env = {k: os.environ.pop(k, None) for k in ("NOTIFY_WEBHOOK_URL", "NOTIFY_WEBHOOK_SECRET")}
        try:
            check("senden ohne Liste: rot", a.main(["senden", "--datei", str(tmp / "fehlt.json")]), 1)
            a.schreibe(aus, [], "kaputt")
            check("senden mit Grund: rot", a.main(["senden", "--datei", str(aus)]), 1)
            a.schreibe(aus, [])
            check("senden ohne neue Events: grün, kein Aufruf", a.main(["senden", "--datei", str(aus)]), 0)
            a.schreibe(aus, [z("Ganz neu")])
            check("senden ohne NOTIFY_WEBHOOK_URL: grün", a.main(["senden", "--datei", str(aus)]), 0)
        finally:
            for k, v in alt_env.items():
                if v is not None:
                    os.environ[k] = v

    # melde: 503 wird wiederholt, eine Zeitüberschreitung nie (die Mails
    # könnten schon in der Warteschlange liegen).
    import urllib.request as _ur
    alt_open, alt_sleep = _ur.urlopen, a.time.sleep
    aufrufe = []

    class Antwort:
        def __enter__(self):
            return self
        def __exit__(self, *x):
            return False
        def read(self):
            return b'{"checkedSubscriptions": 2, "mailsQueued": 2, "collectedForLater": 0}'

    try:
        a.time.sleep = lambda s: None
        def offen_503(req, timeout=0):
            aufrufe.append(req.get_header("X-notify-secret"))
            if len(aufrufe) < 2:
                raise urllib.error.HTTPError(req.full_url, 503, "busy", {}, None)
            return Antwort()
        _ur.urlopen = offen_503
        antwort = a.melde("https://x.example/hook", "geheim", [z("Ganz neu")])
        check("503 einmal wiederholt, dann Antwort gelesen", (len(aufrufe), antwort["mailsQueued"]), (2, 2))
        check("Secret im Kopf", aufrufe[0], "geheim")
        aufrufe.clear()
        def zu_langsam(req, timeout=0):
            aufrufe.append(1)
            raise _socket.timeout("timed out")
        _ur.urlopen = zu_langsam
        try:
            a.melde("https://x.example/hook", None, [z("Ganz neu")])
            fehler = None
        except RuntimeError as exc:
            fehler = str(exc)
        check("Zeitüberschreitung: kein zweiter Aufruf", (len(aufrufe), fehler is not None), (1, True))
    finally:
        _ur.urlopen, a.time.sleep = alt_open, alt_sleep

    wurzel = Path(__file__).resolve().parent.parent
    update = (wurzel / "scripts" / "update_events.py").read_text(encoding="utf-8")
    check("der Datenlauf meldet nicht mehr selbst", "os.environ.get(\"NOTIFY_WEBHOOK_URL\")" in update, False)
    pages = (wurzel / ".github" / "workflows" / "pages.yml").read_text(encoding="utf-8")
    abo = pages[pages.index("\n  abo:"):] if "\n  abo:" in pages else ""
    check("pages.yml: Job „abo“ läuft erst nach dem Deploy", "needs: deploy" in abo, True)
    check("pages.yml: der Job meldet mit den Secrets",
          "abo_meldung.py senden" in abo and "secrets.NOTIFY_WEBHOOK_URL" in abo, True)
    upload = pages.index("upload-pages-artifact")
    check("pages.yml: Vergleich samt deploy.json VOR dem Upload",
          0 < pages.index("abo_meldung.py vergleichen") < upload and "--deploy-json deploy.json" in pages, True)
    check("pages.yml: die Liste liegt nicht in der hochgeladenen Seite", "$RUNNER_TEMP/abo/abo-neu.json" in pages, True)
    check("Secrets nur im Job „abo“", pages.count("secrets.NOTIFY_WEBHOOK"), abo.count("secrets.NOTIFY_WEBHOOK"))


def test_abo_mail() -> None:
    """Die Abo-Mail im Stil der Seite (vom Nutzer am 20.10.2026: „Ja, bitte
    die schönere Mail-Vorlage bauen"): rendert den echten Abschnitt aus
    functions/index.js in node - eine Karte je Veranstaltung, Längen wie
    in der Länge-Spalte (am ganzen Bestand gegen EED.formatLength),
    Datum mit Wochentag, maskierte Daten, Betreff ohne „für Alle neuen
    Events", höchstens MAX_KARTEN Karten."""
    print("\nAbo-Mail (functions/index.js, aboMail):")
    import json as _json
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        print("  (übersprungen: node fehlt)")
        return
    wurzel = Path(__file__).resolve().parent.parent
    quelle = (wurzel / "functions" / "index.js").read_text(encoding="utf-8")
    check("Abschnitt mit Marken in functions/index.js",
          "// ==== Abo-Mail: Anfang ====" in quelle and "// ==== Abo-Mail: Ende ====" in quelle, True)
    check("Einmal-Abos nutzen dieselbe Vorlage", "message: einmalMail(treffer[0])" in quelle, True)

    def lauf(name, tag, km=None, **mehr):
        e = {"name": name, "datum_start": tag, "datum_ende": tag, "standort": "Ulm", "land": "Deutschland",
             "art1": "Laufen", "art2": "Straße", "laenge_km": km, "veranstalter_url": "https://lauf.example/"}
        e.update(mehr)
        return e
    beispiele = [
        lauf("Stadtlauf", "2026-10-25", 21.1), lauf("Stadtlauf", "2026-10-25", 10.0), lauf("Stadtlauf", "2026-10-25", 5.0),
        lauf("Cross <b>X</b>", "2026-11-07", None, art1="Fahrrad", art2="Cyclocross", dauer_h=0.5,
             veranstalter_url="javascript:alert(1)"),
        lauf("Seetriathlon", "2027-06-12", 25.75, art1="Triathlon", art2="Straße", datum_ende="2027-06-13"),
        lauf("Herbstlauf", "2027-09-01", 10.0, datum_vorlaeufig=True, charity=True),
    ]
    wieviele = [lauf(f"Lauf {i}", "2026-12-01", 10.0) for i in range(40)]
    skript = (
        "const fs=require('fs');"
        f"const src=fs.readFileSync({_json.dumps(str(wurzel / 'functions' / 'index.js'))},'utf8');"
        "const tri=src.slice(src.indexOf('const TRIATHLON_FORMAT_IM_LABEL'),src.indexOf('const TRIATHLON_FORMAT_KATEGORIE'));"
        "const teil=src.slice(src.indexOf('// ==== Abo-Mail: Anfang ===='),src.indexOf('// ==== Abo-Mail: Ende ===='));"
        "const defineString=(n,o)=>({value:()=>o.default});"
        "const m=new Function('defineString','SITE_URL',tri+teil+';return {eventKurz,aboMail,einmalMail,laengeText,MAX_KARTEN};')"
        "(defineString,defineString('',{default:'https://endurance-events.de'}));"
        "global.window=global;"
        f"require({_json.dumps(str(wurzel / 'filters.js'))});"
        "global.EnduranceFilters.relativeDays=()=>'';"
        f"require({_json.dumps(str(wurzel / 'event-detail.js'))});"
        "const EED=global.EnduranceDetail;"
        f"const ev=JSON.parse(fs.readFileSync({_json.dumps(str(wurzel / 'events.json'))},'utf8')).filter(e=>!e.pruefung_offen);"
        "const abw=ev.filter(e=>(m.laengeText(m.eventKurz(e))||'–')!==EED.formatLength(e,'de')).map(e=>e.name).slice(0,5);"
        "const b=JSON.parse(process.argv[1]), w=JSON.parse(process.argv[2]);"
        "const kurz=x=>x.map(m.eventKurz);"
        "const mit=m.aboMail('id1',{name:'Laufen · 25 km um Ulm',rhythmus:'woechentlich',token:'t'},kurz(b),0);"
        "const ohne=m.aboMail('id1',{name:'Alle neuen Events',rhythmus:'monatlich'},kurz(b),3);"
        "const eins=m.aboMail('id1',{name:'',rhythmus:'sofort'},kurz(b.slice(0,1)),0);"
        "const viel=m.aboMail('id1',{name:'X',rhythmus:'woechentlich'},kurz(w),0);"
        "const einmal=m.einmalMail(b[0]);"
        "console.log(JSON.stringify({abw,anzahl:ev.length,mit,ohne,eins,viel_karten:(viel.html.match(/Zum Event/g)||[]).length,"
        "viel_html:viel.html.length,viel_text:viel.text.includes('… und 10 weitere Events'),einmal,max:m.MAX_KARTEN}));"
    )
    r = subprocess.run([node, "-e", skript, _json.dumps(beispiele), _json.dumps(wieviele)], capture_output=True, text=True)
    if r.returncode != 0:
        check("Abo-Mail läuft in node", r.stderr.strip()[:300], "")
        return
    d = _json.loads(r.stdout)
    check(f"Längen wie in der Länge-Spalte ({d['anzahl']} sichtbare Zeilen)", d["abw"], [])
    mit, ohne, eins = d["mit"], d["ohne"], d["eins"]
    check("Betreff mit Filter", mit["subject"], "4 neue Events: Laufen · 25 km um Ulm")
    check("Betreff ohne Filter (kein „für Alle neuen Events“)", ohne["subject"], "4+ neue Events auf Endurance Events")
    check("Betreff bei einem Event", eins["subject"], "Neues Event: Stadtlauf")
    html = mit["html"]
    check("eine Karte je Veranstaltung", html.count("Zum Event"), 4)
    check("Längen kürzeste zuerst, Dezimalkomma", "5km &middot; 10km &middot; 21,1km" in html, True)
    check("unter einer Stunde in Minuten", "30min" in html, True)
    check("Triathlon als Format", ">Sprint<" in html, True)
    check("Datum mit Wochentag", "So, 25.10.2026" in html, True)
    check("mehrtägig", "Sa, 12.06. – So, 13.06.2027" in html, True)
    check("vorläufiger Termin nur als Monat mit *", "September 2027*" in html and "Termin noch zu bestätigen" in html, True)
    check("Charity-Marke", "&#9829; Charity" in html, True)
    check("Name maskiert", "Cross &lt;b&gt;X&lt;/b&gt;" in html and "<b>X</b>" not in html, True)
    check("nur http(s) wird verlinkt", "javascript:" in html, False)
    check("Link in die Liste mit Suche", "https://endurance-events.de/events.html?s=Stadtlauf" in html, True)
    check("Abmelde-Link und Impressum", "Abo beenden" in html and "/impressum.html" in html, True)
    check("keine Bilder (kein Lesebeleg)", "<img" in html.lower(), False)
    check("Textfassung mit Link", "https://endurance-events.de/events.html?s=Stadtlauf" in mit["text"], True)
    check("übersprungene Strecken genannt", "3 weitere Strecken" in ohne["html"], True)
    check("höchstens MAX_KARTEN Karten", d["viel_karten"], d["max"])
    check("Rest als Zahl (auch im Text)", d["viel_text"], True)
    check("Mail bleibt unter der Gmail-Grenze (~100 KB)", d["viel_html"] < 90000, True)
    check("Einmal-Abo im selben Gewand", d["einmal"]["subject"] == "Neues Event gefunden: Stadtlauf"
          and "Zum Event" in d["einmal"]["html"], True)


def test_geloescht() -> None:
    """Status „Löschen" (vom Nutzer am 20.10.2026: „löschen soll nur die
    eine Zeile herausnehmen und schauen ob das Event nächstes Jahr passend
    ist"): Der Import setzt `geloescht`, das Gedächtnis nimmt die Strecke aus
    der Vorjahreskopie und merkt sie; die passende Zeile der NÄCHSTEN
    Ausgabe steht wieder in der Korrekturliste - die übrigen Termine einer
    Serie und spätere Jahre nicht, und nach der Prüfung auch nicht mehr."""
    print("\nStatus „Löschen“ (Zeile raus, nächste Ausgabe wird geprüft):")
    import json as _json
    import tempfile
    import clean_events
    import veranstalter_seiten as vs
    from audit_events import offene_pruefungen

    def z(name, datum, km=None, wettbewerb=None, dauer=None, standort="Ladbergen"):
        return dict(name=name, standort=standort, land="Deutschland", lat=52.13, lon=7.74, art1="Laufen",
                    art2="Straße", datum_start=datum, datum_ende=datum, laenge_km=km, dauer_h=dauer,
                    wettbewerb=wettbewerb, veranstalter_url="https://ladbergen-ultra.example/")
    ged = vs.Gedaechtnis({})
    zeilen = [z("Ladbergen Ultramarathon", "2026-10-14", 21.1, "Halbmarathon"),
              z("Ladbergen Ultramarathon", "2026-10-14", 42.2, "Marathon"),
              z("Ladbergen Ultramarathon", "2026-10-14", 44.0, "Ultra 44 km")]
    vs.lernen(zeilen, ged, overrides={})
    k = vs.schluessel_von(zeilen[0])
    bericht = vs.merke_geloescht([zeilen[2]], ged, "2026-10-20")
    eintrag = ged.daten[k]
    check("gelöscht gemerkt, Strecke aus der Vorjahreskopie",
          (len(bericht), [st["laenge_km"] for st in eintrag["strecken"]], len(eintrag["geloescht"]),
           eintrag["geloescht"][0].get("ganz")), (1, [21.1, 42.2], 1, None))
    check("zweimal gemerkt = einmal", vs.merke_geloescht([zeilen[2]], ged, "2026-10-20"), [])
    check("Vorjahreskopie ohne die gelöschte Strecke",
          [r["laenge_km"] for r in vs.manual_events_aus(eintrag, "2027-10-13", "2027-10-13", "2027-09-01")],
          [21.1, 42.2])
    naechste = z("Ladbergen Ultramarathon", "2027-10-13", 44.0, "Ultra 44 km")
    check("nächste Ausgabe derselben Strecke: wieder eine Frage",
          [len(ged.geloescht_fuer(r)) for r in (
              naechste, z("Ladbergen Ultramarathon", "2027-10-13", 43.8),        # neu vermessen
              z("Ladbergen Ultramarathon", "2027-10-13", 42.2),                  # andere Strecke
              z("Ladbergen Ultramarathon", "2026-11-20", 44.0),                  # gleiche Saison
              z("Ladbergen Ultramarathon", "2028-10-11", 44.0))],                # zwei Jahre später
          [1, 1, 0, 0, 0])
    arten = lambda r, prot: [a for e, g in offene_pruefungen([r], "2027-01-01", prot, gedaechtnis=ged)  # noqa: E731
                             for a, _ in g]
    check("Korrekturliste: Grund „geloescht“", "geloescht" in arten(naechste, {}), True)
    check("nach der Prüfung nicht mehr", "geloescht" in arten(naechste, {
        "ladbergen ultramarathon|2027-10-13": {"ergebnis": "quelle_ok"}}), False)
    check("ohne Gedächtnis kein Grund", "geloescht" in [a for e, g in offene_pruefungen([naechste], "2027-01-01", {})
                                                         for a, _ in g], False)
    # Die ganze Veranstaltung gelöscht: nächstes Jahr passt jede Strecke.
    h = z("Hartberg bewegt sich", "2026-10-17", 6.0, standort="Hartberg")
    vs.lernen([h], ged, overrides={})
    vs.merke_geloescht([h], ged, "2026-10-20")
    check("ganze Veranstaltung: ganz, jede Strecke passt",
          (ged.daten[vs.schluessel_von(h)]["geloescht"][0].get("ganz"),
           len(ged.geloescht_fuer(z("Hartberg bewegt sich", "2027-10-16", 10.0, standort="Hartberg")))), (True, 1))
    check("ganz bleibt beim zweiten Aufruf einmal", (vs.merke_geloescht([h], ged, "2026-10-20"),
                                                     len(ged.daten[vs.schluessel_von(h)]["geloescht"])), ([], 1))
    vs.merke_geloescht([], ged, "2028-06-01")
    check("alte Markierungen räumt der Aufruf weg", ("geloescht" in ged.daten[k],
                                                     "geloescht" in ged.daten[vs.schluessel_von(h)]), (False, False))
    check("ohne Eintrag: Warnung, nichts angelegt",
          (vs.merke_geloescht([z("Gibt es nicht", "2026-10-14", 5.0, standort="Nirgendwo")], ged)[0][:1],
           len(ged.daten)), ("⚠", 2))

    alt_laden = clean_events.load_manual_overrides
    clean_events.load_manual_overrides = lambda: {
        "Ladbergen Ultramarathon|2026-10-14|44": {"exclude": True, "geloescht": True},
        "Ladbergen Ultramarathon|2026-10-14|42.2": {"exclude": True},
        "Ladbergen Ultramarathon|2026-10-14|21.1": {"exclude": True, "geloescht": True, "abgesagt": True}}
    try:
        treffer = clean_events.geloeschte_zeilen(zeilen)
    finally:
        clean_events.load_manual_overrides = alt_laden
    check("geloeschte_zeilen: nur Löschen der Korrekturliste", [e["laenge_km"] for e in treffer], [44.0])

    # Der Import setzt die Markierung (braucht openpyxl).
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("  (Import übersprungen: openpyxl fehlt)")
        return
    import korrektur_seite as ks
    import korrekturliste as kl
    kopf = [k for _, k, _ in kl.SPALTEN]
    werte = {k: "" for k in kopf}
    werte.update({"Nr": 1, "Datum Start": "14.10.2026", "Datum Ende": "14.10.2026",
                  "Name (nicht änderbar)": "Ladbergen Ultramarathon", "Wettbewerb / Strecke": "Ultra 44 km",
                  "Sportart": "Laufen", "Kategorie": "Straße", "Distanz (km)": "44", "Ort": "Ladbergen",
                  "Land": "Deutschland", "Veranstalter-Link": "https://ladbergen-ultra.example/"})
    orig = dict(zeilen[2], datum_vorlaeufig=None)
    antwort = {"liste": "events", "id": "e1", "status": "Löschen", "werte": {}, "anmerkung": "",
               "kopf": kopf, "zeile": [werte[k] for k in kopf], "orig": _json.dumps(orig, ensure_ascii=False)}
    pfade = ("MANUAL_OVERRIDES_PATH", "GEPRUEFT_PATH", "EVENTS_JSON_PATH")
    alt = {n: getattr(kl, n) for n in pfade}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for n, inhalt in zip(pfade, ({}, {}, [dict(orig)])):
            (tmp / f"{n}.json").write_text(_json.dumps(inhalt), encoding="utf-8")
            setattr(kl, n, tmp / f"{n}.json")
        try:
            ks.baue_excel("events", [antwort], tmp / "a.xlsx")
            rc = kl.importiere(tmp / "a.xlsx", False, False)
            ov = _json.loads((tmp / "MANUAL_OVERRIDES_PATH.json").read_text(encoding="utf-8"))
        finally:
            for n, v in alt.items():
                setattr(kl, n, v)
    eintrag = ov.get("Ladbergen Ultramarathon|2026-10-14|44", {})
    check("Import: exclude + geloescht, nicht abgesagt",
          (rc, eintrag.get("exclude"), eintrag.get("geloescht"), eintrag.get("abgesagt")), (0, True, True, None))


def test_insta_logos() -> None:
    """Logo-Reels (27.10.2026): am letzten Freitag „monat" (Zahl aller
    Events des nächsten Monats, dann die Logos der großen Events im Takt),
    am 1. „anmeldung" (große Events, deren Anmeldung in dem Monat öffnet);
    die Logo-Suche auf der Veranstalterseite und die Pflichtfelder von
    grosse_events.json."""
    import datetime as _dt
    import io as _io
    import json as _json
    import random as _random
    import re as _re

    print("\nInstagram-Logo-Reels (insta_inhalte monat/anmeldung, insta_logos):")
    import insta_inhalte as I
    import insta_logos as L
    import insta_video as V
    from PIL import Image as _Image

    # --- Logo-Suche: das Logo im Startseiten-Link schlägt den Sponsor
    html = ('<html><body><header><a href="/"><img src="/img/logo.svg" alt="Stadtlauf"></a></header>'
            '<div class="sponsors"><img class="logo" src="/s/garmin_logo.png"></div>'
            '<footer><img src="/f/partner-logo.png"></footer>'
            '<link rel="apple-touch-icon" href="/apple.png"></body></html>')
    k = L.kandidaten(html, "https://www.stadtlauf.de/", "Stadtlauf")
    check("Logo-Suche: Kopf-Logo zuerst, Sponsoren und Fuß nie, Touch-Icon zuletzt",
          [u for _, u in k], ["https://www.stadtlauf.de/img/logo.svg", "https://www.stadtlauf.de/apple.png"])
    check("Logo-Suche: Sponsor im Umfeld (class=sponsors) fliegt raus",
          any("garmin" in u for _, u in k), False)

    # --- Aufbereiten: Rand weg, helles Logo erkannt, zu kleines abgewiesen
    weiss = _Image.new("RGBA", (600, 400), (0, 0, 0, 0))
    weiss.paste((255, 255, 255, 255), (100, 150, 280, 250))     # zwei „Buchstaben" mit Lücke
    weiss.paste((255, 255, 255, 255), (320, 150, 500, 250))
    puffer = _io.BytesIO()
    weiss.save(puffer, "PNG")
    erg = L.aufbereiten(puffer.getvalue())
    check("Aufbereiten: transparenter Rand abgeschnitten, weißes Logo = hell (Navy-Karte)",
          (erg[0].size, erg[1]), ((400, 100), True))
    klein = _io.BytesIO()
    _Image.new("RGB", (40, 40), (200, 0, 0)).save(klein, "PNG")
    check("Aufbereiten: zu kleines Bild ist kein Logo", L.aufbereiten(klein.getvalue()), None)
    check("Slug ohne Umlaute", L.slug("Zürich Marathon – Läufe"), "zuerich-marathon-laeufe")

    # --- Testdaten: sechs große Events im November
    def zeile(name, ort, tag, sport="Laufen", km=42.195):
        return {"name": name, "standort": ort, "land": "Deutschland", "art1": sport, "art2": "Straße",
                "datum_start": tag, "datum_ende": tag, "lat": 50.0, "lon": 8.0, "laenge_km": km,
                "veranstalter_url": "https://x.de/"}
    ev = [zeile(f"Dorflauf {i}", "Hinterwald", f"2026-11-{i + 1:02d}", km=10.0) for i in range(20)]
    ev += [zeile("Mainufer Marathon", "Frankfurt", "2026-11-08"), zeile("Mainufer Halbmarathon", "Frankfurt", "2026-11-08", km=21.1),
           zeile("Domstadt-Lauf", "Köln", "2026-11-15"), zeile("Isar Triathlon", "München", "2026-11-21", "Triathlon", 51.5),
           zeile("Elbe Run", "Hamburg", "2026-11-29"), zeile("Kinder Elbe Run", "Hamburg", "2026-11-29", km=2.0),
           zeile("Saale Lauf", "Halle", "2026-11-01"), zeile("Rhein Lauf", "Bonn", "2026-11-22"),
           zeile("Domstadt-Lauf", "Köln", "2027-04-25")]
    grosse = [
        {"name": "Mainufer Marathon", "suche": ["mainufer"], "land": "Deutschland", "sport": "Laufen", "starter": 20000,
         "anmeldestart": "2026-11-03", "webseite": "https://mainufer.de/"},
        {"name": "Domstadt-Lauf", "suche": ["domstadt-lauf"], "land": "Deutschland", "sport": "Laufen", "starter": 9000,
         "anmeldestart": "2026-11", "webseite": "https://domstadt.de/"},
        {"name": "Isar Triathlon", "suche": ["isar triathlon"], "land": "Deutschland", "sport": "Triathlon", "starter": 1500,
         "anmeldestart": "2026-11-20", "termin": "2027-07-04"},
        {"name": "Elbe Run", "suche": ["elbe run"], "land": "Deutschland", "sport": "Laufen", "starter": 6000,
         "anmeldestart": "2026-12-01", "webseite": "https://elberun.de/"},
        {"name": "Alpen Radmarathon", "suche": ["alpen rad"], "land": "Österreich", "sport": "Fahrrad", "starter": 4000,
         "anmeldestart": "2026-11-10", "termin": "2027-08-29"},
        {"name": "Saale Lauf", "suche": ["saale lauf"], "land": "Deutschland", "sport": "Laufen", "starter": 900},
        {"name": "Rhein Lauf", "suche": ["rhein lauf"], "land": "Deutschland", "sport": "Laufen", "starter": 1200},
    ]
    check("letzter Freitag im Oktober 2026 ist der 30.", [d for d in range(20, 32) if I.letzter_freitag(_dt.date(2026, 10, d))], [30])
    b = I.plane(_dt.date(2026, 10, 30), ev, momente=I.lade_momente(), clips=[], grosse=grosse)
    logos = [f for f in b["folien"] if f["typ"] == "logo"]
    check("monat: letzter Freitag → Logo-Reel für den nächsten Monat", b["format"], "monat")
    check("monat: erste Folie = Zahl ALLER Events im November", (b["folien"][0]["zahl"], b["folien"][0]["titel"]),
          (str(len(I.veranstaltungen(ev, _dt.date(2026, 11, 1), _dt.date(2026, 11, 30)))), "EVENTS IM NOVEMBER 2026"))
    check("monat: je großes Event ein Logo (Halbmarathon zählt zum Marathon, Kinderlauf nie), nach Datum",
          [(f["name"], f["datum"], f["logo_id"]) for f in logos],
          [("Saale Lauf", "1. Nov", "saale-lauf"), ("Mainufer Marathon", "8. Nov", "mainufer-marathon"),
           ("Domstadt-Lauf", "15. Nov", "domstadt-lauf"), ("Isar Triathlon", "21. Nov", "isar-triathlon"),
           ("Rhein Lauf", "22. Nov", "rhein-lauf"), ("Elbe Run", "29. Nov", "elbe-run")])
    alt = I.MONAT_MAX_LOGOS
    I.MONAT_MAX_LOGOS = 5
    try:
        knapp = I.plane(_dt.date(2026, 10, 30), ev, momente=I.lade_momente(), clips=[], grosse=grosse)
    finally:
        I.MONAT_MAX_LOGOS = alt
    check("monat: bei mehr Events als Platz die größten nach Starterzahl (der kleinste fällt)",
          [f["name"] for f in knapp["folien"] if f["typ"] == "logo"],
          ["Mainufer Marathon", "Domstadt-Lauf", "Isar Triathlon", "Rhein Lauf", "Elbe Run"])
    check("monat: ein Schlag je Logo, Ende mit „Check it out“, Schwarzweiß im November",
          (all(f.get("blitz") for f in logos), b["folien"][-1]["typ"], b["folien"][-1]["titel"], b["stil"]),
          (True, "cta", "Check it out!", "sw"))
    plan = V.zeitplan(b["folien"], schlag=0.5)
    check("monat: Logo-Folien stehen genau einen Schlag", {round(e - s, 2) for s, e, f in plan if f["typ"] == "logo"}, {0.5})
    weniger = I.plane(_dt.date(2026, 10, 30), ev, momente=I.lade_momente(), clips=[], grosse=grosse[:5])
    check("monat: unter sechs Logos gilt der normale Plan", weniger["format"] != "monat", True)

    a = I.plane(_dt.date(2026, 11, 1), ev, momente=I.lade_momente(), clips=[], grosse=grosse)
    karten = [f for f in a["folien"] if f["typ"] == "logo"]
    check("anmeldung: am 1. die großen Events, deren Anmeldung in dem Monat öffnet", a["format"], "anmeldung")
    check("anmeldung: Anmeldestart (genau oder nur Monat) und nächster Renntermin, nach Anmeldestart",
          [(f["name"], f["ueber"], f["datum"]) for f in karten],
          [("Domstadt-Lauf", "Anmeldung im November", "15. Nov 2026"),
           ("Mainufer Marathon", "Anmeldung ab 3. Nov", "8. Nov 2026"),
           ("Alpen Radmarathon", "Anmeldung ab 10. Nov", "29. Aug 2027"),
           ("Isar Triathlon", "Anmeldung ab 20. Nov", "21. Nov 2026")])
    check("anmeldung: Caption sagt „ohne Gewähr“", "ohne Gewähr" in a["caption"], True)
    check("anmeldung: unter drei Events gilt der Wochenplan",
          I.plane(_dt.date(2026, 12, 1), ev, momente=I.lade_momente(), clips=[], grosse=grosse)["format"] != "anmeldung", True)
    check("Formate monat, anmeldung und top über --format wählbar", {"monat", "anmeldung", "top"} <= set(I.FORMATE), True)

    # --- Top 5 (jeder zweite Samstag): Kategorie, Rangfolge, Rotation, Karte
    def top_zeilen(name, ort, tag, km, sport="Laufen", art2="Straße", land="Deutschland"):
        return dict(zeile(name, ort, tag, sport, km), art2=art2, land=land)
    tev = []
    tgross = []
    for i, (n, st) in enumerate([("Alpha", 9000), ("Bravo", 7000), ("Charlie", 5000), ("Delta", 3000),
                                 ("Echo", 2000), ("Foxtrot", 1000)]):
        tev.append(top_zeilen(f"{n} Marathon", f"Ort{i}", f"2027-0{i + 3}-1{i}", 42.195))
        tev.append(top_zeilen(f"{n} Marathon", f"Ort{i}", f"2027-0{i + 3}-1{i}", 10.0))
        tgross.append({"name": f"{n} Marathon", "suche": [n.lower() + " marathon"], "land": "Deutschland",
                       "sport": "Laufen", "starter": st})
    tev.append(top_zeilen("Golf Halbmarathon", "OrtG", "2027-05-02", 21.1))
    tgross.append({"name": "Golf Halbmarathon", "suche": ["golf halbmarathon"], "land": "Deutschland",
                   "sport": "Laufen", "starter": 99999})
    kat = {k[0]: k for k in I.TOP_KATEGORIEN}
    check("top: Marathon-Kategorie, größte zuerst, nur Veranstaltungen mit 42,2 km",
          [e["g"]["name"] for e in I.top_kandidaten(_dt.date(2026, 11, 14), tev, tgross, kat["marathon"])][:3],
          ["Alpha Marathon", "Bravo Marathon", "Charlie Marathon"])
    check("top: Halbmarathon-Kategorie ohne Veranstaltungen mit Marathon",
          [e["g"]["name"] for e in I.top_kandidaten(_dt.date(2026, 11, 14), tev, tgross, kat["halbmarathon"])],
          ["Golf Halbmarathon"])
    t = I.format_top(_dt.date(2026, 11, 14), tev, _random.Random(1), tgross)
    ranglog = [(f["aktiv"], f["eintraege"][f["aktiv"]]["rang"]) for f in t["folien"]
               if f["typ"] == "top" and f["aktiv"] is not None]
    check("top: Countdown 5 → 1, dann die Übersicht, Schluss „Check it out“",
          ([r for _, r in ranglog], [f["typ"] for f in t["folien"]][-2:], t["folien"][-2]["aktiv"]),
          ([5, 4, 3, 2, 1], ["top", "cta"], None))
    check("top: die fünf größten, Caption sagt, wie gezählt wird",
          ([e["name"] for e in t["folien"][1]["eintraege"]][::-1],
           "allen Teilnehmenden der Veranstaltung" in t["caption"]),
          (["Alpha Marathon", "Bravo Marathon", "Charlie Marathon", "Delta Marathon", "Echo Marathon"], True))
    samstage = [_dt.date(2026, 10, 31) + _dt.timedelta(days=14 * i) for i in range(4)]
    check("top: jeder zweite Samstag, sonst „saison“",
          [I.format_am(_dt.date(2026, 10, 31) + _dt.timedelta(days=7 * i)) for i in range(4)],
          ["top", "saison", "top", "saison"])
    ganz = I.lade_grosse()
    ev_echt = I.lade_events()
    kats = [I.format_top(d, ev_echt, _random.Random(1), ganz) for d in samstage]
    kats = [k["kategorie"] for k in kats if k]
    check("top: am echten Bestand keine Kategorie zweimal hintereinander",
          all(a != b for a, b in zip(kats, kats[1:])), True)
    for design in I.TOP_DESIGNS:
        for aktiv in (0, None):
            bild = V.folie_bild(dict(t["folien"][1], variante=design, aktiv=aktiv), "Laufen")
            check(f"top: Design {design} zeichnet ({'aktiv' if aktiv == 0 else 'Übersicht'})",
                  bild.getchannel("A").getbbox() is not None, True)
    karte = V.karte_bild(400, 400, [(52.52, 13.40, "1")], 0)
    check("Karte: Länderfläche aus laender.json und Punkt Berlin im Nordosten",
          (karte.getpixel((200, 250))[3] > 0, karte.getpixel((5, 5))[3]), (True, 0))

    # --- Folie ohne Logo im Cache: Namenskarte statt Absturz
    bild = V.folie_bild(dict(logos[0], logo_id="gibt-es-nicht"), "Laufen")
    check("Logo-Folie ohne Logo zeichnet die Namenskarte (Karte weiß, mittig)",
          bild.getpixel(((V.KARTE[0] + V.KARTE[2]) // 2, V.KARTE[1] + 20))[:3], V.WEISS)

    # --- grosse_events.json: Pflichtfelder und Quellen
    wurzel = Path(__file__).resolve().parent.parent
    liste = _json.loads((wurzel / "marketing" / "instagram" / "grosse_events.json").read_text(encoding="utf-8"))["events"]
    fehler = []
    for g in liste:
        n = g.get("name", "?")
        if not g.get("suche") or not all(isinstance(x, str) and x == x.lower() and x for x in g["suche"]):
            fehler.append(f"{n}: suche")
        if g.get("land") not in I.LAND_KURZ or g.get("sport") not in I.SPORTARTEN:
            fehler.append(f"{n}: land/sport")
        if g.get("starter") is None:
            if "ironman" not in n.lower():
                fehler.append(f"{n}: starter fehlt (nur bei IRONMAN ohne veröffentlichte Zahl erlaubt)")
        elif not isinstance(g["starter"], int) or not str(g.get("starter_quelle", "")).startswith("https://"):
            fehler.append(f"{n}: starter/quelle")

        if g.get("anmeldestart") is not None:
            if not _re.fullmatch(r"\d{4}-\d{2}(-\d{2})?", str(g["anmeldestart"])) or not I.anmeldestart(g):
                fehler.append(f"{n}: anmeldestart")
            if not str(g.get("anmeldestart_quelle") or "").startswith("http"):
                fehler.append(f"{n}: anmeldestart ohne Quelle")
        if g.get("webseite") and not g["webseite"].startswith("https://"):
            fehler.append(f"{n}: webseite")
    check("grosse_events.json: Pflichtfelder, Quellen, Datumsformat", fehler, [])
    namen = [L.slug(g["name"]) for g in liste]
    check("grosse_events.json: Namen eindeutig (Logo-Datei je Name)", len(namen), len(set(namen)))
    ordner = wurzel / "marketing" / "instagram" / "logos"
    dateien = {p.stem for p in ordner.glob("*.png")}
    meta = _json.loads((ordner / "logos.json").read_text(encoding="utf-8"))
    check("logos/: jede Datei gehört zu einem großen Event und hat Quelle + Prüfdatum in logos.json",
          (sorted(dateien - set(namen)), sorted(d for d in dateien if not (meta.get(d) or {}).get("geprueft"))), ([], []))
    check("Der tägliche Lauf holt keine Logos aus dem Netz (nur geprüfte Dateien)",
          "logo_fuer" in (wurzel / "scripts" / "insta_lauf.py").read_text(encoding="utf-8"), False)


def main() -> int:
    for test in (test_distanz, test_rundung, test_kategorie, test_land, test_statistik, test_instagram, test_instagram_konten, test_insta_freigabe, test_insta_logos,
                 test_wettbewerbe, test_hoehenprofil, test_offizieller_link,
                 test_duplikate, test_namensvereinheitlichung, test_namensrand,
                 test_vergangene_events, test_zeitrennen, test_kalenderdateien, test_web_data,
                 test_meldungen,
                 test_ortsverzeichnis, test_js_syntax, test_abo_rhythmen,
                 test_override_schluessel, test_suche_uebersetzungen,
                 test_fremde_sportart, test_zwei_rennen_in_einer_zeile,
                 test_koordinaten_widerspruch, test_override_koordinaten,
                 test_zwei_sportarten_im_namen, test_audit_pruefungen, test_walking, test_kinder, test_rundenlaenge_zeitrennen, test_dauer_minuten, test_pruefung_offen, test_korrektur_seite, test_kopierschutz,
                 test_stundenlauf, test_such_vorschlaege,
                 test_nicht_ausdauer, test_staffeln, test_datum_vorlaeufig, test_laufen_weiterleitung, test_veranstalter_links, test_neue_quellen, test_serientermin_im_label,
                 test_schwimmen_regeln, test_schwimmkalender, test_faellt_aus, test_check_again, test_geloescht, test_abo_meldung, test_abo_mail, test_turbosport, test_radsportevents,
                 test_cyclingaustria, test_swimsports, test_fsieben, test_datasport, test_kleine_radquellen, test_kilometerliebe, test_dealgrid, test_marsch_quellen, test_lvpfalz_lck, test_siebter_weg, test_gedaechtnis, test_serien,
                 test_kalender_staging,
                 test_mehrsport_teilstrecken,
                 test_manuelle_events,
                 test_keine_fremden_dateien, test_laender_maske,
                 test_asset_stempel, test_farbschema_skript, test_wartung, test_vorschau_sperre,
                 test_lauf_im_triathlonkalender,
                 test_wettbewerb_zusatz, test_triathlon_format, test_charity_merkmal,
                 test_triathlon_format_kopie):
        test()

    print()
    if _failures:
        print(f"❌ {len(_failures)} Test(s) fehlgeschlagen:")
        for line in _failures:
            print(f"  - {line}")
        return 1
    print("✅ Alle Tests bestanden.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
