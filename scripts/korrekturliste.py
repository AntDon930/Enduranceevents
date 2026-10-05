#!/usr/bin/env python3
"""Korrekturliste als Excel - für die Prüfung durch den Nutzer und zurück.

Warum es dieses Skript gibt
---------------------------
Der Nutzer hat am 05.10.2026 um eine Excel-Liste gebeten: alle Events,
bei denen Angaben fehlen oder die zweifelhaft sind, mit Markierung, WAS
fehlt oder fraglich ist. Er prüft sie und schickt eine korrigierte
Fassung im selben Format zurück. Dieses Skript macht beides:

    python3 scripts/korrekturliste.py export                 # -> korrekturliste.xlsx
    python3 scripts/korrekturliste.py import datei.xlsx --dry-run
    python3 scripts/korrekturliste.py import datei.xlsx      # schreibt die Korrekturen

Was in die Liste kommt
----------------------
1. alle Meldungen von `audit_events.pruefe_event()` für Zeilen, die noch
   NICHT in `geprueft.json` stehen (dieselbe Auswahl wie `--offen`),
2. vorläufige Termine (`datum_vorlaeufig`),
3. Veranstaltungen, die die Einzelprüfung als „unklar" abgelegt hat -
   mit der Begründung aus `geprueft.json`.
Je Strecke eine Zeile; die betroffenen Zellen sind rot (fehlt) oder
orange (fraglich) hinterlegt.

Wie der Rückweg funktioniert
----------------------------
Das Blatt `_original` (ausgeblendet) hält je Nr. die Werte beim Export.
Beim Import zählt nur, was sich GEGENÜBER DIESEM STAND geändert hat - so
ist egal, ob events.json inzwischen ein Datenlauf neu geschrieben hat.
Jede Änderung wird ein Eintrag in `manual_overrides.json` unter dem
genauesten Schlüssel ("<Name>|<Datum>|<km>", sonst
"<Name>|<Datum>|@<Wettbewerb>", sonst "<Name>|<Datum>"), „Löschen"
wird `exclude: true`, und der Status landet in `geprueft.json`. Danach
`clean_events.py` und `build_ics.py` laufen lassen (macht `import`
selbst, außer mit `--dry-run` oder `--ohne-aufraeumen`).
"""
from __future__ import annotations

import argparse
import datetime
import json
import re
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from audit_events import ist_geprueft, lade_geprueft, pruefe_event, GEPRUEFT_PATH  # noqa: E402
from scraper_lib import (EVENTS_JSON_PATH, MANUAL_OVERRIDES_PATH,  # noqa: E402
                         orte_aus_places, override_keys)

try:
    from openpyxl import Workbook, load_workbook  # noqa: E402
except ImportError:
    sys.exit("openpyxl fehlt: pip install openpyxl")
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
STANDARD_DATEI = REPO / "korrekturliste.xlsx"

BLATT = "Events prüfen"
BLATT_ORIGINAL = "_original"

# (Feld in events.json, Spaltenkopf, Breite). Die Reihenfolge ist die
# der Tabelle; "name" ist nicht änderbar (kein Override-Feld - der Name
# ist Teil des Schlüssels).
SPALTEN = [
    ("_nr", "Nr", 7),
    ("datum_start", "Datum Start", 13),
    ("datum_ende", "Datum Ende", 13),
    ("datum_vorlaeufig", "Termin vorläufig", 11),
    ("name", "Name (nicht änderbar)", 40),
    ("wettbewerb", "Wettbewerb / Strecke", 30),
    ("art1", "Sportart", 12),
    ("art2", "Kategorie", 14),
    ("laenge_km", "Distanz (km)", 11),
    ("dauer_h", "Dauer (h)", 9),
    ("standort", "Ort", 20),
    ("land", "Land", 16),
    ("veranstalter_url", "Veranstalter-Link", 45),
    ("lat", "Breite", 10),
    ("lon", "Länge", 10),
    ("_hinweis", "Was fehlt / was ist fraglich", 55),
    ("_status", "Status", 20),
    ("_anmerkung", "Deine Anmerkung", 35),
]
FELDER = [f for f, _, _ in SPALTEN if not f.startswith("_")]
AENDERBAR = [f for f in FELDER if f != "name"]
SPALTE = {f: i + 1 for i, (f, _, _) in enumerate(SPALTEN)}

STATUS_WERTE = ["Stimmt so", "Korrigiert", "Löschen", "Unklar"]
SPORTARTEN = ["Laufen", "Fahrrad", "Schwimmen", "Triathlon"]
KATEGORIEN = ["Straße", "Trail", "Mountainbike", "Cyclecross", "Gravel", "Hindernis",
              "Duathlon", "Bahn", "Backyard Ultra", "Freiwasser", "Cross", "Swimrun",
              "Becken", "Zeitfahren", "Aquathlon", "Indoor"]
LAENDER = ["Deutschland", "Österreich", "Schweiz", "Italien (Südtirol)"]

ROT = PatternFill("solid", fgColor="F8C9C4")       # fehlt
ORANGE = PatternFill("solid", fgColor="FCE2B0")    # fraglich
GRAU = PatternFill("solid", fgColor="EDEDED")      # nicht änderbar
KOPF = PatternFill("solid", fgColor="0B1B33")
SCHRIFT = "Arial"

# Audit-Kategorie -> (Felder, "fehlt" oder "fraglich", Text für die Liste).
# Was hier nicht steht, wird als "fraglich" am Namen markiert.
KATEGORIE_FELDER: dict[str, tuple[tuple[str, ...], str, str]] = {
    "weder Distanz noch Dauer": (("laenge_km", "dauer_h"), "fehlt", "Distanz (oder Dauer) fehlt"),
    "kein Veranstalter-Link": (("veranstalter_url",), "fehlt", "Veranstalter-Link fehlt"),
    "Portallink statt offizieller Seite": (("veranstalter_url",), "fraglich",
                                           "Link führt auf ein Portal, nicht auf die Veranstalterseite"),
    "keine Koordinaten": (("lat", "lon", "standort"), "fehlt",
                          "Kartenposition fehlt (Ort prüfen - Koordinaten ergänze ich dann)"),
    "kein Ort": (("standort",), "fehlt", "Ort fehlt"),
    "kein Land": (("land",), "fehlt", "Land fehlt"),
    "Wochentag Mo-Do": (("datum_start",), "fraglich", "Termin an einem Werktag (Mo-Do) - Datum prüfen"),
    "Veranstaltung dauert über eine Woche": (("datum_start", "datum_ende"), "fraglich",
                                             "dauert über eine Woche - Start/Ende prüfen"),
    "Ende liegt vor dem Start": (("datum_start", "datum_ende"), "fraglich", "Ende liegt vor dem Start"),
    "Datum unlesbar": (("datum_start",), "fraglich", "Datum unlesbar"),
    "Jahreszahl im Namen passt nicht zum Datum": (("datum_start",), "fraglich",
                                                  "Jahreszahl im Namen passt nicht zum Datum"),
    "auffällige Distanz": (("laenge_km",), "fraglich", "ungewöhnliche Distanz"),
    "„Marathon“, aber Distanz passt nicht": (("laenge_km",), "fraglich", "„Marathon“, aber die Distanz passt nicht"),
    "„Halbmarathon“, aber Distanz passt nicht": (("laenge_km",), "fraglich",
                                                 "„Halbmarathon“, aber die Distanz passt nicht"),
    "Zahl im Label weicht von laenge_km ab": (("laenge_km", "wettbewerb"), "fraglich",
                                              "km im Streckennamen ≠ Distanz"),
    "Rundenlänge statt Renndistanz": (("laenge_km",), "fraglich", "evtl. Rundenlänge statt Gesamtdistanz"),
    "Kinderlauf-Label mit Erwachsenendistanz": (("laenge_km", "wettbewerb"), "fraglich",
                                                "Kinderlauf mit Erwachsenendistanz?"),
    "Zeitrennen im Namen, aber eine Distanz gesetzt": (("laenge_km", "dauer_h"), "fraglich",
                                                       "Zeitrennen (Stunden) - Distanz oder Dauer?"),
    "Distanz UND Dauer gesetzt": (("laenge_km", "dauer_h"), "fraglich",
                                  "Distanz UND Dauer gesetzt - nur eins davon"),
    "Koordinaten passen nicht zum Land": (("land", "lat", "lon"), "fraglich",
                                          "Kartenposition liegt nicht im angegebenen Land"),
    "Trail/Cross im Namen, andere Kategorie": (("art2",), "fraglich", "Trail/Cross im Namen, Kategorie passt nicht"),
    "Berglauf im Namen, andere Kategorie": (("art2",), "fraglich", "Berglauf im Namen, Kategorie passt nicht"),
    "Testeintrag der Quelle?": (("name",), "fraglich", "Testeintrag der Quelle? (evtl. löschen)"),
    "Name sieht unsauber aus": (("name",), "fraglich", "Name sieht unsauber aus"),
    "Name komplett in Großbuchstaben": (("name",), "fraglich", "Name komplett in Großbuchstaben"),
    "Name sehr lang (Beschreibung statt Name?)": (("name",), "fraglich", "Name sehr lang"),
}


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def sammle(events: list[dict], heute: str) -> list[tuple[dict, list[tuple[tuple[str, ...], str, str]]]]:
    """Die Zeilen der Liste: (Event, [(Felder, Art, Text), …])."""
    protokoll = lade_geprueft()
    unklar = {k: v for k, v in protokoll.items() if v.get("ergebnis") == "unklar"}
    zeilen = []
    for e in events:
        if (e.get("datum_start") or "") < heute:
            continue
        funde: list[tuple[tuple[str, ...], str, str]] = []
        if not ist_geprueft(protokoll, e):
            for kategorie, _ in pruefe_event(e):
                funde.append(KATEGORIE_FELDER.get(kategorie, (("name",), "fraglich", kategorie)))
        if e.get("datum_vorlaeufig"):
            funde.append((("datum_start", "datum_vorlaeufig"), "fraglich",
                          "Termin nur geschätzt (Veranstalter nennt das Datum noch nicht)"))
        eintrag = unklar.get(f"{(e.get('name') or '').strip()}|{e.get('datum_start')}".casefold())
        if eintrag:
            funde.append((("name",), "fraglich",
                          f"Frühere Prüfung unklar: {eintrag.get('notiz') or 'ohne Begründung'}"))
        if funde:
            zeilen.append((e, funde))

    # Exakt doppelte Zeilen (wie audit_events.pruefe_bestand)
    schluessel = lambda e: ((e.get("name") or "").casefold(), e.get("datum_start"),  # noqa: E731
                            e.get("laenge_km"), (e.get("wettbewerb") or "").casefold())
    anzahl: dict = {}
    for e in events:
        anzahl[schluessel(e)] = anzahl.get(schluessel(e), 0) + 1
    vorhanden = {id(e) for e, _ in zeilen}
    for e in events:
        if anzahl[schluessel(e)] > 1 and (e.get("datum_start") or "") >= heute:
            hinweis = (("name",), "fraglich", "Zeile doppelt vorhanden - eine davon löschen")
            if id(e) in vorhanden:
                next(f for x, f in zeilen if x is e).append(hinweis)
            else:
                zeilen.append((e, [hinweis]))
                vorhanden.add(id(e))

    zeilen.sort(key=lambda z: (z[0].get("datum_start") or "", (z[0].get("name") or "").casefold(),
                               z[0].get("laenge_km") or 0))
    return zeilen


def als_zelle(feld: str, wert):
    if wert is None:
        return None
    if feld in ("datum_start", "datum_ende"):
        try:
            return datetime.date.fromisoformat(wert)
        except ValueError:
            return wert
    if feld == "datum_vorlaeufig":
        return "ja" if wert else None
    return wert


def export(ziel: Path) -> None:
    events = json.loads(EVENTS_JSON_PATH.read_text(encoding="utf-8"))
    heute = datetime.date.today().isoformat()
    zeilen = sammle(events, heute)

    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    duenn = Side(style="thin", color="D0D0D0")
    rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)

    for i, (_, kopf, breite) in enumerate(SPALTEN, start=1):
        c = ws.cell(row=1, column=i, value=kopf)
        c.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        c.fill = KOPF
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.row_dimensions[1].height = 32

    orig = wb.create_sheet(BLATT_ORIGINAL)
    orig.append(["Nr", "Werte beim Export (JSON) - nicht ändern"])

    for nr, (e, funde) in enumerate(zeilen, start=1):
        r = nr + 1
        markiert: dict[str, PatternFill] = {}
        for felder, art, _ in funde:
            for f in felder:
                if markiert.get(f) is ROT:
                    continue
                markiert[f] = ROT if art == "fehlt" else ORANGE
        texte = []
        for _, art, text in funde:
            t = f"{'FEHLT' if art == 'fehlt' else 'FRAGLICH'}: {text}"
            if t not in texte:
                texte.append(t)

        for feld, _, _ in SPALTEN:
            col = SPALTE[feld]
            if feld == "_nr":
                wert = nr
            elif feld == "_hinweis":
                wert = "\n".join(texte)
            elif feld.startswith("_"):
                wert = None
            else:
                wert = als_zelle(feld, e.get(feld))
            c = ws.cell(row=r, column=col, value=wert)
            c.font = Font(name=SCHRIFT, size=10)
            c.border = rahmen
            c.alignment = Alignment(vertical="top", wrap_text=feld in ("_hinweis", "_anmerkung", "name"))
            if feld in ("datum_start", "datum_ende"):
                c.number_format = "DD.MM.YYYY"
            elif feld in ("laenge_km", "dauer_h"):
                c.number_format = "0.0##"
            elif feld in ("lat", "lon"):
                c.number_format = "0.0000"
            if feld == "name" or feld == "_nr" or feld == "_hinweis":
                c.fill = GRAU
            if feld in markiert:
                c.fill = markiert[feld]
            if feld == "veranstalter_url" and wert:
                c.hyperlink = wert
                c.font = Font(name=SCHRIFT, size=10, color="1F4E99", underline="single")
        orig.append([nr, json.dumps({f: e.get(f) for f in FELDER}, ensure_ascii=False)])

    letzte = len(zeilen) + 1
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{letzte}"

    def liste(feld: str, werte: list[str]) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(werte) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Ungültiger Wert",
                            error="Bitte einen Wert aus der Liste wählen.")
        sp = get_column_letter(SPALTE[feld])
        dv.add(f"{sp}2:{sp}{max(letzte, 2)}")
        ws.add_data_validation(dv)

    liste("_status", STATUS_WERTE)
    liste("art1", SPORTARTEN)
    liste("art2", KATEGORIEN)
    liste("land", LAENDER)
    liste("datum_vorlaeufig", ["ja", "nein"])
    ws.cell(row=1, column=SPALTE["name"]).comment = Comment(
        "Der Name ist Teil des Schlüssels und lässt sich hier nicht ändern. "
        "Ist er falsch, bitte in 'Deine Anmerkung' schreiben.", "Endurance Events")

    orig.sheet_state = "hidden"
    # openpyxl speichert Formeln ohne Ergebnis - Excel soll die Übersicht
    # beim Öffnen rechnen.
    wb.calculation.fullCalcOnLoad = True
    _anleitung(wb, len(zeilen), letzte)
    wb.move_sheet("Anleitung", offset=-2)
    wb.active = 1
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    print(f"{len(zeilen)} Zeilen ({len({(e.get('name'), e.get('datum_start')) for e, _ in zeilen})} "
          f"Veranstaltungen) -> {ziel}")


def _anleitung(wb: Workbook, anzahl: int, letzte: int) -> None:
    ws = wb.create_sheet("Anleitung")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    f = lambda **k: Font(name=SCHRIFT, **k)  # noqa: E731
    zeilen = [
        ("Korrekturliste Endurance Events", None, f(bold=True, size=14)),
        (f"Stand {datetime.date.today():%d.%m.%Y} - {anzahl} Strecken, bei denen etwas fehlt "
         f"oder fraglich ist.", None, f(size=10)),
        None,
        ("So geht's", None, f(bold=True, size=11)),
        ("1.", "Im Blatt „Events prüfen“ die farbigen Zellen ansehen. Spalte „Was fehlt / was ist "
               "fraglich“ sagt, worum es geht.", None),
        ("2.", "Falsches direkt in der Zelle überschreiben, Fehlendes eintragen. Formate: Datum "
               "TT.MM.JJJJ, Distanz in km mit Komma (z. B. 21,1), Dauer in Stunden (z. B. 6).", None),
        ("3.", "In „Status“ wählen: Stimmt so / Korrigiert / Löschen / Unklar. Leer = nicht angesehen "
               "(Änderungen in der Zeile übernehme ich trotzdem).", None),
        ("4.", "Alles, was nicht in eine Zelle passt (z. B. falscher Name, Hinweise), in „Deine "
               "Anmerkung“.", None),
        ("5.", "Datei speichern (als .xlsx, gleiche Spalten) und mir zurückschicken.", None),
        None,
        ("Farben", None, f(bold=True, size=11)),
        ("rot", "Angabe fehlt", None),
        ("orange", "Angabe vorhanden, aber fraglich", None),
        ("grau", "nicht änderbar (Nr., Name, Hinweis)", None),
        None,
        ("Bitte nicht", None, f(bold=True, size=11)),
        ("", "Spalten löschen, umbenennen oder umsortieren, die Spalte „Nr“ ändern oder das "
             "ausgeblendete Blatt „_original“ anfassen - daran erkenne ich, was du geändert hast. "
             "Zeilen löschen oder filtern/sortieren ist in Ordnung.", None),
        ("", "Kartenposition (Breite/Länge): meist reicht es, den Ort zu korrigieren - die "
             "Koordinaten ergänze ich beim Einlesen.", None),
        None,
        ("Beispiel", None, f(bold=True, size=11)),
        ("Vorher", "Distanz leer (rot), Hinweis „FEHLT: Distanz (oder Dauer) fehlt“", None),
        ("Nachher", "Distanz 21,1 · Status „Korrigiert“ · Anmerkung „laut Ausschreibung 2027“", None),
        None,
        ("Übersicht", None, f(bold=True, size=11)),
    ]
    for z in zeilen:
        if z is None:
            ws.append([])
            continue
        a, b, font = z
        ws.append([a, b])
        r = ws.max_row
        ws.cell(row=r, column=1).font = font or f(bold=True, size=10)
        ws.cell(row=r, column=2).font = f(size=10)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        if a in ("rot", "orange", "grau"):
            ws.cell(row=r, column=1).fill = {"rot": ROT, "orange": ORANGE, "grau": GRAU}[a]

    h = get_column_letter(SPALTE["_hinweis"])
    s = get_column_letter(SPALTE["_status"])
    bereich_h = f"'{BLATT}'!${h}$2:${h}${letzte}"
    bereich_s = f"'{BLATT}'!${s}$2:${s}${letzte}"
    uebersicht = [
        ("Strecken in der Liste", f"=COUNTA('{BLATT}'!$A$2:$A${letzte})"),
        ("davon mit fehlender Angabe", f'=COUNTIF({bereich_h},"*FEHLT*")'),
        ("Distanz fehlt", f'=COUNTIF({bereich_h},"*Distanz (oder Dauer) fehlt*")'),
        ("Portallink statt Veranstalterseite", f'=COUNTIF({bereich_h},"*Portal*")'),
        ("Termin nur geschätzt", f'=COUNTIF({bereich_h},"*nur geschätzt*")'),
        ("Termin an einem Werktag", f'=COUNTIF({bereich_h},"*Werktag*")'),
        ("Frühere Prüfung unklar", f'=COUNTIF({bereich_h},"*Prüfung unklar*")'),
        ("schon bearbeitet (Status gesetzt)", f"=COUNTA({bereich_s})"),
        ("noch offen", f"=B{{r1}}-B{{r2}}"),
    ]
    erste = ws.max_row + 1
    for i, (text, formel) in enumerate(uebersicht):
        if "{r1}" in formel:
            formel = formel.format(r1=erste, r2=erste + i - 1)
        ws.append([text, formel])
        r = ws.max_row
        ws.cell(row=r, column=1).font = f(size=10)
        ws.cell(row=r, column=2).font = f(size=10, bold=True)
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

DATUM_RE = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{4})\s*$")


def lies_wert(feld: str, wert):
    """Zellwert -> Wert wie in events.json. Wirft ValueError mit Klartext."""
    if isinstance(wert, str):
        wert = wert.strip()
        if wert == "":
            wert = None
    if wert is None:
        return False if feld == "datum_vorlaeufig" else None
    if feld in ("datum_start", "datum_ende"):
        if isinstance(wert, datetime.datetime):
            return wert.date().isoformat()
        if isinstance(wert, datetime.date):
            return wert.isoformat()
        m = DATUM_RE.match(str(wert))
        if m:
            return datetime.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
        return datetime.date.fromisoformat(str(wert)).isoformat()
    if feld == "datum_vorlaeufig":
        s = str(wert).strip().lower()
        if s in ("ja", "x", "true", "1", "wahr"):
            return True
        if s in ("nein", "false", "0", "falsch"):
            return False
        raise ValueError(f"'{wert}' ist weder ja noch nein")
    if feld in ("laenge_km", "dauer_h", "lat", "lon"):
        if isinstance(wert, (int, float)):
            zahl = float(wert)
        else:
            s = re.sub(r"\s*(km|h|std\.?|stunden)$", "", str(wert), flags=re.I).replace(",", ".")
            try:
                zahl = float(s)
            except ValueError:
                raise ValueError(f"'{wert}' ist keine Zahl") from None
        return round(zahl, 2 if feld in ("laenge_km", "dauer_h") else 7)
    if feld == "art1" and wert not in SPORTARTEN:
        raise ValueError(f"Sportart '{wert}' unbekannt ({', '.join(SPORTARTEN)})")
    if feld == "art2" and wert not in KATEGORIEN:
        raise ValueError(f"Kategorie '{wert}' unbekannt")
    if feld == "land" and wert not in LAENDER:
        raise ValueError(f"Land '{wert}' unbekannt ({', '.join(LAENDER)})")
    if feld == "veranstalter_url" and not re.match(r"^https?://", str(wert)):
        wert = "https://" + str(wert)
    return str(wert)


def gleich(feld: str, a, b) -> bool:
    if feld == "datum_vorlaeufig":
        return bool(a) == bool(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < (1e-6 if feld in ("lat", "lon") else 0.005)
    return a == b


def koordinaten(standort: str, land: str | None) -> tuple[float, float] | None:
    """Ort -> Koordinaten aus places.json, nur bei eindeutigem Treffer."""
    if not standort:
        return None
    try:
        orte = orte_aus_places(land if land in LAENDER else None)
    except (ValueError, KeyError):
        return None
    name = standort.strip().lower()
    treffer = (orte.get(name) or orte.get(re.sub(r"\s*\(.*?\)\s*$", "", name))
               or orte.get(name.split("/")[0].strip()) or [])
    if len({(round(t[1], 2), round(t[2], 2)) for t in treffer}) == 1:
        return treffer[0][1], treffer[0][2]
    return None


def importiere(quelle: Path, dry_run: bool, aufraeumen: bool) -> int:
    wb = load_workbook(quelle)
    if BLATT not in wb.sheetnames or BLATT_ORIGINAL not in wb.sheetnames:
        print(f"Blatt '{BLATT}' oder '{BLATT_ORIGINAL}' fehlt - ist das die Korrekturliste?")
        return 1
    ws, orig_ws = wb[BLATT], wb[BLATT_ORIGINAL]
    original = {int(r[0]): json.loads(r[1]) for r in orig_ws.iter_rows(min_row=2, values_only=True)
                if r[0] is not None}

    kopf = [c.value for c in ws[1]]
    erwartet = [k for _, k, _ in SPALTEN]
    if kopf[:len(erwartet)] != erwartet:
        abw = [f"{e!r} statt {k!r}" for e, k in zip(erwartet, kopf) if e != k]
        print("Die Spalten wurden verändert: " + "; ".join(abw[:5]))
        return 1

    overrides = json.loads(MANUAL_OVERRIDES_PATH.read_text(encoding="utf-8"))
    protokoll = json.loads(GEPRUEFT_PATH.read_text(encoding="utf-8"))
    events = json.loads(EVENTS_JSON_PATH.read_text(encoding="utf-8"))
    # Der Schlüssel "<Name>|<Datum>" trifft ALLE Strecken einer Veranstaltung.
    zeilen_je_event: dict[tuple, int] = {}
    for e in events:
        k = ((e.get("name") or "").strip().casefold(), e.get("datum_start"))
        zeilen_je_event[k] = zeilen_je_event.get(k, 0) + 1

    heute = datetime.date.today().isoformat()
    fehler, aenderungen, geloescht, status_zaehler = [], [], [], {}
    neue_overrides: dict[str, dict] = {}
    neues_protokoll: dict[str, dict] = {}

    for zeile in ws.iter_rows(min_row=2, values_only=True):
        if zeile[0] is None:
            continue
        try:
            nr = int(zeile[0])
        except (TypeError, ValueError):
            fehler.append(f"Nr '{zeile[0]}' ist keine Zahl")
            continue
        alt = original.get(nr)
        if alt is None:
            fehler.append(f"Nr {nr}: nicht im Blatt _original")
            continue
        wert = {f: zeile[SPALTE[f] - 1] for f in FELDER}
        status = (zeile[SPALTE["_status"] - 1] or "").strip() if isinstance(
            zeile[SPALTE["_status"] - 1], str) else ""
        anmerkung = (zeile[SPALTE["_anmerkung"] - 1] or "")
        anmerkung = str(anmerkung).strip()
        bez = f"Nr {nr} {alt['name']!r} {alt['datum_start']}"

        neu: dict = {}
        zeilen_fehler = False
        for f in AENDERBAR:
            try:
                v = lies_wert(f, wert[f])
            except (ValueError, TypeError) as ex:
                fehler.append(f"{bez}: {dict((a, b) for a, b, _ in SPALTEN)[f]}: {ex}")
                zeilen_fehler = True
                continue
            if not gleich(f, v, alt.get(f)):
                neu[f] = v
        if zeilen_fehler:
            continue

        # Ort geändert, Koordinaten nicht: aus places.json nachziehen.
        if "standort" in neu and "lat" not in neu and "lon" not in neu:
            pos = koordinaten(neu["standort"], neu.get("land", alt.get("land")))
            if pos:
                neu["lat"], neu["lon"] = pos
            else:
                fehler.append(f"{bez}: Ort '{neu['standort']}' nicht eindeutig gefunden - "
                              f"Koordinaten bitte von Hand nachtragen (Ort wird trotzdem übernommen)")
        # Nur der Start geändert: Bei eintägigen Events wandert das Ende mit,
        # sonst darf es nicht vor dem neuen Start liegen.
        if "datum_start" in neu and "datum_ende" not in neu:
            if alt.get("datum_ende") in (None, alt["datum_start"]):
                neu["datum_ende"] = neu["datum_start"]
            elif alt["datum_ende"] < neu["datum_start"]:
                fehler.append(f"{bez}: neues Startdatum liegt nach dem Enddatum - Ende bitte mit ändern")
                continue
        # datum_vorlaeufig ändert sich nur, wenn die Zelle geändert wurde -
        # ein neues Datum allein sagt nicht, ob es bestätigt ist.
        if neu.get("datum_vorlaeufig") is False:
            neu["datum_vorlaeufig"] = None  # Feld entfernen statt false speichern

        loeschen = status.lower() == "löschen"
        if not neu and not loeschen and not status and not anmerkung:
            continue

        # Fehlende Koordinaten bei einer BEARBEITETEN Zeile aus dem Ort
        # ergänzen - nicht angefasste Zeilen bleiben, wie sie sind.
        if (not loeschen and wert["lat"] is None and wert["lon"] is None
                and alt.get("lat") is None and "standort" not in neu and alt.get("standort")):
            pos = koordinaten(alt["standort"], neu.get("land", alt.get("land")))
            if pos:
                neu["lat"], neu["lon"] = pos

        # Immer den genauesten Schlüssel: Distanz, sonst Wettbewerb, sonst
        # die ganze Veranstaltung (nur wenn sie aus dieser einen Zeile besteht
        # oder keine Strecke einen Namen hat).
        keys = override_keys(alt["name"], alt["datum_start"], alt.get("laenge_km"),
                             alt.get("wettbewerb"))
        schluessel = keys[0]
        if (neu or loeschen) and len(keys) == 1:
            n = zeilen_je_event.get((alt["name"].strip().casefold(), alt["datum_start"]), 0)
            if n > 1:
                fehler.append(f"{bez}: Strecke ohne Distanz und ohne Namen - die Änderung gilt für "
                              f"ALLE {n} Strecken dieser Veranstaltung. Bitte prüfen.")
        if loeschen:
            neue_overrides.setdefault(schluessel, {})["exclude"] = True
            geloescht.append(bez)
        elif neu:
            neue_overrides.setdefault(schluessel, {}).update(neu)
            aenderungen.append(f"{bez}: " + ", ".join(
                f"{f} {alt.get(f)!r} -> {v!r}" for f, v in neu.items()))
        if neu or loeschen:
            notiz = "Korrekturliste des Nutzers" + (f": {anmerkung}" if anmerkung else "")
            neue_overrides[schluessel]["_note"] = notiz

        ergebnis = {"löschen": "entfernt", "unklar": "unklar", "stimmt so": "quelle_ok",
                    "korrigiert": "korrigiert"}.get(status.lower())
        if ergebnis is None and neu:
            ergebnis = "korrigiert"
        if ergebnis:
            status_zaehler[ergebnis] = status_zaehler.get(ergebnis, 0) + 1
            pkey = f"{alt['name'].strip()}|{neu.get('datum_start', alt['datum_start'])}"
            vorher = neues_protokoll.get(pkey, {}).get("ergebnis")
            # Eine Veranstaltung, deren eine Strecke korrigiert und deren
            # andere bestätigt wurde, ist "korrigiert".
            rang = ["quelle_ok", "unklar", "korrigiert", "entfernt"]
            if vorher and rang.index(vorher) > rang.index(ergebnis):
                ergebnis = vorher
            neues_protokoll[pkey] = {
                "am": heute, "ergebnis": ergebnis,
                "quelle": neu.get("veranstalter_url") or alt.get("veranstalter_url") or "",
                "notiz": "Korrekturliste des Nutzers" + (f": {anmerkung}" if anmerkung else ""),
            }
        elif anmerkung:
            fehler.append(f"{bez}: Anmerkung ohne Änderung/Status: {anmerkung}")

    print(f"Änderungen: {len(aenderungen)}, gelöscht: {len(geloescht)}, "
          f"Status: {status_zaehler or '-'}")
    for z in aenderungen:
        print("  ~ " + z)
    for z in geloescht:
        print("  - " + z)
    if fehler:
        print(f"\nHINWEISE ({len(fehler)}) - bitte ansehen:")
        for z in fehler:
            print("  ! " + z)

    if dry_run:
        print("\n(--dry-run: nichts geschrieben)")
        return 0

    for k, v in neue_overrides.items():
        bestehend = next((x for x in overrides if x.casefold() == k.casefold()), k)
        eintrag = overrides.setdefault(bestehend, {})
        alte_notiz = eintrag.get("_note")
        eintrag.update({f: w for f, w in v.items()})
        if alte_notiz and alte_notiz != v.get("_note"):
            eintrag["_note"] = f"{alte_notiz} | {v.get('_note')}"
    protokoll.update(neues_protokoll)
    MANUAL_OVERRIDES_PATH.write_text(json.dumps(overrides, ensure_ascii=False, indent=2) + "\n",
                                     encoding="utf-8")
    GEPRUEFT_PATH.write_text(json.dumps(protokoll, ensure_ascii=False, indent=2) + "\n",
                             encoding="utf-8")
    print(f"\n{len(neue_overrides)} Einträge in {MANUAL_OVERRIDES_PATH.name}, "
          f"{len(neues_protokoll)} in {GEPRUEFT_PATH.name} geschrieben.")

    if aufraeumen and neue_overrides:
        # --no-geocoding: Ohne den Schalter fragt clean_events.py Nominatim
        # (1 Anfrage/s) und braucht weit über zehn Minuten; Koordinaten
        # zieht dieses Skript selbst aus places.json nach.
        for befehl in (["clean_events.py", "--quiet", "--no-geocoding"], ["build_ics.py"]):
            print(f"\n$ python3 scripts/{' '.join(befehl)}")
            subprocess.run([sys.executable, str(REPO / "scripts" / befehl[0]), *befehl[1:]], check=True)
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = parser.add_subparsers(dest="befehl", required=True)
    p_ex = sub.add_parser("export", help="Excel-Liste erzeugen")
    p_ex.add_argument("ziel", nargs="?", type=Path, default=STANDARD_DATEI)
    p_im = sub.add_parser("import", help="korrigierte Excel-Liste einlesen")
    p_im.add_argument("quelle", type=Path)
    p_im.add_argument("--dry-run", action="store_true", help="nur anzeigen, nichts schreiben")
    p_im.add_argument("--ohne-aufraeumen", action="store_true",
                      help="clean_events.py/build_ics.py danach NICHT laufen lassen")
    args = parser.parse_args()
    if args.befehl == "export":
        export(args.ziel)
        return 0
    return importiere(args.quelle, args.dry_run, not args.ohne_aufraeumen)


if __name__ == "__main__":
    sys.exit(main())
